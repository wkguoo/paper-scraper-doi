"""Shared helpers for DOI batch input, cookie checks, and user reports."""

from __future__ import annotations

import csv
import json
import re
from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable


PREVIEW_LIMIT = 200
TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "gbk", "cp936")
DOI_COLUMN_CANDIDATES = ("doi", "DOI", "Doi", "DOI号", "doi号", "DOI號", "doi號")
COLUMN_ALIASES = {
    "title": ("title", "article title", "paper title", "标题", "题名", "文献标题"),
    "authors": ("authors", "author", "作者", "作者列表"),
    "journal": ("journal", "source", "publication", "期刊", "期刊名称"),
    "year": ("year", "publication year", "年份", "发表年份"),
    "date": ("date", "publication date", "日期", "发表日期"),
    "doi": DOI_COLUMN_CANDIDATES,
}
COOKIE_DOMAINS = ("sciencedirect.com", "elsevier.com", "sciencedirectassets.com")


@dataclass(frozen=True)
class DoiRecord:
    row_number: int
    doi: str
    title: str = ""
    authors: str = ""
    journal: str = ""
    year: str = ""
    date: str = ""
    raw_value: str = ""


@dataclass(frozen=True)
class DoiPreviewRow:
    row_number: int
    doi: str
    title: str = ""
    raw_value: str = ""
    status: str = "valid"
    reason: str = ""


@dataclass(frozen=True)
class DoiPreview:
    rows: list[DoiPreviewRow]
    total_doi: int
    total_rows: int
    source: str
    doi_column: str = ""
    encoding: str = ""
    sheet_name: str = ""
    warnings: list[str] = field(default_factory=list)
    status_counts: dict[str, int] = field(default_factory=dict)


@dataclass(frozen=True)
class CookieCheck:
    path: str
    exists: bool
    readable: bool
    valid_json: bool
    cookie_count: int
    relevant_cookie_count: int
    has_sciencedirect_cookie: bool
    is_usable: bool
    message: str


@dataclass(frozen=True)
class PdfDownloadRecord:
    doi: str
    pii: str
    title: str
    status: str
    file: str = ""
    reason: str = ""


@dataclass(frozen=True)
class SupplementDownloadRecord:
    doi: str
    pii: str
    article_title: str
    article_file: str = ""
    supplement_index: int = 0
    supplement_title: str = ""
    source_url: str = ""
    status: str = ""
    file: str = ""
    content_type: str = ""
    size_bytes: int = 0
    reason: str = ""


@dataclass(frozen=True)
class DownloadRunResult:
    pdf_success: int
    pdf_failed: int
    pdf_skipped: int
    pdf_records: list[PdfDownloadRecord] = field(default_factory=list)
    supplement_success: int = 0
    supplement_failed: int = 0
    supplement_skipped: int = 0
    supplement_not_found: int = 0
    supplement_records: list[SupplementDownloadRecord] = field(default_factory=list)

    def __iter__(self):
        yield self.pdf_success
        yield self.pdf_failed
        yield self.pdf_skipped
        yield self.pdf_records


@dataclass(frozen=True)
class RunSummary:
    input_path: str
    output_dir: str
    total_doi: int
    resolved_count: int
    failure_reasons: dict[str, int]
    pdf_success: int
    pdf_failed: int
    pdf_skipped: int
    resolved_path: str = ""
    failed_path: str = ""
    pdf_report_path: str = ""
    cookie_message: str = ""
    retry_input_path: str = ""
    retry_input_count: int = 0
    retry_input_excluded_count: int = 0
    beginner_recommendations: list[str] = field(default_factory=list)
    supplement_requested: bool = False
    supplement_success: int = 0
    supplement_failed: int = 0
    supplement_skipped: int = 0
    supplement_not_found: int = 0
    supplement_report_path: str = ""
    browser_message: str = ""
    download_next_steps: str = ""


@dataclass(frozen=True)
class RunEvent:
    stage: str
    status: str
    timestamp: str = ""
    row_number: int | str = ""
    doi: str = ""
    title: str = ""
    pii: str = ""
    file: str = ""
    reason: str = ""
    counts: dict[str, int] = field(default_factory=dict)


@dataclass(frozen=True)
class RetryInputResult:
    path: str
    row_count: int
    excluded_count: int
    excluded_reasons: dict[str, int] = field(default_factory=dict)


@dataclass(frozen=True)
class ResumeCandidate:
    path: str
    updated_at: str
    has_summary_json: bool
    has_pdf_report: bool
    input_match: bool
    overlap_count: int
    overlap_total: int
    overlap_ratio: float
    score: float
    reason: str = ""


def normalize_column_name(name: object) -> str:
    return re.sub(r"[\s_\-]+", "", str(name or "")).lower()


def find_column(headers: Iterable[str], candidates: Iterable[str]) -> str | None:
    normalized = {normalize_column_name(header): header for header in headers}
    for candidate in candidates:
        key = normalize_column_name(candidate)
        if key in normalized:
            return normalized[key]
    return None


def row_value(row: dict[str, object], field: str) -> str:
    column = find_column(row.keys(), COLUMN_ALIASES[field])
    if not column:
        return ""
    value = row.get(column, "")
    return "" if value is None else str(value).strip()


def clean_doi(doi: object) -> str:
    value = str(doi or "").strip()
    value = re.sub(r"^(https?://)?(dx\.)?doi\.org/", "", value, flags=re.I)
    value = re.sub(r"^doi\s*:\s*", "", value, flags=re.I)
    value = re.split(r"[\]\}\s<>\"']+", value, maxsplit=1)[0]
    value = value.strip().strip(".,;，。；、")
    while value.endswith(")") and value.count(")") > value.count("("):
        value = value[:-1].rstrip().strip(".,;，。；、")
    return value


def extract_doi_from_text(text: object) -> str:
    match = re.search(r"10\.\d{4,9}/[^\s\"'<>\]\}]+", str(text or ""), flags=re.I)
    if not match:
        return ""
    return clean_doi(match.group(0))


def load_doi_records(
    input_path: str | Path,
    doi_column: str | None = None,
    sheet_name: str | None = None,
) -> list[DoiRecord]:
    path = Path(input_path)
    ext = path.suffix.lower()
    if ext == ".csv":
        records, _, _ = _read_delimited_records(path, ",", doi_column)
        return records
    if ext == ".tsv":
        try:
            records, _, _ = _read_delimited_records(path, "\t", doi_column)
            return records
        except ValueError:
            return _read_text_records(path)
    if ext in {".txt", ".md", ".markdown"}:
        return _read_text_records(path)
    if ext in {".xlsx", ".xlsm"}:
        records, _ = _read_xlsx_records(path, doi_column, sheet_name)
        return records
    raise ValueError("仅支持 .csv、.xlsx、.xlsm、.txt、.md、.markdown、.tsv 文件")


def preview_doi_input(
    input_path: str | Path | None = None,
    pasted_text: str = "",
    doi_column: str | None = None,
    sheet_name: str | None = None,
    limit: int = PREVIEW_LIMIT,
) -> DoiPreview:
    if input_path:
        path = Path(input_path)
        ext = path.suffix.lower()
        if ext == ".csv":
            records, encoding, found_column = _read_delimited_records(path, ",", doi_column)
            return _preview_from_records(records, str(path), found_column, encoding, "", limit)
        if ext == ".tsv":
            try:
                records, encoding, found_column = _read_delimited_records(path, "\t", doi_column)
                return _preview_from_records(records, str(path), found_column, encoding, "", limit)
            except ValueError:
                records = _read_text_records(path)
                return _preview_from_records(records, str(path), "逐行扫描", "", "", limit)
        if ext in {".txt", ".md", ".markdown"}:
            records = _read_text_records(path)
            return _preview_from_records(records, str(path), "逐行扫描", "", "", limit)
        if ext in {".xlsx", ".xlsm"}:
            records, found_sheet = _read_xlsx_records(path, doi_column, sheet_name)
            found_column = _find_doi_column_from_records_source(path, doi_column, found_sheet)
            return _preview_from_records(records, str(path), found_column, "", found_sheet, limit)
        raise ValueError("仅支持 .csv、.xlsx、.xlsm、.txt、.md、.markdown、.tsv 文件")

    records = _records_from_pasted_text(pasted_text)
    return _preview_from_records(records, "粘贴内容", "逐行扫描", "", "", limit)


def check_cookie_json(path: str | Path) -> CookieCheck:
    cookie_path = Path(path)
    if not str(path).strip():
        return CookieCheck("", False, False, False, 0, 0, False, False, "未选择 Cookie JSON 文件")
    if not cookie_path.exists():
        return CookieCheck(str(cookie_path), False, False, False, 0, 0, False, False, "找不到 Cookie JSON 文件")

    data = None
    last_decode_error: Exception | None = None
    for encoding in TEXT_ENCODINGS:
        try:
            data = json.loads(cookie_path.read_text(encoding=encoding))
            break
        except UnicodeDecodeError as exc:
            last_decode_error = exc
        except json.JSONDecodeError as exc:
            return CookieCheck(str(cookie_path), True, True, False, 0, 0, False, False, f"Cookie 文件不是有效 JSON: {exc}")
        except OSError as exc:
            return CookieCheck(str(cookie_path), True, False, False, 0, 0, False, False, f"无法读取 Cookie 文件: {exc}")

    if data is None:
        return CookieCheck(
            str(cookie_path),
            True,
            False,
            False,
            0,
            0,
            False,
            False,
            f"无法识别 Cookie 文件编码: {last_decode_error}",
        )

    cookie_items = _cookie_items_from_json(data)
    cookie_count = len(cookie_items)
    relevant_count = sum(1 for item in cookie_items if _is_relevant_cookie(item))
    has_relevant = relevant_count > 0

    if cookie_count == 0:
        message = "Cookie 文件可读，但没有识别到 Cookie 项"
    elif not has_relevant:
        message = f"已识别 {cookie_count} 个 Cookie，但未发现 ScienceDirect/Elsevier 相关域"
    else:
        message = f"已识别 {cookie_count} 个 Cookie，其中 {relevant_count} 个与 ScienceDirect/Elsevier 相关"

    return CookieCheck(
        str(cookie_path),
        True,
        True,
        True,
        cookie_count,
        relevant_count,
        has_relevant,
        has_relevant,
        message,
    )


def write_pdf_download_report(records: list[PdfDownloadRecord], output_dir: str | Path) -> Path:
    path = Path(output_dir) / "pdf_download_report.csv"
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=["doi", "pii", "title", "status", "file", "reason"])
        writer.writeheader()
        for record in records:
            writer.writerow({
                "doi": record.doi,
                "pii": record.pii,
                "title": record.title,
                "status": record.status,
                "file": record.file,
                "reason": record.reason,
            })
    return path


def write_supplement_download_report(records: list[SupplementDownloadRecord], output_dir: str | Path) -> Path:
    path = Path(output_dir) / "supplement_download_report.csv"
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "doi",
        "pii",
        "article_title",
        "article_file",
        "supplement_index",
        "supplement_title",
        "source_url",
        "status",
        "file",
        "content_type",
        "size_bytes",
        "reason",
    ]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for record in records:
            writer.writerow({
                "doi": record.doi,
                "pii": record.pii,
                "article_title": record.article_title,
                "article_file": record.article_file,
                "supplement_index": record.supplement_index,
                "supplement_title": record.supplement_title,
                "source_url": record.source_url,
                "status": record.status,
                "file": record.file,
                "content_type": record.content_type,
                "size_bytes": record.size_bytes,
                "reason": record.reason,
            })
    return path


def collect_retry_input_rows(
    pdf_report_path: str | Path | None = None,
    doi_failed_path: str | Path | None = None,
    selected_dois: set[str] | None = None,
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    selected = {clean_doi(doi).lower() for doi in (selected_dois or set()) if clean_doi(doi)}

    for row in _read_csv_rows(pdf_report_path):
        if str(row.get("status") or "").strip().lower() != "failed":
            continue
        doi = clean_doi(row.get("doi", ""))
        if not doi:
            continue
        key = doi.lower()
        if selected and key not in selected:
            continue
        if key in seen:
            continue
        seen.add(key)
        rows.append({
            "doi": doi,
            "title": str(row.get("title") or "").strip(),
            "source_status": "pdf_failed",
            "reason": str(row.get("reason") or "").strip(),
        })

    ignored_reasons = {"DOI 为空", "重复 DOI，已跳过"}
    for row in _read_csv_rows(doi_failed_path):
        reason = str(row.get("reason") or "").strip()
        if reason in ignored_reasons:
            continue
        doi = clean_doi(row.get("doi", ""))
        if not doi:
            continue
        key = doi.lower()
        if selected and key not in selected:
            continue
        if key in seen:
            continue
        seen.add(key)
        rows.append({
            "doi": doi,
            "title": str(row.get("title") or "").strip(),
            "source_status": "resolve_failed",
            "reason": reason,
        })

    return rows


def write_retry_input_csv(
    rows: list[dict[str, str]],
    output_dir: str | Path,
    timestamp: str | None = None,
) -> Path:
    stamp = timestamp or datetime.now().strftime("%Y%m%d_%H%M%S")
    path = Path(output_dir) / f"retry_failed_doi_{stamp}.csv"
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=["doi", "title", "source_status", "reason"])
        writer.writeheader()
        for row in rows:
            writer.writerow({
                "doi": row.get("doi", ""),
                "title": row.get("title", ""),
                "source_status": row.get("source_status", ""),
                "reason": row.get("reason", ""),
            })
    return path


def write_retry_input_from_reports(
    pdf_report_path: str | Path | None = None,
    doi_failed_path: str | Path | None = None,
    output_dir: str | Path | None = None,
    timestamp: str | None = None,
    selected_dois: set[str] | None = None,
) -> RetryInputResult:
    rows = collect_retry_input_rows(
        pdf_report_path,
        doi_failed_path,
        selected_dois=selected_dois,
    )
    excluded_reasons = _retry_exclusion_counts(
        pdf_report_path,
        doi_failed_path,
        selected_dois=selected_dois,
    )
    if not rows:
        return RetryInputResult("", 0, sum(excluded_reasons.values()), dict(excluded_reasons))

    target_dir = Path(output_dir) if output_dir else _first_existing_parent(pdf_report_path, doi_failed_path)
    path = write_retry_input_csv(rows, target_dir, timestamp=timestamp)
    return RetryInputResult(str(path), len(rows), sum(excluded_reasons.values()), dict(excluded_reasons))


def write_run_event(path: str | Path, event: RunEvent) -> Path:
    event_path = Path(path)
    event_path.parent.mkdir(parents=True, exist_ok=True)
    payload = asdict(event)
    if not payload["timestamp"]:
        payload["timestamp"] = datetime.now().isoformat(timespec="seconds")
    with event_path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(payload, ensure_ascii=False, sort_keys=True) + "\n")
    return event_path


def safe_write_run_event(path: str | Path | None, event: RunEvent) -> None:
    if not path:
        return
    try:
        write_run_event(path, event)
    except Exception:
        # Structured telemetry must never interrupt a long download task.
        return


def read_run_events(path: str | Path) -> list[dict[str, object]]:
    event_path = Path(path)
    if not event_path.exists():
        return []
    events: list[dict[str, object]] = []
    with event_path.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            events.append(json.loads(line))
    return events


def write_run_summary_json(
    summary: RunSummary,
    event_path: str | Path | None = None,
) -> Path:
    output_dir = Path(summary.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "run_summary.json"
    payload = {
        "input_path": summary.input_path,
        "output_dir": summary.output_dir,
        "total_doi": summary.total_doi,
        "resolved_count": summary.resolved_count,
        "resolve_failed_count": sum(summary.failure_reasons.values()),
        "failure_reasons": summary.failure_reasons,
        "pdf_success": summary.pdf_success,
        "pdf_failed": summary.pdf_failed,
        "pdf_skipped": summary.pdf_skipped,
        "resolved_path": summary.resolved_path,
        "failed_path": summary.failed_path,
        "pdf_report_path": summary.pdf_report_path,
        "cookie_message": summary.cookie_message,
        "retry_input_path": summary.retry_input_path,
        "retry_input_count": summary.retry_input_count,
        "retry_input_excluded_count": summary.retry_input_excluded_count,
        "beginner_recommendations": summary.beginner_recommendations,
        "supplement_requested": summary.supplement_requested,
        "supplement_success": summary.supplement_success,
        "supplement_failed": summary.supplement_failed,
        "supplement_skipped": summary.supplement_skipped,
        "supplement_not_found": summary.supplement_not_found,
        "supplement_report_path": summary.supplement_report_path,
        "browser_message": summary.browser_message,
        "download_next_steps": summary.download_next_steps,
        "event_path": str(event_path) if event_path else "",
        "written_at": datetime.now().isoformat(timespec="seconds"),
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def write_run_summary(summary: RunSummary) -> Path:
    output_dir = Path(summary.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "run_summary.txt"
    lines = [
        "DOI 批量任务报告",
        "=" * 18,
        f"输入来源: {summary.input_path or '粘贴内容/临时文件'}",
        f"输出目录: {summary.output_dir}",
        f"识别 DOI 数: {summary.total_doi}",
        f"成功解析数: {summary.resolved_count}",
        f"解析失败数: {sum(summary.failure_reasons.values())}",
        "",
        "解析失败原因:",
    ]
    if summary.failure_reasons:
        for reason, count in sorted(summary.failure_reasons.items(), key=lambda item: (-item[1], item[0])):
            lines.append(f"- {reason}: {count}")
    else:
        lines.append("- 无")

    lines.extend([
        "",
        "PDF 下载:",
        f"- 成功: {summary.pdf_success}",
        f"- 失败: {summary.pdf_failed}",
        f"- 跳过: {summary.pdf_skipped}",
        "",
        "补充材料下载:",
        f"- 是否请求: {'是' if summary.supplement_requested else '否'}",
    ])
    if summary.supplement_requested:
        lines.extend([
            f"- 成功: {summary.supplement_success}",
            f"- 失败: {summary.supplement_failed}",
            f"- 跳过: {summary.supplement_skipped}",
            f"- 未发现: {summary.supplement_not_found}",
        ])
    else:
        lines.append("- 未请求：PDF 下载未启用或补充材料下载已关闭。")
    lines.extend([
        "",
        "输出文件:",
        f"- 解析成功表: {summary.resolved_path or '未生成'}",
        f"- DOI 失败报告: {summary.failed_path or '未生成'}",
        f"- PDF 下载报告: {summary.pdf_report_path or '未生成'}",
        f"- Supplement 下载报告: {summary.supplement_report_path or '未生成'}",
    ])
    if summary.retry_input_path:
        lines.append(f"- 重试输入表: {summary.retry_input_path}")
    if summary.cookie_message:
        lines.extend(["", f"Cookie 检查: {summary.cookie_message}"])
    lines.extend(["", "下一步建议:"])
    if summary.beginner_recommendations:
        lines.append("")
        lines.append("小白下一步建议:")
        for recommendation in summary.beginner_recommendations:
            lines.append(f"- {recommendation}")
    if summary.retry_input_path:
        lines.append(f"- 已生成重试输入 {summary.retry_input_count} 条；请预览确认后再手动运行。")
    if summary.pdf_failed:
        lines.append("- 若 PDF 大量失败，优先检查 Cookie 是否过期、机构权限是否可访问 PDF、Edge/调试浏览器中是否出现验证码或限速提示。")
    if summary.supplement_failed:
        lines.append("- 若补充材料失败，先查看 supplement_download_report.csv；常见原因是附件链接返回登录页、权限不足或远端响应不是附件文件。")
    if summary.failure_reasons:
        lines.append("- 先查看 doi_batch_failed.csv，非 ScienceDirect DOI 或重复 DOI 不会中断整个任务。")
    if not summary.failure_reasons and not summary.pdf_failed and not summary.supplement_failed:
        lines.append("- 任务完成，无需处理。")

    if summary.browser_message:
        lines.extend(["", f"Browser: {summary.browser_message}"])
    if summary.download_next_steps:
        for step in summary.download_next_steps.splitlines():
            step = step.strip()
            if step:
                lines.append(f"- {step}")

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def failure_reason_counts(failures: Iterable[dict[str, object]]) -> dict[str, int]:
    counter: Counter[str] = Counter()
    for item in failures:
        reason = str(item.get("reason") or "未知原因")
        counter[reason] += 1
    return dict(counter)


def load_resume_success_dois(output_dir: str | Path | None) -> set[str]:
    if not output_dir:
        return set()
    pdf_report = Path(output_dir) / "pdf_download_report.csv"
    success: set[str] = set()
    for row in _read_csv_rows(pdf_report):
        if str(row.get("status") or "").strip().lower() != "success":
            continue
        doi = clean_doi(row.get("doi", ""))
        if doi:
            success.add(doi.lower())
    return success


def find_resume_candidates(
    output_root: str | Path | None,
    input_path: str | Path | None = None,
    current_dois: Iterable[str] | None = None,
    limit: int = 5,
) -> list[ResumeCandidate]:
    if not output_root:
        return []
    root = Path(output_root)
    if not root.exists() or not root.is_dir():
        return []

    input_name = Path(input_path).name.lower() if input_path else ""
    current = {clean_doi(doi).lower() for doi in (current_dois or []) if clean_doi(doi)}
    candidates: list[ResumeCandidate] = []
    for child in root.iterdir():
        if not child.is_dir():
            continue
        summary_path = child / "run_summary.json"
        pdf_report_path = child / "pdf_download_report.csv"
        if not summary_path.exists() and not pdf_report_path.exists():
            continue

        summary_data = _read_json_object(summary_path)
        summary_input = str(summary_data.get("input_path") or "") if summary_data else ""
        input_match = bool(input_name and summary_input and Path(summary_input).name.lower() == input_name)
        old_dois = _collect_report_dois(child)
        overlap_count = len(current & old_dois) if current and old_dois else 0
        overlap_total = min(len(current), len(old_dois)) if current and old_dois else 0
        overlap_ratio = (overlap_count / overlap_total) if overlap_total else 0.0
        updated_ts = _candidate_mtime(child, summary_path, pdf_report_path)
        updated_at = datetime.fromtimestamp(updated_ts).isoformat(timespec="seconds")
        score = (
            (100.0 if input_match else 0.0)
            + overlap_ratio * 50.0
            + (10.0 if pdf_report_path.exists() else 0.0)
            + (5.0 if summary_path.exists() else 0.0)
            + min(updated_ts / 10_000_000_000, 1.0)
        )
        reasons = []
        if input_match:
            reasons.append("输入文件名一致")
        if overlap_count:
            reasons.append(f"DOI 重合 {overlap_count} 条")
        if not reasons:
            reasons.append("最近历史结果")
        candidates.append(ResumeCandidate(
            path=str(child),
            updated_at=updated_at,
            has_summary_json=summary_path.exists(),
            has_pdf_report=pdf_report_path.exists(),
            input_match=input_match,
            overlap_count=overlap_count,
            overlap_total=overlap_total,
            overlap_ratio=overlap_ratio,
            score=score,
            reason="；".join(reasons),
        ))
    candidates.sort(key=lambda item: (item.score, item.updated_at), reverse=True)
    return candidates[:limit]


def filter_records_for_resume(
    records: Iterable[dict[str, object]],
    success_dois: set[str],
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    kept: list[dict[str, object]] = []
    skipped: list[dict[str, object]] = []
    normalized_success = {doi.lower() for doi in success_dois}
    for record in records:
        doi = clean_doi(record.get("doi", ""))
        if doi and doi.lower() in normalized_success:
            skipped.append({
                "row_number": record.get("row_number", ""),
                "doi": doi,
                "title": record.get("title", ""),
                "reason": "已在历史结果中成功下载 PDF，断点恢复跳过",
            })
        else:
            kept.append(dict(record))
    return kept, skipped


def collect_failure_table_rows(
    pdf_report_path: str | Path | None = None,
    doi_failed_path: str | Path | None = None,
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()

    for row in _read_csv_rows(pdf_report_path):
        status = str(row.get("status") or "").strip().lower()
        if status not in {"failed", "skipped"}:
            continue
        doi = clean_doi(row.get("doi", ""))
        if not doi:
            continue
        kind = "PDF失败" if status == "failed" else "跳过"
        key = (kind, doi.lower())
        if key in seen:
            continue
        seen.add(key)
        rows.append({
            "kind": kind,
            "doi": doi,
            "title": str(row.get("title") or "").strip(),
            "status": status,
            "reason": str(row.get("reason") or "").strip(),
            "source": "pdf_download_report.csv",
        })

    ignored_reasons = {"DOI 为空", "重复 DOI，已跳过"}
    for row in _read_csv_rows(doi_failed_path):
        doi = clean_doi(row.get("doi", ""))
        reason = str(row.get("reason") or "").strip()
        if not doi or reason in ignored_reasons:
            continue
        key = ("解析失败", doi.lower())
        if key in seen:
            continue
        seen.add(key)
        rows.append({
            "kind": "解析失败",
            "doi": doi,
            "title": str(row.get("title") or "").strip(),
            "status": "failed",
            "reason": reason,
            "source": "doi_batch_failed.csv",
        })

    return rows


def _read_delimited_records(path: Path, delimiter: str, doi_column: str | None) -> tuple[list[DoiRecord], str, str]:
    last_error: Exception | None = None
    for encoding in TEXT_ENCODINGS:
        try:
            with path.open("r", newline="", encoding=encoding) as f:
                reader = csv.DictReader(f, delimiter=delimiter)
                headers = [header or "" for header in (reader.fieldnames or [])]
                found_column = _find_doi_column(headers, doi_column)
                records = [_record_from_row(row_number, row, found_column) for row_number, row in enumerate(reader, start=2)]
                return records, encoding, found_column
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"无法识别表格编码，请另存为 UTF-8；最后一次错误: {last_error}")


def _read_text_records(path: Path) -> list[DoiRecord]:
    last_error: Exception | None = None
    for encoding in TEXT_ENCODINGS:
        try:
            records: list[DoiRecord] = []
            with path.open("r", encoding=encoding) as f:
                for line_number, line in enumerate(f, start=1):
                    doi = extract_doi_from_text(line)
                    if doi:
                        records.append(DoiRecord(row_number=line_number, doi=doi, raw_value=line.strip()))
            return records
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"无法识别文本编码，请另存为 UTF-8；最后一次错误: {last_error}")


def _read_xlsx_records(path: Path, doi_column: str | None, sheet_name: str | None) -> tuple[list[DoiRecord], str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("读取 xlsx 需要安装 openpyxl") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        requested_sheet = (sheet_name or "").strip()
        if requested_sheet and requested_sheet not in wb.sheetnames:
            raise ValueError(f"找不到工作表：{requested_sheet}。可用工作表：{', '.join(wb.sheetnames)}")
        found_sheet = requested_sheet or wb.sheetnames[0]
        ws = wb[found_sheet]
        rows_iter = ws.iter_rows(values_only=True)
        headers_raw = next(rows_iter, None)
        if not headers_raw:
            return [], found_sheet
        headers = [str(header).strip() if header is not None else "" for header in headers_raw]
        found_column = _find_doi_column(headers, doi_column)
        records: list[DoiRecord] = []
        for row_number, values in enumerate(rows_iter, start=2):
            row = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))}
            records.append(_record_from_row(row_number, row, found_column))
        return records, found_sheet
    finally:
        wb.close()


def _find_doi_column(headers: list[str], doi_column: str | None) -> str:
    requested = (doi_column or "").strip()
    if requested:
        found = find_column(headers, (requested,))
        if found:
            return found
        raise ValueError(f"找不到 DOI 列：{requested}。现有列：{', '.join(headers)}")
    found = find_column(headers, DOI_COLUMN_CANDIDATES)
    if found:
        return found
    raise ValueError(f"未识别 DOI 列。请填写 DOI 列名。现有列：{', '.join(headers)}")


def _record_from_row(row_number: int, row: dict[str, object], doi_column: str) -> DoiRecord:
    raw_value = "" if row.get(doi_column, "") is None else str(row.get(doi_column, "")).strip()
    return DoiRecord(
        row_number=row_number,
        doi=clean_doi(raw_value),
        title=row_value(row, "title"),
        authors=row_value(row, "authors"),
        journal=row_value(row, "journal"),
        year=row_value(row, "year"),
        date=row_value(row, "date"),
        raw_value=raw_value,
    )


def _records_from_pasted_text(text: str) -> list[DoiRecord]:
    records: list[DoiRecord] = []
    for line_number, line in enumerate((text or "").splitlines(), start=1):
        doi = extract_doi_from_text(line)
        if doi:
            records.append(DoiRecord(row_number=line_number, doi=doi, raw_value=line.strip()))
    return records


def _preview_from_records(
    records: list[DoiRecord],
    source: str,
    doi_column: str,
    encoding: str,
    sheet_name: str,
    limit: int,
) -> DoiPreview:
    status_counts = {"valid": 0, "empty": 0, "invalid": 0, "duplicate": 0}
    seen: set[str] = set()
    preview_rows: list[DoiPreviewRow] = []

    for record in records:
        raw_value = record.raw_value if record.raw_value != "" else record.doi
        normalized_doi = extract_doi_from_text(record.doi) or extract_doi_from_text(raw_value)
        if not str(raw_value or record.doi).strip():
            status = "empty"
            reason = "DOI 为空"
        elif not normalized_doi:
            status = "invalid"
            reason = "未识别到 DOI"
        elif normalized_doi.lower() in seen:
            status = "duplicate"
            reason = "重复 DOI"
        else:
            status = "valid"
            reason = ""
            seen.add(normalized_doi.lower())
        status_counts[status] += 1
        preview_rows.append(DoiPreviewRow(
            row_number=record.row_number,
            doi=normalized_doi,
            title=record.title,
            raw_value=raw_value,
            status=status,
            reason=reason,
        ))

    rows = preview_rows[:limit]
    return DoiPreview(
        rows=rows,
        total_doi=status_counts["valid"],
        total_rows=len(records),
        source=source,
        doi_column=doi_column,
        encoding=encoding,
        sheet_name=sheet_name,
        status_counts=status_counts,
    )


def _read_csv_rows(path: str | Path | None) -> list[dict[str, str]]:
    if not path:
        return []
    csv_path = Path(path)
    if not csv_path.exists():
        return []
    last_error: Exception | None = None
    for encoding in TEXT_ENCODINGS:
        try:
            with csv_path.open("r", newline="", encoding=encoding) as f:
                return [
                    {str(key or ""): "" if value is None else str(value) for key, value in row.items()}
                    for row in csv.DictReader(f)
                ]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"无法识别 CSV 编码，请另存为 UTF-8；最后一次错误: {last_error}")


def _first_existing_parent(*paths: str | Path | None) -> Path:
    for path in paths:
        if path:
            candidate = Path(path)
            if candidate.parent.exists():
                return candidate.parent
    return Path.cwd()


def _retry_exclusion_counts(
    pdf_report_path: str | Path | None = None,
    doi_failed_path: str | Path | None = None,
    selected_dois: set[str] | None = None,
) -> Counter[str]:
    counter: Counter[str] = Counter()
    selected = {clean_doi(doi).lower() for doi in (selected_dois or set()) if clean_doi(doi)}

    for row in _read_csv_rows(pdf_report_path):
        status = str(row.get("status") or "").strip().lower()
        doi = clean_doi(row.get("doi", ""))
        key = doi.lower()
        if status == "failed":
            if not doi:
                counter["PDF失败 DOI 为空"] += 1
            elif selected and key not in selected:
                counter["未选中"] += 1
        elif status == "skipped":
            counter[str(row.get("reason") or "PDF 跳过").strip() or "PDF 跳过"] += 1

    ignored_reasons = {"DOI 为空", "重复 DOI，已跳过"}
    for row in _read_csv_rows(doi_failed_path):
        reason = str(row.get("reason") or "").strip()
        doi = clean_doi(row.get("doi", ""))
        key = doi.lower()
        if not doi:
            counter[reason or "DOI 为空"] += 1
        elif reason in ignored_reasons:
            counter[reason] += 1
        elif selected and key not in selected:
            counter["未选中"] += 1
    return counter


def _read_json_object(path: Path) -> dict[str, object]:
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}
    return data if isinstance(data, dict) else {}


def _collect_report_dois(output_dir: Path) -> set[str]:
    dois: set[str] = set()
    for filename in ("pdf_download_report.csv", "doi_batch_failed.csv"):
        for row in _read_csv_rows(output_dir / filename):
            doi = clean_doi(row.get("doi", ""))
            if doi:
                dois.add(doi.lower())
    return dois


def _candidate_mtime(output_dir: Path, *paths: Path) -> float:
    mtimes = [output_dir.stat().st_mtime]
    for path in paths:
        if path.exists():
            mtimes.append(path.stat().st_mtime)
    return max(mtimes)


def _find_doi_column_from_records_source(path: Path, doi_column: str | None, sheet_name: str) -> str:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[sheet_name]
            headers_raw = next(ws.iter_rows(values_only=True), None)
            headers = [str(header).strip() if header is not None else "" for header in (headers_raw or [])]
            return _find_doi_column(headers, doi_column)
        finally:
            wb.close()
    except Exception:
        return doi_column or ""


def _cookie_items_from_json(data: object) -> list[dict[str, object]]:
    if isinstance(data, dict) and isinstance(data.get("cookies"), list):
        data = data["cookies"]
    if isinstance(data, list):
        return [item for item in data if isinstance(item, dict)]
    if isinstance(data, dict):
        return [{"name": key, "value": value, "domain": ""} for key, value in data.items()]
    return []


def _is_relevant_cookie(item: dict[str, object]) -> bool:
    domain = str(item.get("domain") or item.get("Domain") or "").lower()
    name = str(item.get("name") or item.get("Name") or "").lower()
    return any(cookie_domain in domain for cookie_domain in COOKIE_DOMAINS) or name in {"euid", "sdmsession"}
