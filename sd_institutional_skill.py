from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import tempfile
import time
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable
from urllib.parse import quote
from urllib.request import Request, urlopen

from doi_batch_utils import (
    COLUMN_ALIASES,
    DownloadRunResult,
    DoiRecord,
    PdfDownloadRecord,
    RunSummary,
    SupplementDownloadRecord,
    TEXT_ENCODINGS,
    check_cookie_json,
    clean_doi,
    extract_doi_from_text,
    failure_reason_counts,
    find_column,
    load_doi_records,
    row_value,
    write_pdf_download_report,
    write_pdf_bytes_atomic,
    write_run_summary,
    write_run_summary_json,
    write_supplement_download_report,
)
from paper_automation.deduplicator import deduplicate_candidates
from paper_automation.elsevier_api import ElsevierApiClient, ElsevierApiResult, ElsevierAttachment
from paper_automation.file_manager import enrich_row_metadata_for_delivery, make_pdf_filename
from paper_automation.metadata_resolver import JsonGetter, MetadataResolver, SearchProvider, semantic_scholar_search_provider
from paper_automation.models import MetadataResult, PaperCandidate
from paper_automation.parser import has_extra_bibliographic_signal, is_probable_paper_title, parse_mixed_text
from paper_automation.pdf_validation import is_pdf_bytes, is_valid_pdf
from sd_scraper import ScienceDirectScraper
from sd_supplements import SAFE_EXTENSIONS, SupplementCandidate, make_supplement_filename, supplement_status_counts


SUPPORTED_INPUT_EXTENSIONS = {".xlsx", ".xlsm", ".csv", ".tsv", ".txt", ".md", ".markdown"}
TABULAR_INPUT_EXTENSIONS = {".xlsx", ".xlsm", ".csv", ".tsv"}
METADATA_FIELDS = ("title", "authors", "journal", "year", "date")
IGNORED_DIR_NAMES = {
    ".git",
    ".venv",
    "venv",
    "env",
    "results",
    "dist",
    "pdfs",
    "__pycache__",
    "build",
}
COOKIE_DOMAINS = ("sciencedirect.com", "elsevier.com", "sciencedirectassets.com")
COOKIE_CACHE_NAME = "sciencedirect_cookies.json"
DEFAULT_LOGIN_WAIT_SECONDS = 600
DEFAULT_METADATA_CONFIDENCE = 0.92


@dataclass(frozen=True)
class SourceEntry:
    entry_id: int
    source: str
    row_number: int
    raw_text: str
    doi: str = ""
    title: str = ""
    authors: str = ""
    journal: str = ""
    year: str = ""
    date: str = ""
    initial_status: str = "recognized"
    initial_reason: str = ""


@dataclass(frozen=True)
class IntakeRow:
    source: str
    row_number: int
    input_doi: str
    doi: str = ""
    input_title: str = ""
    title: str = ""
    authors: str = ""
    journal: str = ""
    year: str = ""
    date: str = ""
    raw_value: str = ""
    metadata_source: str = ""
    match_basis: str = ""
    confidence: str = ""
    status: str = "valid"
    reason: str = ""
    review_hint: str = ""


@dataclass(frozen=True)
class IntakeResult:
    all_rows: list[IntakeRow]
    unique_rows: list[IntakeRow]
    preview_path: Path
    merged_input_path: Path
    status_counts: dict[str, int]

    @property
    def valid_count(self) -> int:
        return len(self.unique_rows)


@dataclass
class ElsevierApiAttemptRecord:
    """Sanitised audit row; credential values and response bodies are forbidden."""

    doi: str
    status: str
    http_status: int | None = None
    api_key_present: bool = False
    insttoken_present: bool = False
    full_xml_received: bool = False
    attachment_eid: str = ""
    main_eid_present: bool = False
    pdf_size_bytes: int = 0
    pdf_valid: bool = False
    browser_fallback: bool = False
    reason: str = ""


@dataclass
class ElsevierApiPhaseResult:
    resolved_records: list[dict]
    pdf_records: list[PdfDownloadRecord]
    supplement_records: list[SupplementDownloadRecord]
    fallback_rows: list[IntakeRow]
    attempts: list[ElsevierApiAttemptRecord]


def main(argv: list[str] | None = None) -> int:
    parser = build_arg_parser()
    args = parser.parse_args(argv)

    if not args.text and not args.input and not args.folder:
        parser.error("需要提供 --text、--input 或 --folder 中的至少一种输入")

    output_root = choose_output_root(args.out) if args.choose_out else Path(args.out).expanduser().resolve()
    run_name = args.run_name or f"sd_skill_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    run_dir = output_root / run_name
    run_dir.mkdir(parents=True, exist_ok=True)
    auth_dir = output_root / "_auth"
    auth_dir.mkdir(parents=True, exist_ok=True)
    preflight_only = args.preflight or args.beginner

    try:
        intake = build_intake(
            texts=args.text or [],
            input_paths=[Path(path) for path in args.input or []],
            folder_paths=[Path(path) for path in args.folder or []],
            output_dir=run_dir,
            doi_column=args.doi_column,
            sheet_name=args.sheet,
            resolve_metadata=not preflight_only,
            resolve_title_only_files=args.resolve_title_only,
            email=args.email,
            min_confidence=args.min_confidence,
            search_provider=semantic_scholar_search_provider(email=args.email) if args.auto_web_search else None,
            max_search_candidates=args.max_search_candidates,
        )
    except Exception as exc:
        print(f"[错误] DOI 输入整理失败: {exc}", flush=True)
        return 2

    print("[输入整理] 完成", flush=True)
    print(f"- 预览表: {intake.preview_path}", flush=True)
    print(f"- 合并输入表: {intake.merged_input_path}", flush=True)
    print(f"- 有效唯一 DOI: {intake.valid_count}", flush=True)
    print(f"- 状态统计: {dict(intake.status_counts)}", flush=True)

    if intake.valid_count == 0:
        failed_path = write_intake_failed_report(intake.all_rows, run_dir / "doi_batch_failed.csv")
        pdf_report_path = write_pdf_download_report([], run_dir)
        summary = RunSummary(
            input_path=str(intake.merged_input_path),
            output_dir=str(run_dir),
            total_doi=0,
            resolved_count=0,
            failure_reasons=intake_failure_reasons(intake.all_rows),
            pdf_success=0,
            pdf_failed=0,
            pdf_skipped=0,
            resolved_path="",
            failed_path=str(failed_path),
            pdf_report_path=str(pdf_report_path),
            cookie_message="",
            beginner_recommendations=build_beginner_recommendations(intake, preflight_only=preflight_only),
            supplement_requested=False,
        )
        summary_path = write_run_summary(summary)
        summary_json_path = write_run_summary_json(summary)
        print("[结束] 没有可处理的有效 DOI。", flush=True)
        print(f"- DOI failure report: {failed_path}", flush=True)
        print(f"- PDF report: {pdf_report_path}", flush=True)
        print(f"- Run summary: {summary_path}", flush=True)
        print(f"- Run summary JSON: {summary_json_path}", flush=True)
        return 1

    if preflight_only:
        failed_path = write_intake_failed_report(intake.all_rows, run_dir / "doi_batch_failed.csv")
        preflight_pdf_records = [
            PdfDownloadRecord(
                doi=row.doi,
                pii="",
                title=row.title,
                status="not_requested",
                reason="preflight",
            )
            for row in intake.unique_rows
        ]
        pdf_report_path = write_pdf_download_report(preflight_pdf_records, run_dir)
        summary = RunSummary(
            input_path=str(intake.merged_input_path),
            output_dir=str(run_dir),
            total_doi=intake.valid_count,
            resolved_count=0,
            failure_reasons=intake_failure_reasons(intake.all_rows),
            pdf_success=0,
            pdf_failed=0,
            pdf_skipped=len(intake.unique_rows),
            resolved_path="",
            failed_path=str(failed_path),
            pdf_report_path=str(pdf_report_path),
            cookie_message=cookie_status_message(auth_dir / COOKIE_CACHE_NAME),
            beginner_recommendations=build_beginner_recommendations(
                intake,
                preflight_only=True,
                auto_web_search=args.auto_web_search,
            ),
            supplement_requested=False,
        )
        summary_path = write_run_summary(summary)
        summary_json_path = write_run_summary_json(summary)
        print("[预检] 完成；未联网解析元数据、未解析 ScienceDirect PII，未下载 PDF。", flush=True)
        print(f"- 输出目录: {run_dir}", flush=True)
        print(f"- 可进入后续解析的 DOI: {intake.valid_count}", flush=True)
        print(f"- 需复核/排除: {sum(intake_failure_reasons(intake.all_rows).values())}", flush=True)
        print(f"- PDF 明细: {pdf_report_path}", flush=True)
        print(f"- 任务摘要: {summary_path}", flush=True)
        print(f"- JSON 摘要: {summary_json_path}", flush=True)
        return 0

    download_pdfs = not args.dry_run and not args.no_download_pdfs
    # Supplements are on by default; --no-download-supplements applies to both
    # the API and browser paths.
    download_supplements = download_pdfs and (
        False
        if bool(getattr(args, "no_download_supplements", False))
        else True
        if getattr(args, "download_supplements", None) is None
        else bool(args.download_supplements)
    )
    cookie_cache_path = auth_dir / COOKIE_CACHE_NAME
    scraper = None
    cookie_message = explicit_cookie_status_message(args.cookies) if args.cookies else cookie_status_message(cookie_cache_path)
    api_phase = ElsevierApiPhaseResult([], [], [], [], [])
    browser_results: list[dict] = []
    browser_failures: list[dict] = []
    browser_pdf_records: list[PdfDownloadRecord] = []
    browser_supplement_records: list[SupplementDownloadRecord] = []

    if download_pdfs:
        api_phase = run_elsevier_api_phase(
            intake.unique_rows,
            run_dir=run_dir,
            email=args.email,
            download_supplements=download_supplements,
        )
        if api_phase.fallback_rows and args.api_only:
            attempt_by_doi = {_doi_key(item.doi): item for item in api_phase.attempts}
            browser_failures = []
            for row in api_phase.fallback_rows:
                attempt = attempt_by_doi.get(_doi_key(row.doi))
                status = attempt.status if attempt is not None else "missing_api_attempt"
                detail = attempt.reason if attempt is not None else "missing_api_attempt"
                browser_failures.append(
                    {
                        "row_number": row.row_number,
                        "doi": row.doi,
                        "reason": f"api_only_{status}:{detail or status}",
                    }
                )
                if attempt is not None:
                    attempt.browser_fallback = False
                    attempt.reason = f"api_only_no_fallback:{detail or status}"
            cookie_message = "API-only：未读取 Cookie，未创建或启动浏览器。"
            print(
                f"[API-only] {len(api_phase.fallback_rows)} 条 API 未成功；"
                "已记录失败，禁止浏览器兜底。",
                flush=True,
            )
        elif api_phase.fallback_rows:
            browser_input_path = write_merged_input(
                api_phase.fallback_rows,
                run_dir / "elsevier_api_browser_fallback.csv",
            )
            scraper = make_scraper(cookie_cache_path, browser_exe=args.browser_exe, cookies_path=args.cookies)
            browser_results, browser_failures = scraper.resolve_doi_batch(str(browser_input_path))
            if browser_results:
                cache_devtools_cookies(scraper, cookie_cache_path)
                try:
                    download_result = scraper.download_pdfs_devtools(
                        browser_results,
                        str(run_dir),
                        login_wait_seconds=args.login_wait_seconds,
                        interactive_login=False,
                        download_supplements=download_supplements,
                        session_break_seconds=float(getattr(args, "session_break_seconds", 60.0) or 60.0),
                        session_break_every=int(getattr(args, "session_break_every", 8) or 8),
                        resume=True,
                    )
                except Exception as exc:
                    browser_pdf_records = [
                        PdfDownloadRecord(
                            doi=item.get("doi", ""),
                            pii=item.get("pii", ""),
                            title=item.get("title", ""),
                            status="failed",
                            reason=f"download_exception_{type(exc).__name__}",
                        )
                        for item in browser_results
                    ]
                    print("[警告] PDF 下载器异常；已为每篇文献写入可恢复的失败记录。", flush=True)
                else:
                    if download_result:
                        _, _, _, browser_pdf_records = download_result
                        if isinstance(download_result, DownloadRunResult):
                            browser_supplement_records = download_result.supplement_records
                    else:
                        browser_pdf_records = [
                            PdfDownloadRecord(
                                doi=item.get("doi", ""),
                                pii=item.get("pii", ""),
                                title=item.get("title", ""),
                                status="failed",
                                reason="PDF 下载流程未返回状态",
                            )
                            for item in browser_results
                        ]
                    cache_devtools_cookies(scraper, cookie_cache_path)
                    cookie_message = cookie_status_message(cookie_cache_path)
            else:
                print("[提示] API 失败项未解析到可供浏览器下载的 ScienceDirect 记录。", flush=True)
        else:
            cookie_message = "Elsevier API 已完成全部主 PDF；未读取或启动浏览器 Cookie 流程。"
    else:
        # Dry runs keep the existing DOI-to-PII resolution behaviour and do not
        # make Elsevier full-text API requests.
        scraper = ScienceDirectScraper(browser_exe=args.browser_exe)
        browser_results, browser_failures = scraper.resolve_doi_batch(str(intake.merged_input_path))
        reason = "dry_run" if args.dry_run else "no_download_pdfs"
        browser_pdf_records = [
            PdfDownloadRecord(
                doi=item.get("doi", ""),
                pii=item.get("pii", ""),
                title=item.get("title", ""),
                status="not_requested",
                reason=reason,
            )
            for item in browser_results
        ]

    results = merge_resolved_records(intake.unique_rows, api_phase.resolved_records, browser_results)
    failures = browser_failures
    pdf_records = merge_pdf_records(
        intake.unique_rows,
        api_phase.pdf_records,
        browser_pdf_records,
        browser_failures,
        download_requested=download_pdfs,
    )
    supplement_records = [*api_phase.supplement_records, *browser_supplement_records]
    pdf_success = sum(record.status == "success" for record in pdf_records)
    pdf_failed = sum(record.status == "failed" for record in pdf_records)
    pdf_skipped = sum(record.status in {"skipped", "not_requested"} for record in pdf_records)
    supplement_success, supplement_failed, supplement_skipped, supplement_not_found = supplement_status_counts(
        supplement_records
    )
    api_attempt_report_path = write_elsevier_api_attempt_report(api_phase.attempts, run_dir)

    resolved_path = ""
    if results:
        resolved_path = (
            scraper.save_to_xlsx(results, "doi_batch_resolved.xlsx", str(run_dir))
            if scraper is not None and hasattr(scraper, "save_to_xlsx")
            else save_resolved_results(results, run_dir / "doi_batch_resolved.xlsx")
        )
    else:
        print("[提示] 未解析到任何 ScienceDirect DOI。", flush=True)
    failed_path = (
        scraper.save_failed_doi_report(failures, "doi_batch_failed.csv", str(run_dir))
        if scraper is not None and hasattr(scraper, "save_failed_doi_report")
        else write_resolve_failed_report(failures, run_dir / "doi_batch_failed.csv")
    )

    supplement_report_path = ""

    pdf_report_path = write_pdf_download_report(pdf_records, run_dir)
    if download_supplements and results:
        supplement_report_path = str(write_supplement_download_report(supplement_records, run_dir))
    summary = RunSummary(
        input_path=str(intake.merged_input_path),
        output_dir=str(run_dir),
        total_doi=intake.valid_count,
        resolved_count=len(results),
        failure_reasons=failure_reason_counts(failures),
        pdf_success=pdf_success,
        pdf_failed=pdf_failed,
        pdf_skipped=pdf_skipped,
        resolved_path=resolved_path,
        failed_path=failed_path,
        pdf_report_path=str(pdf_report_path),
        cookie_message=cookie_message,
        supplement_success=supplement_success,
        supplement_failed=supplement_failed,
        supplement_skipped=supplement_skipped,
        supplement_not_found=supplement_not_found,
        supplement_report_path=supplement_report_path,
        supplement_requested=download_supplements and bool(results),
        beginner_recommendations=build_beginner_recommendations(
            intake,
            failure_reasons=failure_reason_counts(failures),
            pdf_failed=pdf_failed,
            preflight_only=False,
            auto_web_search=args.auto_web_search,
        ),
        browser_message=getattr(scraper, "last_browser_message", "") if scraper is not None else "",
        download_next_steps=getattr(scraper, "last_download_next_steps", "") if scraper is not None else "",
    )
    summary_path = write_run_summary(summary)
    summary_json_path = write_run_summary_json(summary)

    print("[报告] 完成", flush=True)
    print(f"- 输出目录: {run_dir}", flush=True)
    print(f"- 解析成功: {len(results)}", flush=True)
    print(f"- PDF 成功/失败/跳过: {pdf_success}/{pdf_failed}/{pdf_skipped}", flush=True)
    if download_supplements:
        print(
            f"- 补充材料 成功/失败/跳过/未发现: "
            f"{supplement_success}/{supplement_failed}/{supplement_skipped}/{supplement_not_found}",
            flush=True,
        )
    print(f"- PDF 明细: {pdf_report_path}", flush=True)
    print(f"- Elsevier API 脱敏审计: {api_attempt_report_path}", flush=True)
    if supplement_report_path:
        print(f"- 补充材料明细: {supplement_report_path}", flush=True)
    print(f"- 任务摘要: {summary_path}", flush=True)
    print(f"- JSON 摘要: {summary_json_path}", flush=True)
    return 0


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "ScienceDirect institutional adapter / preflight helper (internal & compatibility). "
            "For new literature downloads prefer: paper_batch.py start ..."
        ),
        epilog="Default user entry is paper_batch.py or Codex skill paper-download, not this script alone.",
    )
    parser.add_argument("--text", action="append", help="DOI or pasted paper text; may be repeated")
    parser.add_argument("--input", action="append", help="Input file path; may be repeated")
    parser.add_argument("--folder", action="append", help="Folder to recursively scan for DOI files; may be repeated")
    parser.add_argument("--out", default="results", help="Output root directory; default: results")
    parser.add_argument("--run-name", help="Optional run directory name under --out")
    parser.add_argument("--doi-column", help="DOI column name for tabular files")
    parser.add_argument("--sheet", help="Excel sheet name for all Excel inputs")
    parser.add_argument("--email", default=os.environ.get("PAPER_SKILL_EMAIL", ""), help="Email for polite Crossref/OpenAlex API use")
    parser.add_argument("--min-confidence", type=float, default=DEFAULT_METADATA_CONFIDENCE, help="Minimum title-match confidence for title-only automatic DOI use")
    parser.add_argument("--beginner", action="store_true", help="Run a beginner-friendly preflight workflow and write next-step guidance")
    parser.add_argument("--preflight", action="store_true", help="Only prepare and review input; do not resolve ScienceDirect PII or download PDFs")
    parser.add_argument("--auto-web-search", action="store_true", help="Use optional Semantic Scholar search fallback for title-only metadata resolution")
    parser.add_argument("--max-search-candidates", type=int, default=5, help="Maximum optional search candidates to inspect when --auto-web-search is enabled")
    parser.add_argument("--resolve-title-only", action="store_true", help="For file/folder inputs, also try to resolve rows without explicit DOI by title")
    parser.add_argument("--choose-out", action="store_true", help="Open a Windows folder picker for the output root when available")
    parser.add_argument("--dry-run", action="store_true", help="Resolve DOI metadata but do not download PDFs")
    parser.add_argument("--no-download-pdfs", action="store_true", help="Skip PDF downloads after DOI resolution")
    parser.add_argument(
        "--api-only",
        action="store_true",
        help="Use Elsevier API only; never create a browser fallback",
    )
    parser.add_argument(
        "--download-supplements",
        action="store_true",
        default=None,
        help="Also download ScienceDirect supplementary files (default: on)",
    )
    parser.add_argument(
        "--no-download-supplements",
        action="store_true",
        help="Do not download ScienceDirect supplementary files",
    )
    parser.add_argument(
        "--session-break-seconds",
        type=float,
        default=60.0,
        help="Fixed rest seconds after every N successful SD downloads (default 60; was 150)",
    )
    parser.add_argument(
        "--session-break-every",
        type=int,
        default=8,
        help="Rest after this many successful SD downloads (default 8)",
    )
    parser.add_argument("--cookies", help="Explicit Cookie JSON file to use before cached/browser cookies")
    parser.add_argument(
        "--login-wait-seconds",
        type=int,
        default=DEFAULT_LOGIN_WAIT_SECONDS,
        help="Seconds to wait for browser institutional login when needed",
    )
    parser.add_argument(
        "--browser-exe",
        help="Browser executable path for institutional login/download (defaults to Chrome, then Edge)",
    )
    return parser


def build_intake(
    texts: Iterable[str],
    input_paths: Iterable[Path],
    folder_paths: Iterable[Path],
    output_dir: Path,
    doi_column: str | None = None,
    sheet_name: str | None = None,
    resolve_metadata: bool = False,
    resolve_title_only_files: bool = False,
    email: str = "",
    min_confidence: float = DEFAULT_METADATA_CONFIDENCE,
    http_json: JsonGetter | None = None,
    search_provider: SearchProvider | None = None,
    max_search_candidates: int = 5,
) -> IntakeResult:
    output_dir.mkdir(parents=True, exist_ok=True)
    entries: list[SourceEntry] = []
    next_id = 1

    for index, text in enumerate(texts, start=1):
        source = f"pasted_text_{index}"
        new_entries, next_id = entries_from_mixed_text(text, source, next_id)
        entries.extend(new_entries)

    for input_path in input_paths:
        new_entries, next_id = load_entries_from_file(
            input_path,
            next_id,
            doi_column=doi_column,
            sheet_name=sheet_name,
            resolve_title_only=resolve_title_only_files,
        )
        entries.extend(new_entries)

    for folder_path in folder_paths:
        for file_path in iter_input_files(folder_path):
            new_entries, next_id = load_entries_from_file(
                file_path,
                next_id,
                doi_column=doi_column,
                sheet_name=sheet_name,
                resolve_title_only=resolve_title_only_files,
            )
            entries.extend(new_entries)

    print("[输入整理] 本地识别 DOI/题名，先写入预览表...", flush=True)
    all_rows, unique_rows, counts = classify_entries(
        entries,
        resolve_metadata=False,
        email=email,
        min_confidence=min_confidence,
        http_json=http_json,
        search_provider=search_provider,
        max_search_candidates=max_search_candidates,
    )
    preview_path = output_dir / "doi_intake_preview.csv"
    merged_path = output_dir / "merged_doi_input.csv"
    write_intake_preview(all_rows, preview_path)
    write_merged_input(unique_rows, merged_path)
    print(f"[输入整理] 本地预览已写出: {preview_path}", flush=True)
    print(f"[输入整理] 本地合并表已写出: {merged_path}", flush=True)

    if resolve_metadata and has_title_only_metadata_candidates(entries):
        print("[输入整理] 检测到题名-only/短引用候选，开始联网元数据增强...", flush=True)
        all_rows, unique_rows, counts = classify_entries(
            entries,
            resolve_metadata=True,
            email=email,
            min_confidence=min_confidence,
            http_json=http_json,
            search_provider=search_provider,
            max_search_candidates=max_search_candidates,
        )
        write_intake_preview(all_rows, preview_path)
        write_merged_input(unique_rows, merged_path)
        print("[输入整理] 联网元数据增强完成，预览表已更新。", flush=True)
    elif resolve_metadata:
        print("[输入整理] 未发现需要题名匹配的候选，跳过联网元数据增强。", flush=True)

    return IntakeResult(
        all_rows=all_rows,
        unique_rows=unique_rows,
        preview_path=preview_path,
        merged_input_path=merged_path,
        status_counts=dict(counts),
    )


def has_title_only_metadata_candidates(entries: Iterable[SourceEntry]) -> bool:
    return any(
        entry.initial_status == "recognized" and not entry.doi and bool(entry.title)
        for entry in entries
    )


def entries_from_mixed_text(text: str, source: str, next_id: int) -> tuple[list[SourceEntry], int]:
    entries: list[SourceEntry] = []
    for candidate in parse_mixed_text(text):
        entries.append(SourceEntry(
            entry_id=next_id,
            source=source,
            row_number=candidate.source_index,
            raw_text=candidate.raw_text,
            doi=candidate.doi,
            title=candidate.title,
            initial_status=candidate.status,
            initial_reason=candidate.reason,
        ))
        next_id += 1
    return entries, next_id


def entries_from_explicit_doi_text(text: str, source: str, next_id: int) -> tuple[list[SourceEntry], int]:
    entries: list[SourceEntry] = []
    for candidate in parse_mixed_text(text):
        if not candidate.doi:
            continue
        entries.append(SourceEntry(
            entry_id=next_id,
            source=source,
            row_number=candidate.source_index,
            raw_text=candidate.raw_text,
            doi=candidate.doi,
            title=candidate.title,
            initial_status=candidate.status,
            initial_reason=candidate.reason,
        ))
        next_id += 1
    return entries, next_id


def iter_input_files(folder: Path) -> Iterable[Path]:
    root = folder.expanduser().resolve()
    if not root.exists():
        raise FileNotFoundError(f"找不到文件夹: {root}")
    if not root.is_dir():
        raise NotADirectoryError(f"不是文件夹: {root}")

    for current_root, dirnames, filenames in os.walk(root):
        dirnames[:] = [
            name for name in dirnames
            if name.lower() not in IGNORED_DIR_NAMES
        ]
        for filename in sorted(filenames):
            path = Path(current_root) / filename
            if should_use_input_file(path):
                yield path


def should_use_input_file(path: Path) -> bool:
    name = path.name.lower()
    if name.startswith("cookies") and path.suffix.lower() == ".json":
        return False
    if name in {"cookie.json", COOKIE_CACHE_NAME.lower()}:
        return False
    return path.suffix.lower() in SUPPORTED_INPUT_EXTENSIONS


def load_entries_from_file(
    path: Path,
    next_id: int,
    doi_column: str | None = None,
    sheet_name: str | None = None,
    resolve_title_only: bool = False,
) -> tuple[list[SourceEntry], int]:
    file_path = path.expanduser().resolve()
    if not file_path.exists():
        raise FileNotFoundError(f"找不到输入文件: {file_path}")
    if not should_use_input_file(file_path):
        print(f"[跳过] 不支持或不应扫描的输入文件: {file_path}")
        return [], next_id

    if file_path.suffix.lower() in {".txt", ".md", ".markdown"}:
        text = read_text_file(file_path)
        if resolve_title_only:
            return entries_from_mixed_text(text, str(file_path), next_id)
        return entries_from_explicit_doi_text(text, str(file_path), next_id)

    if resolve_title_only:
        records = load_tabular_records(file_path, doi_column=doi_column, sheet_name=sheet_name)
    else:
        records = load_explicit_doi_records(file_path, doi_column=doi_column, sheet_name=sheet_name)
    entries: list[SourceEntry] = []
    for record in records:
        normalized = extract_doi_from_text(record.doi) or extract_doi_from_text(record.raw_value)
        raw_value = record.raw_value if record.raw_value != "" else record.doi
        context_text = " ".join(
            value for value in [record.title, record.authors, record.journal, record.year, record.date]
            if str(value or "").strip()
        )
        if not str(raw_value or record.doi or record.title).strip():
            initial_status = "empty"
            initial_reason = "DOI 为空"
        elif record.doi and not normalized:
            initial_status = "invalid"
            initial_reason = "未识别到 DOI"
        elif not normalized and record.title and not is_probable_paper_title(record.title):
            initial_status = "needs_review"
            initial_reason = "not_probable_title"
        elif not normalized and record.title and not has_extra_bibliographic_signal(context_text, record.title):
            initial_status = "needs_review"
            initial_reason = "insufficient_bibliographic_context"
        elif not normalized and not record.title:
            initial_status = "needs_review"
            initial_reason = "not_probable_title"
        else:
            initial_status = "recognized"
            initial_reason = ""
        entries.append(SourceEntry(
            entry_id=next_id,
            source=str(file_path),
            row_number=record.row_number,
            raw_text=raw_value,
            doi=normalized,
            title=record.title,
            authors=record.authors,
            journal=record.journal,
            year=record.year,
            date=record.date,
            initial_status=initial_status,
            initial_reason=initial_reason,
        ))
        next_id += 1
    return entries, next_id


def load_tabular_records(
    path: Path,
    doi_column: str | None = None,
    sheet_name: str | None = None,
) -> list[DoiRecord]:
    if doi_column:
        return load_doi_records(path, doi_column=doi_column, sheet_name=sheet_name)

    headers = read_tabular_headers(path, sheet_name=sheet_name)
    if find_column(headers, COLUMN_ALIASES["doi"]):
        return load_doi_records(path, doi_column=doi_column, sheet_name=sheet_name)
    if find_column(headers, COLUMN_ALIASES["title"]):
        return load_title_only_records(path, sheet_name=sheet_name)
    return load_doi_records(path, doi_column=doi_column, sheet_name=sheet_name)


def load_explicit_doi_records(
    path: Path,
    doi_column: str | None = None,
    sheet_name: str | None = None,
) -> list[DoiRecord]:
    try:
        return load_doi_records(path, doi_column=doi_column, sheet_name=sheet_name)
    except ValueError:
        if doi_column:
            raise
        return scan_tabular_rows_for_explicit_dois(path, sheet_name=sheet_name)


def scan_tabular_rows_for_explicit_dois(path: Path, sheet_name: str | None = None) -> list[DoiRecord]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return scan_delimited_rows_for_explicit_dois(path, delimiter=",")
    if suffix == ".tsv":
        return scan_delimited_rows_for_explicit_dois(path, delimiter="\t")
    if suffix in {".xlsx", ".xlsm"}:
        return scan_xlsx_rows_for_explicit_dois(path, sheet_name=sheet_name)
    raise ValueError(f"Unsupported tabular input: {path}")


def scan_delimited_rows_for_explicit_dois(path: Path, delimiter: str) -> list[DoiRecord]:
    last_error: Exception | None = None
    for encoding in TEXT_ENCODINGS:
        try:
            records: list[DoiRecord] = []
            with path.open("r", newline="", encoding=encoding) as f:
                reader = csv.DictReader(f, delimiter=delimiter)
                for row_number, row in enumerate(reader, start=2):
                    normalized_row = {str(key or ""): value for key, value in row.items()}
                    if extract_doi_from_text(row_text(normalized_row)):
                        records.append(record_from_explicit_doi_row(row_number, normalized_row))
            return records
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"Unable to detect tabular encoding for {path}: {last_error}")


def scan_xlsx_rows_for_explicit_dois(path: Path, sheet_name: str | None = None) -> list[DoiRecord]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Reading xlsx inputs requires openpyxl") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        requested_sheet = (sheet_name or "").strip()
        if requested_sheet and requested_sheet not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {requested_sheet}. Available sheets: {', '.join(wb.sheetnames)}")
        found_sheet = requested_sheet or wb.sheetnames[0]
        ws = wb[found_sheet]
        rows_iter = ws.iter_rows(values_only=True)
        headers_raw = next(rows_iter, None)
        if not headers_raw:
            return []
        headers = [str(header).strip() if header is not None else "" for header in headers_raw]
        records: list[DoiRecord] = []
        for row_number, values in enumerate(rows_iter, start=2):
            row = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))}
            if extract_doi_from_text(row_text(row)):
                records.append(record_from_explicit_doi_row(row_number, row))
        return records
    finally:
        wb.close()


def record_from_explicit_doi_row(row_number: int, row: dict[str, object]) -> DoiRecord:
    raw_value = row_text(row)
    return DoiRecord(
        row_number=row_number,
        doi=extract_doi_from_text(raw_value),
        title=row_value(row, "title"),
        authors=row_value(row, "authors"),
        journal=row_value(row, "journal"),
        year=row_value(row, "year"),
        date=row_value(row, "date"),
        raw_value=raw_value,
    )


def row_text(row: dict[str, object]) -> str:
    return " | ".join(
        str(value).strip()
        for value in row.values()
        if value is not None and str(value).strip()
    )


def read_tabular_headers(path: Path, sheet_name: str | None = None) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_delimited_headers(path, delimiter=",")
    if suffix == ".tsv":
        return read_delimited_headers(path, delimiter="\t")
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx_headers(path, sheet_name=sheet_name)
    return []


def read_delimited_headers(path: Path, delimiter: str) -> list[str]:
    last_error: Exception | None = None
    for encoding in TEXT_ENCODINGS:
        try:
            with path.open("r", newline="", encoding=encoding) as f:
                reader = csv.DictReader(f, delimiter=delimiter)
                return [header or "" for header in (reader.fieldnames or [])]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"Unable to detect tabular encoding for {path}: {last_error}")


def read_xlsx_headers(path: Path, sheet_name: str | None = None) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Reading xlsx inputs requires openpyxl") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        requested_sheet = (sheet_name or "").strip()
        if requested_sheet and requested_sheet not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {requested_sheet}. Available sheets: {', '.join(wb.sheetnames)}")
        found_sheet = requested_sheet or wb.sheetnames[0]
        ws = wb[found_sheet]
        headers_raw = next(ws.iter_rows(values_only=True), None)
        if not headers_raw:
            return []
        return [str(header).strip() if header is not None else "" for header in headers_raw]
    finally:
        wb.close()


def load_title_only_records(path: Path, sheet_name: str | None = None) -> list[DoiRecord]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_delimited_metadata_records(path, delimiter=",")
    if suffix == ".tsv":
        return read_delimited_metadata_records(path, delimiter="\t")
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx_metadata_records(path, sheet_name=sheet_name)
    raise ValueError(f"Unsupported title-only tabular input: {path}")


def read_delimited_metadata_records(path: Path, delimiter: str) -> list[DoiRecord]:
    last_error: Exception | None = None
    for encoding in TEXT_ENCODINGS:
        try:
            with path.open("r", newline="", encoding=encoding) as f:
                reader = csv.DictReader(f, delimiter=delimiter)
                headers = [header or "" for header in (reader.fieldnames or [])]
                rows = (
                    (row_number, {str(key or ""): value for key, value in row.items()})
                    for row_number, row in enumerate(reader, start=2)
                )
                return metadata_records_from_rows(headers, rows)
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"Unable to detect tabular encoding for {path}: {last_error}")


def read_xlsx_metadata_records(path: Path, sheet_name: str | None = None) -> list[DoiRecord]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Reading xlsx inputs requires openpyxl") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        requested_sheet = (sheet_name or "").strip()
        if requested_sheet and requested_sheet not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {requested_sheet}. Available sheets: {', '.join(wb.sheetnames)}")
        found_sheet = requested_sheet or wb.sheetnames[0]
        ws = wb[found_sheet]
        rows_iter = ws.iter_rows(values_only=True)
        headers_raw = next(rows_iter, None)
        if not headers_raw:
            return []
        headers = [str(header).strip() if header is not None else "" for header in headers_raw]
        rows = (
            (
                row_number,
                {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))},
            )
            for row_number, values in enumerate(rows_iter, start=2)
        )
        return metadata_records_from_rows(headers, rows)
    finally:
        wb.close()


def metadata_records_from_rows(
    headers: list[str],
    rows: Iterable[tuple[int, dict[str, object]]],
) -> list[DoiRecord]:
    if not has_title_column(headers):
        raise ValueError(f"No title column found for title-only input. Existing columns: {', '.join(headers)}")

    records: list[DoiRecord] = []
    for row_number, row in rows:
        doi_raw = row_value(row, "doi")
        title = row_value(row, "title")
        authors = row_value(row, "authors")
        journal = row_value(row, "journal")
        year = row_value(row, "year")
        date = row_value(row, "date")
        raw_value = doi_raw or title or authors or journal or year or date
        records.append(DoiRecord(
            row_number=row_number,
            doi=clean_doi(doi_raw),
            title=title,
            authors=authors,
            journal=journal,
            year=year,
            date=date,
            raw_value=raw_value,
        ))
    return records


def has_title_column(headers: list[str]) -> bool:
    return bool(find_column(headers, COLUMN_ALIASES["title"]))


def read_text_file(path: Path) -> str:
    last_error: Exception | None = None
    for encoding in TEXT_ENCODINGS:
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError as exc:
            last_error = exc
    raise UnicodeDecodeError("text", b"", 0, 1, f"无法识别文本编码: {last_error}")


def classify_entries(
    entries: list[SourceEntry],
    *,
    resolve_metadata: bool,
    email: str,
    min_confidence: float,
    http_json: JsonGetter | None,
    search_provider: SearchProvider | None = None,
    max_search_candidates: int = 5,
) -> tuple[list[IntakeRow], list[IntakeRow], Counter]:
    counts: Counter = Counter()
    recognized = [
        entry for entry in entries
        if entry.initial_status == "recognized" and (entry.doi or entry.title)
    ]
    candidates = [
        PaperCandidate(entry.entry_id, entry.raw_text, entry.doi, entry.title)
        for entry in recognized
    ]
    deduped = deduplicate_candidates(candidates)
    duplicate_by_id = {item.source_index: item for item in deduped.duplicates}
    unique_ids = {item.source_index for item in deduped.unique}
    resolver = (
        MetadataResolver(
            email=email,
            http_json=http_json,
            search_provider=search_provider,
            max_search_candidates=max_search_candidates,
        )
        if resolve_metadata
        else None
    )

    all_rows: list[IntakeRow] = []
    unique_rows: list[IntakeRow] = []
    resolved_doi_owner: dict[str, int] = {}
    metadata_total = sum(1 for item in deduped.unique if not item.doi and item.title) if resolver else 0
    metadata_index = 0

    for entry in entries:
        if entry.initial_status in {"empty", "invalid", "needs_review"}:
            row = row_from_entry(
                entry,
                doi="",
                title=entry.title,
                status=entry.initial_status,
                reason=entry.initial_reason or entry.initial_status,
            )
        elif entry.entry_id in duplicate_by_id:
            duplicate = duplicate_by_id[entry.entry_id]
            duplicate_label = {
                "duplicate_doi": "重复 DOI",
                "duplicate_title": "重复题名",
                "possible_duplicate_title": "疑似重复题名",
            }.get(duplicate.reason, duplicate.reason)
            row = row_from_entry(
                entry,
                doi=entry.doi,
                title=entry.title,
                status="duplicate",
                reason=f"{duplicate_label}: duplicate_of={duplicate.duplicate_of}; score={duplicate.score:.3f}",
            )
        elif entry.entry_id in unique_ids:
            if resolver is not None and not entry.doi and entry.title:
                metadata_index += 1
                print(
                    f"[元数据增强] [{metadata_index}/{metadata_total}] "
                    f"按题名匹配 DOI: {entry.title[:80]}",
                    flush=True,
                )
            metadata = resolve_entry_metadata(entry, resolver, min_confidence)
            row = row_from_metadata(entry, metadata, min_confidence)
        else:
            row = row_from_entry(entry, doi="", title=entry.title, status="needs_review", reason="not_selected")
        if row.status == "valid" and row.doi:
            doi_key = row.doi.lower()
            if doi_key in resolved_doi_owner:
                duplicate_reason = f"重复 DOI: duplicate_of={resolved_doi_owner[doi_key]}"
                row = IntakeRow(**{
                    **row.__dict__,
                    "status": "duplicate",
                    "reason": duplicate_reason,
                    "review_hint": intake_review_hint("duplicate", duplicate_reason, row.doi),
                })
            else:
                resolved_doi_owner[doi_key] = entry.entry_id
        all_rows.append(row)
        counts[row.status] += 1
        if row.status == "valid" and row.doi:
            unique_rows.append(row)

    for key in ("valid", "duplicate", "empty", "invalid", "needs_review"):
        counts.setdefault(key, 0)
    return all_rows, unique_rows, counts


def resolve_entry_metadata(
    entry: SourceEntry,
    resolver: MetadataResolver | None,
    min_confidence: float,
) -> MetadataResult:
    candidate = PaperCandidate(entry.entry_id, entry.raw_text, entry.doi, entry.title)
    if resolver is None:
        return MetadataResult(
            source_index=entry.entry_id,
            query_title=entry.title,
            doi=entry.doi,
            title=entry.title,
            authors=split_authors(entry.authors),
            journal=entry.journal,
            year=entry.year,
            confidence=1.0 if entry.doi else 0.0,
            source="input",
            match_basis="input_doi" if entry.doi else "",
        )
    metadata = resolver.resolve_one(candidate)
    if entry.doi and not metadata.doi:
        metadata.doi = entry.doi
    if entry.title and not metadata.title:
        metadata.title = entry.title
    if entry.authors and not metadata.authors:
        metadata.authors = split_authors(entry.authors)
    if entry.journal and not metadata.journal:
        metadata.journal = entry.journal
    if entry.year and not metadata.year:
        metadata.year = entry.year
    if entry.doi:
        metadata.confidence = max(metadata.confidence, 1.0)
    elif metadata.doi and not metadata.confidence:
        metadata.confidence = min_confidence
    return metadata


def row_from_metadata(entry: SourceEntry, metadata: MetadataResult, min_confidence: float) -> IntakeRow:
    doi = metadata.doi or entry.doi
    is_title_only = not entry.doi
    if is_title_only and (not doi or metadata.confidence < min_confidence):
        return row_from_entry(
            entry,
            doi=doi,
            title=metadata.title or entry.title,
            status="needs_review",
            reason=f"metadata_confidence_below_threshold:{metadata.confidence:.3f}",
            metadata=metadata,
        )
    if not doi:
        return row_from_entry(
            entry,
            doi="",
            title=metadata.title or entry.title,
            status="needs_review",
            reason=metadata.reason or "metadata_not_found",
            metadata=metadata,
        )
    return row_from_entry(
        entry,
        doi=doi,
        title=metadata.title or entry.title,
        status="valid",
        metadata=metadata,
    )


def row_from_entry(
    entry: SourceEntry,
    *,
    doi: str,
    title: str,
    status: str,
    reason: str = "",
    metadata: MetadataResult | None = None,
) -> IntakeRow:
    authors = entry.authors
    journal = entry.journal
    year = entry.year
    metadata_source = ""
    match_basis = ""
    confidence = ""
    if metadata is not None:
        authors = "; ".join(metadata.authors) if metadata.authors else authors
        journal = metadata.journal or journal
        year = metadata.year or year
        metadata_source = metadata.source
        match_basis = metadata.match_basis
        confidence = f"{metadata.confidence:.3f}"
    normalized_doi = clean_doi(doi).lower() if doi else ""
    review_hint = intake_review_hint(status, reason, normalized_doi)
    return IntakeRow(
        source=entry.source,
        row_number=entry.row_number,
        input_doi=entry.doi,
        doi=normalized_doi,
        input_title=entry.title,
        title=title,
        authors=authors,
        journal=journal,
        year=year,
        date=entry.date,
        raw_value=entry.raw_text,
        metadata_source=metadata_source,
        match_basis=match_basis,
        confidence=confidence,
        status=status,
        reason=reason,
        review_hint=review_hint,
    )


def intake_review_hint(status: str, reason: str, doi: str = "") -> str:
    reason_value = str(reason or "")
    doi_value = clean_doi(doi).lower()
    if status == "valid":
        if doi_value.startswith("10.1016/"):
            return "可进入 ScienceDirect 解析；正式下载前仍需机构权限。"
        return "疑似非 ScienceDirect/Elsevier DOI；可能无法通过 ScienceDirect 下载。"
    if status == "duplicate":
        return "重复项，程序只保留首次识别记录。"
    if status == "empty":
        return "空行或无可用信息，可忽略。"
    if status == "invalid":
        return "请检查 DOI 格式，或重新复制完整 DOI 后重跑。"
    if status == "needs_review":
        if reason_value.startswith("metadata_confidence_below_threshold"):
            return "题名匹配置信度低，请补 DOI 或完整期刊、年份、卷页后重跑。"
        if reason_value == "insufficient_bibliographic_context":
            return "信息不足，请补 DOI、完整题名、期刊、年份或卷页后重跑。"
        if reason_value == "not_probable_title":
            return "不像完整论文题名，请补 DOI 或删除说明性文字后重跑。"
        if reason_value == "metadata_not_found":
            return "公开元数据未找到，请补 DOI 或更完整引用后重跑。"
        return "请人工确认 DOI 后重跑。"
    return ""


def split_authors(authors: str) -> list[str]:
    return [part.strip() for part in str(authors or "").split(";") if part.strip()]


def run_elsevier_api_phase(
    rows: list[IntakeRow],
    *,
    run_dir: Path,
    email: str,
    download_supplements: bool,
    client: ElsevierApiClient | None = None,
) -> ElsevierApiPhaseResult:
    """Try Elsevier before any browser/cookie work and return mergeable rows."""

    api_client = client or ElsevierApiClient()
    resolved_records: list[dict] = []
    pdf_records: list[PdfDownloadRecord] = []
    supplement_records: list[SupplementDownloadRecord] = []
    fallback_rows: list[IntakeRow] = []
    attempts: list[ElsevierApiAttemptRecord] = []
    circuit_status = ""

    for row in rows:
        if circuit_status:
            result = ElsevierApiResult(
                status=circuit_status,
                doi=row.doi,
                reason=f"api_circuit_open_after_{circuit_status}",
            )
        else:
            result = api_client.download_article(row.doi, include_supplements=download_supplements)

        final_status = result.status
        final_reason = result.reason
        article: dict | None = None
        pdf_size = 0
        pdf_valid = False

        if result.status == "success" and is_pdf_bytes(result.pdf_bytes):
            article = _article_from_api_result(row, result, email=email)
            target_name = _api_pdf_filename(article)
            target_path = run_dir / "pdfs" / target_name
            try:
                if target_path.is_symlink():
                    raise OSError("unsafe_symlink")
                if not is_valid_pdf(target_path):
                    write_pdf_bytes_atomic(target_path, result.pdf_bytes)
                pdf_valid = is_valid_pdf(target_path)
                if not pdf_valid:
                    raise ValueError("published_pdf_invalid")
                pdf_size = target_path.stat().st_size
            except OSError:
                final_status = "network_error"
                final_reason = "pdf_atomic_write_error"
            except (TypeError, ValueError):
                final_status = "invalid_pdf"
                final_reason = "published_pdf_invalid"
            else:
                article["file"] = target_name
                resolved_records.append(article)
                pdf_records.append(
                    PdfDownloadRecord(
                        doi=row.doi,
                        pii=result.pii,
                        title=article.get("title", ""),
                        status="success",
                        file=target_name,
                    )
                )
                if download_supplements:
                    supplement_records.extend(
                        download_elsevier_api_supplements(
                            api_client,
                            result,
                            article=article,
                            article_file=target_name,
                            run_dir=run_dir,
                        )
                    )
        elif result.status == "success":
            final_status = "invalid_pdf"
            final_reason = "client_success_without_valid_pdf"

        browser_fallback = final_status != "success"
        if browser_fallback:
            fallback_rows.append(row)
        attempts.append(
            ElsevierApiAttemptRecord(
                doi=row.doi,
                status=final_status,
                http_status=result.http_status,
                api_key_present=bool(api_client.api_key),
                insttoken_present=bool(api_client.insttoken),
                full_xml_received=result.full_xml_received,
                attachment_eid=result.attachment_eid,
                main_eid_present=bool(result.attachment_eid),
                pdf_size_bytes=pdf_size,
                pdf_valid=pdf_valid,
                browser_fallback=browser_fallback,
                reason=final_reason,
            )
        )
        if not circuit_status:
            if result.status in {"unauthorized", "rate_limited"}:
                circuit_status = result.status
            elif result.supplement_status in {"unauthorized", "rate_limited"}:
                circuit_status = result.supplement_status

    return ElsevierApiPhaseResult(
        resolved_records=resolved_records,
        pdf_records=pdf_records,
        supplement_records=supplement_records,
        fallback_rows=fallback_rows,
        attempts=attempts,
    )


def _article_from_api_result(row: IntakeRow, result: ElsevierApiResult, *, email: str) -> dict:
    article = {
        "source_index": row.row_number,
        "title": row.title or result.title,
        "authors": row.authors or "; ".join(result.authors),
        "journal": row.journal or result.journal,
        "volume": "",
        "issue": "",
        "year": row.year or result.year,
        "date": row.date,
        "doi": row.doi,
        "abstract": "",
        "article_type": "",
        "open_access": "",
        "url": f"https://doi.org/{quote(row.doi)}",
        "pdf_url": "",
        "pii": result.pii,
    }
    if not article["title"] or not article["authors"] or not article["year"]:
        article = enrich_row_metadata_for_delivery(article, email=email)
    # API metadata is authoritative enough to fill any field that remained
    # empty after the existing DOI enrichment helper ran.
    article["title"] = article.get("title") or result.title or row.input_title or row.doi
    article["authors"] = article.get("authors") or "; ".join(result.authors)
    article["journal"] = article.get("journal") or result.journal
    article["year"] = article.get("year") or result.year
    return article


def _author_list(value: object) -> list[str]:
    text = str(value or "").replace(" | ", ";").replace("|", ";")
    return [part.strip() for part in text.split(";") if part.strip()]


def _api_pdf_filename(article: dict) -> str:
    metadata = MetadataResult(
        source_index=int(article.get("source_index") or 0),
        query_title=str(article.get("title") or ""),
        doi=str(article.get("doi") or ""),
        title=str(article.get("title") or ""),
        authors=_author_list(article.get("authors")),
        journal=str(article.get("journal") or ""),
        year=str(article.get("year") or ""),
        source="elsevier_api",
    )
    return make_pdf_filename(metadata)


def download_elsevier_api_supplements(
    client: ElsevierApiClient,
    result: ElsevierApiResult,
    *,
    article: dict,
    article_file: str,
    run_dir: Path,
) -> list[SupplementDownloadRecord]:
    if result.supplement_status != "success":
        return [
            SupplementDownloadRecord(
                doi=result.doi,
                pii=result.pii,
                article_title=str(article.get("title") or ""),
                article_file=article_file,
                status="failed",
                reason=f"elsevier_api_{result.supplement_status}:{result.supplement_reason or result.supplement_status}",
            )
        ]
    if not result.supplements:
        return [
            SupplementDownloadRecord(
                doi=result.doi,
                pii=result.pii,
                article_title=str(article.get("title") or ""),
                article_file=article_file,
                status="not_found",
                reason="elsevier_api_no_supplement_objects",
            )
        ]

    article_stem = Path(article_file).stem
    supplement_dir = run_dir / "supplements" / article_stem
    records: list[SupplementDownloadRecord] = []
    for index, attachment in enumerate(result.supplements, start=1):
        raw_label = Path(attachment.filename or attachment.eid).name
        candidate = SupplementCandidate(
            url=attachment.api_url,
            title=Path(raw_label).stem or "Supplementary material",
        )
        filename = make_supplement_filename(index, candidate, attachment.mime_type)
        suffix = Path(filename).suffix.lower()
        if suffix not in SAFE_EXTENSIONS:
            records.append(
                _api_supplement_record(
                    result,
                    article,
                    article_file,
                    attachment,
                    index,
                    "failed",
                    reason="unsupported_attachment_extension",
                )
            )
            continue

        target_path = supplement_dir / filename
        if _valid_existing_attachment(target_path, suffix):
            records.append(
                _api_supplement_record(
                    result,
                    article,
                    article_file,
                    attachment,
                    index,
                    "skipped",
                    file=str(Path("supplements") / article_stem / filename),
                    content_type=attachment.mime_type,
                    size_bytes=target_path.stat().st_size,
                    reason="existing_valid_attachment",
                )
            )
            continue

        # Keep the temporary filename out of the already long article-stem
        # directory.  On Windows a valid final attachment path can otherwise
        # cross MAX_PATH only while mkstemp adds its random suffix.  A sibling
        # staging directory remains on the same volume, so os.replace keeps the
        # existing atomic-publish semantics.
        staging_dir = run_dir / ".elsevier_api_tmp"
        fd: int | None = None
        temp_path: Path | None = None
        try:
            supplement_dir.mkdir(parents=True, exist_ok=True)
            staging_dir.mkdir(parents=True, exist_ok=True)
            fd, temp_name = tempfile.mkstemp(
                prefix=".attachment.",
                suffix=".tmp",
                dir=str(staging_dir),
            )
            temp_path = Path(temp_name)
            with os.fdopen(fd, "wb") as sink:
                fd = None
                object_result = client.stream_attachment(attachment, sink)
                sink.flush()
                os.fsync(sink.fileno())
            if object_result.status != "success":
                records.append(
                    _api_supplement_record(
                        result,
                        article,
                        article_file,
                        attachment,
                        index,
                        "failed",
                        content_type=object_result.content_type or attachment.mime_type,
                        reason=f"elsevier_api_{object_result.status}:{object_result.reason or object_result.status}",
                    )
                )
                continue
            if suffix == ".pdf":
                valid = is_valid_pdf(temp_path)
            else:
                valid = temp_path.is_file() and not temp_path.is_symlink() and temp_path.stat().st_size > 0
            if not valid:
                records.append(
                    _api_supplement_record(
                        result,
                        article,
                        article_file,
                        attachment,
                        index,
                        "failed",
                        content_type=object_result.content_type or attachment.mime_type,
                        reason="invalid_attachment_content",
                    )
                )
                continue
            os.replace(temp_path, target_path)
            records.append(
                _api_supplement_record(
                    result,
                    article,
                    article_file,
                    attachment,
                    index,
                    "success",
                    file=str(Path("supplements") / article_stem / filename),
                    content_type=object_result.content_type or attachment.mime_type,
                    size_bytes=target_path.stat().st_size,
                )
            )
        except Exception as exc:
            records.append(
                _api_supplement_record(
                    result,
                    article,
                    article_file,
                    attachment,
                    index,
                    "failed",
                    reason=f"attachment_exception_{type(exc).__name__}",
                )
            )
        finally:
            if fd is not None:
                try:
                    os.close(fd)
                except OSError:
                    pass
            if temp_path is not None:
                try:
                    temp_path.unlink()
                except FileNotFoundError:
                    pass
    return records


def _valid_existing_attachment(path: Path, suffix: str) -> bool:
    if path.is_symlink() or not path.is_file():
        return False
    return is_valid_pdf(path) if suffix == ".pdf" else path.stat().st_size > 0


def _api_supplement_record(
    result: ElsevierApiResult,
    article: dict,
    article_file: str,
    attachment: ElsevierAttachment,
    index: int,
    status: str,
    *,
    file: str = "",
    content_type: str = "",
    size_bytes: int = 0,
    reason: str = "",
) -> SupplementDownloadRecord:
    return SupplementDownloadRecord(
        doi=result.doi,
        pii=result.pii,
        article_title=str(article.get("title") or ""),
        article_file=article_file,
        supplement_index=index,
        supplement_title=attachment.filename or attachment.eid,
        source_url=attachment.api_url,
        status=status,
        file=file,
        content_type=content_type or attachment.mime_type,
        size_bytes=size_bytes,
        reason=reason,
    )


def merge_resolved_records(
    input_rows: list[IntakeRow],
    api_records: list[dict],
    browser_records: list[dict],
) -> list[dict]:
    by_doi = {
        _doi_key(record.get("doi", "")): record
        for record in [*api_records, *browser_records]
        if _doi_key(record.get("doi", ""))
    }
    return [by_doi[_doi_key(row.doi)] for row in input_rows if _doi_key(row.doi) in by_doi]


def merge_pdf_records(
    input_rows: list[IntakeRow],
    api_records: list[PdfDownloadRecord],
    browser_records: list[PdfDownloadRecord],
    browser_failures: list[dict],
    *,
    download_requested: bool,
) -> list[PdfDownloadRecord]:
    by_doi = {
        _doi_key(record.doi): record
        for record in [*api_records, *browser_records]
        if _doi_key(record.doi)
    }
    failure_by_doi = {
        _doi_key(failure.get("doi", "")): failure
        for failure in browser_failures
        if _doi_key(failure.get("doi", ""))
    }
    merged: list[PdfDownloadRecord] = []
    for row in input_rows:
        key = _doi_key(row.doi)
        record = by_doi.get(key)
        if record is not None:
            merged.append(record)
            continue
        failure = failure_by_doi.get(key, {})
        merged.append(
            PdfDownloadRecord(
                doi=row.doi,
                pii="",
                title=row.title,
                status="failed" if download_requested or failure else "not_requested",
                reason=str(failure.get("reason") or ("browser_result_missing" if download_requested else "not_requested")),
            )
        )
    return merged


def _doi_key(value: object) -> str:
    return str(value or "").strip().casefold()


def write_elsevier_api_attempt_report(records: list[ElsevierApiAttemptRecord], output_dir: Path) -> Path:
    path = output_dir / "elsevier_api_attempts.csv"
    fieldnames = [
        "doi",
        "status",
        "http_status",
        "api_key_present",
        "insttoken_present",
        "full_xml_received",
        "attachment_eid",
        "main_eid_present",
        "pdf_size_bytes",
        "pdf_valid",
        "browser_fallback",
        "reason",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for record in records:
            writer.writerow({name: getattr(record, name) for name in fieldnames})
    return path


def write_resolve_failed_report(failures: list[dict], path: Path) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=["row_number", "doi", "reason"], extrasaction="ignore")
        writer.writeheader()
        writer.writerows(failures)
    return str(path)


def save_resolved_results(results: list[dict], path: Path) -> str:
    fields = ScienceDirectScraper.FIELDS
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "papers"
        worksheet.append(fields)
        for item in results:
            worksheet.append([item.get(field, "") for field in fields])
        workbook.save(path)
        return str(path)
    except ImportError:
        csv_path = path.with_suffix(".csv")
        with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(results)
        return str(csv_path)


def write_intake_preview(rows: list[IntakeRow], path: Path) -> Path:
    fieldnames = [
        "source",
        "row_number",
        "input_doi",
        "doi",
        "input_title",
        "title",
        "authors",
        "journal",
        "year",
        "date",
        "raw_value",
        "metadata_source",
        "match_basis",
        "confidence",
        "status",
        "reason",
        "review_hint",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row.__dict__)
    return path


def write_merged_input(rows: list[IntakeRow], path: Path) -> Path:
    fieldnames = ["source", "row_number", "title", "authors", "journal", "year", "date", "doi", "raw_value"]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            if not row.doi:
                continue
            writer.writerow({
                "source": row.source,
                "row_number": row.row_number,
                "title": row.title,
                "authors": row.authors,
                "journal": row.journal,
                "year": row.year,
                "date": row.date,
                "doi": row.doi,
                "raw_value": row.raw_value,
            })
    return path


def write_intake_failed_report(rows: list[IntakeRow], path: Path) -> Path:
    fieldnames = ["row_number", "doi", "reason"]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            if row.status == "valid" and row.doi:
                continue
            writer.writerow({
                "row_number": row.row_number,
                "doi": row.input_doi or row.doi,
                "reason": f"{row.status}: {row.reason or row.status}",
            })
    return path


def intake_failure_reasons(rows: list[IntakeRow]) -> dict[str, int]:
    counts: Counter[str] = Counter()
    for row in rows:
        if row.status == "valid" and row.doi:
            continue
        counts[row.reason or row.status] += 1
    return dict(counts)


def build_beginner_recommendations(
    intake: IntakeResult,
    failure_reasons: dict[str, int] | None = None,
    pdf_failed: int = 0,
    *,
    preflight_only: bool = False,
    auto_web_search: bool = False,
) -> list[str]:
    recommendations: list[str] = []
    valid_count = intake.status_counts.get("valid", 0)
    review_count = intake.status_counts.get("needs_review", 0)
    duplicate_count = intake.status_counts.get("duplicate", 0)
    non_sciencedirect_count = sum(
        1 for row in intake.unique_rows
        if row.doi and not row.doi.lower().startswith("10.1016/")
    )

    if preflight_only:
        recommendations.append("这是预检结果；确认 doi_intake_preview.csv 后，去掉 --preflight/--beginner 再正式下载。")
    if valid_count:
        recommendations.append(f"已有 {valid_count} 条唯一 DOI 可进入后续解析；正式下载仍依赖机构权限。")
    if review_count:
        recommendations.append(f"有 {review_count} 条需要人工复核；优先按 review_hint 补 DOI、完整题名、期刊、年份或卷页。")
    if duplicate_count:
        recommendations.append(f"发现 {duplicate_count} 条重复输入；程序只保留首次识别记录。")
    if non_sciencedirect_count:
        recommendations.append(f"有 {non_sciencedirect_count} 条 DOI 疑似不是 ScienceDirect/Elsevier，可改用 OA 流程。")
    if not auto_web_search and review_count and not preflight_only:
        recommendations.append("若题名或短引用较多，可重跑正式解析时加 --resolve-title-only --auto-web-search 尝试公开学术搜索补 DOI。")
    if pdf_failed:
        recommendations.append("PDF 失败时先看 pdf_download_report.csv；常见原因是 cookie 过期、无机构权限、验证码或限速。")
    if failure_reasons:
        recommendations.append("解析失败时先按失败原因分组处理，可按失败原因分组后分批处理。")
    if not recommendations:
        recommendations.append("当前没有需要处理的异常项。")
    return recommendations


def choose_output_root(default_out: str, dialog_func: Callable[[], str] | None = None) -> Path:
    if dialog_func is None:
        try:
            from tkinter import Tk, filedialog
        except Exception as exc:
            print(f"[提示] 无法打开文件夹选择器，使用默认输出目录: {default_out} ({exc})")
            return Path(default_out).expanduser().resolve()

        def dialog_func() -> str:
            root = Tk()
            root.withdraw()
            try:
                return filedialog.askdirectory(title="选择 ScienceDirect PDF 输出目录") or ""
            finally:
                root.destroy()

    try:
        selected = dialog_func()
    except Exception as exc:
        print(f"[提示] 文件夹选择器不可用，使用默认输出目录: {default_out} ({exc})")
        selected = ""
    return Path(selected or default_out).expanduser().resolve()


def make_scraper(
    cookie_cache_path: Path,
    browser_exe: str | None = None,
    cookies_path: str | None = None,
) -> ScienceDirectScraper:
    browser_kwargs = {"browser_exe": browser_exe} if browser_exe else {}
    if cookies_path:
        cookie_check = check_cookie_json(cookies_path)
        print(f"[Cookie] {cookie_check.message}")
        if cookie_check.is_usable:
            return ScienceDirectScraper(
                cookies_file=str(Path(cookies_path).expanduser().resolve()),
                use_browser_cookies=False,
                **browser_kwargs,
            )
        print("[Cookie] 显式 Cookie 文件不可用，将回退到缓存或本机浏览器 Cookie。")
    if cookie_cache_path.exists() and cookie_cache_path.stat().st_size > 0:
        cookie_check = check_cookie_json(cookie_cache_path)
        if cookie_check.is_usable:
            return ScienceDirectScraper(
                cookies_file=str(cookie_cache_path),
                use_browser_cookies=False,
                **browser_kwargs,
            )
        print(f"[Cookie] 缓存不可用，将改从本机浏览器读取: {cookie_check.message}")
    return ScienceDirectScraper(use_browser_cookies=True, **browser_kwargs)


def explicit_cookie_status_message(cookies_path: str | None) -> str:
    if not cookies_path:
        return ""
    return check_cookie_json(cookies_path).message


def cookie_status_message(cookie_cache_path: Path) -> str:
    count = 0
    if cookie_cache_path.exists():
        try:
            data = json.loads(cookie_cache_path.read_text(encoding="utf-8"))
            if isinstance(data, list):
                count = len(data)
        except Exception:
            return f"Cookie 缓存存在但不可读: {cookie_cache_path}"
    if count:
        return f"已缓存 {count} 个 ScienceDirect/Elsevier 相关 Cookie: {cookie_cache_path}"
    return "未发现可用 Cookie 缓存；下载时将尝试从本机浏览器/调试会话获取"


def cache_devtools_cookies(
    scraper: ScienceDirectScraper,
    cookie_cache_path: Path,
    *,
    opener: Callable = urlopen,
    websocket_connect: Callable | None = None,
) -> int:
    if websocket_connect is None:
        try:
            import websocket
        except ImportError:
            return 0
        websocket_connect = websocket.create_connection

    if not scraper._is_chrome_debug_ready():
        scraper._launch_chrome_with_debug()
    if not scraper._is_chrome_debug_ready():
        return 0

    try:
        cookies = extract_devtools_cookies(
            debug_port=scraper.CHROME_DBG_PORT,
            opener=opener,
            websocket_connect=websocket_connect,
        )
    except Exception:
        return 0

    relevant = [cookie for cookie in cookies if is_relevant_cookie(cookie)]
    if not relevant:
        return 0
    write_cookie_cache(relevant, cookie_cache_path)
    return len(relevant)


def extract_devtools_cookies(
    *,
    debug_port: int,
    opener: Callable = urlopen,
    websocket_connect: Callable,
) -> list[dict[str, object]]:
    base_url = f"http://127.0.0.1:{debug_port}"
    tab = None
    ws = None
    try:
        req = Request(f"{base_url}/json/new?{quote('about:blank', safe=':/?&=%')}", method="PUT")
        with opener(req, timeout=20) as response:
            tab = json.loads(response.read())
        ws = websocket_connect(tab["webSocketDebuggerUrl"], timeout=30, suppress_origin=True)
        ws.send(json.dumps({"id": 1, "method": "Network.getAllCookies"}))
        deadline = time.time() + 20
        while time.time() < deadline:
            msg = json.loads(ws.recv())
            if msg.get("id") == 1:
                return msg.get("result", {}).get("cookies", []) or []
    finally:
        if ws is not None:
            try:
                ws.close()
            except Exception:
                pass
        if tab:
            try:
                opener(Request(f"{base_url}/json/close/{tab['id']}"), timeout=10)
            except Exception:
                pass
    return []


def is_relevant_cookie(cookie: dict[str, object]) -> bool:
    domain = str(cookie.get("domain") or "").lower()
    name = str(cookie.get("name") or "").lower()
    return any(cookie_domain in domain for cookie_domain in COOKIE_DOMAINS) or name in {"euid", "sdmsession"}


def write_cookie_cache(cookies: list[dict[str, object]], path: Path) -> Path:
    safe_fields = (
        "name",
        "value",
        "domain",
        "path",
        "expires",
        "httpOnly",
        "secure",
        "sameSite",
    )
    serializable = []
    for cookie in cookies:
        item = {field: cookie[field] for field in safe_fields if field in cookie}
        if item.get("name") and item.get("value"):
            serializable.append(item)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(serializable, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


if __name__ == "__main__":
    sys.exit(main())
