from __future__ import annotations

import csv
import errno
import hashlib
import inspect
import json
import math
import os
import shutil
import stat
import tempfile
import time
from contextlib import contextmanager
from dataclasses import asdict, dataclass, is_dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterator

if os.name == "nt":
    import msvcrt
else:
    import fcntl


_WINDOWS_INVALID_CHARS = frozenset('<>:"|?*')
_WINDOWS_RESERVED_NAMES = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    *(f"COM{number}" for number in range(1, 10)),
    *(f"LPT{number}" for number in range(1, 10)),
    "COM¹",
    "COM²",
    "COM³",
    "LPT¹",
    "LPT²",
    "LPT³",
}

_LOCK_POLL_SECONDS = 0.05
_STATE_REPLACE_TIMEOUT_SECONDS = 2.0


@dataclass(frozen=True)
class BatchPaths:
    root: Path
    pdfs: Path
    reports: Path
    working: Path
    state: Path
    normalized_input: Path
    manual_retry: Path
    zotero_fallback: Path
    zotero_results: Path


NORMALIZED_FIELDS = [
    "task_id",
    "source_index",
    "input_doi",
    "input_title",
    "doi",
    "title",
    "authors",
    "journal",
    "year",
    "publisher",
    "status",
    "source",
    "file",
    "reason",
]

SUCCESS_STATUSES = {
    "oa_downloaded",
    "institutional_downloaded",
    "zotero_existing_pdf",
    "zotero_downloaded",
    "manual_imported",
}
ZOTERO_RESULT_FIELDS = [
    "task_id",
    "zotero_item_id",
    "attachment_path",
    "status",
    "reason",
]
ZOTERO_SUCCESS = {
    "existing_pdf": "zotero_existing_pdf",
    "downloaded": "zotero_downloaded",
}
ZOTERO_FAILURE_STATUSES = {
    "no_pdf",
    "not_found",
    "metadata_uncertain",
    "zotero_unavailable",
    "no_attachment",
    "download_failed",
    "zotero_api_unavailable",
    "user_cancelled",
    "job_expired",
    "job_id_conflict",
    "plugin_error",
}
ZOTERO_INPUT_STATUSES = set(ZOTERO_SUCCESS) | ZOTERO_FAILURE_STATUSES
FINAL_MANIFEST_FIELDS = [*NORMALIZED_FIELDS, "zotero_item_id"]
FINAL_REPORT_FILENAMES = (
    "final_manifest.csv",
    "final_manifest.xlsx",
    "failed.csv",
    "run_summary.txt",
    "batch_status.csv",
    "batch_status.json",
)
# User-facing delivery (keep this simple — one input, one inventory, and three
# predictable folders; reports/cache remain internal to the run directory).
USER_DELIVERY_DIR_NAME = "结果"
USER_INVENTORY_NAME = "下载清单.csv"
USER_PDF_DIR_NAME = "pdf"
USER_MD_DIR_NAME = "md"
USER_SUPPLEMENT_DIR_NAME = "补充材料"
USER_INVENTORY_FIELDS = [
    "序号",
    "状态",
    "DOI",
    "题名",
    "作者",
    "年份",
    "期刊",
    "下载来源",
    "结果文件",
    "补充材料",
    "失败原因",
    "task_id",
]
_STAGE_SUCCESS_STATUSES = {"downloaded", "oa_downloaded", "institutional_downloaded"}
_INSTITUTIONAL_SOURCES = {"sciencedirect", "non_elsevier", "institutional"}
_INPUT_SNAPSHOT_DIR_NAME = "input_source"


def user_delivery_dir(paths: BatchPaths) -> Path:
    """Return the only user-facing delivery directory for a batch."""

    return paths.root / USER_DELIVERY_DIR_NAME


def user_inventory_path(paths: BatchPaths) -> Path:
    """Return the user-facing inventory path inside ``结果``."""

    return user_delivery_dir(paths) / USER_INVENTORY_NAME


def user_pdf_dir(paths: BatchPaths) -> Path:
    """Return the user-facing article-PDF directory."""

    return user_delivery_dir(paths) / USER_PDF_DIR_NAME


def user_md_dir(paths: BatchPaths) -> Path:
    """Return the reserved directory for later PDF-to-Markdown conversion."""

    return user_delivery_dir(paths) / USER_MD_DIR_NAME


def user_supplement_dir(paths: BatchPaths) -> Path:
    """Return the optional user-facing supplementary-material directory."""

    return user_delivery_dir(paths) / USER_SUPPLEMENT_DIR_NAME


@dataclass(frozen=True)
class BatchRunResult:
    paths: BatchPaths
    total_count: int
    success_count: int
    failed_count: int
    manual_retry_count: int
    zotero_fallback_count: int


def _validate_path_component(value: str, *, error: str, clean_spaces: bool = False) -> str:
    raw = str(value)
    path = Path(raw)
    if (
        not raw
        or path.is_absolute()
        or bool(path.drive)
        or "/" in raw
        or "\\" in raw
        or ".." in raw
        or any(character in _WINDOWS_INVALID_CHARS or ord(character) < 32 for character in raw)
    ):
        raise ValueError(error)

    cleaned = "_".join(raw.strip().split()) if clean_spaces else raw
    if (
        not cleaned
        or cleaned in {".", ".."}
        or cleaned != cleaned.rstrip(". ")
        or cleaned.split(".", 1)[0].upper() in _WINDOWS_RESERVED_NAMES
    ):
        raise ValueError(error)
    return cleaned


def _batch_paths_for_root(root: Path) -> BatchPaths:
    """Build BatchPaths for an existing or newly created run root."""
    root = root.expanduser().resolve()
    pdfs, reports, working = root / "pdfs", root / "reports", root / "working"
    for path in (pdfs, reports, working):
        path.mkdir(parents=True, exist_ok=True)
    return BatchPaths(
        root=root,
        pdfs=pdfs,
        reports=reports,
        working=working,
        state=working / "batch_state.json",
        normalized_input=working / "normalized_input.csv",
        manual_retry=working / "manual_retry.csv",
        zotero_fallback=working / "zotero_fallback.csv",
        zotero_results=working / "zotero_results.csv",
    )


def default_run_name_from_input(
    input_path: str | Path | None,
    input_text: str | None = None,
) -> str:
    """Stable folder name from input file stem (or generic fallback)."""
    if input_path is not None:
        stem = Path(input_path).expanduser().stem.strip() or "paper_batch"
        try:
            return _validate_path_component(stem, error="invalid_run_name", clean_spaces=True)
        except ValueError:
            return "paper_batch"
    if input_text and str(input_text).strip():
        return "paper_batch_text"
    return "paper_batch"


def create_batch_paths(
    output_root: str | Path,
    run_name: str | None = None,
    now: datetime | None = None,
    *,
    fixed: bool = False,
) -> BatchPaths:
    """Create batch paths under output_root.

    - fixed=False (legacy): always ``{name}_{timestamp}`` unique folder.
    - fixed=True: use ``{name}`` only (reuse same folder; no new timestamp).
      Requires a run_name (or defaults to paper_batch).
    """
    prefix = "paper_batch" if run_name is None else _validate_path_component(
        run_name,
        error="invalid_run_name",
        clean_spaces=True,
    )
    output = Path(output_root).expanduser().resolve()
    output.mkdir(parents=True, exist_ok=True)

    if fixed:
        root = output / prefix
        if root.exists() and not root.is_dir():
            raise ValueError("fixed_run_path_not_directory")
        root.mkdir(parents=True, exist_ok=True)
        return _batch_paths_for_root(root)

    stamp = (now or datetime.now()).strftime("%Y%m%d_%H%M%S")
    base_name = f"{prefix}_{stamp}"
    counter = 1
    while True:
        name = base_name if counter == 1 else f"{base_name}_{counter}"
        root = output / name
        try:
            root.mkdir()
        except FileExistsError:
            counter += 1
        else:
            break

    return _batch_paths_for_root(root)


def _state_path(run_dir: str | Path) -> Path:
    return Path(run_dir).expanduser().resolve() / "working" / "batch_state.json"


def _replace_state_file(temporary: Path, destination: Path) -> None:
    deadline = time.monotonic() + _STATE_REPLACE_TIMEOUT_SECONDS
    while True:
        try:
            os.replace(temporary, destination)
        except OSError as exc:
            retryable = os.name == "nt" and (
                isinstance(exc, PermissionError)
                or getattr(exc, "winerror", None) in {5, 32, 33}
            )
            remaining = deadline - time.monotonic()
            if not retryable or remaining <= 0:
                raise
            time.sleep(min(_LOCK_POLL_SECONDS, remaining))
        else:
            return


def _write_batch_state_file(path: Path, payload: dict) -> Path:
    serialized = json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
    path.parent.mkdir(parents=True, exist_ok=True)
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f"{path.name}.",
        suffix=".tmp",
        dir=path.parent,
    )
    temporary = Path(temporary_name)
    descriptor_open = True
    try:
        handle = os.fdopen(file_descriptor, "w", encoding="utf-8", newline="\n")
        descriptor_open = False
        with handle:
            handle.write(serialized)
            handle.flush()
            os.fsync(handle.fileno())
        _replace_state_file(temporary, path)
    finally:
        if descriptor_open:
            try:
                os.close(file_descriptor)
            except OSError:
                pass
        temporary.unlink(missing_ok=True)
    return path


def save_batch_state(paths: BatchPaths, payload: dict) -> Path:
    """Durably replace one snapshot; read-modify-write callers must hold the lock."""

    return _write_batch_state_file(paths.state, payload)


def load_batch_state(run_dir: str | Path) -> dict:
    """Read one snapshot; read-modify-write callers must hold ``batch_state_lock``."""

    path = _state_path(run_dir)
    return json.loads(path.read_text(encoding="utf-8"))


def _try_file_lock(handle) -> None:
    handle.seek(0)
    if os.name == "nt":
        msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
    else:
        fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)


def _release_file_lock(handle) -> None:
    handle.seek(0)
    if os.name == "nt":
        msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
    else:
        fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def _is_lock_contention(exc: OSError) -> bool:
    return (
        isinstance(exc, BlockingIOError)
        or exc.errno in {errno.EACCES, errno.EAGAIN}
        or getattr(exc, "winerror", None) in {32, 33}
    )


@contextmanager
def _file_lock(
    lock_path: Path,
    *,
    timeout: float,
    invalid_timeout_error: str,
    timeout_error: str,
) -> Iterator[None]:
    timeout_seconds = float(timeout)
    if not math.isfinite(timeout_seconds) or timeout_seconds < 0:
        raise ValueError(invalid_timeout_error)
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    handle = lock_path.open("a+b")
    acquired = False
    try:
        deadline = time.monotonic() + timeout_seconds
        while True:
            try:
                _try_file_lock(handle)
            except OSError as exc:
                if not _is_lock_contention(exc):
                    raise
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    raise TimeoutError(timeout_error) from exc
                time.sleep(min(_LOCK_POLL_SECONDS, remaining))
            else:
                acquired = True
                break
        yield
    finally:
        try:
            if acquired:
                _release_file_lock(handle)
        finally:
            handle.close()


@contextmanager
def batch_state_lock(
    run_dir: str | Path,
    *,
    timeout: float = 10.0,
) -> Iterator[None]:
    """Exclusively lock one run's state; this context manager is not re-entrant."""

    lock_path = _state_path(run_dir).with_suffix(".lock")
    with _file_lock(
        lock_path,
        timeout=timeout,
        invalid_timeout_error="invalid_batch_state_lock_timeout",
        timeout_error="batch_state_lock_timeout",
    ):
        yield


def claim_manual_retry(
    run_dir: str | Path,
    *,
    timeout: float = 10.0,
    validate_state: Callable[[dict], object] | None = None,
    validate_before_claim: Callable[[dict], object] | None = None,
) -> tuple[bool, dict]:
    """Atomically claim a run's one allowed manual retry and return its state."""

    with batch_state_lock(run_dir, timeout=timeout):
        state = load_batch_state(run_dir)
        if validate_state is not None:
            validate_state(state)
        if state.get("manual_retry_used"):
            return False, state
        if (
            validate_before_claim is not None
            and validate_before_claim(state) is False
        ):
            return False, state
        state["manual_retry_used"] = True
        _write_batch_state_file(_state_path(run_dir), state)
        return True, state


def is_valid_pdf(path: str | Path, minimum_size: int | None = None) -> bool:
    """Validate a local PDF via shared header + %%EOF checks."""
    from paper_automation.pdf_validation import DEFAULT_MINIMUM_SIZE, is_valid_pdf as _is_valid_pdf

    return _is_valid_pdf(path, minimum_size if minimum_size is not None else DEFAULT_MINIMUM_SIZE)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _same_pdf_content(path: Path, expected_hash: str) -> bool:
    return (
        not path.is_symlink()
        and is_valid_pdf(path)
        and _sha256(path) == expected_hash
    )


def _path_exists(path: Path) -> bool:
    return path.exists() or path.is_symlink()


@contextmanager
def _pdf_publish_lock(destination: Path, *, timeout: float) -> Iterator[None]:
    with _file_lock(
        destination / ".pdf_publish.lock",
        timeout=timeout,
        invalid_timeout_error="invalid_pdf_publish_lock_timeout",
        timeout_error="pdf_publish_lock_timeout",
    ):
        yield


def _cleanup_stale_pdf_snapshots(destination: Path) -> None:
    # Snapshot creation also holds the destination lock, so any snapshot visible
    # to the current lock owner was abandoned by a process that no longer owns it.
    for snapshot in destination.glob(".pdf_snapshot_*.tmp"):
        snapshot.unlink(missing_ok=True)


def _publish_pdf_snapshot(snapshot: Path, target: Path) -> bool:
    try:
        os.link(snapshot, target)
    except FileExistsError:
        return False
    except OSError as exc:
        raise OSError(f"hard_link_publish_failed: {exc}") from exc
    return True


def _snapshot_pdf_handle(source_handle, destination: Path) -> tuple[Path, str]:
    file_descriptor, snapshot_name = tempfile.mkstemp(
        prefix=".pdf_snapshot_",
        suffix=".tmp",
        dir=destination,
    )
    snapshot = Path(snapshot_name)
    digest = hashlib.sha256()
    descriptor_open = True
    try:
        snapshot_handle = os.fdopen(file_descriptor, "wb")
        descriptor_open = False
        with snapshot_handle:
            for chunk in iter(lambda: source_handle.read(1024 * 1024), b""):
                snapshot_handle.write(chunk)
                digest.update(chunk)
            snapshot_handle.flush()
            os.fsync(snapshot_handle.fileno())
        if not is_valid_pdf(snapshot):
            raise ValueError("not_pdf_response")
        return snapshot, digest.hexdigest()
    except BaseException:
        snapshot.unlink(missing_ok=True)
        raise
    finally:
        if descriptor_open:
            try:
                os.close(file_descriptor)
            except OSError:
                pass


def _snapshot_pdf_source(source: Path, destination: Path) -> tuple[Path, str]:
    with source.open("rb") as source_handle:
        return _snapshot_pdf_handle(source_handle, destination)


def _publish_verified_pdf_snapshot(
    snapshot: Path,
    destination: Path,
    filename: str,
    source_hash: str,
    *,
    reuse_any_hash: bool = False,
) -> Path:
    """Publish an already validated private snapshot without reopening its source."""

    if reuse_any_hash:
        for candidate in sorted(destination.glob("*.pdf"), key=lambda path: path.name.casefold()):
            if _same_pdf_content(candidate, source_hash):
                return candidate

    target = destination / filename
    while True:
        if _path_exists(target):
            if _same_pdf_content(target, source_hash):
                return target
            break
        if _publish_pdf_snapshot(snapshot, target):
            return target

    counter = 1
    while True:
        suffix = "" if counter == 1 else f"_{counter}"
        candidate = target.with_name(
            f"{target.stem}_{source_hash[:8]}{suffix}{target.suffix}"
        )
        if _path_exists(candidate):
            if _same_pdf_content(candidate, source_hash):
                return candidate
            counter += 1
            continue
        if _publish_pdf_snapshot(snapshot, candidate):
            return candidate


def copy_pdf_safely(
    source: str | Path,
    destination_dir: str | Path,
    filename: str,
    *,
    lock_timeout: float = 10.0,
) -> Path:
    safe_filename = _validate_path_component(filename, error="invalid_filename")
    source_path = Path(source).expanduser().resolve()
    if not source_path.is_file():
        raise ValueError("not_pdf_response")
    # Keep the caller's absolute spelling for the returned path. Windows may
    # resolve the same directory as either a long path or an 8.3 short path
    # (for example RUNNER~1); returning the worker-local spelling makes
    # otherwise identical paths compare differently across processes.
    display_destination = Path(destination_dir).expanduser()
    if not display_destination.is_absolute():
        display_destination = Path.cwd() / display_destination
    destination = display_destination.resolve()
    destination.mkdir(parents=True, exist_ok=True)
    with _pdf_publish_lock(destination, timeout=lock_timeout):
        _cleanup_stale_pdf_snapshots(destination)
        snapshot = None
        try:
            snapshot, source_hash = _snapshot_pdf_source(source_path, destination)
            published = _publish_verified_pdf_snapshot(
                snapshot,
                destination,
                safe_filename,
                source_hash,
            )
            return display_destination / published.name
        finally:
            if snapshot is not None:
                snapshot.unlink(missing_ok=True)


def run_oa_stage(rows: list[dict], output_dir: Path, options: Any):
    """Late import keeps the Task 2 PDF helper free of a circular import."""

    from .batch_stages import run_oa_stage as stage

    return stage(rows, output_dir, options)


def run_sciencedirect_stage(input_path: Path, output_dir: Path, options: Any):
    from .batch_stages import run_sciencedirect_stage as stage

    return stage(input_path, output_dir, options)


def run_non_elsevier_stage(input_path: Path, output_dir: Path, options: Any):
    from .batch_stages import run_non_elsevier_stage as stage

    return stage(input_path, output_dir, options)


def _batch_options_type():
    from .batch_stages import BatchOptions

    return BatchOptions


def _stage_input_writer(rows: list[dict], path: Path) -> Path:
    from .batch_stages import write_stage_input

    return write_stage_input(rows, path)


def _split_institutional_rows(rows: list[dict]) -> tuple[list[dict], list[dict]]:
    from .batch_stages import split_institutional_rows

    return split_institutional_rows(rows)


def _needs_manual_retry(status: str, reason: str) -> bool:
    from .batch_stages import needs_manual_retry

    return needs_manual_retry(status, reason)


def _normalise_doi(value: object) -> str:
    from doi_batch_utils import clean_doi

    return clean_doi(str(value or "")).lower()


def _empty_normalized_row() -> dict[str, str]:
    return {field: "" for field in NORMALIZED_FIELDS}


def _normalise_row_mapping(row: object) -> dict:
    if is_dataclass(row):
        raw = asdict(row)
    elif isinstance(row, dict):
        raw = dict(row)
    else:
        raise ValueError("invalid_batch_row")
    result = _empty_normalized_row()
    result.update({key: value for key, value in raw.items() if key in result or key.startswith("_") or key == "fixture_pdf"})
    for field in NORMALIZED_FIELDS:
        result[field] = str(result.get(field, "") or "")
    return result


def _write_csv_rows(path: Path, rows: list[dict]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=NORMALIZED_FIELDS)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in NORMALIZED_FIELDS})
    return path


def _atomic_report_file(path: Path, suffix: str, writer: Callable[[Path], None]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.stem}.",
        suffix=suffix,
        dir=path.parent,
    )
    temporary = Path(temporary_name)
    try:
        os.close(descriptor)
        writer(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    return path


def _report_row(row: dict, fields: list[str]) -> dict[str, str]:
    return {field: str(row.get(field, "") or "") for field in fields}


def _csv_safe_value(value: object) -> str:
    text = str(value or "")
    stripped = text.lstrip()
    return f"'{text}" if stripped[:1] in {"=", "+", "-", "@"} else text


def _write_report_csv(path: Path, fields: list[str], rows: list[dict]) -> Path:
    normalized_rows = [
        {field: _csv_safe_value(row.get(field, "")) for field in fields}
        for row in rows
    ]

    def write(temporary: Path) -> None:
        with temporary.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(normalized_rows)
            handle.flush()
            os.fsync(handle.fileno())

    return _atomic_report_file(path, ".csv.tmp", write)


def _write_report_text(path: Path, content: str) -> Path:
    def write(temporary: Path) -> None:
        with temporary.open("w", encoding="utf-8", newline="\n") as handle:
            handle.write(content)
            handle.flush()
            os.fsync(handle.fileno())

    return _atomic_report_file(path, ".txt.tmp", write)


def _write_report_json(path: Path, payload: dict[str, object]) -> Path:
    content = json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
    return _write_report_text(path, content)


def _write_final_manifest_xlsx(path: Path, rows: list[dict]) -> Path:
    manifest_rows = [_report_row(row, FINAL_MANIFEST_FIELDS) for row in rows]

    def write(temporary: Path) -> None:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell

        workbook = Workbook(write_only=True)
        try:
            worksheet = workbook.create_sheet("final_manifest")
            for values in [FINAL_MANIFEST_FIELDS, *(
                [row[field] for field in FINAL_MANIFEST_FIELDS] for row in manifest_rows
            )]:
                cells = []
                for value in values:
                    cell = WriteOnlyCell(worksheet, value=str(value or ""))
                    cell.data_type = "s"
                    cells.append(cell)
                worksheet.append(cells)
            workbook.save(temporary)
        finally:
            workbook.close()
        with temporary.open("r+b") as handle:
            handle.flush()
            os.fsync(handle.fileno())

    return _atomic_report_file(path, ".xlsx.tmp", write)


def _copy_report_backup(source: Path, destination: Path) -> None:
    with source.open("rb") as source_handle, destination.open("xb") as target_handle:
        for chunk in iter(lambda: source_handle.read(1024 * 1024), b""):
            target_handle.write(chunk)
        target_handle.flush()
        os.fsync(target_handle.fileno())


def _cleanup_report_transaction_dir(directory: Path) -> None:
    for filename in FINAL_REPORT_FILENAMES:
        (directory / filename).unlink(missing_ok=True)
    try:
        directory.rmdir()
    except OSError:
        pass


def _publish_report_set(paths: BatchPaths, staging: Path) -> None:
    backup = Path(tempfile.mkdtemp(prefix=".final_reports_backup_", dir=paths.reports))
    existed: dict[str, bool] = {}
    try:
        for filename in FINAL_REPORT_FILENAMES:
            target = paths.reports / filename
            existed[filename] = target.exists() or target.is_symlink()
            if existed[filename]:
                _copy_report_backup(target, backup / filename)
        try:
            for filename in FINAL_REPORT_FILENAMES:
                os.replace(staging / filename, paths.reports / filename)
        except Exception as publish_error:
            rollback_errors: list[str] = []
            for filename in FINAL_REPORT_FILENAMES:
                target = paths.reports / filename
                try:
                    if existed[filename]:
                        os.replace(backup / filename, target)
                    else:
                        target.unlink(missing_ok=True)
                except OSError as rollback_error:
                    rollback_errors.append(f"{filename}:{rollback_error}")
            if rollback_errors:
                raise RuntimeError(
                    "final_report_rollback_failed:" + ";".join(rollback_errors)
                ) from publish_error
            raise
    finally:
        _cleanup_report_transaction_dir(backup)


def _tabular_doi_cells(input_path: Path) -> dict[str, str] | None:
    if input_path.suffix.lower() not in {".csv", ".tsv", ".xlsx", ".xlsm"}:
        return None
    from doi_batch_utils import COLUMN_ALIASES, find_column
    from sd_institutional_skill import load_tabular_records, read_tabular_headers

    headers = read_tabular_headers(input_path)
    has_doi_column = find_column(headers, COLUMN_ALIASES["doi"]) is not None
    records = load_tabular_records(input_path)
    return {
        str(record.row_number): str(record.raw_value or "") if has_doi_column else ""
        for record in records
    }


def normalize_input(
    *,
    input_text: str | None,
    input_path: str | Path | None,
    paths: BatchPaths,
    options: Any,
) -> list[dict]:
    """Use the existing intake contract without ever modifying the source input."""

    if (input_text is None) == (input_path is None):
        raise ValueError("exactly_one_input_required")
    from paper_automation.parser import extract_dois
    from sd_institutional_skill import build_intake

    path_values = [] if input_path is None else [Path(input_path).expanduser().resolve()]
    tabular_doi_cells = None if not path_values else _tabular_doi_cells(path_values[0])
    # Opt1: skip title-only networked metadata unless explicitly enabled.
    resolve_title = bool(getattr(options, "resolve_title_metadata", False))
    intake = build_intake(
        texts=[] if input_text is None else [input_text],
        input_paths=path_values,
        folder_paths=[],
        output_dir=paths.working / "intake",
        resolve_metadata=resolve_title,
        resolve_title_only_files=resolve_title,
        min_confidence=0.92,
        email=str(getattr(options, "email", "") or ""),
    )

    from .failure_routing import is_noise_title_line
    from paper_automation.parser import is_probable_paper_title

    rows: list[dict] = []
    seen_dois: set[str] = set()
    seen_titles: set[str] = set()
    pending_title_for_doi: str = ""
    for intake_row in intake.all_rows:
        raw = asdict(intake_row)
        title = str(raw.get("title", "") or "").strip()
        raw_value = str(raw.get("raw_value", "") or "").strip()
        intake_status = str(raw.get("status", "") or "").strip().lower()
        intake_reason = str(raw.get("reason", "") or "").strip()
        if intake_status == "duplicate":
            continue
        # B1.1: drop markdown section headers / notes early.
        display_text = title or raw_value
        if is_noise_title_line(display_text, intake_reason) and not extract_dois(
            f"{raw_value} {raw.get('input_doi', '')} {raw.get('doi', '')}"
        ):
            continue
        extracted_dois: list[str] = []
        source_index = str(raw.get("row_number", "") or "")
        explicit_doi_cell = "" if tabular_doi_cells is None else tabular_doi_cells.get(source_index, "")
        doi_not_from_column = bool(
            tabular_doi_cells is not None
            and not explicit_doi_cell.strip()
            and str(raw.get("input_doi", "") or "").strip()
        )
        if tabular_doi_cells is None:
            doi_sources = (raw.get("raw_value", ""), raw.get("input_doi", ""), raw.get("doi", ""))
        elif explicit_doi_cell.strip():
            doi_sources = (explicit_doi_cell,)
        elif intake_status == "valid" and not doi_not_from_column:
            doi_sources = (raw.get("doi", ""),)
        else:
            doi_sources = ()
        for value in doi_sources:
            for doi in extract_dois(value):
                normalized_doi = _normalise_doi(doi)
                if normalized_doi and normalized_doi not in extracted_dois:
                    extracted_dois.append(normalized_doi)

        if not extracted_dois:
            # A1: default DOI-only intake — drop title-only noise unless title metadata mode.
            if doi_not_from_column:
                # Safety audit: DOI appeared outside the DOI column.
                extracted_dois = [""]
            elif not resolve_title:
                # Remember a plausible title so a following DOI-only line can inherit it.
                if (
                    title
                    and is_probable_paper_title(title)
                    and not is_noise_title_line(title, intake_reason)
                ):
                    pending_title_for_doi = title
                continue
            else:
                if not title or is_noise_title_line(title, intake_reason):
                    continue
                if intake_reason in {"not_probable_title"} and not is_probable_paper_title(title):
                    continue
                title_key = " ".join(title.lower().split())
                if title_key and title_key in seen_titles:
                    continue
                if title_key:
                    seen_titles.add(title_key)
                extracted_dois = [""]
                pending_title_for_doi = ""

        for doi in extracted_dois:
            if doi:
                if doi in seen_dois:
                    continue
                seen_dois.add(doi)
                if not title and pending_title_for_doi:
                    title = pending_title_for_doi
                    pending_title_for_doi = ""
            is_pending = intake_status == "valid" and bool(doi)
            row = _empty_normalized_row()
            row.update({
                "source_index": source_index,
                "input_doi": explicit_doi_cell if explicit_doi_cell else str(raw.get("input_doi", "") or ""),
                "input_title": str(raw.get("input_title", "") or "") or title,
                "doi": doi,
                "title": title,
                "authors": str(raw.get("authors", "") or ""),
                "journal": str(raw.get("journal", "") or ""),
                "year": str(raw.get("year", "") or ""),
                "status": "pending" if is_pending else "metadata_uncertain",
                "source": str(raw.get("source", "") or "input"),
                "reason": (
                    "doi_not_from_doi_column"
                    if doi_not_from_column
                    else "" if is_pending
                    else str(raw.get("reason", "") or "metadata_uncertain")
                ),
            })
            rows.append(row)

    if not rows:
        raise ValueError("empty_input")
    for index, row in enumerate(rows, start=1):
        row["task_id"] = f"paper-{index:04d}"
    _snapshot_input_source(paths, input_text=input_text, input_path=input_path)
    _write_csv_rows(paths.normalized_input, rows)
    return rows


def _snapshot_input_source(
    paths: BatchPaths,
    *,
    input_text: str | None,
    input_path: str | Path | None,
) -> Path:
    """Keep one immutable copy of the user input for the final delivery.

    The original source is never edited.  A fixed batch can be resumed after
    the original file has moved or been renamed because the snapshot lives in
    the internal ``working`` directory and is copied into ``结果`` at publish
    time.
    """

    snapshot_dir = paths.working / _INPUT_SNAPSHOT_DIR_NAME
    snapshot_dir.mkdir(parents=True, exist_ok=True)
    if input_path is not None:
        source = Path(input_path).expanduser().resolve()
        if not source.is_file():
            raise ValueError("input_file_missing")
        destination = snapshot_dir / source.name
        if destination.exists():
            if destination.read_bytes() != source.read_bytes():
                raise ValueError("input_snapshot_conflict")
        else:
            shutil.copy2(source, destination)
        return destination

    destination = snapshot_dir / "输入清单.txt"
    payload = str(input_text or "")
    if destination.exists():
        if destination.read_text(encoding="utf-8") != payload:
            raise ValueError("input_snapshot_conflict")
    else:
        destination.write_text(payload, encoding="utf-8", newline="")
    return destination


def _is_successful_status(status: object) -> bool:
    value = str(status or "").strip().lower()
    return value in SUCCESS_STATUSES


def _is_terminal_status(status: object) -> bool:
    return _is_successful_status(status) or str(status or "").strip().lower() == "duplicate"


def _update_as_mapping(update: object) -> dict:
    result = _normalise_row_mapping(update)
    if not result["task_id"]:
        raise ValueError("gateway_task_id_missing")
    return result


def _validate_stage_updates(rows: list[dict], updates: object) -> list[dict]:
    if not isinstance(updates, list):
        raise ValueError("invalid_gateway_updates")
    known_ids = {str(row.get("task_id", "")) for row in rows}
    if len(known_ids) != len(rows) or "" in known_ids:
        raise ValueError("invalid_state_task_ids")
    normalized_updates: list[dict] = []
    seen_ids: set[str] = set()
    for update in updates:
        normalized = _update_as_mapping(update)
        task_id = normalized["task_id"]
        if task_id not in known_ids:
            raise ValueError("gateway_task_id_unknown")
        if task_id in seen_ids:
            raise ValueError("gateway_task_id_duplicate")
        seen_ids.add(task_id)
        normalized_updates.append(normalized)
    return normalized_updates


def _copy_successful_pdf(row: dict, paths: BatchPaths, *, email: str = "") -> dict:
    status = str(row.get("status", "") or "").strip().lower()
    if status not in _STAGE_SUCCESS_STATUSES:
        return row
    result = dict(row)
    if status == "downloaded":
        source = str(row.get("source", "") or "").strip().lower()
        if source == "oa":
            canonical_status = "oa_downloaded"
        elif source in _INSTITUTIONAL_SOURCES:
            canonical_status = "institutional_downloaded"
        else:
            result.update(
                status="invalid_download_source",
                file="",
                reason=f"invalid_download_source:{source or 'missing'}",
            )
            return result
    else:
        canonical_status = status
    source_file = str(row.get("file", "") or "").strip()
    source_path = Path(source_file).expanduser() if source_file else None
    try:
        valid_source = bool(
            source_path is not None
            and not source_path.is_symlink()
            and is_valid_pdf(source_path)
        )
    except OSError:
        valid_source = False
    if not valid_source:
        result.update(status="not_pdf_response", file="", reason="not_pdf_response")
        return result
    try:
        # Mandatory delivery naming: 年份-作者-题名 at first publish (not a later rename).
        from .file_manager import enrich_row_metadata_for_delivery

        result = enrich_row_metadata_for_delivery(result, email=email)
        preferred_name = _filename_for_batch_row(result)
        stage_name = _canonical_stage_pdf_name(source_path.name)
        if _has_delivery_placeholders(preferred_name) and stage_name:
            # The ScienceDirect adapter may have richer API/publisher metadata
            # than the original DOI-only batch row.  Its file was already
            # written with make_pdf_filename() and validated above, so preserve
            # that canonical name instead of regressing to 0000-Unknown.
            preferred_name = stage_name
        copied = copy_pdf_safely(source_path, paths.pdfs, preferred_name)
    except (OSError, ValueError) as exc:
        result.update(
            status="not_pdf_response",
            file="",
            reason=f"not_pdf_response:{type(exc).__name__}",
        )
        return result
    result["status"] = canonical_status
    result["file"] = str(copied)
    result["reason"] = ""
    return result


def _has_delivery_placeholders(filename: str) -> bool:
    lowered = str(filename or "").casefold()
    return lowered.startswith("0000-") or "-unknown-" in lowered


def _canonical_stage_pdf_name(filename: str) -> str:
    """Accept only a safe 年份-作者-题名.pdf stage filename."""

    name = str(filename or "").strip()
    try:
        safe_name = _validate_path_component(name, error="invalid_filename")
    except ValueError:
        return ""
    if safe_name != name or Path(name).suffix.casefold() != ".pdf":
        return ""
    parts = Path(name).stem.split("-", 2)
    if len(parts) != 3:
        return ""
    year, author, title = parts
    if len(year) != 4 or not year.isdigit() or not year.startswith(("19", "20")):
        return ""
    if not author or author.casefold() == "unknown" or not title:
        return ""
    return name


def _merge_stage_rows(
    rows: list[dict],
    updates: object,
    paths: BatchPaths,
    *,
    email: str = "",
) -> list[dict]:
    normalized_updates = _validate_stage_updates(rows, updates)
    updates_by_id = {update["task_id"]: update for update in normalized_updates}
    merged_rows: list[dict] = []
    for existing in rows:
        task_id = str(existing.get("task_id", ""))
        update = updates_by_id.get(task_id)
        if update is None or _is_terminal_status(existing.get("status", "")):
            merged_rows.append(dict(existing))
            continue
        merged = dict(existing)
        merged.update(update)
        merged["task_id"] = task_id
        merged_rows.append(_copy_successful_pdf(merged, paths, email=email))
    return merged_rows


def _apply_stage_updates(
    state: dict,
    updates: object,
    paths: BatchPaths,
    *,
    assume_locked: bool = False,
) -> None:
    if not assume_locked:
        with batch_state_lock(paths.root):
            latest = load_batch_state(paths.root)
            _validate_state(latest, expected_run_dir=paths.root)
            _apply_stage_updates(latest, updates, paths, assume_locked=True)
            state.clear()
            state.update(latest)
        return

    _validate_state(state, expected_run_dir=paths.root)
    normalized_updates = _validate_stage_updates(state["rows"], updates)
    email = str((state.get("options") or {}).get("email", "") or "")
    for update in normalized_updates:
        state["rows"] = _merge_stage_rows(state["rows"], [update], paths, email=email)
        save_batch_state(paths, state)


def _required_task_ids(rows: list[dict]) -> list[str]:
    return [
        str(row.get("task_id", ""))
        for row in rows
        if str(row.get("status", "")).strip().lower() not in {"metadata_uncertain", "duplicate"}
        and not _is_successful_status(row.get("status", ""))
    ]


def _mark_required_failures(
    state: dict,
    paths: BatchPaths,
    task_ids: set[str],
    status: str,
    reason: str,
) -> None:
    updates = []
    for row in state["rows"]:
        task_id = str(row.get("task_id", ""))
        if task_id not in task_ids or _is_terminal_status(row.get("status", "")):
            continue
        updates.append({**row, "status": status, "file": "", "reason": reason})
    _apply_stage_updates(state, updates, paths)


def _accepts_on_updates(method: Callable[..., object]) -> bool:
    parameters = inspect.signature(method).parameters.values()
    return any(
        parameter.name == "on_updates" or parameter.kind == inspect.Parameter.VAR_KEYWORD
        for parameter in parameters
    )


def _run_gateway_with_state(
    *,
    runner: object,
    method_name: str,
    rows: list[dict],
    paths: BatchPaths,
    options: Any,
    state: dict,
    required_ids: set[str],
) -> None:
    method = getattr(runner, method_name)
    callback_error: Exception | None = None
    callback_applied_ids: set[str] = set()

    def persist_updates(updates: list[dict]) -> None:
        nonlocal callback_error
        try:
            normalized_updates = _validate_stage_updates(state["rows"], updates)
            _apply_stage_updates(state, normalized_updates, paths)
            callback_applied_ids.update(
                update["task_id"] for update in normalized_updates
            )
        except Exception as exc:
            callback_error = exc
            raise

    try:
        if _accepts_on_updates(method):
            returned_updates = method(rows, paths, options, on_updates=persist_updates)
        else:
            returned_updates = method(rows, paths, options)
    except Exception as exc:
        if callback_error is exc or (
            isinstance(exc, ValueError)
            and str(exc).startswith(("gateway_task_id", "invalid_gateway_updates"))
        ):
            raise
        failure = f"gateway_exception_{type(exc).__name__}"
        _mark_required_failures(state, paths, required_ids, failure, failure)
        return

    normalized_updates = _validate_stage_updates(state["rows"], returned_updates)
    _apply_stage_updates(state, normalized_updates, paths)
    returned_ids = {update["task_id"] for update in normalized_updates}
    missing_ids = required_ids - returned_ids - callback_applied_ids
    if missing_ids:
        _mark_required_failures(
            state,
            paths,
            missing_ids,
            "missing_stage_update",
            "missing_stage_update",
        )


def _pending_rows(rows: list[dict], *, manual_retry_used: bool) -> tuple[list[dict], list[dict]]:
    from .failure_routing import is_zotero_eligible

    manual_rows: list[dict] = []
    fallback_rows: list[dict] = []
    for row in rows:
        if _is_terminal_status(row.get("status", "")):
            continue
        if not manual_retry_used and _needs_manual_retry(row.get("status", ""), row.get("reason", "")):
            manual_rows.append(row)
            continue
        # A2/B3.1: only DOI-bearing, bridge-eligible failures enter Zotero fallback.
        if is_zotero_eligible(
            row.get("status", ""),
            row.get("reason", ""),
            doi=row.get("doi") or row.get("input_doi") or "",
        ):
            fallback_rows.append(row)
    return manual_rows, fallback_rows


def _write_pending_files(paths: BatchPaths, rows: list[dict], *, manual_retry_used: bool = False) -> None:
    manual_rows, fallback_rows = _pending_rows(rows, manual_retry_used=manual_retry_used)
    _write_csv_rows(paths.manual_retry, manual_rows)
    _write_csv_rows(paths.zotero_fallback, fallback_rows)


def _serialize_options(options: Any) -> dict[str, object]:
    if not is_dataclass(options):
        raise ValueError("invalid_batch_options")
    return _validate_options_data(asdict(options))


def _validate_options_data(data: object) -> dict[str, object]:
    expected_fields = {
        "email",
        "cookies",
        "browser_exe",
        "login_wait_seconds",
        "debug_port",
        "throttle_seconds",
        "skip_manual_retry",
        "download_supplements",
        "smart_route",
        "session_break_seconds",
        "session_break_every",
        "resolve_title_metadata",
        "circuit_breaker_threshold",
        "auto_oa_recovery",
        "iucr_short_try",
    }
    if not isinstance(data, dict):
        raise ValueError("invalid_batch_options")
    # Backward compatible with older batch_state snapshots (missing optimization fields).
    payload = dict(data)
    if "skip_manual_retry" not in payload:
        payload["skip_manual_retry"] = True
    payload.setdefault("download_supplements", True)
    payload.setdefault("smart_route", True)
    payload.setdefault("session_break_seconds", 60.0)
    payload.setdefault("session_break_every", 8)
    payload.setdefault("resolve_title_metadata", False)
    payload.setdefault("circuit_breaker_threshold", 3)
    payload.setdefault("auto_oa_recovery", True)
    payload.setdefault("iucr_short_try", True)
    if set(payload) != expected_fields:
        # Ignore unknown keys from future versions; require all expected after defaults.
        payload = {key: payload[key] for key in expected_fields if key in payload}
        for key in expected_fields:
            payload.setdefault(key, {
                "email": "",
                "cookies": "",
                "browser_exe": "",
                "login_wait_seconds": 0,
                "debug_port": 9333,
                "throttle_seconds": 1.0,
                "skip_manual_retry": True,
                "download_supplements": True,
                "smart_route": True,
                "session_break_seconds": 60.0,
                "session_break_every": 8,
                "resolve_title_metadata": False,
                "circuit_breaker_threshold": 3,
                "auto_oa_recovery": True,
                "iucr_short_try": True,
            }[key])
        if set(payload) != expected_fields:
            raise ValueError("invalid_batch_options")
    if not isinstance(payload["email"], str) or not isinstance(payload["browser_exe"], str):
        raise ValueError("invalid_batch_options")
    if not isinstance(payload["cookies"], str):
        raise ValueError("cookies_must_be_path")
    for flag in (
        "skip_manual_retry",
        "download_supplements",
        "smart_route",
        "resolve_title_metadata",
        "auto_oa_recovery",
        "iucr_short_try",
    ):
        if type(payload[flag]) is not bool:
            raise ValueError("invalid_batch_options")
    cookie_path = payload["cookies"].strip()
    lowered_cookie = cookie_path.lower()
    if cookie_path and (
        any(character in cookie_path for character in ("\r", "\n", "=", ";"))
        or "cookie:" in lowered_cookie
        or "://" in cookie_path
        or cookie_path.lstrip().startswith(("{", "["))
        or Path(cookie_path).suffix.lower() != ".json"
    ):
        raise ValueError("cookies_must_be_path")

    login_wait_seconds = payload["login_wait_seconds"]
    if type(login_wait_seconds) is not int or login_wait_seconds < 0:
        raise ValueError("invalid_login_wait_seconds")
    debug_port = payload["debug_port"]
    if type(debug_port) is not int or not 1 <= debug_port <= 65535:
        raise ValueError("invalid_debug_port")
    throttle_seconds = payload["throttle_seconds"]
    if (
        isinstance(throttle_seconds, bool)
        or not isinstance(throttle_seconds, (int, float))
        or not math.isfinite(float(throttle_seconds))
        or float(throttle_seconds) < 0
    ):
        raise ValueError("invalid_throttle_seconds")
    session_break_seconds = payload["session_break_seconds"]
    if (
        isinstance(session_break_seconds, bool)
        or not isinstance(session_break_seconds, (int, float))
        or not math.isfinite(float(session_break_seconds))
        or float(session_break_seconds) < 0
    ):
        raise ValueError("invalid_session_break_seconds")
    session_break_every = payload["session_break_every"]
    if type(session_break_every) is not int or session_break_every < 1:
        raise ValueError("invalid_session_break_every")
    circuit_breaker_threshold = payload["circuit_breaker_threshold"]
    if type(circuit_breaker_threshold) is not int or circuit_breaker_threshold < 1:
        raise ValueError("invalid_circuit_breaker_threshold")
    return {
        "email": payload["email"],
        "cookies": cookie_path,
        "browser_exe": payload["browser_exe"],
        "login_wait_seconds": login_wait_seconds,
        "debug_port": debug_port,
        "throttle_seconds": float(throttle_seconds),
        "skip_manual_retry": payload["skip_manual_retry"],
        "download_supplements": payload["download_supplements"],
        "smart_route": payload["smart_route"],
        "session_break_seconds": float(session_break_seconds),
        "session_break_every": session_break_every,
        "resolve_title_metadata": payload["resolve_title_metadata"],
        "circuit_breaker_threshold": circuit_breaker_threshold,
        "auto_oa_recovery": payload["auto_oa_recovery"],
        "iucr_short_try": payload["iucr_short_try"],
    }


def _options_from_state(state: dict):
    saved = state.get("options")
    try:
        validated = _validate_options_data(saved)
    except ValueError as exc:
        raise ValueError("invalid_batch_state_options") from exc
    return _batch_options_type()(**validated)


def _validate_state(
    state: object,
    *,
    expected_run_dir: str | Path | None = None,
) -> dict:
    if not isinstance(state, dict):
        raise ValueError("invalid_batch_state")
    if state.get("version") != 1 or not isinstance(state.get("run_dir"), str):
        raise ValueError("invalid_batch_state")
    if not isinstance(state.get("manual_retry_used"), bool) or not isinstance(state.get("rows"), list):
        raise ValueError("invalid_batch_state")
    _validate_stage_updates([_normalise_row_mapping(row) for row in state["rows"]], [])
    _options_from_state(state)
    if expected_run_dir is not None:
        saved_run_dir = Path(state["run_dir"]).expanduser()
        if not saved_run_dir.is_absolute():
            raise ValueError("invalid_batch_state_run_dir")
        try:
            saved_root = saved_run_dir.resolve()
            expected_root = Path(expected_run_dir).expanduser().resolve()
        except OSError as exc:
            raise ValueError("invalid_batch_state_run_dir") from exc
        if saved_root != expected_root:
            raise ValueError("invalid_batch_state_run_dir")
    return state


def _paths_from_run_dir(run_dir: str | Path) -> BatchPaths:
    root = Path(run_dir).expanduser().resolve()
    working = root / "working"
    return BatchPaths(
        root=root,
        pdfs=root / "pdfs",
        reports=root / "reports",
        working=working,
        state=working / "batch_state.json",
        normalized_input=working / "normalized_input.csv",
        manual_retry=working / "manual_retry.csv",
        zotero_fallback=working / "zotero_fallback.csv",
        zotero_results=working / "zotero_results.csv",
    )


def paths_from_run_dir(run_dir: str | Path) -> BatchPaths:
    """Return the standard paths for an existing batch without changing it."""

    return _paths_from_run_dir(run_dir)


def _read_csv_rows(path: Path) -> list[dict]:
    if not path.is_file():
        raise ValueError("pending_file_missing")
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        return [{key: str(value or "") for key, value in row.items()} for row in csv.DictReader(handle)]


def _read_zotero_results(path: Path, state_rows: list[dict]) -> list[dict[str, str]]:
    """Read and fully validate the handoff before changing any batch state."""

    if not path.is_file():
        raise ValueError("zotero_results_file_missing")
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        try:
            records = list(csv.reader(handle, strict=True))
        except csv.Error as exc:
            raise ValueError("zotero_results_row_invalid") from exc
        if not records or records[0] != ZOTERO_RESULT_FIELDS:
            raise ValueError("zotero_results_fields_invalid")
        raw_rows = records[1:]

    known_ids = {str(row.get("task_id", "") or "") for row in state_rows}
    if not known_ids or "" in known_ids or len(known_ids) != len(state_rows):
        raise ValueError("invalid_state_task_ids")
    results: list[dict[str, str]] = []
    seen_ids: set[str] = set()
    for values in raw_rows:
        if len(values) != len(ZOTERO_RESULT_FIELDS) or not any(
            str(value or "").strip() for value in values
        ):
            raise ValueError("zotero_results_row_invalid")
        result = {
            field: str(value or "").strip()
            for field, value in zip(ZOTERO_RESULT_FIELDS, values)
        }
        result["status"] = result["status"].lower()
        task_id = result["task_id"]
        if not task_id:
            raise ValueError("zotero_result_task_id_missing")
        if task_id in seen_ids:
            raise ValueError("zotero_result_task_id_duplicate")
        if task_id not in known_ids:
            raise ValueError("zotero_result_task_id_unknown")
        if result["status"] not in ZOTERO_INPUT_STATUSES:
            raise ValueError("zotero_result_status_invalid")
        seen_ids.add(task_id)
        results.append(result)
    return results


def _zotero_attachment_absolute_path(path_value: str) -> Path:
    value = str(path_value or "").strip()
    if not value or "://" in value or value.lower().startswith(("zotero:", "file:")):
        raise ValueError("zotero_attachment_not_local")
    source = Path(value).expanduser()
    if not source.is_absolute():
        raise ValueError("zotero_attachment_not_absolute")
    return source


def _stat_is_reparse_point(details: os.stat_result) -> bool:
    return bool(
        os.name == "nt"
        and getattr(details, "st_file_attributes", 0)
        & stat.FILE_ATTRIBUTE_REPARSE_POINT
    )


def _validate_zotero_attachment_chain(source: Path) -> tuple[Path, os.stat_result]:
    """Validate every lexical path component without following a reparse point."""

    current = Path(source.anchor)
    try:
        for component in source.parts[1:]:
            if component in {"", "."}:
                continue
            current = current.parent if component == ".." else current / component
            details = current.lstat()
            if (
                current.is_symlink()
                or stat.S_ISLNK(details.st_mode)
                or _stat_is_reparse_point(details)
            ):
                raise ValueError("zotero_attachment_reparse_point")
        resolved = source.resolve(strict=True)
        final_details = current.lstat()
        if (
            current.is_symlink()
            or stat.S_ISLNK(final_details.st_mode)
            or _stat_is_reparse_point(final_details)
        ):
            raise ValueError("zotero_attachment_reparse_point")
        if not stat.S_ISREG(final_details.st_mode):
            raise ValueError("not_pdf_response")
    except OSError as exc:
        raise ValueError("not_pdf_response") from exc
    return resolved, final_details


def _local_zotero_attachment(path_value: str) -> Path:
    source = _zotero_attachment_absolute_path(path_value)
    resolved, _ = _validate_zotero_attachment_chain(source)
    if not is_valid_pdf(resolved):
        raise ValueError("not_pdf_response")
    return resolved


def _same_file_identity(left: os.stat_result, right: os.stat_result) -> bool:
    return (left.st_dev, left.st_ino) == (right.st_dev, right.st_ino)


def _stable_open_file(before: os.stat_result, after: os.stat_result) -> bool:
    return (
        _same_file_identity(before, after)
        and stat.S_ISREG(before.st_mode)
        and stat.S_ISREG(after.st_mode)
        and not _stat_is_reparse_point(before)
        and not _stat_is_reparse_point(after)
        and before.st_size == after.st_size
        and before.st_mtime_ns == after.st_mtime_ns
    )


def _revalidate_open_zotero_path(
    source: Path,
    expected_resolved: Path,
    opened_details: os.stat_result,
) -> None:
    try:
        resolved, path_details = _validate_zotero_attachment_chain(source)
    except ValueError as exc:
        if str(exc) == "zotero_attachment_reparse_point":
            raise
        raise ValueError("zotero_attachment_changed") from exc
    if (
        resolved != expected_resolved
        or not _same_file_identity(opened_details, path_details)
        or not stat.S_ISREG(opened_details.st_mode)
        or _stat_is_reparse_point(opened_details)
    ):
        raise ValueError("zotero_attachment_changed")


def _snapshot_zotero_attachment(
    path_value: str,
    destination: Path,
    expected_source: Path,
    verified_hash: str,
) -> tuple[Path, str]:
    """Copy from one validated source handle into a private, fsynced snapshot."""

    source = _zotero_attachment_absolute_path(path_value)
    resolved, _ = _validate_zotero_attachment_chain(source)
    if resolved != expected_source:
        raise ValueError("zotero_attachment_changed")

    snapshot = None
    try:
        with resolved.open("rb") as source_handle:
            opened_before = os.fstat(source_handle.fileno())
            _revalidate_open_zotero_path(source, resolved, opened_before)
            snapshot, snapshot_hash = _snapshot_pdf_handle(source_handle, destination)
            opened_after = os.fstat(source_handle.fileno())
            if not _stable_open_file(opened_before, opened_after):
                raise ValueError("zotero_attachment_changed")
            _revalidate_open_zotero_path(source, resolved, opened_after)
        if snapshot_hash != verified_hash:
            raise ValueError("zotero_attachment_changed")
        return snapshot, snapshot_hash
    except BaseException:
        if snapshot is not None:
            snapshot.unlink(missing_ok=True)
        raise


def _filename_for_batch_row(row: dict) -> str:
    from .file_manager import make_pdf_filename
    from .models import MetadataResult

    try:
        source_index = int(str(row.get("source_index", "") or "0"))
    except ValueError:
        source_index = 0
    authors = [
        part.strip()
        for part in str(row.get("authors", "") or "").replace(";", "|").split("|")
        if part.strip()
    ]
    metadata = MetadataResult(
        source_index=source_index,
        query_title=str(row.get("input_title", "") or row.get("title", "") or ""),
        doi=str(row.get("doi", "") or ""),
        title=str(row.get("title", "") or ""),
        authors=authors,
        journal=str(row.get("journal", "") or ""),
        year=str(row.get("year", "") or ""),
        publisher=str(row.get("publisher", "") or ""),
    )
    return make_pdf_filename(metadata)


def _copy_zotero_attachment(
    row: dict,
    attachment_path: str,
    paths: BatchPaths,
    *,
    email: str = "",
) -> Path:
    source = _local_zotero_attachment(attachment_path)
    source_hash = _sha256(source)
    verified_source = _local_zotero_attachment(attachment_path)
    verified_hash = _sha256(verified_source)
    if source != verified_source or source_hash != verified_hash:
        raise ValueError("zotero_attachment_changed")
    from .file_manager import enrich_row_metadata_for_delivery

    enriched = enrich_row_metadata_for_delivery(row, email=email)
    safe_filename = _validate_path_component(
        _filename_for_batch_row(enriched),
        error="invalid_filename",
    )
    destination = paths.pdfs.expanduser().resolve()
    destination.mkdir(parents=True, exist_ok=True)
    with _pdf_publish_lock(destination, timeout=10.0):
        _cleanup_stale_pdf_snapshots(destination)
        snapshot = None
        try:
            snapshot, snapshot_hash = _snapshot_zotero_attachment(
                attachment_path,
                destination,
                verified_source,
                verified_hash,
            )
            return _publish_verified_pdf_snapshot(
                snapshot,
                destination,
                safe_filename,
                snapshot_hash,
                reuse_any_hash=True,
            )
        finally:
            if snapshot is not None:
                snapshot.unlink(missing_ok=True)


def _metadata_uncertain_audit(reason: object) -> str:
    detail = str(reason or "").strip()
    return "metadata_uncertain" if not detail else f"metadata_uncertain; {detail}"


def _validate_manual_retry_preclaim(state: dict, paths: BatchPaths) -> list[dict]:
    if not paths.manual_retry.is_file():
        raise ValueError("manual_retry_file_missing")
    with paths.manual_retry.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        fieldnames = reader.fieldnames or []
        if not set(NORMALIZED_FIELDS).issubset(fieldnames):
            raise ValueError("manual_retry_fields_invalid")
        rows = [{key: str(value or "") for key, value in row.items()} for row in reader]

    state_by_id = {str(row.get("task_id", "")): row for row in state["rows"]}
    state_manual_ids = {
        task_id
        for task_id, row in state_by_id.items()
        if _needs_manual_retry(row.get("status", ""), row.get("reason", ""))
    }
    seen_ids: set[str] = set()
    for row in rows:
        task_id = row.get("task_id", "")
        if not task_id or task_id in seen_ids:
            raise ValueError("manual_retry_task_id_invalid")
        seen_ids.add(task_id)
        state_row = state_by_id.get(task_id)
        if state_row is None:
            raise ValueError("manual_retry_unknown_task_id")
        if not _needs_manual_retry(row.get("status", ""), row.get("reason", "")):
            raise ValueError("manual_retry_status_invalid")
        if not _needs_manual_retry(state_row.get("status", ""), state_row.get("reason", "")):
            raise ValueError("manual_retry_state_not_pending")
        if str(row.get("status", "")).strip().lower() != str(
            state_row.get("status", "")
        ).strip().lower():
            raise ValueError("manual_retry_status_mismatch")
    if seen_ids != state_manual_ids:
        raise ValueError("manual_retry_task_ids_mismatch")
    return rows


def _status_label_zh(status: object) -> str:
    value = str(status or "").strip().lower()
    if value in SUCCESS_STATUSES:
        return "成功"
    if value == "duplicate":
        return "重复"
    return "失败"


def _source_label_zh(source: object) -> str:
    value = str(source or "").strip().lower()
    labels = {
        "oa": "开放获取",
        "sciencedirect": "ScienceDirect",
        "non_elsevier": "机构（非Elsevier）",
        "institutional": "机构",
        "zotero": "Zotero",
        "manual_import": "外部补入",
        "oa_direct": "开放获取",
    }
    return labels.get(value, str(source or "").strip() or "")


def _safe_delivery_name(name: str, *, fallback: str) -> str:
    cleaned = str(name or "").strip()
    if not cleaned:
        cleaned = fallback
    cleaned = cleaned.replace("\\", "_").replace("/", "_")
    cleaned = "".join(
        "_" if (ch in _WINDOWS_INVALID_CHARS or ord(ch) < 32) else ch
        for ch in cleaned
    ).strip(" .")
    if not cleaned or cleaned in {".", ".."}:
        cleaned = fallback
    stem = cleaned.split(".", 1)[0].upper()
    if stem in _WINDOWS_RESERVED_NAMES:
        cleaned = f"_{cleaned}"
    return cleaned[:180] or fallback


def _safe_input_delivery_name(name: str) -> str:
    """Keep the source input name while avoiding reserved delivery names."""

    cleaned = _safe_delivery_name(name, fallback="输入清单.txt")
    reserved = {
        USER_INVENTORY_NAME.casefold(),
        USER_PDF_DIR_NAME.casefold(),
        USER_MD_DIR_NAME.casefold(),
        USER_SUPPLEMENT_DIR_NAME.casefold(),
    }
    if cleaned.casefold() in reserved:
        cleaned = f"输入清单_{cleaned}"
    return cleaned


def _iter_regular_files(root: Path) -> list[Path]:
    if not root.is_dir():
        return []
    return sorted(
        (
            path
            for path in root.rglob("*")
            if path.is_file() and not path.is_symlink()
        ),
        key=lambda path: str(path).casefold(),
    )


def _copy_input_to_delivery(
    paths: BatchPaths,
    delivery_target: Path,
    staging_results: Path,
    rows: list[dict],
) -> Path:
    """Copy exactly one input list into the user-facing result directory.

    New runs use the immutable input snapshot.  Older runs are supported by
    falling back to an existing non-PDF file in ``结果`` and finally the
    normalized internal CSV.
    """

    snapshot_files = _iter_regular_files(paths.working / _INPUT_SNAPSHOT_DIR_NAME)
    source: Path | None = snapshot_files[0] if snapshot_files else None
    if source is None:
        legacy_candidates = [
            path
            for path in (
                delivery_target.iterdir() if delivery_target.is_dir() else []
            )
            if path.is_file()
            and not path.is_symlink()
            and path.name.casefold() != USER_INVENTORY_NAME.casefold()
            and path.suffix.lower() != ".pdf"
        ]
        source = sorted(legacy_candidates, key=lambda path: path.name.casefold())[0] if legacy_candidates else None
    if source is None and paths.normalized_input.is_file():
        source = paths.normalized_input
    if source is None:
        # Synthetic/legacy states may not have retained their original input.
        # Publish a reproducible normalized input rather than silently omitting
        # the promised input-list slot.
        destination = staging_results / "输入清单.csv"
        _write_csv_rows(destination, rows)
        return destination

    preferred = (
        "输入清单.csv"
        if source == paths.normalized_input
        else _safe_input_delivery_name(source.name)
    )
    destination = staging_results / preferred
    shutil.copy2(source, destination)
    return destination


def _delivery_pdf_candidates(delivery_target: Path) -> list[Path]:
    """Find current and legacy user-facing article PDFs without supplements."""

    candidates: list[Path] = []
    for root in (delivery_target / USER_PDF_DIR_NAME, delivery_target):
        if not root.is_dir():
            continue
        for path in root.iterdir():
            if (
                path.is_file()
                and not path.is_symlink()
                and path.suffix.lower() == ".pdf"
                and path not in candidates
            ):
                candidates.append(path)
    return sorted(candidates, key=lambda path: path.name.casefold())


def _copy_existing_user_tree(source: Path, destination: Path) -> bool:
    """Copy regular files from a user-owned tree and report whether any exist."""

    files = _iter_regular_files(source)
    if not files:
        return False
    for source_file in files:
        relative = source_file.relative_to(source)
        target = destination / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source_file, target)
    return True


def _collect_supplement_dirs(row: dict, paths: BatchPaths) -> list[Path]:
    """Locate supplement folders written by stage downloaders for this paper."""

    candidates: list[Path] = []
    file_value = str(row.get("file", "") or "").strip()
    stems: list[str] = []
    if file_value:
        path = Path(file_value)
        stems.append(path.stem)
        # Stage copies may keep paper-0001 while delivery uses year_author names.
    task_id = str(row.get("task_id", "") or "").strip()
    if task_id:
        stems.append(task_id)
    doi = str(row.get("doi", "") or row.get("input_doi", "") or "").strip().lower()
    if doi:
        stems.append(doi.replace("/", "_"))
        stems.append(doi.rsplit("/", 1)[-1])

    search_roots = [
        paths.reports / "sciencedirect" / "supplements",
        paths.reports / "sciencedirect" / "pdfs" / "supplements",
        paths.reports / "oa" / "supplements",
        paths.reports / "non_elsevier_institutional" / "supplements",
        paths.pdfs / "supplements",
        paths.root / "supplements",
        paths.root / USER_DELIVERY_DIR_NAME / USER_SUPPLEMENT_DIR_NAME,
    ]
    seen: set[str] = set()
    for root in search_roots:
        if not root.is_dir():
            continue
        for stem in stems:
            if not stem:
                continue
            folder = root / stem
            key = str(folder).lower()
            if key in seen:
                continue
            if folder.is_dir() and any(folder.iterdir()):
                candidates.append(folder)
                seen.add(key)
    return candidates


def _copy_tree_files(source_dir: Path, destination_dir: Path) -> list[str]:
    destination_dir.mkdir(parents=True, exist_ok=True)
    relative_names: list[str] = []
    for source in sorted(source_dir.rglob("*")):
        if not source.is_file() or source.is_symlink():
            continue
        rel = source.relative_to(source_dir)
        target = destination_dir / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
        relative_names.append(str(rel).replace("\\", "/"))
    return relative_names


def _replace_user_directory(target: Path, staging: Path) -> None:
    backup: Path | None = None
    if target.exists() or target.is_symlink():
        backup = target.with_name(f".{target.name}.bak_{time.time_ns()}")
        if backup.exists() or backup.is_symlink():
            if backup.is_dir() and not backup.is_symlink():
                shutil.rmtree(backup, ignore_errors=True)
            else:
                backup.unlink(missing_ok=True)
        target.replace(backup)
    try:
        staging.replace(target)
    except Exception:
        if backup is not None and backup.exists() and not target.exists():
            backup.replace(target)
        raise
    if backup is not None:
        if backup.is_dir() and not backup.is_symlink():
            shutil.rmtree(backup, ignore_errors=True)
        else:
            backup.unlink(missing_ok=True)


def _file_sha256(path: Path, *, chunk: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            block = handle.read(chunk)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest()


def publish_user_delivery(paths: BatchPaths, rows: list[dict]) -> Path:
    """Publish the five-item user package under ``结果``.

    The package contains the original input snapshot, ``下载清单.csv``,
    ``pdf`` and ``md`` directories, plus ``补充材料`` only when files exist.
    Internal ``pdfs``/``reports``/``working`` remain outside the package.

    **A1 merge-safe:** existing PDFs, supplementary files, and Markdown files
    in the user package are preserved by content copy when the batch is
    republished. Legacy PDFs directly under ``结果`` are migrated into
    ``结果/pdf``.
    """

    paths.root.mkdir(parents=True, exist_ok=True)
    # Stage beside the batch instead of inside it.  The published layout is
    # unchanged, but this saves the run-directory component from every nested
    # supplement path and avoids Windows MAX_PATH failures before the final
    # same-volume atomic directory replace.
    staging_root = Path(
        tempfile.mkdtemp(prefix=".delivery_", dir=paths.root.parent)
    )
    staging_results = staging_root / USER_DELIVERY_DIR_NAME
    staging_results.mkdir(parents=True, exist_ok=True)
    staging_pdf_dir = staging_results / USER_PDF_DIR_NAME
    staging_pdf_dir.mkdir(parents=True, exist_ok=True)
    staging_md_dir = staging_results / USER_MD_DIR_NAME
    staging_md_dir.mkdir(parents=True, exist_ok=True)
    inventory_rows: list[dict[str, str]] = []
    used_names: set[str] = set()
    staged_hashes: set[str] = set()

    delivery_target = user_delivery_dir(paths)
    prior_orphans = _delivery_pdf_candidates(delivery_target)
    legacy_supplement_dirs = [
        path
        for path in (delivery_target.iterdir() if delivery_target.is_dir() else [])
        if path.is_dir()
        and not path.is_symlink()
        and path.name.casefold().endswith("_supplements")
    ]
    supplement_staged = _copy_existing_user_tree(
        delivery_target / USER_SUPPLEMENT_DIR_NAME,
        staging_results / USER_SUPPLEMENT_DIR_NAME,
    )
    try:
        _copy_existing_user_tree(delivery_target / USER_MD_DIR_NAME, staging_md_dir)
        _copy_input_to_delivery(paths, delivery_target, staging_results, rows)

        for index, row in enumerate(rows, start=1):
            status = str(row.get("status", "") or "").strip()
            label = _status_label_zh(status)
            reason = ""
            if label == "失败":
                reason = str(row.get("reason", "") or "").strip() or status or "unknown_failure"
            elif label == "重复":
                reason = str(row.get("reason", "") or "").strip() or "duplicate"

            result_rel = ""
            supplements_rel = ""
            if label == "成功":
                source_file = str(row.get("file", "") or "").strip()
                source_path = Path(source_file).expanduser() if source_file else None
                if source_path is not None and not source_path.is_absolute():
                    source_path = paths.root / source_path
                if source_path is not None and source_path.is_file() and not source_path.is_symlink():
                    preferred = _safe_delivery_name(
                        source_path.name if source_path.suffix.lower() == ".pdf" else f"{source_path.stem}.pdf",
                        fallback=f"{row.get('task_id') or f'paper-{index:04d}'}.pdf",
                    )
                    if not preferred.lower().endswith(".pdf"):
                        preferred = f"{preferred}.pdf"
                    base = preferred
                    counter = 2
                    while preferred.lower() in used_names:
                        preferred = f"{Path(base).stem}_{counter}.pdf"
                        counter += 1
                    used_names.add(preferred.lower())
                    target_pdf = staging_pdf_dir / preferred
                    shutil.copy2(source_path, target_pdf)
                    try:
                        staged_hashes.add(_file_sha256(target_pdf))
                    except OSError:
                        pass
                    result_rel = f"{USER_DELIVERY_DIR_NAME}/{USER_PDF_DIR_NAME}/{preferred}"

                    supplement_dirs = _collect_supplement_dirs(row, paths)
                    preferred_supplement = (
                        delivery_target
                        / USER_SUPPLEMENT_DIR_NAME
                        / Path(preferred).stem
                    )
                    legacy_supplement = delivery_target / f"{Path(preferred).stem}_supplements"
                    supplement_dirs.extend(
                        path
                        for path in (preferred_supplement, legacy_supplement)
                        if path.is_dir() and path not in supplement_dirs
                    )
                    if supplement_dirs:
                        suppl_dest = staging_results / USER_SUPPLEMENT_DIR_NAME / Path(preferred).stem
                        copied_names: list[str] = []
                        for folder in supplement_dirs:
                            copied_names.extend(_copy_tree_files(folder, suppl_dest))
                        if copied_names or _iter_regular_files(suppl_dest):
                            supplement_staged = True
                            supplements_rel = (
                                f"{USER_DELIVERY_DIR_NAME}/{USER_SUPPLEMENT_DIR_NAME}/"
                                f"{Path(preferred).stem}"
                            )
                else:
                    label = "失败"
                    reason = str(row.get("reason", "") or "").strip() or "missing_result_pdf"

            inventory_rows.append(
                {
                    "序号": str(index),
                    "状态": label,
                    "DOI": str(row.get("doi", "") or row.get("input_doi", "") or ""),
                    "题名": str(row.get("title", "") or row.get("input_title", "") or ""),
                    "作者": str(row.get("authors", "") or ""),
                    "年份": str(row.get("year", "") or ""),
                    "期刊": str(row.get("journal", "") or ""),
                    "下载来源": _source_label_zh(row.get("source", "")),
                    "结果文件": result_rel,
                    "补充材料": supplements_rel,
                    "失败原因": reason,
                    "task_id": str(row.get("task_id", "") or ""),
                }
            )

        # Preserve legacy per-paper supplement directories while migrating them
        # into one global user-facing directory.
        for legacy in legacy_supplement_dirs:
            stem = legacy.name[: -len("_supplements")]
            if _copy_existing_user_tree(
                legacy,
                staging_results / USER_SUPPLEMENT_DIR_NAME / stem,
            ):
                supplement_staged = True

        # A1: re-attach orphan PDFs not already staged by content hash.
        orphan_index = 0
        for orphan in prior_orphans:
            try:
                digest = _file_sha256(orphan)
            except OSError:
                continue
            if digest in staged_hashes:
                continue
            preferred = _safe_delivery_name(
                orphan.name,
                fallback=f"external_{orphan_index + 1:04d}.pdf",
            )
            if not preferred.lower().endswith(".pdf"):
                preferred = f"{preferred}.pdf"
            base = preferred
            counter = 2
            while preferred.lower() in used_names:
                preferred = f"{Path(base).stem}_{counter}.pdf"
                counter += 1
            used_names.add(preferred.lower())
            shutil.copy2(orphan, staging_pdf_dir / preferred)
            staged_hashes.add(digest)
            orphan_index += 1
            inventory_rows.append(
                {
                    "序号": str(len(inventory_rows) + 1),
                    "状态": "外部补入",
                    "DOI": "",
                    "题名": Path(preferred).stem,
                    "作者": "",
                    "年份": "",
                    "期刊": "",
                    "下载来源": "外部补入",
                    "结果文件": f"{USER_DELIVERY_DIR_NAME}/{USER_PDF_DIR_NAME}/{preferred}",
                    "补充材料": "",
                    "失败原因": "",
                    "task_id": "",
                }
            )

        if not supplement_staged:
            shutil.rmtree(staging_results / USER_SUPPLEMENT_DIR_NAME, ignore_errors=True)

        inventory_staging = staging_results / USER_INVENTORY_NAME
        _write_report_csv(inventory_staging, USER_INVENTORY_FIELDS, inventory_rows)

        _replace_user_directory(delivery_target, staging_results)
        return user_inventory_path(paths)
    finally:
        shutil.rmtree(staging_root, ignore_errors=True)


def write_final_reports(paths: BatchPaths, rows: list[dict]) -> None:
    """Write stable final reports from the current state, with or without Zotero data."""

    manifest_rows = [dict(row) for row in rows]
    failed_rows = [
        row for row in manifest_rows
        if not _is_successful_status(row.get("status", ""))
        and str(row.get("status", "")).strip().lower() != "duplicate"
    ]
    duplicate_count = sum(
        1 for row in manifest_rows
        if str(row.get("status", "")).strip().lower() == "duplicate"
    )
    success_count = sum(
        1 for row in manifest_rows if _is_successful_status(row.get("status", ""))
    )
    failure_count = len(failed_rows)
    failure_status_counts: dict[str, int] = {}
    for row in failed_rows:
        status = str(row.get("status", "") or "").strip() or "missing_status"
        failure_status_counts[status] = failure_status_counts.get(status, 0) + 1
    summary_lines = [
        f"input_count: {len(manifest_rows)}",
        f"success_count: {success_count}",
        f"failure_count: {failure_count}",
        f"user_inventory: {user_inventory_path(paths)}",
        f"user_results_directory: {user_delivery_dir(paths)}",
        f"final_pdf_directory: {user_pdf_dir(paths)}",
        f"duplicate_terminal_rows_excluded: {duplicate_count}",
        "failure_status_counts:",
    ]
    summary_lines.extend(
        f"{status}: {count}" for status, count in sorted(failure_status_counts.items())
    )
    summary_lines.append("failed_tasks:")
    summary_lines.extend(
        "{task_id}\t{status}\t{reason}".format(
            task_id=str(row.get("task_id", "") or ""),
            status=str(row.get("status", "") or ""),
            reason=str(row.get("reason", "") or ""),
        )
        for row in failed_rows
    )
    summary = {
        "total_count": len(manifest_rows),
        "success_count": success_count,
        "failed_count": failure_count,
        "user_inventory": str(user_inventory_path(paths)),
        "user_results_directory": str(user_delivery_dir(paths)),
    }
    try:
        paths.reports.mkdir(parents=True, exist_ok=True)
        paths.working.mkdir(parents=True, exist_ok=True)
        with _file_lock(
            paths.working / "final_reports.lock",
            timeout=10.0,
            invalid_timeout_error="invalid_final_report_lock_timeout",
            timeout_error="final_report_lock_timeout",
        ):
            staging = Path(
                tempfile.mkdtemp(prefix=".final_reports_staging_", dir=paths.reports)
            )
            try:
                _write_report_csv(staging / "final_manifest.csv", FINAL_MANIFEST_FIELDS, manifest_rows)
                _write_final_manifest_xlsx(staging / "final_manifest.xlsx", manifest_rows)
                _write_report_csv(staging / "failed.csv", FINAL_MANIFEST_FIELDS, failed_rows)
                _write_report_text(staging / "run_summary.txt", "\n".join(summary_lines) + "\n")
                # Keep Task 4 report names available for start/resume callers.
                _write_report_csv(staging / "batch_status.csv", NORMALIZED_FIELDS, manifest_rows)
                _write_report_json(staging / "batch_status.json", summary)
                _publish_report_set(paths, staging)
            finally:
                _cleanup_report_transaction_dir(staging)
        # User-facing package: one input, one inventory, and predictable folders.
        publish_user_delivery(paths, manifest_rows)
    except Exception as exc:
        raise RuntimeError(f"final_report_write_failed:{type(exc).__name__}:{exc}") from exc


def _write_latest_state_outputs(
    paths: BatchPaths,
    state: dict,
    *,
    pending_manual_retry_used: bool | None,
    assume_locked: bool = False,
) -> None:
    """Publish pending/final reports from the latest disk state under its writer lock."""

    if not assume_locked:
        with batch_state_lock(paths.root):
            _write_latest_state_outputs(
                paths,
                state,
                pending_manual_retry_used=pending_manual_retry_used,
                assume_locked=True,
            )
        return

    latest = load_batch_state(paths.root)
    _validate_state(latest, expected_run_dir=paths.root)
    if pending_manual_retry_used is not None:
        _write_pending_files(
            paths,
            latest["rows"],
            manual_retry_used=pending_manual_retry_used,
        )
    write_final_reports(paths, latest["rows"])
    state.clear()
    state.update(latest)


def finalize_batch(
    run_dir: str | Path,
    zotero_results: str | Path,
) -> BatchRunResult:
    """Reconcile one fully validated Zotero export without changing Zotero attachments."""

    paths = _paths_from_run_dir(run_dir)
    state = load_batch_state(paths.root)
    _validate_state(state, expected_run_dir=paths.root)
    results_path = Path(zotero_results)
    _read_zotero_results(results_path, state["rows"])
    with batch_state_lock(paths.root):
        state = load_batch_state(paths.root)
        _validate_state(state, expected_run_dir=paths.root)
        results = _read_zotero_results(results_path, state["rows"])
        results_by_id = {result["task_id"]: result for result in results}
        for index, existing in enumerate(state["rows"]):
            task_id = str(existing.get("task_id", "") or "")
            result = results_by_id.get(task_id)
            if result is None or _is_terminal_status(existing.get("status", "")):
                continue
            updated = dict(existing)
            zotero_status = result["status"]
            if zotero_status in ZOTERO_SUCCESS:
                if not result["zotero_item_id"]:
                    updated.update(
                        status="zotero_item_id_missing",
                        source="zotero",
                        file="",
                        reason="zotero_item_id_missing",
                        zotero_item_id="",
                    )
                else:
                    try:
                        email = str((state.get("options") or {}).get("email", "") or "")
                        target = _copy_zotero_attachment(
                            updated,
                            result["attachment_path"],
                            paths,
                            email=email,
                        )
                    except (OSError, ValueError) as exc:
                        updated.update(
                            status="not_pdf_response",
                            source="zotero",
                            file="",
                            reason=f"not_pdf_response:{exc}",
                            zotero_item_id=result["zotero_item_id"],
                        )
                    else:
                        previous_status = str(existing.get("status", "") or "").strip().lower()
                        updated.update(
                            status=ZOTERO_SUCCESS[zotero_status],
                            source="zotero",
                            file=str(target),
                            reason=(
                                _metadata_uncertain_audit(existing.get("reason", ""))
                                if previous_status == "metadata_uncertain" else ""
                            ),
                            zotero_item_id=result["zotero_item_id"],
                        )
            else:
                updated.update(
                    status=zotero_status,
                    source="zotero",
                    file="",
                    reason=result["reason"],
                    zotero_item_id=result["zotero_item_id"],
                )
            state["rows"][index] = updated
            save_batch_state(paths, state)
        _write_latest_state_outputs(
            paths,
            state,
            pending_manual_retry_used=bool(state.get("manual_retry_used")),
            assume_locked=True,
        )
    return _result_from_state(paths, state)


def _result_from_state(paths: BatchPaths, state: dict) -> BatchRunResult:
    rows = state["rows"]
    manual_rows, fallback_rows = _pending_rows(
        rows,
        manual_retry_used=bool(state.get("manual_retry_used")),
    )
    return BatchRunResult(
        paths=paths,
        total_count=len(rows),
        success_count=sum(1 for row in rows if _is_successful_status(row.get("status", ""))),
        failed_count=sum(1 for row in rows if not _is_terminal_status(row.get("status", ""))),
        manual_retry_count=len(manual_rows),
        zotero_fallback_count=len(fallback_rows),
    )


def result_from_state(paths: BatchPaths, state: dict) -> BatchRunResult:
    """Validate and summarize an existing batch state without changing it."""

    if not isinstance(paths, BatchPaths) or paths != _paths_from_run_dir(paths.root):
        raise ValueError("invalid_batch_paths")
    _validate_state(state, expected_run_dir=paths.root)
    return _result_from_state(paths, state)


class DefaultStageGateway:
    def run_initial(
        self,
        rows: list[dict],
        paths: BatchPaths,
        options: Any,
        *,
        on_updates: Callable[[list[dict]], None] | None = None,
    ) -> list[dict]:
        runnable = [dict(row) for row in rows if str(row.get("status", "")).lower() == "pending"]
        if not runnable:
            return []

        from .batch_stages import is_elsevier_doi, is_gold_oa_doi, split_route_rows

        # Download order: Elsevier first, then everything else.
        gold_oa, elsevier_rows, other_rows = split_route_rows(runnable)
        smart_route = bool(getattr(options, "smart_route", True))
        print(
            f"[download_order] 1) Elsevier={len(elsevier_rows)}  "
            f"2) other={len(gold_oa) + len(other_rows)} "
            f"(gold_oa={len(gold_oa)}, non_elsevier={len(other_rows)}; "
            f"smart_route={'on' if smart_route else 'off'})",
            flush=True,
        )

        updates: list[dict] = []
        working = [dict(row) for row in rows]
        runnable_ids = {str(row.get("task_id", "")) for row in runnable}

        def _pending_from(pool: list[dict]) -> list[dict]:
            pool_ids = {str(row.get("task_id", "")) for row in pool}
            return [
                row
                for row in working
                if str(row.get("task_id", "")) in pool_ids
                and str(row.get("task_id", "")) in runnable_ids
                and not _is_terminal_status(row.get("status", ""))
            ]

        def _apply_local(stage_updates: list[dict]) -> None:
            nonlocal working, updates
            if not stage_updates:
                return
            if on_updates is not None:
                on_updates(stage_updates)
            updates.extend(stage_updates)
            email = str(getattr(options, "email", "") or "")
            working = _merge_stage_rows(working, stage_updates, paths, email=email)

        # ---- Phase 1: Elsevier / ScienceDirect first ----
        elsevier_pending = _pending_from(elsevier_rows)
        if elsevier_pending:
            print(f"[download_order] phase1 Elsevier start n={len(elsevier_pending)}", flush=True)
            stage_input = _stage_input_writer(
                elsevier_pending, paths.working / "sciencedirect_input.csv"
            )
            sd_updates = [
                _update_as_mapping(row)
                for row in run_sciencedirect_stage(stage_input, paths.reports, options)
            ]
            _apply_local(sd_updates)

        # ---- Phase 2: non-Elsevier (OA for gold / all remaining, then adapters) ----
        non_elsevier_pool = gold_oa + other_rows
        non_elsevier_pending = _pending_from(non_elsevier_pool)
        if non_elsevier_pending:
            if smart_route:
                oa_candidates = [
                    row
                    for row in non_elsevier_pending
                    if is_gold_oa_doi(str(row.get("doi", "") or ""))
                ]
            else:
                # Legacy: try OA for all remaining non-Elsevier rows.
                oa_candidates = list(non_elsevier_pending)

            if oa_candidates:
                print(f"[download_order] phase2 OA start n={len(oa_candidates)}", flush=True)
                oa_updates = [
                    _update_as_mapping(row)
                    for row in run_oa_stage(oa_candidates, paths.reports, options)
                ]
                _apply_local(oa_updates)

            still_pending = _pending_from(non_elsevier_pool)
            # Never send Elsevier leftovers here; only non-Elsevier.
            still_pending = [
                row
                for row in still_pending
                if not is_elsevier_doi(str(row.get("doi", "") or ""))
            ]
            if still_pending:
                print(
                    f"[download_order] phase2 non-Elsevier institutional start n={len(still_pending)}",
                    flush=True,
                )
                stage_input = _stage_input_writer(
                    still_pending, paths.working / "non_elsevier_input.csv"
                )
                other_updates = [
                    _update_as_mapping(row)
                    for row in run_non_elsevier_stage(stage_input, paths.reports, options)
                ]
                _apply_local(other_updates)

        return updates

    def run_retry(
        self,
        rows: list[dict],
        paths: BatchPaths,
        options: Any,
        *,
        on_updates: Callable[[list[dict]], None] | None = None,
    ) -> list[dict]:
        manual_rows = [
            row for row in rows
            if _needs_manual_retry(row.get("status", ""), row.get("reason", ""))
        ]
        return self._run_institutional_stages(
            manual_rows,
            paths,
            options,
            on_updates=on_updates,
            sd_name="sciencedirect_retry_input.csv",
            other_name="non_elsevier_retry_input.csv",
        )

    def run_institutional_only(
        self,
        rows: list[dict],
        paths: BatchPaths,
        options: Any,
        *,
        on_updates: Callable[[list[dict]], None] | None = None,
    ) -> list[dict]:
        """Institutional stages only (skip OA). Used by retry_failed_batch."""

        runnable = [
            dict(row)
            for row in rows
            if str(row.get("status", "")).lower() == "pending"
        ]
        if not runnable:
            return []
        return self._run_institutional_stages(
            runnable,
            paths,
            options,
            on_updates=on_updates,
            sd_name="sciencedirect_retry_failed_input.csv",
            other_name="non_elsevier_retry_failed_input.csv",
        )

    def _run_institutional_stages(
        self,
        rows: list[dict],
        paths: BatchPaths,
        options: Any,
        *,
        on_updates: Callable[[list[dict]], None] | None,
        sd_name: str,
        other_name: str,
    ) -> list[dict]:
        science_direct, other = _split_institutional_rows(rows)
        updates: list[dict] = []
        if science_direct:
            stage_input = _stage_input_writer(science_direct, paths.working / sd_name)
            stage_updates = [
                _update_as_mapping(row)
                for row in run_sciencedirect_stage(stage_input, paths.reports, options)
            ]
            if on_updates is not None:
                on_updates(stage_updates)
            updates.extend(stage_updates)
        if other:
            stage_input = _stage_input_writer(other, paths.working / other_name)
            stage_updates = [
                _update_as_mapping(row)
                for row in run_non_elsevier_stage(stage_input, paths.reports, options)
            ]
            if on_updates is not None:
                on_updates(stage_updates)
            updates.extend(stage_updates)
        return updates


def run_post_download_ladder(
    paths: BatchPaths,
    state: dict,
    *,
    options: Any | None = None,
    pending_manual_retry_used: bool | None = None,
) -> BatchRunResult:
    """A2: after download stages — limited OA recovery → zotero_fallback → delivery.

    Called from ``start_batch`` and ``retry_failed_batch`` so fixed-run continues
    get the same failure ladder as a fresh start.
    """
    selected_options = options or _options_from_state(state)
    if bool(getattr(selected_options, "auto_oa_recovery", True)):
        try:
            from .oa_recovery import run_limited_oa_recovery_on_batch

            recovered = run_limited_oa_recovery_on_batch(
                paths.root,
                email=str(getattr(selected_options, "email", "") or ""),
                require_oa_signal=True,
            )
            if recovered:
                tried = sum(1 for item in recovered if item.status != "skipped")
                ok = sum(1 for item in recovered if item.status == "oa_downloaded")
                print(
                    f"[OA补救/直下] 尝试 {tried} 条（unsupported/not_pdf 必试，其余需 OA 信号）；"
                    f"成功 {ok}",
                    flush=True,
                )
                with batch_state_lock(paths.root):
                    state.clear()
                    state.update(load_batch_state(paths.root))
        except Exception as exc:
            print(f"[OA补救] 跳过（{type(exc).__name__}: {exc})", flush=True)

    manual_flag = (
        bool(state.get("manual_retry_used"))
        if pending_manual_retry_used is None
        else pending_manual_retry_used
    )
    _write_latest_state_outputs(
        paths,
        state,
        pending_manual_retry_used=manual_flag,
    )
    result = _result_from_state(paths, state)
    print(
        f"[交付] 请查看: {paths.root}\n"
        f"  - {user_inventory_path(paths)}\n"
        f"  - {user_delivery_dir(paths)}\\",
        flush=True,
    )
    if result.zotero_fallback_count > 0:
        print(
            f"[阶梯] 仍有 {result.zotero_fallback_count} 条可走 Zotero "
            f"(paper_batch.py zotero --run-dir \"{paths.root}\")",
            flush=True,
        )
    return result


def start_batch(
    *,
    input_text: str | None,
    input_path: str | Path | None,
    output_root: str | Path,
    run_name: str | None = None,
    options: Any | None = None,
    gateway=None,
    normalizer=None,
    now: datetime | None = None,
    doi_preflight: bool | None = None,
    fixed_run: bool = True,
    fresh: bool = False,
) -> BatchRunResult:
    """Start or continue a batch.

    By default uses a **fixed** folder under ``output_root`` named by ``run_name``
    (or the input file stem), reusing it on later starts so intermediate
    timestamped folders are not created. Pass ``fresh=True`` or
    ``fixed_run=False`` for the legacy unique ``name_timestamp`` folders.
    """
    if (input_text is None) == (input_path is None):
        raise ValueError("exactly_one_input_required")
    selected_options = options or _batch_options_type()()
    serialized_options = _serialize_options(selected_options)
    resolved_name = run_name or default_run_name_from_input(input_path, input_text)

    # Fixed delivery folder (default): one stable path per job.
    use_fixed = bool(fixed_run) and not bool(fresh)
    if use_fixed:
        paths = create_batch_paths(output_root, run_name=resolved_name, now=now, fixed=True)
        print(f"[fixed_run] 批次目录（固定）: {paths.root}", flush=True)
        if paths.state.is_file():
            # Reuse existing batch: continue unresolved rows only (keep successes).
            print("[fixed_run] 发现已有 batch_state，续跑未完成项（成功项保留）…", flush=True)
            return retry_failed_batch(
                paths.root,
                gateway=gateway,
                options=selected_options,
                skip_oa=False,
            )
    else:
        paths = create_batch_paths(
            output_root,
            run_name=resolved_name,
            now=now,
            fixed=False,
        )
        print(f"[run] 新建时间戳批次目录: {paths.root}", flush=True)

    normalize = normalizer or normalize_input
    rows = [_normalise_row_mapping(row) for row in normalize(
        input_text=input_text,
        input_path=input_path,
        paths=paths,
        options=selected_options,
    )]
    # A4: DOI preflight ON by default (fail-open on network errors).
    if doi_preflight is None:
        doi_preflight = True
    if doi_preflight:
        rows = _apply_doi_preflight(rows, email=str(getattr(selected_options, "email", "") or ""), paths=paths)
    _validate_stage_updates(rows, [])
    _write_csv_rows(paths.normalized_input, rows)
    skip_manual = bool(serialized_options.get("skip_manual_retry", True))
    state = {
        "version": 1,
        "run_dir": str(paths.root),
        # skip_manual_retry: treat as already used so all failures go to zotero_fallback.
        "manual_retry_used": skip_manual,
        "options": serialized_options,
        "rows": rows,
    }
    with batch_state_lock(paths.root):
        save_batch_state(paths, state)
    runner = gateway or DefaultStageGateway()
    _run_gateway_with_state(
        runner=runner,
        method_name="run_initial",
        rows=rows,
        paths=paths,
        options=selected_options,
        state=state,
        required_ids=set(_required_task_ids(rows)),
    )
    # A2: institutional → OA recovery → zotero_fallback + delivery.
    return run_post_download_ladder(
        paths,
        state,
        options=selected_options,
        pending_manual_retry_used=bool(state.get("manual_retry_used")),
    )


def retry_failed_batch(
    run_dir: str | Path,
    *,
    gateway=None,
    options: Any | None = None,
    skip_oa: bool = True,
    retry_all: bool = False,
) -> BatchRunResult:
    """Re-run unresolved rows in an existing batch (same run-dir; optimization #5).

    By default skips OA and only re-runs institutional stages for network-class
    failures (A5). Use retry_all=True to restore the previous broader set.
    """

    from .failure_routing import is_retry_eligible

    paths = _paths_from_run_dir(run_dir)
    with batch_state_lock(paths.root):
        state = load_batch_state(paths.root)
        _validate_state(state, expected_run_dir=paths.root)
        selected_options = options or _options_from_state(state)
        if options is not None:
            state["options"] = _serialize_options(selected_options)
        failed_rows = [
            row
            for row in state["rows"]
            if not _is_terminal_status(row.get("status", ""))
            and is_retry_eligible(
                row.get("status", ""),
                row.get("reason", ""),
                retry_all=retry_all,
            )
        ]
        if not failed_rows:
            return _result_from_state(paths, state)
        # Reset unresolved rows to pending so stages will process them.
        for row in failed_rows:
            row["status"] = "pending"
            row["source"] = "retry_failed"
            row["file"] = ""
            if str(row.get("reason", "") or "").startswith("doi_"):
                pass
            else:
                row["reason"] = "retry_failed_reset"
        # Opt1: do not re-run DOI preflight on retry by default (avoids hangs).
        save_batch_state(paths, state)

    runner = gateway or DefaultStageGateway()
    method = "run_institutional_only" if skip_oa else "run_initial"
    retry_ids = {
        str(row.get("task_id", ""))
        for row in state["rows"]
        if str(row.get("status", "")).lower() == "pending"
    }
    if not retry_ids:
        _write_latest_state_outputs(
            paths,
            state,
            pending_manual_retry_used=bool(state.get("manual_retry_used")),
        )
        return _result_from_state(paths, state)
    pending_rows = [row for row in state["rows"] if str(row.get("task_id", "")) in retry_ids]
    _run_gateway_with_state(
        runner=runner,
        method_name=method,
        rows=pending_rows if method == "run_institutional_only" else state["rows"],
        paths=paths,
        options=selected_options,
        state=state,
        required_ids=retry_ids,
    )
    # A2: even institutional-only retries still run limited OA (unsupported/not_pdf)
    # then rebuild zotero_fallback + delivery package.
    return run_post_download_ladder(
        paths,
        state,
        options=selected_options,
        pending_manual_retry_used=bool(state.get("manual_retry_used")),
    )


def _apply_doi_preflight(rows: list[dict], *, email: str, paths: BatchPaths) -> list[dict]:
    try:
        from .doi_preflight import preflight_rows
    except Exception:
        return rows
    try:
        updated, changes = preflight_rows(rows, email=email)
    except Exception as exc:
        print(f"[DOI预检] 跳过（{type(exc).__name__}）")
        return rows
    if changes:
        print(f"[DOI预检] 校正/补全 {len(changes)} 条：")
        for change in changes[:20]:
            print(
                f"  - {change.task_id}: {change.action} "
                f"{change.old_doi or '(无)'} -> {change.new_doi or '(无)'} ({change.detail})"
            )
        if len(changes) > 20:
            print(f"  ... 另有 {len(changes) - 20} 条")
        try:
            log_path = paths.working / "doi_preflight_changes.csv"
            import csv as _csv
            with log_path.open("w", newline="", encoding="utf-8-sig") as handle:
                writer = _csv.DictWriter(
                    handle,
                    fieldnames=["task_id", "action", "old_doi", "new_doi", "detail"],
                )
                writer.writeheader()
                for change in changes:
                    writer.writerow(
                        {
                            "task_id": change.task_id,
                            "action": change.action,
                            "old_doi": change.old_doi,
                            "new_doi": change.new_doi,
                            "detail": change.detail,
                        }
                    )
        except OSError:
            pass
    else:
        print("[DOI预检] 无需校正")
    return updated


def resume_batch(
    run_dir: str | Path,
    *,
    gateway=None,
    options: Any | None = None,
) -> BatchRunResult:
    paths = _paths_from_run_dir(run_dir)
    retry_csv_rows: list[dict] = []

    def validate_state(state: dict) -> None:
        _validate_state(state, expected_run_dir=paths.root)

    def validate_before_claim(state: dict) -> bool:
        validated_rows = _validate_manual_retry_preclaim(state, paths)
        if options is not None:
            _serialize_options(options)
        if not validated_rows:
            return False
        retry_csv_rows.extend(validated_rows)
        return True

    claimed, state = claim_manual_retry(
        run_dir,
        validate_state=validate_state,
        validate_before_claim=validate_before_claim,
    )
    if not claimed:
        return _result_from_state(paths, state)
    selected_options = options or _options_from_state(state)
    retry_ids = {row.get("task_id", "") for row in retry_csv_rows}
    retry_rows = [row for row in state["rows"] if str(row.get("task_id", "")) in retry_ids]
    if retry_rows:
        runner = gateway or DefaultStageGateway()
        _run_gateway_with_state(
            runner=runner,
            method_name="run_retry",
            rows=retry_rows,
            paths=paths,
            options=selected_options,
            state=state,
            required_ids=set(retry_ids),
        )
    _write_latest_state_outputs(
        paths,
        state,
        pending_manual_retry_used=True,
    )
    return _result_from_state(paths, state)
