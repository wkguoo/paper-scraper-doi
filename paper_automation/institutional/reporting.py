from __future__ import annotations

import csv
import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from doi_batch_utils import TEXT_ENCODINGS, clean_doi

from .models import InstitutionalReportRow


REPORT_FIELDS = [
    "row_number",
    "doi",
    "title",
    "journal",
    "year",
    "publisher",
    "adapter",
    "status",
    "file",
    "reason",
    "landing_url",
    "final_landing_url",
    "pdf_url",
    "metadata_source",
]
MERGE_FIELDS = [
    "institutional_adapter",
    "institutional_status",
    "institutional_pdf_file",
    "institutional_reason",
    "institutional_landing_url",
    "institutional_final_landing_url",
    "institutional_pdf_url",
]


def ensure_output_dirs(output_dir: str | Path) -> dict[str, Path]:
    root = Path(output_dir) / "non_elsevier_institutional"
    pdfs = root / "pdfs"
    root.mkdir(parents=True, exist_ok=True)
    pdfs.mkdir(parents=True, exist_ok=True)
    return {"root": root, "pdfs": pdfs}


def write_report(rows: list[InstitutionalReportRow], output_dir: str | Path) -> Path:
    path = ensure_output_dirs(output_dir)["root"] / "institutional_pdf_download_report.csv"
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=REPORT_FIELDS)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "row_number": row.row_number,
                    "doi": row.doi,
                    "title": row.title,
                    "journal": row.journal,
                    "year": row.year,
                    "publisher": row.publisher,
                    "adapter": row.adapter,
                    "status": row.status,
                    "file": row.file,
                    "reason": row.reason,
                    "landing_url": row.landing_url,
                    "final_landing_url": row.final_landing_url,
                    "pdf_url": row.pdf_url,
                    "metadata_source": row.metadata_source,
                }
            )
    return path


def write_run_summary(
    rows: list[InstitutionalReportRow],
    output_dir: str | Path,
    report_path: str | Path,
    manifest_update_path: str = "",
) -> Path:
    directories = ensure_output_dirs(output_dir)
    status_counts = Counter(row.status for row in rows)
    payload = {
        "output_dir": str(directories["root"]),
        "pdf_dir": str(directories["pdfs"]),
        "total_input": len(rows),
        "resolved_count": sum(1 for row in rows if row.doi),
        "downloaded_count": status_counts.get("pdf_downloaded", 0),
        "failed_count": len(rows) - status_counts.get("pdf_downloaded", 0),
        "status_counts": dict(status_counts),
        "report_path": str(report_path),
        "manifest_update_path": manifest_update_path,
        "written_at": datetime.now().isoformat(timespec="seconds"),
    }
    path = directories["root"] / "run_summary.json"
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def merge_manifest(
    manifest_path: str | Path,
    rows: list[InstitutionalReportRow],
    output_dir: str | Path,
) -> Path:
    existing_rows = read_tabular_rows(manifest_path)
    updates = {clean_doi(row.doi).lower(): row for row in rows if clean_doi(row.doi)}
    fieldnames = list(existing_rows[0].keys()) if existing_rows else ["doi"]
    for field in MERGE_FIELDS:
        if field not in fieldnames:
            fieldnames.append(field)
    merged_rows: list[dict[str, str]] = []
    seen: set[str] = set()
    for row in existing_rows:
        doi_key = clean_doi(row.get("doi", "")).lower()
        update = updates.get(doi_key)
        merged = dict(row)
        if update:
            merged.update(
                {
                    "institutional_adapter": update.adapter,
                    "institutional_status": update.status,
                    "institutional_pdf_file": update.file,
                    "institutional_reason": update.reason,
                    "institutional_landing_url": update.landing_url,
                    "institutional_final_landing_url": update.final_landing_url,
                    "institutional_pdf_url": update.pdf_url,
                }
            )
            seen.add(doi_key)
        merged_rows.append(merged)
    for doi_key, update in updates.items():
        if doi_key in seen:
            continue
        appended = {field: "" for field in fieldnames}
        appended["doi"] = update.doi
        appended["title"] = update.title
        appended["institutional_adapter"] = update.adapter
        appended["institutional_status"] = update.status
        appended["institutional_pdf_file"] = update.file
        appended["institutional_reason"] = update.reason
        appended["institutional_landing_url"] = update.landing_url
        appended["institutional_final_landing_url"] = update.final_landing_url
        appended["institutional_pdf_url"] = update.pdf_url
        merged_rows.append(appended)
    path = ensure_output_dirs(output_dir)["root"] / "final_manifest_institutional_update.csv"
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(merged_rows)
    return path


def read_tabular_rows(path: str | Path) -> list[dict[str, str]]:
    target = Path(path)
    suffix = target.suffix.lower()
    if suffix == ".csv":
        return _read_csv_rows(target)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx_rows(target)
    raise ValueError("仅支持 CSV/XLSX/XLSM manifest 合并")


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    last_error: UnicodeDecodeError | None = None
    for encoding in TEXT_ENCODINGS:
        try:
            with path.open("r", newline="", encoding=encoding) as handle:
                return [{key: str(value or "") for key, value in row.items()} for row in csv.DictReader(handle)]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"无法解码 CSV 文件: {path} ({last_error})")


def _read_xlsx_rows(path: Path) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("缺少 openpyxl，无法读取 XLSX manifest") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "").strip() for value in rows[0]]
    result: list[dict[str, str]] = []
    for values in rows[1:]:
        result.append({headers[index]: str(value or "") for index, value in enumerate(values)})
    return result
