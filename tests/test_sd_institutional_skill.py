from __future__ import annotations

import csv
import json
import sys
import tempfile
import types
import unittest
from pathlib import Path
from urllib.parse import unquote_plus
from unittest.mock import MagicMock, patch


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


class InstitutionalSkillIntakeTests(unittest.TestCase):
    def test_main_writes_empty_reports_when_no_valid_doi(self) -> None:
        from sd_institutional_skill import main

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_path = root / "invalid.csv"
            with input_path.open("w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=["title", "doi"])
                writer.writeheader()
                writer.writerow({"title": "", "doi": "not-a-doi"})

            exit_code = main([
                "--input",
                str(input_path),
                "--out",
                str(root),
                "--run-name",
                "no_valid",
                "--dry-run",
            ])

            run_dir = root / "no_valid"
            failed_path = run_dir / "doi_batch_failed.csv"
            pdf_report_path = run_dir / "pdf_download_report.csv"
            summary_path = run_dir / "run_summary.txt"

            self.assertNotEqual(exit_code, 0)
            self.assertTrue(failed_path.exists())
            self.assertTrue(pdf_report_path.exists())
            self.assertTrue(summary_path.exists())
            with failed_path.open("r", encoding="utf-8-sig") as f:
                failed_rows = list(csv.DictReader(f))
            self.assertEqual(len(failed_rows), 1)
            self.assertIn("invalid", failed_rows[0]["reason"])
            with pdf_report_path.open("r", encoding="utf-8-sig") as f:
                self.assertEqual(list(csv.DictReader(f)), [])

            summary_text = summary_path.read_text(encoding="utf-8")
            self.assertIn(str(run_dir / "merged_doi_input.csv"), summary_text)
            self.assertIn(str(failed_path), summary_text)
            self.assertIn(str(pdf_report_path), summary_text)
            self.assertIn("识别 DOI 数: 0", summary_text)
            self.assertIn("成功解析数: 0", summary_text)
            self.assertIn("解析失败数: 1", summary_text)
            self.assertIn("- 成功: 0", summary_text)
            self.assertIn("- 失败: 0", summary_text)
            self.assertIn("- 跳过: 0", summary_text)

    def test_beginner_preflight_writes_review_hints_without_sciencedirect_resolution(self) -> None:
        from sd_institutional_skill import main
        from sd_scraper import ScienceDirectScraper
        from paper_automation.models import MetadataResult

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with patch.object(
                ScienceDirectScraper,
                "resolve_doi_batch",
                side_effect=AssertionError("preflight must not resolve ScienceDirect DOI"),
            ), patch(
                "sd_institutional_skill.MetadataResolver.resolve_one",
                return_value=MetadataResult(
                    source_index=1,
                    query_title="",
                    doi="10.1016/j.actamat.2024.119999",
                    title="Resolved DOI paper",
                    journal="Acta Materialia",
                    year="2024",
                    confidence=1.0,
                    source="crossref",
                    match_basis="input_doi",
                ),
            ):
                exit_code = main([
                    "--text",
                    "DOI: 10.1016/j.actamat.2024.119999\n"
                    "P8: Ti-Mo beta-Ti stress-induced martensitic transformation 是否包含 SXRD",
                    "--out",
                    str(root),
                    "--run-name",
                    "beginner_preflight",
                    "--beginner",
                    "--preflight",
                ])

            run_dir = root / "beginner_preflight"
            preview_text = (run_dir / "doi_intake_preview.csv").read_text(encoding="utf-8-sig")
            summary_text = (run_dir / "run_summary.txt").read_text(encoding="utf-8")
            pdf_report_text = (run_dir / "pdf_download_report.csv").read_text(encoding="utf-8-sig")

        self.assertEqual(exit_code, 0)
        self.assertIn("review_hint", preview_text)
        self.assertIn("请补 DOI", preview_text)
        self.assertIn("小白下一步建议", summary_text)
        self.assertIn("preflight", pdf_report_text)

    def test_main_downloads_supplements_by_default_and_can_disable_them(self) -> None:
        from doi_batch_utils import DownloadRunResult, PdfDownloadRecord, SupplementDownloadRecord
        from paper_automation.models import MetadataResult
        from sd_institutional_skill import main

        calls: list[bool] = []

        class FakeScraper:
            def resolve_doi_batch(self, _input_path: str):
                return [
                    {
                        "doi": "10.1016/j.actamat.2024.119999",
                        "pii": "S1359645424000012",
                        "title": "Paper with supplements",
                        "authors": "Zhang Wei",
                        "year": "2024",
                    }
                ], []

            def save_to_xlsx(self, _results: list[dict], filename: str, output_dir: str) -> str:
                path = Path(output_dir) / filename
                path.write_text("placeholder", encoding="utf-8")
                return str(path)

            def save_failed_doi_report(self, _failures: list[dict], filename: str, output_dir: str) -> str:
                path = Path(output_dir) / filename
                with path.open("w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.DictWriter(f, fieldnames=["row_number", "doi", "reason"])
                    writer.writeheader()
                return str(path)

            def download_pdfs_devtools(self, results: list[dict], output_dir: str, **kwargs: object) -> DownloadRunResult:
                download_supplements = bool(kwargs["download_supplements"])
                calls.append(download_supplements)
                article = results[0]
                supplement_records = []
                if download_supplements:
                    supplement_records.append(
                        SupplementDownloadRecord(
                            doi=article["doi"],
                            pii=article["pii"],
                            article_title=article["title"],
                            article_file="2024_Zhang_Paper with supplements_12345678.pdf",
                            supplement_index=0,
                            status="not_found",
                            reason="未发现补充材料",
                        )
                    )
                return DownloadRunResult(
                    pdf_success=1,
                    pdf_failed=0,
                    pdf_skipped=0,
                    pdf_records=[
                        PdfDownloadRecord(
                            doi=article["doi"],
                            pii=article["pii"],
                            title=article["title"],
                            status="success",
                            file="2024_Zhang_Paper with supplements_12345678.pdf",
                        )
                    ],
                    supplement_success=0,
                    supplement_failed=0,
                    supplement_skipped=0,
                    supplement_not_found=1 if download_supplements else 0,
                    supplement_records=supplement_records,
                )

        def fake_make_scraper(_cookie_cache_path: Path, browser_exe: str | None = None) -> FakeScraper:
            self.assertIsNone(browser_exe)
            return FakeScraper()

        def fake_resolve_one(_self: object, _candidate: object) -> MetadataResult:
            return MetadataResult(
                source_index=1,
                query_title="",
                doi="10.1016/j.actamat.2024.119999",
                title="Paper with supplements",
                confidence=1.0,
                source="crossref",
            )

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_path = root / "papers.csv"
            input_path.write_text("doi\n10.1016/j.actamat.2024.119999\n", encoding="utf-8")
            with patch("sd_institutional_skill.make_scraper", side_effect=fake_make_scraper), patch(
                "sd_institutional_skill.cache_devtools_cookies",
                return_value=0,
            ), patch("sd_institutional_skill.MetadataResolver.resolve_one", fake_resolve_one):
                default_exit = main([
                    "--input",
                    str(input_path),
                    "--out",
                    str(root),
                    "--run-name",
                    "default_supplements",
                ])
                disabled_exit = main([
                    "--input",
                    str(input_path),
                    "--out",
                    str(root),
                    "--run-name",
                    "disabled_supplements",
                    "--no-download-supplements",
                ])

            default_summary = json.loads((root / "default_supplements" / "run_summary.json").read_text(encoding="utf-8"))
            disabled_summary = json.loads((root / "disabled_supplements" / "run_summary.json").read_text(encoding="utf-8"))
            default_report_exists = (root / "default_supplements" / "supplement_download_report.csv").exists()

        self.assertEqual(default_exit, 0)
        self.assertEqual(disabled_exit, 0)
        self.assertEqual(calls, [True, False])
        self.assertTrue(default_report_exists)
        self.assertTrue(default_summary["supplement_requested"])
        self.assertEqual(default_summary["supplement_not_found"], 1)
        self.assertFalse(disabled_summary["supplement_requested"])
        self.assertEqual(disabled_summary["supplement_report_path"], "")
        self.assertFalse((root / "disabled_supplements" / "supplement_download_report.csv").exists())

    def test_main_marks_supplements_not_requested_for_non_pdf_runs(self) -> None:
        from paper_automation.models import MetadataResult
        from sd_institutional_skill import main

        class FakeScraper:
            def resolve_doi_batch(self, _input_path: str):
                return [
                    {
                        "doi": "10.1016/j.actamat.2024.119999",
                        "pii": "S1359645424000012",
                        "title": "Dry run paper",
                        "authors": "Zhang Wei",
                        "year": "2024",
                    }
                ], []

            def save_to_xlsx(self, _results: list[dict], filename: str, output_dir: str) -> str:
                path = Path(output_dir) / filename
                path.write_text("placeholder", encoding="utf-8")
                return str(path)

            def save_failed_doi_report(self, _failures: list[dict], filename: str, output_dir: str) -> str:
                path = Path(output_dir) / filename
                with path.open("w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.DictWriter(f, fieldnames=["row_number", "doi", "reason"])
                    writer.writeheader()
                return str(path)

            def download_pdfs_devtools(self, *_args: object, **_kwargs: object) -> object:
                raise AssertionError("PDF and supplement downloads must not run")

        def fake_resolve_one(_self: object, _candidate: object) -> MetadataResult:
            return MetadataResult(
                source_index=1,
                query_title="",
                doi="10.1016/j.actamat.2024.119999",
                title="Dry run paper",
                confidence=1.0,
                source="crossref",
            )

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_path = root / "papers.csv"
            input_path.write_text("doi\n10.1016/j.actamat.2024.119999\n", encoding="utf-8")
            scenarios = [
                ("dry_run", ["--dry-run"]),
                ("no_pdf", ["--no-download-pdfs"]),
                ("preflight", ["--preflight"]),
            ]
            with patch("sd_institutional_skill.ScienceDirectScraper", side_effect=lambda *a, **k: FakeScraper()), patch(
                "sd_institutional_skill.MetadataResolver.resolve_one",
                fake_resolve_one,
            ):
                exit_codes = {
                    name: main([
                        "--input",
                        str(input_path),
                        "--out",
                        str(root),
                        "--run-name",
                        name,
                        *extra_args,
                    ])
                    for name, extra_args in scenarios
                }
            summaries = {
                name: json.loads((root / name / "run_summary.json").read_text(encoding="utf-8"))
                for name, _extra_args in scenarios
            }
            report_paths = {
                name: root / name / "supplement_download_report.csv"
                for name, _extra_args in scenarios
            }

        self.assertEqual(exit_codes, {"dry_run": 0, "no_pdf": 0, "preflight": 0})
        for name, summary in summaries.items():
            with self.subTest(name=name):
                self.assertFalse(summary["supplement_requested"])
                self.assertEqual(summary["supplement_report_path"], "")
                self.assertFalse(report_paths[name].exists())

    def test_pdf_failure_records_supplement_skipped_without_downloader(self) -> None:
        from sd_scraper import ScienceDirectScraper

        class FakeResponse:
            def __init__(self, payload: dict[str, object]) -> None:
                self.payload = payload

            def __enter__(self) -> "FakeResponse":
                return self

            def __exit__(self, *_args: object) -> None:
                return None

            def read(self) -> bytes:
                return json.dumps(self.payload).encode("utf-8")

        class FakeWebSocket:
            def __init__(self) -> None:
                self.last_message: dict[str, object] = {}

            def settimeout(self, _timeout: int) -> None:
                return None

            def send(self, payload: str) -> None:
                self.last_message = json.loads(payload)

            def recv(self) -> str:
                method = self.last_message.get("method")
                message_id = self.last_message.get("id", 1)
                if method == "Runtime.evaluate":
                    params = self.last_message.get("params") or {}
                    expression = params.get("expression") if isinstance(params, dict) else ""
                    value = (
                        "https://pdf.sciencedirectassets.com/main.pdf"
                        if expression == "location.href"
                        else "<html><body>article</body></html>"
                    )
                    return json.dumps({"id": message_id, "result": {"result": {"value": value}}})
                if method == "Network.getCookies":
                    return json.dumps({"id": message_id, "result": {"cookies": []}})
                return json.dumps({"method": "Page.loadEventFired"})

            def close(self) -> None:
                return None

        tab_counter = {"value": 0}

        def fake_urlopen(request: object, timeout: int = 0) -> FakeResponse:
            url = getattr(request, "full_url", str(request))
            if "/json/new" in url:
                tab_counter["value"] += 1
                tab_id = f"tab-{tab_counter['value']}"
                return FakeResponse({"id": tab_id, "webSocketDebuggerUrl": f"ws://{tab_id}"})
            if "/json/close/" in url:
                return FakeResponse({})
            raise AssertionError(f"Unexpected URL: {url}")

        article = {
            "doi": "10.1016/j.actamat.2024.119999",
            "pii": "S1359645424000012",
            "title": "Failed PDF paper",
            "authors": "Zhang Wei",
            "year": "2024",
        }
        expected_filename = ScienceDirectScraper._make_pdf_filename(1, article)
        fake_websocket_module = types.SimpleNamespace(create_connection=lambda *_args, **_kwargs: FakeWebSocket())
        scraper = ScienceDirectScraper()
        scraper._is_chrome_debug_ready = lambda: True
        scraper._launch_chrome_with_debug = lambda: None
        supplement_downloader = MagicMock(return_value=[])

        with tempfile.TemporaryDirectory() as tmp, patch.dict(sys.modules, {"websocket": fake_websocket_module}), patch(
            "urllib.request.urlopen",
            side_effect=fake_urlopen,
        ), patch("sd_scraper._dt_capture_pdf", return_value=(None, "synthetic PDF failure")), patch(
            "sd_scraper.download_supplements_for_article",
            supplement_downloader,
        ), patch("sd_scraper.time.sleep", return_value=None), patch("sd_scraper.random.uniform", return_value=0):
            result = scraper.download_pdfs_devtools(
                [article],
                tmp,
                interactive_login=False,
                login_wait_seconds=0,
                download_supplements=True,
            )

        supplement_downloader.assert_not_called()
        self.assertEqual(result.pdf_failed, 1)
        self.assertEqual(result.supplement_skipped, 1)
        self.assertEqual(len(result.supplement_records), 1)
        record = result.supplement_records[0]
        self.assertEqual(record.status, "skipped")
        self.assertEqual(record.supplement_index, 0)
        self.assertEqual(record.article_file, expected_filename)
        self.assertIn("PDF download failed", record.reason)

    def test_english_cli_writes_supplement_report_from_download_result(self) -> None:
        from doi_batch_utils import DownloadRunResult, PdfDownloadRecord, SupplementDownloadRecord
        import sd_scraper_en

        calls: list[bool] = []

        class FakeScraper:
            def __init__(self, *_args: object, **_kwargs: object) -> None:
                return None

            def search_by_keyword(
                self,
                _query: str,
                _count: int,
                _sort: str,
                _date: str,
                _article_type: str,
            ) -> list[dict]:
                return [
                    {
                        "doi": "10.1016/j.actamat.2024.119999",
                        "pii": "S1359645424000012",
                        "title": "English CLI paper",
                        "authors": "Zhang Wei",
                        "year": "2024",
                    }
                ]

            def save_to_xlsx(self, _results: list[dict], filename: str, output_dir: str) -> str:
                path = Path(output_dir) / filename
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text("placeholder", encoding="utf-8")
                return str(path)

            def download_pdfs_devtools(self, results: list[dict], output_dir: str, **kwargs: object) -> DownloadRunResult:
                calls.append(bool(kwargs["download_supplements"]))
                article = results[0]
                return DownloadRunResult(
                    pdf_success=1,
                    pdf_failed=0,
                    pdf_skipped=0,
                    pdf_records=[
                        PdfDownloadRecord(
                            doi=article["doi"],
                            pii=article["pii"],
                            title=article["title"],
                            status="success",
                            file="2024_Zhang_English CLI paper_12345678.pdf",
                        )
                    ],
                    supplement_success=1,
                    supplement_records=[
                        SupplementDownloadRecord(
                            doi=article["doi"],
                            pii=article["pii"],
                            article_title=article["title"],
                            article_file="2024_Zhang_English CLI paper_12345678.pdf",
                            supplement_index=1,
                            supplement_title="Supplementary data",
                            status="success",
                            file="supplements/item.zip",
                        )
                    ],
                )

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            argv = [
                "sd_scraper_en.py",
                "-m",
                "keyword",
                "-q",
                "alloy",
                "--download-pdfs",
                "--output",
                str(root),
                "--filename",
                "en_case",
            ]
            with patch.object(sd_scraper_en, "ScienceDirectScraper", FakeScraper), patch.object(sys, "argv", argv):
                sd_scraper_en.main()

            report_path = root / "supplement_download_report.csv"
            self.assertTrue(report_path.exists())
            with report_path.open("r", encoding="utf-8-sig") as f:
                rows = list(csv.DictReader(f))

        self.assertEqual(calls, [True])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["status"], "success")

    def test_english_devtools_uses_legacy_existing_pdf_filename_for_supplements(self) -> None:
        from doi_batch_utils import SupplementDownloadRecord
        import sd_scraper_en

        class FakeResponse:
            def __init__(self, payload: dict[str, object]) -> None:
                self.payload = payload

            def __enter__(self) -> "FakeResponse":
                return self

            def __exit__(self, *_args: object) -> None:
                return None

            def read(self) -> bytes:
                return json.dumps(self.payload).encode("utf-8")

        class FakeWebSocket:
            def __init__(self) -> None:
                self.last_message: dict[str, object] = {}

            def settimeout(self, _timeout: int) -> None:
                return None

            def send(self, payload: str) -> None:
                self.last_message = json.loads(payload)

            def recv(self) -> str:
                method = self.last_message.get("method")
                message_id = self.last_message.get("id", 1)
                if method == "Runtime.evaluate":
                    params = self.last_message.get("params") or {}
                    expression = params.get("expression") if isinstance(params, dict) else ""
                    value = (
                        "https://pdf.sciencedirectassets.com/main.pdf"
                        if expression == "location.href"
                        else "<html><body>supplement</body></html>"
                    )
                    return json.dumps({"id": message_id, "result": {"result": {"value": value}}})
                if method == "Network.getCookies":
                    return json.dumps({"id": message_id, "result": {"cookies": []}})
                return json.dumps({"method": "Page.loadEventFired"})

            def close(self) -> None:
                return None

        tab_counter = {"value": 0}

        def fake_urlopen(request: object, timeout: int = 0) -> FakeResponse:
            url = getattr(request, "full_url", str(request))
            if "/json/new" in url:
                tab_counter["value"] += 1
                tab_id = f"tab-{tab_counter['value']}"
                return FakeResponse({"id": tab_id, "webSocketDebuggerUrl": f"ws://{tab_id}"})
            if "/json/close/" in url:
                return FakeResponse({})
            raise AssertionError(f"Unexpected URL: {url}")

        article = {
            "doi": "10.1016/j.actamat.2024.119999",
            "pii": "S1359645424000012",
            "title": "Legacy PDF paper",
            "authors": "Zhang Wei",
            "year": "2024",
        }
        scraper = sd_scraper_en.ScienceDirectScraper()
        scraper._is_chrome_debug_ready = lambda: True
        scraper._launch_chrome_with_debug = lambda: None
        legacy_filename = scraper._make_legacy_english_pdf_filename(1, article)
        seen_article_files: list[str] = []

        def fake_download_supplements(
            *,
            article: dict,
            article_index: int,
            article_file: str,
            article_html: str,
            article_url: str,
            output_dir: str,
            session: object,
            headers: dict[str, str],
        ) -> list[SupplementDownloadRecord]:
            seen_article_files.append(article_file)
            return [
                SupplementDownloadRecord(
                    doi=article["doi"],
                    pii=article["pii"],
                    article_title=article["title"],
                    article_file=article_file,
                    supplement_index=0,
                    status="not_found",
                )
            ]

        fake_websocket_module = types.SimpleNamespace(create_connection=lambda *_args, **_kwargs: FakeWebSocket())
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pdf_dir = root / "pdfs"
            pdf_dir.mkdir()
            (pdf_dir / legacy_filename).write_bytes(b"%PDF legacy")
            with patch.dict(sys.modules, {"websocket": fake_websocket_module}), patch(
                "urllib.request.urlopen",
                side_effect=fake_urlopen,
            ), patch("sd_scraper_en.download_supplements_for_article", side_effect=fake_download_supplements), patch(
                "sd_scraper_en._dt_capture_pdf",
                side_effect=AssertionError("existing PDF skip must not fetch PDF"),
            ), patch("sd_scraper_en.time.sleep", return_value=None), patch(
                "sd_scraper_en.random.uniform",
                return_value=0,
            ):
                result = scraper.download_pdfs_devtools([article], str(root), download_supplements=True)

        self.assertEqual(result.pdf_skipped, 1)
        self.assertEqual(result.pdf_records[0].file, legacy_filename)
        self.assertEqual(seen_article_files, [legacy_filename])
        self.assertEqual(result.supplement_records[0].article_file, legacy_filename)

    def test_main_passes_browser_exe_to_dry_run_scraper(self) -> None:
        import sd_institutional_skill

        constructed: list[dict[str, object]] = []

        class FakeScraper:
            def __init__(self, **kwargs: object) -> None:
                constructed.append(kwargs)
                self.last_browser_message = "Edge: fake"
                self.last_download_next_steps = ""

            def resolve_doi_batch(self, _input_path: str) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
                return [], []

            def save_failed_doi_report(self, _failures: list[dict[str, str]], filename: str, output_dir: str) -> str:
                path = Path(output_dir) / filename
                path.write_text("row_number,doi,reason\n", encoding="utf-8-sig")
                return str(path)

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with patch.object(sd_institutional_skill, "ScienceDirectScraper", FakeScraper):
                exit_code = sd_institutional_skill.main([
                    "--text",
                    "10.1016/j.actamat.2026.121972",
                    "--out",
                    str(root),
                    "--run-name",
                    "dry",
                    "--dry-run",
                    "--browser-exe",
                    r"C:\Program Files (x86)\Microsoft\Edge Beta\Application\msedge.exe",
                ])

        self.assertEqual(exit_code, 0)
        self.assertEqual(
            constructed,
            [{"browser_exe": r"C:\Program Files (x86)\Microsoft\Edge Beta\Application\msedge.exe"}],
        )
    def test_build_intake_merges_text_files_and_ignores_output_dirs(self) -> None:
        from sd_institutional_skill import build_intake, iter_input_files

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_dir = root / "sources"
            source_dir.mkdir()
            csv_path = source_dir / "papers.csv"
            with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
                writer = csv.DictWriter(f, fieldnames=["title", "doi"])
                writer.writeheader()
                writer.writerow({"title": "A", "doi": "10.1016/j.actamat.2024.119999"})
                writer.writerow({"title": "Duplicate", "doi": "https://doi.org/10.1016/j.actamat.2024.119999"})
                writer.writerow({"title": "Invalid", "doi": "not-a-doi"})
                writer.writerow({"title": "Empty", "doi": ""})

            ignored = source_dir / "results"
            ignored.mkdir()
            (ignored / "ignored.txt").write_text("10.1016/j.ignored.2024.1", encoding="utf-8")
            (source_dir / "cookies.json").write_text("[]", encoding="utf-8")

            files = [path.name for path in iter_input_files(source_dir)]
            self.assertEqual(files, ["papers.csv"])

            output_dir = root / "out"
            result = build_intake(
                texts=["extra DOI: 10.1016/j.scriptamat.2024.115000"],
                input_paths=[],
                folder_paths=[source_dir],
                output_dir=output_dir,
            )

            self.assertEqual(result.status_counts["valid"], 2)
            self.assertEqual(result.status_counts["duplicate"], 1)
            self.assertEqual(result.status_counts["invalid"], 1)
            self.assertEqual(result.status_counts["needs_review"], 1)
            self.assertEqual(result.valid_count, 2)

            preview_text = result.preview_path.read_text(encoding="utf-8-sig")
            merged_text = result.merged_input_path.read_text(encoding="utf-8-sig")
            self.assertIn("重复 DOI", preview_text)
            self.assertIn("未识别到 DOI", preview_text)
            self.assertIn("not_probable_title", preview_text)
            self.assertIn("10.1016/j.scriptamat.2024.115000", merged_text)
            self.assertNotIn("10.1016/j.ignored.2024.1", merged_text)

    def test_build_intake_resolves_title_only_text_to_doi_metadata(self) -> None:
        from sd_institutional_skill import build_intake

        def fake_json(url: str, headers: dict[str, str] | None = None, timeout: int = 20) -> dict:
            if "api.crossref.org/works?" in url:
                return {
                    "message": {
                        "items": [
                            {
                                "DOI": "10.1016/j.actamat.2024.119999",
                                "title": ["Additive manufacturing of gamma-TiAl alloys"],
                                "author": [{"family": "Zhang", "given": "Wei"}],
                                "container-title": ["Acta Materialia"],
                                "published-print": {"date-parts": [[2024]]},
                                "publisher": "Elsevier",
                                "URL": "https://doi.org/10.1016/j.actamat.2024.119999",
                            }
                        ]
                    }
                }
            return {}

        with tempfile.TemporaryDirectory() as tmp:
            result = build_intake(
                texts=[
                    "Additive manufacturing of gamma-TiAl alloys. Acta Materialia, 2024\n"
                    "DOI: 10.1016/j.actamat.2024.119999"
                ],
                input_paths=[],
                folder_paths=[],
                output_dir=Path(tmp),
                resolve_metadata=True,
                http_json=fake_json,
            )
            merged_text = result.merged_input_path.read_text(encoding="utf-8-sig")
            preview_text = result.preview_path.read_text(encoding="utf-8-sig")

        self.assertEqual(result.valid_count, 1)
        self.assertEqual(result.status_counts["duplicate"], 1)
        self.assertIn("10.1016/j.actamat.2024.119999", merged_text)
        self.assertIn("Zhang Wei", preview_text)
        self.assertIn("Acta Materialia", preview_text)
        self.assertIn("重复 DOI", preview_text)

    def test_build_intake_duplicate_after_metadata_resolution_rewrites_review_hint(self) -> None:
        from paper_automation.models import MetadataResult
        from sd_institutional_skill import SourceEntry, classify_entries

        entries = [
            SourceEntry(
                1,
                "text",
                1,
                "Austenite precipitation thermodynamics in steels. Acta Materialia, 2024",
                title="Austenite precipitation thermodynamics in steels",
            ),
            SourceEntry(
                2,
                "text",
                2,
                "Hydrogen trapping behavior in titanium alloys. Acta Materialia, 2024",
                title="Hydrogen trapping behavior in titanium alloys",
            ),
        ]
        resolved = [
            MetadataResult(
                source_index=1,
                query_title="Austenite precipitation thermodynamics in steels",
                doi="10.1016/j.actamat.2024.119999",
                title="Austenite precipitation thermodynamics in steels",
                journal="Acta Materialia",
                year="2024",
                confidence=0.95,
                source="crossref",
                match_basis="title_similarity",
            ),
            MetadataResult(
                source_index=2,
                query_title="Hydrogen trapping behavior in titanium alloys",
                doi="10.1016/j.actamat.2024.119999",
                title="Hydrogen trapping behavior in titanium alloys",
                journal="Acta Materialia",
                year="2024",
                confidence=0.95,
                source="crossref",
                match_basis="title_similarity",
            ),
        ]

        with patch("sd_institutional_skill.MetadataResolver.resolve_one", side_effect=resolved) as resolve_one:
            rows, unique_rows, counts = classify_entries(
                entries,
                resolve_metadata=True,
                email="",
                min_confidence=0.65,
                http_json=lambda *_args, **_kwargs: {},
            )

        self.assertEqual(resolve_one.call_count, 2)
        self.assertEqual(len(unique_rows), 1)
        self.assertEqual(counts["duplicate"], 1)
        self.assertEqual(rows[1].status, "duplicate")
        self.assertEqual(rows[1].review_hint, "重复项，程序只保留首次识别记录。")
        self.assertNotIn("可进入 ScienceDirect", rows[1].review_hint)

    def test_title_only_resolution_requires_extra_bibliographic_signal(self) -> None:
        from sd_institutional_skill import build_intake

        calls: list[str] = []

        def fake_json(url: str, headers: dict[str, str] | None = None, timeout: int = 20) -> dict:
            calls.append(url)
            return {
                "message": {
                    "items": [
                        {
                            "DOI": "10.1016/j.wrong.2025.1",
                            "title": ["Wrongly matched paper"],
                        }
                    ]
                }
            }

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            csv_path = root / "title_only.csv"
            with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
                writer = csv.DictWriter(f, fieldnames=["title", "notes"])
                writer.writeheader()
                writer.writerow({
                    "title": "Microstructure evolution in laser powder bed fused nickel superalloys",
                    "notes": "only a standalone title, no year author or journal",
                })
                writer.writerow({
                    "title": "",
                    "notes": "P8: Ti-Mo beta-Ti stress-induced martensitic transformation 是否包含 SXRD",
                })
                writer.writerow({
                    "title": "P8: Ti-Mo beta-Ti stress-induced martensitic transformation 是否包含 SXRD",
                    "notes": "remark copied into title column",
                })

            result = build_intake(
                texts=[],
                input_paths=[csv_path],
                folder_paths=[],
                output_dir=root / "out",
                resolve_metadata=True,
                resolve_title_only_files=True,
                http_json=fake_json,
            )
            preview_text = result.preview_path.read_text(encoding="utf-8-sig")
            merged_text = result.merged_input_path.read_text(encoding="utf-8-sig")

        self.assertEqual(calls, [])
        self.assertEqual(result.valid_count, 0)
        self.assertEqual(result.status_counts["needs_review"], 2)
        self.assertEqual(result.status_counts["empty"], 1)
        self.assertIn("insufficient_bibliographic_context", preview_text)
        self.assertIn("not_probable_title", preview_text)
        self.assertNotIn("10.1016/j.wrong.2025.1", merged_text)

    def test_build_intake_title_only_csv_resolves_or_needs_review(self) -> None:
        from sd_institutional_skill import build_intake

        def fake_json(url: str, headers: dict[str, str] | None = None, timeout: int = 20) -> dict:
            decoded_url = unquote_plus(url)
            if "api.crossref.org/works?" in url and "Resolvable title only paper" in decoded_url:
                return {
                    "message": {
                        "items": [
                            {
                                "DOI": "10.1016/j.actamat.2025.120001",
                                "title": ["Resolvable title only paper"],
                                "author": [{"family": "Chen", "given": "Ming"}],
                                "container-title": ["Acta Materialia"],
                                "published-print": {"date-parts": [[2025]]},
                                "publisher": "Elsevier",
                                "URL": "https://doi.org/10.1016/j.actamat.2025.120001",
                            }
                        ]
                    }
                }
            return {}

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            csv_path = root / "title_only.csv"
            with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
                writer = csv.DictWriter(f, fieldnames=["paper title", "author", "source", "publication year"])
                writer.writeheader()
                writer.writerow({
                    "paper title": "Resolvable title only paper",
                    "author": "Chen Ming",
                    "source": "Acta Materialia",
                    "publication year": "2025",
                })
                writer.writerow({
                    "paper title": "Unmatched title only paper",
                    "author": "Li Qiang",
                    "source": "Scripta Materialia",
                    "publication year": "2024",
                })

            result = build_intake(
                texts=[],
                input_paths=[csv_path],
                folder_paths=[],
                output_dir=root / "out",
                resolve_metadata=True,
                resolve_title_only_files=True,
                http_json=fake_json,
            )

            preview_text = result.preview_path.read_text(encoding="utf-8-sig")
            merged_text = result.merged_input_path.read_text(encoding="utf-8-sig")

        self.assertEqual(result.valid_count, 1)
        self.assertEqual(result.status_counts["needs_review"], 1)
        self.assertIn("10.1016/j.actamat.2025.120001", merged_text)
        self.assertIn("Resolvable title only paper", preview_text)
        self.assertIn("Unmatched title only paper", preview_text)
        self.assertIn("metadata_confidence_below_threshold", preview_text)

    def test_build_intake_explicit_missing_doi_column_still_fails(self) -> None:
        from sd_institutional_skill import build_intake

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            csv_path = root / "title_only.csv"
            with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
                writer = csv.DictWriter(f, fieldnames=["title"])
                writer.writeheader()
                writer.writerow({"title": "Title only should not override requested DOI column"})

            with self.assertRaises(ValueError):
                build_intake(
                    texts=[],
                    input_paths=[csv_path],
                    folder_paths=[],
                    output_dir=root / "out",
                    doi_column="missing_doi",
                    resolve_metadata=True,
                    http_json=lambda *_args, **_kwargs: {},
                )

    def test_build_intake_title_only_xlsx_resolves_when_openpyxl_available(self) -> None:
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        from sd_institutional_skill import build_intake

        def fake_json(url: str, headers: dict[str, str] | None = None, timeout: int = 20) -> dict:
            decoded_url = unquote_plus(url)
            if "api.crossref.org/works?" in url and "XLSX title only paper" in decoded_url:
                return {
                    "message": {
                        "items": [
                            {
                                "DOI": "10.1016/j.scriptamat.2025.116001",
                                "title": ["XLSX title only paper"],
                                "author": [{"family": "Wang", "given": "Yu"}],
                                "container-title": ["Scripta Materialia"],
                                "published-online": {"date-parts": [[2025]]},
                                "publisher": "Elsevier",
                                "URL": "https://doi.org/10.1016/j.scriptamat.2025.116001",
                            }
                        ]
                    }
                }
            return {}

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            xlsx_path = root / "title_only.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Papers"
            ws.append(["title", "authors", "journal", "year"])
            ws.append(["XLSX title only paper", "Wang Yu", "Scripta Materialia", "2025"])
            wb.save(xlsx_path)
            wb.close()

            result = build_intake(
                texts=[],
                input_paths=[xlsx_path],
                folder_paths=[],
                output_dir=root / "out",
                resolve_metadata=True,
                resolve_title_only_files=True,
                http_json=fake_json,
            )

            merged_text = result.merged_input_path.read_text(encoding="utf-8-sig")
            preview_text = result.preview_path.read_text(encoding="utf-8-sig")

        self.assertEqual(result.valid_count, 1)
        self.assertIn("10.1016/j.scriptamat.2025.116001", merged_text)
        self.assertIn("XLSX title only paper", preview_text)

    def test_file_inputs_do_not_resolve_title_only_rows_by_default(self) -> None:
        from sd_institutional_skill import build_intake

        calls: list[str] = []

        def fake_json(url: str, headers: dict[str, str] | None = None, timeout: int = 20) -> dict:
            if "api.crossref.org/works?" in url:
                calls.append(url)
                return {
                    "message": {
                        "items": [
                            {
                                "DOI": "10.4028/www.scientific.net/msf.849.219",
                                "title": ["Ti-Mo beta titanium transformation"],
                                "publisher": "Scientific.Net",
                            }
                        ]
                    }
                }
            return {}

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            md_path = root / "papers.md"
            md_path.write_text(
                "| P8 | Remarkable contribution of stress-induced martensitic transformation in Ti-Mo beta-Ti alloys. "
                "[10.1016/j.scriptamat.2024.116254](https://doi.org/10.1016/j.scriptamat.2024.116254) |\n"
                "- P8: Ti-Mo beta-Ti stress-induced martensitic transformation 是否包含 SXRD。\n",
                encoding="utf-8",
            )

            result = build_intake(
                texts=[],
                input_paths=[md_path],
                folder_paths=[],
                output_dir=root / "out",
                resolve_metadata=True,
                http_json=fake_json,
            )

            merged_text = result.merged_input_path.read_text(encoding="utf-8-sig")
            preview_text = result.preview_path.read_text(encoding="utf-8-sig")

        self.assertEqual(calls, [])
        self.assertEqual(result.valid_count, 1)
        self.assertIn("10.1016/j.scriptamat.2024.116254", merged_text)
        self.assertNotIn("10.4028/www.scientific.net/msf.849.219", preview_text)
        self.assertNotIn("10.4028/www.scientific.net/msf.849.219", merged_text)

    def test_tabular_inputs_scan_explicit_doi_without_title_only_resolution_by_default(self) -> None:
        from sd_institutional_skill import build_intake

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            csv_path = root / "mixed.csv"
            with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
                writer = csv.DictWriter(f, fieldnames=["title", "notes"])
                writer.writeheader()
                writer.writerow({
                    "title": "Old Acta paper",
                    "notes": "DOI: 10.1016/S1359-6454(00)00218-4",
                })
                writer.writerow({
                    "title": "Ti-Mo beta-Ti stress-induced martensitic transformation",
                    "notes": "no explicit DOI",
                })

            result = build_intake(
                texts=[],
                input_paths=[csv_path],
                folder_paths=[],
                output_dir=root / "out",
                resolve_metadata=True,
                http_json=lambda *_args, **_kwargs: {},
            )

            merged_text = result.merged_input_path.read_text(encoding="utf-8-sig")

        self.assertEqual(result.valid_count, 1)
        self.assertIn("10.1016/s1359-6454(00)00218-4", merged_text)
        self.assertNotIn("Ti-Mo beta-Ti stress-induced", merged_text)

    def test_choose_output_root_falls_back_when_dialog_fails(self) -> None:
        from sd_institutional_skill import choose_output_root

        def failing_dialog() -> str:
            raise RuntimeError("no display")

        with tempfile.TemporaryDirectory() as tmp:
            selected = choose_output_root(tmp, dialog_func=failing_dialog)

        self.assertEqual(selected, Path(tmp).resolve())

    def test_build_intake_reads_gbk_csv_and_requested_excel_sheet(self) -> None:
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        from sd_institutional_skill import build_intake

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            csv_path = root / "gbk.csv"
            with csv_path.open("w", newline="", encoding="gbk") as f:
                writer = csv.DictWriter(f, fieldnames=["题名", "DOI号"])
                writer.writeheader()
                writer.writerow({"题名": "中文题名", "DOI号": "10.1016/j.jallcom.2025.177000"})

            xlsx_path = root / "papers.xlsx"
            wb = Workbook()
            first = wb.active
            first.title = "Wrong"
            first.append(["title", "doi"])
            first.append(["Wrong", "10.1016/j.wrong.2024.1"])
            target = wb.create_sheet("Target")
            target.append(["标题", "DOI"])
            target.append(["Right", "10.1016/j.right.2024.2"])
            wb.save(xlsx_path)

            result = build_intake(
                texts=[],
                input_paths=[csv_path, xlsx_path],
                folder_paths=[],
                output_dir=root / "out",
                sheet_name="Target",
            )

            self.assertEqual(result.valid_count, 2)
            merged_text = result.merged_input_path.read_text(encoding="utf-8-sig")
            self.assertIn("中文题名", merged_text)
            self.assertIn("10.1016/j.right.2024.2", merged_text)
            self.assertNotIn("10.1016/j.wrong.2024.1", merged_text)


class InstitutionalSkillCookieTests(unittest.TestCase):
    def test_scraper_reads_edge_cookies_before_chrome_cookies(self) -> None:
        import sd_scraper

        class Cookie:
            def __init__(self, name: str, value: str) -> None:
                self.name = name
                self.value = value

        chrome_called = False

        def fake_edge(domain_name: str) -> list[Cookie]:
            self.assertEqual(domain_name, ".sciencedirect.com")
            return [Cookie("SDMSESSION", "edge-secret")]

        def fake_chrome(domain_name: str) -> list[Cookie]:
            nonlocal chrome_called
            chrome_called = True
            return [Cookie("SDMSESSION", "chrome-secret")]

        with patch.object(sd_scraper, "HAS_BROWSER_COOKIE3", True), \
                patch.object(sd_scraper.browser_cookie3, "edge", fake_edge), \
                patch.object(sd_scraper.browser_cookie3, "chrome", fake_chrome):
            scraper = sd_scraper.ScienceDirectScraper(use_browser_cookies=True)

        self.assertFalse(chrome_called)
        self.assertEqual(scraper._cookie_dict["SDMSESSION"], "edge-secret")

    def test_extract_devtools_cookies_and_cache_filters_relevant_domains(self) -> None:
        from sd_institutional_skill import cache_devtools_cookies

        class FakeResponse:
            def __init__(self, payload: dict[str, object]) -> None:
                self.payload = payload

            def __enter__(self) -> "FakeResponse":
                return self

            def __exit__(self, *args: object) -> None:
                return None

            def read(self) -> bytes:
                return json.dumps(self.payload).encode("utf-8")

        class FakeWebSocket:
            def send(self, _payload: str) -> None:
                return None

            def recv(self) -> str:
                return json.dumps({
                    "id": 1,
                    "result": {
                        "cookies": [
                            {"domain": ".sciencedirect.com", "name": "SDMSESSION", "value": "secret"},
                            {"domain": ".elsevier.com", "name": "EUID", "value": "secret2"},
                            {"domain": ".example.com", "name": "other", "value": "nope"},
                        ]
                    },
                })

            def close(self) -> None:
                return None

        class FakeScraper:
            CHROME_DBG_PORT = 9222

            def _is_chrome_debug_ready(self) -> bool:
                return True

            def _launch_chrome_with_debug(self) -> None:
                raise AssertionError("Chrome should not be launched when already ready")

        def fake_opener(request: object, timeout: int = 0) -> FakeResponse:
            url = getattr(request, "full_url", str(request))
            if "/json/new" in url:
                return FakeResponse({"id": "tab-1", "webSocketDebuggerUrl": "ws://tab-1"})
            if "/json/close/" in url:
                return FakeResponse({})
            raise AssertionError(f"Unexpected URL: {url}")

        def fake_ws_connect(*_args: object, **_kwargs: object) -> FakeWebSocket:
            return FakeWebSocket()

        with tempfile.TemporaryDirectory() as tmp:
            cache_path = Path(tmp) / "_auth" / "sciencedirect_cookies.json"
            count = cache_devtools_cookies(
                FakeScraper(),
                cache_path,
                opener=fake_opener,
                websocket_connect=fake_ws_connect,
            )
            data = json.loads(cache_path.read_text(encoding="utf-8"))

        self.assertEqual(count, 2)
        self.assertEqual([item["domain"] for item in data], [".sciencedirect.com", ".elsevier.com"])
        self.assertNotIn(".example.com", json.dumps(data))

    def test_cookie_status_message_never_includes_cookie_values(self) -> None:
        from sd_institutional_skill import cookie_status_message, write_cookie_cache

        with tempfile.TemporaryDirectory() as tmp:
            cache_path = Path(tmp) / "sciencedirect_cookies.json"
            write_cookie_cache([
                {"domain": ".sciencedirect.com", "name": "SDMSESSION", "value": "super-secret"}
            ], cache_path)
            message = cookie_status_message(cache_path)

        self.assertIn("已缓存 1 个", message)
        self.assertNotIn("super-secret", message)

    def test_make_scraper_ignores_invalid_cookie_cache(self) -> None:
        from sd_institutional_skill import make_scraper

        with tempfile.TemporaryDirectory() as tmp:
            cache_path = Path(tmp) / "sciencedirect_cookies.json"
            cache_path.write_text("[]", encoding="utf-8")
            with patch("sd_institutional_skill.ScienceDirectScraper") as factory:
                make_scraper(cache_path)

        factory.assert_called_once_with(use_browser_cookies=True)


class InstitutionalSkillFilenameTests(unittest.TestCase):
    def test_resolve_doi_to_article_fills_sciencedirect_citation_metadata(self) -> None:
        from sd_scraper import ScienceDirectScraper

        class FakeResponse:
            url = "https://www.sciencedirect.com/science/article/pii/S1359645424000012"
            text = """
            <html><head>
              <meta name="citation_title" content="High entropy alloy deformation mechanisms">
              <meta name="citation_author" content="Zhang Wei">
              <meta name="citation_author" content="Li Qiang">
              <meta name="citation_publication_date" content="2024/05/10">
              <meta name="citation_journal_title" content="Acta Materialia">
            </head></html>
            """

        class FakeSession:
            def get(self, *_args, **_kwargs):
                return FakeResponse()

        scraper = ScienceDirectScraper()
        scraper.session = FakeSession()
        article, error = scraper._resolve_doi_to_article({
            "doi": "10.1016/j.actamat.2024.119999",
            "title": "",
            "authors": "",
            "journal": "",
            "year": "",
            "date": "",
        })

        self.assertEqual(error, "")
        self.assertIsNotNone(article)
        assert article is not None
        self.assertEqual(article["title"], "High entropy alloy deformation mechanisms")
        self.assertEqual(article["authors"], "Zhang Wei; Li Qiang")
        self.assertEqual(article["journal"], "Acta Materialia")
        self.assertEqual(article["year"], "2024")

    def test_sciencedirect_filename_uses_metadata_and_doi_hash(self) -> None:
        from sd_scraper import ScienceDirectScraper

        filename = ScienceDirectScraper._make_pdf_filename(
            1,
            {
                "year": "2024",
                "authors": "Zhang Wei; Li Qiang",
                "title": "A/B:C*D? gamma-TiAl alloy <test>",
                "doi": "10.1016/j.actamat.2024.119999",
            },
        )

        self.assertTrue(filename.startswith("2024_Zhang_A B C D gamma-TiAl alloy test_"))
        self.assertTrue(filename.endswith(".pdf"))
        self.assertNotRegex(filename, r'[<>:"/\\|?*]')

    def test_sciencedirect_filename_uses_doi_fallback_without_unknown_labels(self) -> None:
        from sd_scraper import ScienceDirectScraper

        filename = ScienceDirectScraper._make_pdf_filename(
            1,
            {
                "pii": "S1359645424000012",
                "doi": "10.1016/j.actamat.2024.119999",
            },
        )

        self.assertTrue(filename.startswith("undated_no-author_DOI 10.1016 j.actamat.2024.119999_"))
        self.assertNotIn("unknown-year", filename)
        self.assertNotIn("Unknown", filename)
        self.assertTrue(filename.endswith(".pdf"))


class ScienceDirectPdfAccessUrlTests(unittest.TestCase):
    def test_pdf_access_url_accepts_only_real_assets_or_viewer_wrappers(self) -> None:
        from sd_scraper import is_sciencedirect_pdf_access_url

        accepted = [
            "https://pdf.sciencedirectassets.com/271429/1-s2.0-S1359645424000012-main.pdf",
            (
                "chrome-extension://mhjfbmdgcfjbbpaeojofohoefgiehjai/index.html"
                "?src=https%3A%2F%2Fpdf.sciencedirectassets.com%2F271429%2Fmain.pdf"
            ),
            (
                "edge://pdf-viewer/index.html"
                "?src=https%3A%2F%2Fpdf.sciencedirectassets.com%2F271429%2Fmain.pdf"
            ),
            (
                "chrome://pdf-viewer/index.html"
                "?src=https%3A%2F%2Fpdf.sciencedirectassets.com%2F271429%2Fmain.pdf"
            ),
        ]
        rejected = [
            "",
            "http://pdf.sciencedirectassets.com/271429/main.pdf",
            "ftp://pdf.sciencedirectassets.com/271429/main.pdf",
            "https://www.sciencedirect.com/science/article/pii/S1359645424000012/pdfft",
            "https://www.sciencedirect.com/science/article/pii/S1359645424000012",
            "https://id.elsevier.com/as/authorization.oauth2?client_id=sciencedirect",
            "https://cas.example.edu/login?service=https%3A%2F%2Fwww.sciencedirect.com%2F",
            "https://example.com/not-a-real-pdf-route",
            (
                "chrome://settings/?q="
                "https%3A%2F%2Fpdf.sciencedirectassets.com%2F271429%2Fmain.pdf"
            ),
            (
                "chrome-extension://aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa/index.html"
                "?src=https%3A%2F%2Fpdf.sciencedirectassets.com%2F271429%2Fmain.pdf"
            ),
            (
                "chrome-extension://mhjfbmdgcfjbbpaeojofohoefgiehjai/index.html"
                "?src=https%3A%2F%2Fwww.sciencedirect.com%2Fscience%2Farticle%2Fpii%2FS1359645424000012%2Fpdfft"
            ),
        ]

        for url in accepted:
            with self.subTest(url=url):
                self.assertTrue(is_sciencedirect_pdf_access_url(url))
        for url in rejected:
            with self.subTest(url=url):
                self.assertFalse(is_sciencedirect_pdf_access_url(url))


if __name__ == "__main__":
    unittest.main()
