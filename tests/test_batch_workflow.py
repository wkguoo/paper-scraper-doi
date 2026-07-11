from __future__ import annotations

import errno
import hashlib
import json
import multiprocessing
import os
import shutil
import sys
import tempfile
import traceback
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


def _hold_batch_state_lock_worker(run_dir, acquired, release, result_queue) -> None:
    try:
        from paper_automation.batch_workflow import batch_state_lock

        with batch_state_lock(run_dir, timeout=5.0):
            acquired.set()
            if not release.wait(10.0):
                raise TimeoutError("test_release_timeout")
        result_queue.put(("released", ""))
    except BaseException as exc:
        result_queue.put(("error", f"{type(exc).__name__}: {exc}"))


def _attempt_batch_state_lock_worker(run_dir, result_queue) -> None:
    try:
        from paper_automation.batch_workflow import batch_state_lock

        with batch_state_lock(run_dir, timeout=0.3):
            result_queue.put(("acquired", ""))
    except TimeoutError as exc:
        result_queue.put(("timeout", str(exc)))
    except BaseException as exc:
        result_queue.put(("error", f"{type(exc).__name__}: {exc}"))


def _claim_manual_retry_worker(run_dir, start, result_queue) -> None:
    try:
        from paper_automation.batch_workflow import claim_manual_retry

        if not start.wait(10.0):
            raise TimeoutError("test_start_timeout")
        claimed, state = claim_manual_retry(run_dir, timeout=5.0)
        result_queue.put(("ok", claimed, state["manual_retry_used"]))
    except BaseException as exc:
        result_queue.put(("error", f"{type(exc).__name__}: {exc}", False))


def _concurrent_state_save_worker(paths, payload, start, replace_barrier, result_queue) -> None:
    try:
        from paper_automation.batch_workflow import save_batch_state

        if not start.wait(10.0):
            raise TimeoutError("test_start_timeout")
        original_replace = os.replace
        temporary_name = ""
        first_replace = True

        def coordinated_replace(source, destination):
            nonlocal first_replace, temporary_name
            temporary_name = Path(source).name
            if first_replace:
                first_replace = False
                replace_barrier.wait(timeout=10.0)
            return original_replace(source, destination)

        with patch.object(os, "replace", new=coordinated_replace):
            save_batch_state(paths, payload)
        result_queue.put(("ok", payload["writer"], temporary_name))
    except BaseException as exc:
        result_queue.put(("error", f"{type(exc).__name__}: {exc}"))


def _copy_pdf_worker(source, destination, filename, start, started, result_queue) -> None:
    try:
        from paper_automation.batch_workflow import copy_pdf_safely

        if not start.wait(10.0):
            raise TimeoutError("test_start_timeout")
        started.set()
        copied = copy_pdf_safely(source, destination, filename)
        result_queue.put(("ok", str(copied)))
    except BaseException as exc:
        result_queue.put(
            ("error", f"{type(exc).__name__}: {exc}", traceback.format_exc())
        )


def _controlled_copy_pdf_worker(
    source,
    destination,
    filename,
    publish_reached,
    release_publish,
    phase_queue,
    result_queue,
) -> None:
    """Pause at the old partial write or the new atomic publish boundary."""

    try:
        from paper_automation.batch_workflow import copy_pdf_safely

        original_copyfileobj = shutil.copyfileobj
        original_link = os.link
        phase_reported = False

        def report_phase(phase):
            nonlocal phase_reported
            if phase_reported:
                return
            phase_reported = True
            phase_queue.put(phase)
            publish_reached.set()
            if not release_publish.wait(10.0):
                raise TimeoutError("test_publish_release_timeout")

        def controlled_copyfileobj(source_handle, target_handle, length=0):
            prefix = source_handle.read(5)
            target_handle.write(prefix)
            target_handle.flush()
            os.fsync(target_handle.fileno())
            report_phase("direct_write")
            return original_copyfileobj(source_handle, target_handle, length)

        def controlled_link(source_path, target_path, *args, **kwargs):
            report_phase("hard_link")
            return original_link(source_path, target_path, *args, **kwargs)

        with patch.object(shutil, "copyfileobj", new=controlled_copyfileobj), patch.object(
            os,
            "link",
            new=controlled_link,
        ):
            copied = copy_pdf_safely(source, destination, filename)
        result_queue.put(("ok", str(copied)))
    except BaseException as exc:
        result_queue.put(("error", f"{type(exc).__name__}: {exc}"))


class BatchFileTests(unittest.TestCase):
    def _join_workers(self, *workers) -> None:
        for worker in workers:
            if worker is None:
                continue
            worker.join(10.0)
            if worker.is_alive():
                worker.terminate()
                worker.join(5.0)
            self.assertFalse(worker.is_alive(), f"worker did not stop: {worker.name}")

    def test_paths_and_state_are_created_under_new_run_directory(self) -> None:
        from paper_automation.batch_workflow import create_batch_paths, load_batch_state, save_batch_state

        with tempfile.TemporaryDirectory() as tmp:
            paths = create_batch_paths(Path(tmp), now=datetime(2026, 7, 10, 17, 0, 0))
            save_batch_state(paths, {"version": 1, "rows": []})

            self.assertEqual(paths.root.name, "paper_batch_20260710_170000")
            self.assertTrue(paths.pdfs.is_dir())
            self.assertTrue(paths.reports.is_dir())
            self.assertEqual(load_batch_state(paths.root)["version"], 1)

    def test_state_save_fsyncs_unique_temp_before_atomic_replace(self) -> None:
        import paper_automation.batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            paths = workflow.create_batch_paths(Path(tmp))
            events = []
            real_fsync = os.fsync
            real_replace = os.replace

            def recording_fsync(file_descriptor):
                events.append(("fsync", ""))
                return real_fsync(file_descriptor)

            def recording_replace(source, destination):
                events.append(("replace", Path(source).name))
                return real_replace(source, destination)

            with patch.object(os, "fsync", side_effect=recording_fsync), patch.object(
                os,
                "replace",
                side_effect=recording_replace,
            ):
                workflow.save_batch_state(paths, {"version": 1, "rows": []})

            self.assertGreaterEqual(len(events), 2)
            self.assertEqual(events[0][0], "fsync")
            self.assertEqual(events[1][0], "replace")
            self.assertTrue(events[1][1].startswith("batch_state.json."))
            self.assertNotEqual(events[1][1], "batch_state.json.tmp")
            self.assertFalse(list(paths.working.glob("batch_state.json.*.tmp")))

    def test_concurrent_state_saves_remain_valid_and_leave_no_temp_files(self) -> None:
        from paper_automation.batch_workflow import create_batch_paths, load_batch_state

        context = multiprocessing.get_context("spawn")
        with tempfile.TemporaryDirectory() as tmp:
            paths = create_batch_paths(Path(tmp))
            start = context.Event()
            replace_barrier = context.Barrier(2)
            result_queue = context.Queue()
            payloads = [
                {"version": 1, "writer": "one", "rows": [{"value": "a" * 2048}]},
                {"version": 1, "writer": "two", "rows": [{"value": "b" * 2048}]},
            ]
            workers = [
                context.Process(
                    target=_concurrent_state_save_worker,
                    args=(paths, payload, start, replace_barrier, result_queue),
                )
                for payload in payloads
            ]
            for worker in workers:
                worker.start()
            start.set()
            try:
                results = [result_queue.get(timeout=15.0) for _ in workers]
            finally:
                self._join_workers(*workers)

            self.assertEqual(
                sorted(result[0] for result in results),
                ["ok", "ok"],
                results,
            )
            self.assertEqual(len({result[2] for result in results}), 2)
            self.assertTrue(
                all(result[2].startswith("batch_state.json.") for result in results)
            )
            self.assertIn(load_batch_state(paths.root)["writer"], {"one", "two"})
            self.assertFalse(paths.state.with_suffix(".json.tmp").exists())
            self.assertFalse(list(paths.working.glob("batch_state.json.*.tmp")))

    def test_batch_state_lock_excludes_second_process_until_timeout(self) -> None:
        import paper_automation.batch_workflow as workflow

        self.assertTrue(hasattr(workflow, "batch_state_lock"))
        context = multiprocessing.get_context("spawn")
        with tempfile.TemporaryDirectory() as tmp:
            paths = workflow.create_batch_paths(Path(tmp))
            acquired = context.Event()
            release = context.Event()
            holder_result = context.Queue()
            contender_result = context.Queue()
            holder = context.Process(
                target=_hold_batch_state_lock_worker,
                args=(paths.root, acquired, release, holder_result),
            )
            contender = None
            holder.start()
            try:
                self.assertTrue(acquired.wait(10.0), "holder did not acquire state lock")
                contender = context.Process(
                    target=_attempt_batch_state_lock_worker,
                    args=(paths.root, contender_result),
                )
                contender.start()
                self.assertEqual(contender_result.get(timeout=10.0)[0], "timeout")
            finally:
                release.set()
                self._join_workers(holder, contender)

            self.assertEqual(holder_result.get(timeout=5.0)[0], "released")
            self.assertEqual(holder.exitcode, 0)
            self.assertEqual(contender.exitcode, 0)

    def test_batch_state_lock_is_released_after_context_exception(self) -> None:
        import paper_automation.batch_workflow as workflow

        self.assertTrue(hasattr(workflow, "batch_state_lock"))
        with tempfile.TemporaryDirectory() as tmp:
            paths = workflow.create_batch_paths(Path(tmp))
            with self.assertRaisesRegex(RuntimeError, "synthetic"):
                with workflow.batch_state_lock(paths.root, timeout=1.0):
                    raise RuntimeError("synthetic")

            with workflow.batch_state_lock(paths.root, timeout=1.0):
                pass

    def test_batch_state_lock_rejects_negative_and_non_finite_timeouts(self) -> None:
        import paper_automation.batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            paths = workflow.create_batch_paths(Path(tmp))
            for timeout in (-1.0, float("nan"), float("inf"), float("-inf")):
                with self.subTest(timeout=timeout):
                    with self.assertRaisesRegex(
                        ValueError,
                        "invalid_batch_state_lock_timeout",
                    ):
                        with workflow.batch_state_lock(paths.root, timeout=timeout):
                            self.fail("invalid timeout acquired the state lock")

    def test_two_processes_cannot_claim_the_same_manual_retry(self) -> None:
        import paper_automation.batch_workflow as workflow

        self.assertTrue(hasattr(workflow, "claim_manual_retry"))
        context = multiprocessing.get_context("spawn")
        with tempfile.TemporaryDirectory() as tmp:
            paths = workflow.create_batch_paths(Path(tmp))
            workflow.save_batch_state(
                paths,
                {"version": 1, "manual_retry_used": False, "rows": []},
            )
            start = context.Event()
            result_queue = context.Queue()
            workers = [
                context.Process(
                    target=_claim_manual_retry_worker,
                    args=(paths.root, start, result_queue),
                )
                for _ in range(2)
            ]
            for worker in workers:
                worker.start()
            start.set()
            try:
                results = [result_queue.get(timeout=15.0) for _ in workers]
            finally:
                self._join_workers(*workers)

            self.assertEqual(
                [result[0] for result in results],
                ["ok", "ok"],
                results,
            )
            self.assertEqual(sorted(result[1] for result in results), [False, True])
            self.assertTrue(all(result[2] for result in results))
            self.assertTrue(workflow.load_batch_state(paths.root)["manual_retry_used"])

    def test_custom_run_name_is_cleaned_and_always_timestamped(self) -> None:
        from paper_automation.batch_workflow import create_batch_paths

        with tempfile.TemporaryDirectory() as tmp:
            paths = create_batch_paths(
                Path(tmp),
                run_name="Ti alloy batch",
                now=datetime(2026, 7, 10, 17, 0, 0),
            )

            self.assertEqual(paths.root.name, "Ti_alloy_batch_20260710_170000")

    def test_same_timestamp_creates_unique_run_directories(self) -> None:
        from paper_automation.batch_workflow import create_batch_paths

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            now = datetime(2026, 7, 10, 17, 0, 0)

            first = create_batch_paths(root, now=now)
            second = create_batch_paths(root, now=now)

            self.assertEqual(first.root.name, "paper_batch_20260710_170000")
            self.assertEqual(second.root.name, "paper_batch_20260710_170000_2")
            self.assertNotEqual(first.root, second.root)

    def test_run_name_rejects_absolute_and_traversal_paths(self) -> None:
        from paper_automation.batch_workflow import create_batch_paths

        with tempfile.TemporaryDirectory() as tmp:
            sandbox = Path(tmp)
            output_root = sandbox / "output"
            invalid_names = [
                str(sandbox / "absolute"),
                "..",
                "../escape",
                r"..\escape",
            ]

            for run_name in invalid_names:
                with self.subTest(run_name=run_name):
                    with self.assertRaisesRegex(ValueError, "invalid_run_name"):
                        create_batch_paths(output_root, run_name=run_name)

    def test_run_name_rejects_superscript_windows_device_names(self) -> None:
        from paper_automation.batch_workflow import create_batch_paths

        with tempfile.TemporaryDirectory() as tmp:
            for run_name in ("COM¹", "COM²", "COM³", "LPT¹", "LPT²", "LPT³"):
                with self.subTest(run_name=run_name):
                    with self.assertRaisesRegex(ValueError, "invalid_run_name"):
                        create_batch_paths(Path(tmp), run_name=run_name)

    def test_pdf_validation_and_copy_are_non_destructive_and_deduplicated(self) -> None:
        from paper_automation.batch_workflow import copy_pdf_safely, is_valid_pdf

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.pdf"
            source.write_bytes(b"%PDF-1.7\nvalid payload")
            destination = root / "pdfs"

            first = copy_pdf_safely(source, destination, "paper.pdf")
            second = copy_pdf_safely(source, destination, "paper.pdf")

            self.assertTrue(is_valid_pdf(first))
            self.assertEqual(first, second)
            self.assertTrue(source.exists())
            self.assertEqual(len(list(destination.glob("*.pdf"))), 1)

    def test_same_pdf_content_logic_never_reuses_a_symlink(self) -> None:
        import paper_automation.batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "paper.pdf"
            payload = b"%PDF-1.7\nlogic symlink payload"
            pdf.write_bytes(payload)
            expected_hash = hashlib.sha256(payload).hexdigest()

            with patch.object(Path, "is_symlink", return_value=True):
                reusable = workflow._same_pdf_content(pdf, expected_hash)

            self.assertFalse(reusable)

    def test_copy_treats_real_symlinks_as_occupied_and_publishes_independent_pdf(
        self,
    ) -> None:
        import paper_automation.batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.pdf"
            payload = b"%PDF-1.7\nreal symlink payload"
            source.write_bytes(payload)
            destination = root / "pdfs"
            destination.mkdir()
            short_hash = hashlib.sha256(payload).hexdigest()[:8]
            target = destination / "paper.pdf"
            hash_candidate = destination / f"paper_{short_hash}.pdf"
            try:
                target.symlink_to(source)
                hash_candidate.symlink_to(source)
            except (OSError, NotImplementedError) as exc:
                permission_error = (
                    isinstance(exc, PermissionError)
                    or getattr(exc, "errno", None) in {errno.EACCES, errno.EPERM}
                    or getattr(exc, "winerror", None) in {5, 1314}
                )
                if permission_error or isinstance(exc, NotImplementedError):
                    self.skipTest(f"file symlinks unavailable: {exc}")
                raise

            copied = workflow.copy_pdf_safely(source, destination, "paper.pdf")

            self.assertEqual(copied.name, f"paper_{short_hash}_2.pdf")
            self.assertTrue(target.is_symlink())
            self.assertTrue(hash_candidate.is_symlink())
            self.assertFalse(copied.is_symlink())
            self.assertFalse(os.path.samefile(source, copied))
            self.assertEqual(copied.read_bytes(), payload)
            self.assertTrue(workflow.is_valid_pdf(copied))

            source.write_bytes(b"%PDF-1.7\nsource changed later")
            self.assertEqual(copied.read_bytes(), payload)

    def test_same_pdf_concurrent_copies_publish_one_complete_target(self) -> None:
        from paper_automation.batch_workflow import is_valid_pdf

        context = multiprocessing.get_context("spawn")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.pdf"
            source.write_bytes(b"%PDF-1.7\nconcurrent identical payload")
            destination = root / "pdfs"
            publish_reached = context.Event()
            release_publish = context.Event()
            phase_queue = context.Queue()
            result_queue = context.Queue()
            second_start = context.Event()
            second_started = context.Event()
            first = context.Process(
                target=_controlled_copy_pdf_worker,
                args=(
                    source,
                    destination,
                    "paper.pdf",
                    publish_reached,
                    release_publish,
                    phase_queue,
                    result_queue,
                ),
            )
            second = context.Process(
                target=_copy_pdf_worker,
                args=(
                    source,
                    destination,
                    "paper.pdf",
                    second_start,
                    second_started,
                    result_queue,
                ),
            )
            first.start()
            results = []
            try:
                self.assertTrue(
                    publish_reached.wait(10.0),
                    "first worker did not reach its publish boundary",
                )
                phase = phase_queue.get(timeout=5.0)
                second.start()
                second_start.set()
                self.assertTrue(second_started.wait(10.0), "second worker did not start")
                if phase == "direct_write":
                    results.append(result_queue.get(timeout=10.0))
                release_publish.set()
                while len(results) < 2:
                    results.append(result_queue.get(timeout=10.0))
            finally:
                release_publish.set()
                self._join_workers(first, second)

            self.assertEqual([result[0] for result in results], ["ok", "ok"], results)
            copied_paths = [Path(result[1]) for result in results]
            delivered = list(destination.glob("*.pdf"))
            self.assertEqual(copied_paths[0], copied_paths[1])
            self.assertEqual(len(delivered), 1)
            self.assertTrue(is_valid_pdf(delivered[0]))

    def test_terminated_publish_leaves_no_partial_pdf_and_next_copy_recovers(self) -> None:
        from paper_automation.batch_workflow import copy_pdf_safely, is_valid_pdf

        context = multiprocessing.get_context("spawn")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.pdf"
            payload = b"%PDF-1.7\nrecoverable payload after termination"
            source.write_bytes(payload)
            destination = root / "pdfs"
            publish_reached = context.Event()
            release_publish = context.Event()
            phase_queue = context.Queue()
            result_queue = context.Queue()
            worker = context.Process(
                target=_controlled_copy_pdf_worker,
                args=(
                    source,
                    destination,
                    "paper.pdf",
                    publish_reached,
                    release_publish,
                    phase_queue,
                    result_queue,
                ),
            )
            worker.start()
            try:
                self.assertTrue(
                    publish_reached.wait(10.0),
                    "worker did not reach its publish boundary",
                )
                phase_queue.get(timeout=5.0)
                worker.terminate()
                worker.join(10.0)
                self.assertFalse(worker.is_alive())
            finally:
                if worker.is_alive():
                    worker.terminate()
                    worker.join(5.0)

            delivered_before_recovery = list(destination.glob("*.pdf"))
            recovered = copy_pdf_safely(source, destination, "paper.pdf")
            delivered_after_recovery = list(destination.glob("*.pdf"))

            self.assertEqual(delivered_before_recovery, [])
            self.assertEqual(delivered_after_recovery, [recovered])
            self.assertEqual(recovered.read_bytes(), payload)
            self.assertTrue(is_valid_pdf(recovered))
            self.assertFalse(list(destination.glob(".pdf_snapshot_*.tmp")))

    def test_pdf_publish_lock_has_a_bounded_timeout(self) -> None:
        import paper_automation.batch_workflow as workflow

        context = multiprocessing.get_context("spawn")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.pdf"
            source.write_bytes(b"%PDF-1.7\npublish lock timeout payload")
            destination = root / "pdfs"
            publish_reached = context.Event()
            release_publish = context.Event()
            phase_queue = context.Queue()
            result_queue = context.Queue()
            worker = context.Process(
                target=_controlled_copy_pdf_worker,
                args=(
                    source,
                    destination,
                    "paper.pdf",
                    publish_reached,
                    release_publish,
                    phase_queue,
                    result_queue,
                ),
            )
            worker.start()
            try:
                self.assertTrue(publish_reached.wait(10.0))
                self.assertEqual(phase_queue.get(timeout=5.0), "hard_link")
                with self.assertRaisesRegex(TimeoutError, "pdf_publish_lock_timeout"):
                    workflow.copy_pdf_safely(
                        source,
                        destination,
                        "paper.pdf",
                        lock_timeout=0.1,
                    )
            finally:
                release_publish.set()
                self._join_workers(worker)

            self.assertEqual(result_queue.get(timeout=5.0)[0], "ok")

    def test_copy_rejects_negative_and_non_finite_lock_timeouts(self) -> None:
        import paper_automation.batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.pdf"
            source.write_bytes(b"%PDF-1.7\ninvalid timeout payload")
            destination = root / "pdfs"

            for index, timeout in enumerate(
                (-1.0, float("nan"), float("inf"), float("-inf")),
                start=1,
            ):
                with self.subTest(timeout=timeout):
                    with self.assertRaisesRegex(
                        ValueError,
                        "invalid_pdf_publish_lock_timeout",
                    ):
                        workflow.copy_pdf_safely(
                            source,
                            destination,
                            f"paper_{index}.pdf",
                            lock_timeout=timeout,
                        )

    def test_hard_link_publish_failure_never_exposes_a_formal_pdf(self) -> None:
        import paper_automation.batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.pdf"
            source.write_bytes(b"%PDF-1.7\nhard link failure payload")
            destination = root / "pdfs"

            with patch.object(
                os,
                "link",
                side_effect=OSError("synthetic hard-link failure"),
            ):
                with self.assertRaisesRegex(OSError, "hard_link_publish_failed"):
                    workflow.copy_pdf_safely(source, destination, "paper.pdf")

            self.assertFalse(list(destination.glob("*.pdf")))
            self.assertFalse(list(destination.glob(".pdf_snapshot_*.tmp")))

    def test_copy_rejects_absolute_and_traversal_filenames(self) -> None:
        from paper_automation.batch_workflow import copy_pdf_safely

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.pdf"
            source.write_bytes(b"%PDF-1.7\nvalid payload")
            destination = root / "pdfs"
            (destination / "nested").mkdir(parents=True)
            invalid_names = [
                str(root / "outside.pdf"),
                "../outside.pdf",
                r"..\outside.pdf",
                "nested/paper.pdf",
                r"nested\paper.pdf",
            ]

            for filename in invalid_names:
                with self.subTest(filename=filename):
                    with self.assertRaisesRegex(ValueError, "invalid_filename"):
                        copy_pdf_safely(source, destination, filename)

    def test_copy_rejects_superscript_windows_device_filenames(self) -> None:
        from paper_automation.batch_workflow import copy_pdf_safely

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.pdf"
            source.write_bytes(b"%PDF-1.7\nvalid payload")
            destination = root / "pdfs"

            for reserved in ("COM¹", "COM²", "COM³", "LPT¹", "LPT²", "LPT³"):
                filename = f"{reserved}.pdf"
                with self.subTest(filename=filename):
                    try:
                        copy_pdf_safely(source, destination, filename)
                    except Exception as exc:
                        self.assertIsInstance(exc, ValueError)
                        self.assertRegex(str(exc), "invalid_filename")
                    else:
                        self.fail("reserved Windows device filename was accepted")

    def test_different_pdf_content_uses_hash_suffix_without_overwriting(self) -> None:
        from paper_automation.batch_workflow import copy_pdf_safely

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first_source = root / "first.pdf"
            second_source = root / "second.pdf"
            first_payload = b"%PDF-1.7\nfirst payload"
            second_payload = b"%PDF-1.7\nsecond payload"
            first_source.write_bytes(first_payload)
            second_source.write_bytes(second_payload)
            destination = root / "pdfs"

            first = copy_pdf_safely(first_source, destination, "paper.pdf")
            second = copy_pdf_safely(second_source, destination, "paper.pdf")

            short_hash = hashlib.sha256(second_payload).hexdigest()[:8]
            self.assertEqual(first.name, "paper.pdf")
            self.assertEqual(second.name, f"paper_{short_hash}.pdf")
            self.assertEqual(first.read_bytes(), first_payload)
            self.assertEqual(second.read_bytes(), second_payload)

    def test_different_pdf_content_concurrent_copies_publish_unique_targets(self) -> None:
        from paper_automation.batch_workflow import is_valid_pdf

        context = multiprocessing.get_context("spawn")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            sources = [root / "first.pdf", root / "second.pdf"]
            payloads = [
                b"%PDF-1.7\nconcurrent first payload",
                b"%PDF-1.7\nconcurrent second payload",
            ]
            for source, payload in zip(sources, payloads):
                source.write_bytes(payload)
            destination = root / "pdfs"
            start = context.Event()
            result_queue = context.Queue()
            started_events = [context.Event(), context.Event()]
            workers = [
                context.Process(
                    target=_copy_pdf_worker,
                    args=(
                        source,
                        destination,
                        "paper.pdf",
                        start,
                        started,
                        result_queue,
                    ),
                )
                for source, started in zip(sources, started_events)
            ]
            for worker in workers:
                worker.start()
            start.set()
            try:
                self.assertTrue(all(event.wait(10.0) for event in started_events))
                results = [result_queue.get(timeout=15.0) for _ in workers]
            finally:
                self._join_workers(*workers)

            self.assertEqual([result[0] for result in results], ["ok", "ok"], results)
            copied_paths = {Path(result[1]) for result in results}
            self.assertEqual(len(copied_paths), 2)
            self.assertEqual({path.read_bytes() for path in copied_paths}, set(payloads))
            self.assertTrue(all(is_valid_pdf(path) for path in copied_paths))
            self.assertEqual(set(destination.glob("*.pdf")), copied_paths)

    def test_occupied_hash_candidate_gets_numeric_suffix_without_overwriting(self) -> None:
        from paper_automation.batch_workflow import copy_pdf_safely

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.pdf"
            source_payload = b"%PDF-1.7\nnew source payload"
            requested_payload = b"%PDF-1.7\nrequested target payload"
            occupied_payload = b"%PDF-1.7\noccupied hash payload"
            source.write_bytes(source_payload)
            destination = root / "pdfs"
            destination.mkdir()
            requested = destination / "paper.pdf"
            requested.write_bytes(requested_payload)
            short_hash = hashlib.sha256(source_payload).hexdigest()[:8]
            occupied = destination / f"paper_{short_hash}.pdf"
            occupied.write_bytes(occupied_payload)

            copied = copy_pdf_safely(source, destination, "paper.pdf")

            self.assertEqual(copied.name, f"paper_{short_hash}_2.pdf")
            self.assertEqual(requested.read_bytes(), requested_payload)
            self.assertEqual(occupied.read_bytes(), occupied_payload)
            self.assertEqual(copied.read_bytes(), source_payload)

    def test_copy_uses_verified_snapshot_if_source_changes_after_first_read(self) -> None:
        import paper_automation.batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.pdf"
            source_payload = b"%PDF-1.7\nstable snapshot payload"
            source.write_bytes(source_payload)
            destination = root / "pdfs"
            destination.mkdir()
            (destination / "paper.pdf").write_bytes(b"%PDF-1.7\nexisting target")
            expected_hash = hashlib.sha256(source_payload).hexdigest()[:8]
            original_copyfileobj = shutil.copyfileobj
            original_link = os.link
            mutation_seen = False

            def mutate_original_once():
                nonlocal mutation_seen
                if not mutation_seen:
                    source.write_bytes(b"")
                    mutation_seen = True

            def mutate_original_then_copy(source_handle, target_handle, length=0):
                mutate_original_once()
                return original_copyfileobj(source_handle, target_handle, length)

            def mutate_original_then_link(source_path, target_path, *args, **kwargs):
                mutate_original_once()
                return original_link(source_path, target_path, *args, **kwargs)

            with patch.object(
                shutil,
                "copyfileobj",
                side_effect=mutate_original_then_copy,
            ), patch.object(os, "link", side_effect=mutate_original_then_link):
                copied = workflow.copy_pdf_safely(source, destination, "paper.pdf")

            self.assertTrue(mutation_seen)
            self.assertEqual(copied.name, f"paper_{expected_hash}.pdf")
            self.assertEqual(copied.read_bytes(), source_payload)
            self.assertTrue(workflow.is_valid_pdf(copied))
            self.assertFalse(list(destination.glob(".pdf_snapshot_*.tmp")))

    def test_non_pdf_is_rejected(self) -> None:
        from paper_automation.batch_workflow import copy_pdf_safely

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "login.pdf"
            source.write_text("<html>login</html>", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "not_pdf_response"):
                copy_pdf_safely(source, root / "pdfs", "paper.pdf")
            self.assertFalse(list((root / "pdfs").glob(".pdf_snapshot_*.tmp")))


class BatchStageTests(unittest.TestCase):
    def test_stage_result_keeps_normalized_fields(self) -> None:
        from paper_automation.batch_stages import StageResult

        result = StageResult(
            task_id="paper-0001",
            doi="10.1016/example",
            title="Example",
            status="downloaded",
            file="paper.pdf",
            reason="",
            source="oa",
        )

        self.assertEqual(result.source, "oa")
        self.assertEqual(result.task_id, "paper-0001")

    def test_route_splits_sciencedirect_from_other_publishers(self) -> None:
        from paper_automation.batch_stages import split_institutional_rows

        rows = [
            {"task_id": "paper-0001", "doi": "10.1016/j.actamat.2024.1"},
            {"task_id": "paper-0002", "doi": "10.1038/s41467-020-1"},
            {"task_id": "paper-0003", "doi": ""},
        ]

        science_direct, other = split_institutional_rows(rows)

        self.assertEqual([row["task_id"] for row in science_direct], ["paper-0001"])
        self.assertEqual([row["task_id"] for row in other], ["paper-0002", "paper-0003"])

    def test_manual_retry_statuses_are_explicit(self) -> None:
        from paper_automation.batch_stages import needs_manual_retry

        self.assertTrue(needs_manual_retry("auth_required", ""))
        self.assertTrue(needs_manual_retry("failed", "captcha_required"))
        self.assertTrue(needs_manual_retry("failed", "turnstile_detected"))
        self.assertTrue(needs_manual_retry("failed", "institutional login required"))
        self.assertFalse(needs_manual_retry("unsupported_publisher", ""))

    def test_write_stage_input_preserves_rows_as_utf8_sig_csv(self) -> None:
        import csv

        from paper_automation.batch_stages import write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            path = write_stage_input(
                [{"task_id": "paper-0001", "doi": "10.1016/a", "title": "中文标题"}],
                Path(tmp) / "working" / "input.csv",
            )

            self.assertTrue(path.read_bytes().startswith(b"\xef\xbb\xbf"))
            with path.open("r", newline="", encoding="utf-8-sig") as handle:
                rows = list(csv.DictReader(handle))

        self.assertEqual(rows, [{"task_id": "paper-0001", "doi": "10.1016/a", "title": "中文标题", "authors": "", "journal": "", "year": ""}])

    def test_multi_doi_stage_row_is_explicit_failure_and_never_sent_downstream(self) -> None:
        from types import SimpleNamespace

        from paper_automation.batch_stages import (
            BatchOptions,
            run_non_elsevier_stage,
            run_oa_stage,
            run_sciencedirect_stage,
            write_stage_input,
        )
        from paper_automation.parser import parse_mixed_text

        row = {
            "task_id": "paper-0001",
            "doi": "10.1016/first 10.1038/second",
            "title": "One source row with two DOI values",
        }
        parsed = parse_mixed_text(row["doi"])
        self.assertEqual([candidate.doi for candidate in parsed], ["10.1016/first", "10.1038/second"])
        self.assertEqual([candidate.source_index for candidate in parsed], [1, 1])

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with patch(
                "paper_automation.batch_stages.run_workflow",
                return_value=SimpleNamespace(manifest_csv="", output_dir=""),
            ) as oa_lower:
                oa_results = run_oa_stage([row], root / "oa-output", BatchOptions())

            input_path = write_stage_input([row], root / "input.csv")
            with patch("paper_automation.batch_stages.sd_main", return_value=0) as sd_lower:
                sd_results = run_sciencedirect_stage(input_path, root / "sd-output", BatchOptions())
            with patch(
                "paper_automation.batch_stages.run_institutional_workflow",
                return_value=SimpleNamespace(report_path=""),
            ) as institutional_lower:
                institutional_results = run_non_elsevier_stage(
                    input_path,
                    root / "institutional-output",
                    BatchOptions(),
                )

        for source, results in (
            ("oa", oa_results),
            ("sciencedirect", sd_results),
            ("non_elsevier", institutional_results),
        ):
            with self.subTest(source=source):
                self.assertEqual(len(results), 1)
                self.assertEqual(results[0].task_id, "paper-0001")
                self.assertEqual(results[0].status, "failed")
                self.assertEqual(results[0].reason, "multiple_dois_in_stage_input")
                self.assertEqual(results[0].source, source)

        oa_lower.assert_not_called()
        sd_lower.assert_not_called()
        institutional_lower.assert_not_called()

    def test_multi_doi_failure_preserves_owner_and_duplicate_layout(self) -> None:
        import csv
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions, run_oa_stage

        with tempfile.TemporaryDirectory() as tmp:
            run_dir = Path(tmp) / "oa_run"
            valid_pdf = run_dir / "pdfs" / "owner.pdf"
            valid_pdf.parent.mkdir(parents=True)
            valid_pdf.write_bytes(b"%PDF-1.7\nowner fixture")
            manifest = run_dir / "metadata" / "manifest.csv"
            manifest.parent.mkdir(parents=True)
            received_lines: list[str] = []

            def fake_workflow(input_text: str, *args, **kwargs):
                received_lines.extend(input_text.splitlines())
                with manifest.open("w", newline="", encoding="utf-8-sig") as handle:
                    writer = csv.DictWriter(handle, fieldnames=["source_index", "doi", "title", "download_status", "file", "reason"])
                    writer.writeheader()
                    writer.writerow({"source_index": "1", "doi": "10.1000/owner", "title": "Owner", "download_status": "downloaded", "file": "owner.pdf"})
                return SimpleNamespace(manifest_csv=str(manifest), output_dir=str(run_dir))

            with patch("paper_automation.batch_stages.run_workflow", side_effect=fake_workflow):
                results = run_oa_stage(
                    [
                        {"task_id": "paper-0001", "doi": "10.1000/first 10.1000/second", "title": "Multi"},
                        {"task_id": "paper-0002", "doi": "10.1000/owner", "title": "Owner"},
                        {"task_id": "paper-0003", "doi": "https://doi.org/10.1000/OWNER", "title": "Duplicate"},
                    ],
                    run_dir.parent,
                    BatchOptions(),
                )

        self.assertEqual(received_lines, ["10.1000/owner"])
        self.assertEqual([result.task_id for result in results], ["paper-0001", "paper-0002", "paper-0003"])
        self.assertEqual([result.status for result in results], ["failed", "downloaded", "duplicate"])
        self.assertEqual(results[0].reason, "multiple_dois_in_stage_input")
        self.assertEqual(results[2].reason, "duplicate_stage_input")

    def test_oa_adapter_maps_manifest_rows_in_input_order(self) -> None:
        import csv
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions, run_oa_stage

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            run_dir = root / "oa_run"
            valid_pdf = run_dir / "pdfs" / "oa.pdf"
            valid_pdf.parent.mkdir(parents=True)
            valid_pdf.write_bytes(b"%PDF-1.7\nvalid fixture")
            manifest = run_dir / "metadata" / "final_manifest.csv"
            manifest.parent.mkdir(parents=True)
            with manifest.open("w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.DictWriter(handle, fieldnames=["source_index", "doi", "title", "download_status", "file", "reason"])
                writer.writeheader()
                writer.writerow({"source_index": "2", "doi": "10.1000/b", "title": "Resolved metadata B", "download_status": "failed", "reason": "no_legal_open_pdf"})
                writer.writerow({"source_index": "1", "doi": "10.1000/a", "title": "Resolved metadata A", "download_status": "downloaded", "file": "oa.pdf"})

            with patch(
                "paper_automation.batch_stages.run_workflow",
                return_value=SimpleNamespace(manifest_csv=str(manifest), output_dir=str(run_dir)),
            ) as mocked:
                results = run_oa_stage(
                    [
                        {"task_id": "paper-0001", "doi": "https://doi.org/10.1000/a", "title": "Original title A"},
                        {"task_id": "paper-0002", "doi": "https://doi.org/10.1000/b", "title": "Original title B"},
                    ],
                    root,
                    BatchOptions(email="researcher@example.edu"),
                )

        self.assertEqual([result.task_id for result in results], ["paper-0001", "paper-0002"])
        self.assertEqual([result.status for result in results], ["downloaded", "failed"])
        self.assertEqual([result.doi for result in results], ["10.1000/a", "10.1000/b"])
        self.assertEqual([result.title for result in results], ["Resolved metadata A", "Resolved metadata B"])
        self.assertEqual(results[0].file, str(valid_pdf.resolve()))
        self.assertEqual(results[1].reason, "no_legal_open_pdf")
        self.assertEqual(mocked.call_args.kwargs["email"], "researcher@example.edu")

    def test_oa_relative_pdf_falls_back_to_manifest_run_directory(self) -> None:
        import csv
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions, run_oa_stage

        with tempfile.TemporaryDirectory() as tmp:
            run_dir = Path(tmp) / "oa_run"
            valid_pdf = run_dir / "pdfs" / "fallback.pdf"
            valid_pdf.parent.mkdir(parents=True)
            valid_pdf.write_bytes(b"%PDF-1.7\nfallback fixture")
            manifest = run_dir / "metadata" / "final_manifest.csv"
            manifest.parent.mkdir(parents=True)
            with manifest.open("w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.DictWriter(handle, fieldnames=["source_index", "doi", "title", "download_status", "file", "reason"])
                writer.writeheader()
                writer.writerow({"source_index": "1", "doi": "10.1000/a", "title": "A", "download_status": "downloaded", "file": "fallback.pdf"})

            with patch(
                "paper_automation.batch_stages.run_workflow",
                return_value=SimpleNamespace(manifest_csv=str(manifest), output_dir=None),
            ):
                results = run_oa_stage(
                    [{"task_id": "paper-0001", "doi": "10.1000/a", "title": "A"}],
                    run_dir.parent,
                    BatchOptions(),
                )

        self.assertEqual(results[0].status, "downloaded")
        self.assertEqual(results[0].file, str(valid_pdf.resolve()))

    def test_oa_adapter_falls_back_to_clean_doi_and_normalized_title(self) -> None:
        import csv
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions, run_oa_stage

        with tempfile.TemporaryDirectory() as tmp:
            run_dir = Path(tmp) / "oa_run"
            manifest = run_dir / "metadata" / "final_manifest.csv"
            manifest.parent.mkdir(parents=True)
            with manifest.open("w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.DictWriter(handle, fieldnames=["source_index", "doi", "title", "download_status", "file", "reason"])
                writer.writeheader()
                writer.writerow({"doi": "10.1000/url", "title": "Resolved URL title", "download_status": "failed", "reason": "no_legal_open_pdf"})
                writer.writerow({"doi": "", "title": "Normalized   Title", "download_status": "needs_review", "reason": "title_only"})

            with patch(
                "paper_automation.batch_stages.run_workflow",
                return_value=SimpleNamespace(manifest_csv=str(manifest), output_dir=str(run_dir)),
            ):
                results = run_oa_stage(
                    [
                        {"task_id": "paper-0001", "doi": "https://doi.org/10.1000/url", "title": "Original URL title"},
                        {"task_id": "paper-0002", "doi": "", "title": " normalized title "},
                    ],
                    run_dir.parent,
                    BatchOptions(),
                )

        self.assertEqual([result.task_id for result in results], ["paper-0001", "paper-0002"])
        self.assertEqual([result.reason for result in results], ["no_legal_open_pdf", "title_only"])

    def test_oa_source_index_ignores_empty_stage_lines(self) -> None:
        import csv
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions, run_oa_stage

        with tempfile.TemporaryDirectory() as tmp:
            run_dir = Path(tmp) / "oa_run"
            valid_pdf = run_dir / "pdfs" / "after-empty.pdf"
            valid_pdf.parent.mkdir(parents=True)
            valid_pdf.write_bytes(b"%PDF-1.7\nafter empty fixture")
            manifest = run_dir / "metadata" / "final_manifest.csv"
            manifest.parent.mkdir(parents=True)
            with manifest.open("w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.DictWriter(handle, fieldnames=["source_index", "doi", "title", "download_status", "file", "reason"])
                writer.writeheader()
                writer.writerow({"source_index": "1", "doi": "10.1000/after-empty", "title": "Resolved", "download_status": "downloaded", "file": "after-empty.pdf"})

            with patch(
                "paper_automation.batch_stages.run_workflow",
                return_value=SimpleNamespace(manifest_csv=str(manifest), output_dir=str(run_dir)),
            ) as mocked:
                results = run_oa_stage(
                    [
                        {"task_id": "paper-0001", "doi": "", "title": ""},
                        {"task_id": "paper-0002", "doi": "10.1000/after-empty", "title": "Original"},
                    ],
                    run_dir.parent,
                    BatchOptions(),
                )

        self.assertEqual(mocked.call_args.args[0], "10.1000/after-empty")
        self.assertEqual([result.task_id for result in results], ["paper-0001", "paper-0002"])
        self.assertEqual(results[0].reason, "missing_stage_report_row")
        self.assertEqual(results[1].status, "downloaded")
        self.assertEqual(results[1].file, str(valid_pdf.resolve()))

    def test_oa_missing_manifest_path_values_are_safe_failures(self) -> None:
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions, run_oa_stage

        variants = {
            "none": SimpleNamespace(manifest_csv=None, output_dir=""),
            "empty": SimpleNamespace(manifest_csv="", output_dir=""),
            "missing_attribute": SimpleNamespace(output_dir=""),
        }
        for name, workflow_result in variants.items():
            with self.subTest(name=name), tempfile.TemporaryDirectory() as tmp, patch(
                "paper_automation.batch_stages.run_workflow",
                return_value=workflow_result,
            ):
                caught = None
                try:
                    results = run_oa_stage(
                        [{"task_id": "paper-0001", "doi": "10.1000/a", "title": "A"}],
                        Path(tmp),
                        BatchOptions(),
                    )
                except Exception as exc:
                    caught = exc
                    results = []

                self.assertIsNone(caught, f"unexpected {type(caught).__name__}: {caught}")
                self.assertEqual(results[0].status, "failed")
                self.assertEqual(results[0].reason, "stage_report_missing")

    def test_sciencedirect_adapter_maps_success_auth_and_cli_options(self) -> None:
        import csv

        from paper_automation.batch_stages import BatchOptions, run_sciencedirect_stage, write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            valid_pdf = root / "sciencedirect" / "pdfs" / "download.pdf"
            valid_pdf.parent.mkdir(parents=True)
            valid_pdf.write_bytes(b"%PDF-1.7\nfixture")
            input_path = write_stage_input(
                [
                    {"task_id": "paper-0001", "doi": "10.1016/a", "title": "A"},
                    {"task_id": "paper-0002", "doi": "10.1016/b", "title": "B"},
                ],
                root / "input.csv",
            )

            def fake_sd_main(argv: list[str]) -> int:
                report_dir = Path(argv[argv.index("--out") + 1]) / "sciencedirect"
                report_dir.mkdir(parents=True, exist_ok=True)
                with (report_dir / "pdf_download_report.csv").open("w", newline="", encoding="utf-8-sig") as handle:
                    fields = ["doi", "pii", "title", "status", "file", "reason", "manual_pdf_url", "manual_status", "manual_reason"]
                    writer = csv.DictWriter(handle, fieldnames=fields)
                    writer.writeheader()
                    writer.writerow({"doi": "10.1016/a", "title": "A", "status": "success", "file": "download.pdf"})
                    writer.writerow({"doi": "10.1016/b", "title": "B", "status": "failed", "reason": "auth_required"})
                return 0

            with patch("paper_automation.batch_stages.sd_main", side_effect=fake_sd_main) as mocked:
                results = run_sciencedirect_stage(
                    input_path,
                    root,
                    BatchOptions(
                        email="researcher@example.edu",
                        cookies="cookies.json",
                        browser_exe="C:/Browser/browser.exe",
                        login_wait_seconds=30,
                    ),
                )

        argv = mocked.call_args.args[0]
        self.assertEqual([result.task_id for result in results], ["paper-0001", "paper-0002"])
        self.assertEqual(results[0].status, "downloaded")
        self.assertEqual(results[0].file, str(valid_pdf.resolve()))
        self.assertEqual(results[1].reason, "auth_required")
        self.assertIn("--run-name", argv)
        self.assertEqual(argv[argv.index("--run-name") + 1], "sciencedirect")
        self.assertIn("--no-download-supplements", argv)
        self.assertEqual(argv[argv.index("--cookies") + 1], "cookies.json")
        self.assertEqual(argv[argv.index("--browser-exe") + 1], "C:/Browser/browser.exe")
        self.assertEqual(argv[argv.index("--login-wait-seconds") + 1], "30")

    def test_sciencedirect_adapter_rejects_invalid_pdf_and_missing_stage_rows(self) -> None:
        import csv

        from paper_automation.batch_stages import BatchOptions, run_sciencedirect_stage, write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            invalid_pdf = root / "login.pdf"
            invalid_pdf.write_text("<html>login</html>", encoding="utf-8")
            input_path = write_stage_input(
                [
                    {"task_id": "paper-0001", "doi": "10.1016/a", "title": "A"},
                    {"task_id": "paper-0002", "doi": "10.1016/b", "title": "B"},
                ],
                root / "input.csv",
            )

            def fake_sd_main(argv: list[str]) -> int:
                report_dir = Path(argv[argv.index("--out") + 1]) / "sciencedirect"
                report_dir.mkdir(parents=True, exist_ok=True)
                with (report_dir / "pdf_download_report.csv").open("w", newline="", encoding="utf-8-sig") as handle:
                    fields = ["doi", "pii", "title", "status", "file", "reason", "manual_pdf_url", "manual_status", "manual_reason"]
                    writer = csv.DictWriter(handle, fieldnames=fields)
                    writer.writeheader()
                    writer.writerow({"doi": "10.1016/a", "title": "A", "status": "success", "file": str(invalid_pdf)})
                return 0

            with patch("paper_automation.batch_stages.sd_main", side_effect=fake_sd_main):
                results = run_sciencedirect_stage(input_path, root, BatchOptions())

        self.assertEqual(results[0].status, "failed")
        self.assertEqual(results[0].reason, "invalid_pdf")
        self.assertEqual(results[1].status, "failed")
        self.assertEqual(results[1].reason, "missing_stage_report_row")

    def test_sciencedirect_adapter_reports_nonzero_exit_without_report(self) -> None:
        from paper_automation.batch_stages import BatchOptions, run_sciencedirect_stage, write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_path = write_stage_input(
                [{"task_id": "paper-0001", "doi": "10.1016/a", "title": "A"}],
                root / "input.csv",
            )
            with patch("paper_automation.batch_stages.sd_main", return_value=3):
                results = run_sciencedirect_stage(input_path, root, BatchOptions())

        self.assertEqual(results[0].status, "failed")
        self.assertEqual(results[0].reason, "stage_exit_code_3")

    def test_sciencedirect_nonzero_exit_preserves_reported_rows(self) -> None:
        import csv

        from paper_automation.batch_stages import BatchOptions, run_sciencedirect_stage, write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            valid_pdf = root / "sciencedirect" / "pdfs" / "partial.pdf"
            valid_pdf.parent.mkdir(parents=True)
            valid_pdf.write_bytes(b"%PDF-1.7\npartial fixture")
            input_path = write_stage_input(
                [
                    {"task_id": "paper-0001", "doi": "10.1016/a", "title": "A"},
                    {"task_id": "paper-0002", "doi": "10.1016/b", "title": "B"},
                ],
                root / "input.csv",
            )

            def fake_sd_main(argv: list[str]) -> int:
                report_dir = Path(argv[argv.index("--out") + 1]) / "sciencedirect"
                with (report_dir / "pdf_download_report.csv").open("w", newline="", encoding="utf-8-sig") as handle:
                    fields = ["doi", "pii", "title", "status", "file", "reason", "manual_pdf_url", "manual_status", "manual_reason"]
                    writer = csv.DictWriter(handle, fieldnames=fields)
                    writer.writeheader()
                    writer.writerow({"doi": "10.1016/a", "title": "A", "status": "success", "file": "partial.pdf"})
                return 7

            with patch("paper_automation.batch_stages.sd_main", side_effect=fake_sd_main):
                results = run_sciencedirect_stage(input_path, root, BatchOptions())

        self.assertEqual([result.task_id for result in results], ["paper-0001", "paper-0002"])
        self.assertEqual(results[0].status, "downloaded")
        self.assertEqual(results[0].file, str(valid_pdf.resolve()))
        self.assertEqual(results[1].status, "failed")
        self.assertEqual(results[1].reason, "stage_exit_code_7")

    def test_sciencedirect_adapter_reports_missing_report(self) -> None:
        from paper_automation.batch_stages import BatchOptions, run_sciencedirect_stage, write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_path = write_stage_input(
                [{"task_id": "paper-0001", "doi": "10.1016/a", "title": "A"}],
                root / "input.csv",
            )
            with patch("paper_automation.batch_stages.sd_main", return_value=0):
                results = run_sciencedirect_stage(input_path, root, BatchOptions())

        self.assertEqual(results[0].status, "failed")
        self.assertEqual(results[0].reason, "stage_report_missing")

    def test_non_elsevier_adapter_maps_report_and_options(self) -> None:
        import csv
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions, run_non_elsevier_stage, write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            valid_pdf = root / "non_elsevier_institutional" / "pdfs" / "institutional.pdf"
            valid_pdf.parent.mkdir(parents=True)
            valid_pdf.write_bytes(b"%PDF-1.7\ninstitutional fixture")
            input_path = write_stage_input(
                [
                    {"task_id": "paper-0001", "doi": "10.1038/a", "title": "A"},
                    {"task_id": "paper-0002", "doi": "", "title": "Unknown"},
                ],
                root / "input.csv",
            )

            def fake_workflow(*args, **kwargs):
                report = root / "non_elsevier_institutional" / "institutional_pdf_download_report.csv"
                report.parent.mkdir(parents=True, exist_ok=True)
                with report.open("w", newline="", encoding="utf-8-sig") as handle:
                    fields = ["row_number", "doi", "title", "journal", "year", "publisher", "adapter", "status", "file", "reason", "landing_url", "final_landing_url", "pdf_url", "metadata_source"]
                    writer = csv.DictWriter(handle, fieldnames=fields)
                    writer.writeheader()
                    writer.writerow({"row_number": "2", "doi": "", "title": "Unknown", "status": "unsupported_publisher", "reason": "no_supported_adapter"})
                    writer.writerow({"row_number": "1", "doi": "10.1038/a", "title": "A", "status": "pdf_downloaded", "file": "institutional.pdf"})
                return SimpleNamespace(report_path=str(report))

            with patch("paper_automation.batch_stages.run_institutional_workflow", side_effect=fake_workflow) as mocked:
                results = run_non_elsevier_stage(
                    input_path,
                    root,
                    BatchOptions(browser_exe="C:/Browser/browser.exe", debug_port=9444, login_wait_seconds=20, throttle_seconds=0.0),
                )

        self.assertEqual([result.task_id for result in results], ["paper-0001", "paper-0002"])
        self.assertEqual([result.status for result in results], ["downloaded", "unsupported_publisher"])
        self.assertEqual(results[0].file, str(valid_pdf.resolve()))
        self.assertEqual(results[1].reason, "no_supported_adapter")
        self.assertEqual(mocked.call_args.kwargs["browser_exe"], "C:/Browser/browser.exe")
        self.assertEqual(mocked.call_args.kwargs["debug_port"], 9444)
        self.assertEqual(mocked.call_args.kwargs["login_wait_seconds"], 20)
        self.assertEqual(mocked.call_args.kwargs["throttle_seconds"], 0.0)

    def test_stage_normalizes_absolute_symlink_to_final_pdf_when_supported(self) -> None:
        import csv

        from paper_automation.batch_stages import BatchOptions, run_sciencedirect_stage, write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            external = root / "external.pdf"
            external.write_bytes(b"%PDF-1.7\nexternal fixture")
            linked = root / "linked.pdf"
            try:
                os.symlink(external, linked)
            except OSError as exc:
                self.skipTest(f"symlink_not_available: {exc}")
            input_path = write_stage_input(
                [{"task_id": "paper-0001", "doi": "10.1016/a", "title": "A"}],
                root / "input.csv",
            )

            def fake_sd_main(argv: list[str]) -> int:
                report_dir = Path(argv[argv.index("--out") + 1]) / "sciencedirect"
                report_dir.mkdir(parents=True, exist_ok=True)
                with (report_dir / "pdf_download_report.csv").open("w", newline="", encoding="utf-8-sig") as handle:
                    fields = ["doi", "pii", "title", "status", "file", "reason", "manual_pdf_url", "manual_status", "manual_reason"]
                    writer = csv.DictWriter(handle, fieldnames=fields)
                    writer.writeheader()
                    writer.writerow({"doi": "10.1016/a", "title": "A", "status": "success", "file": str(linked)})
                return 0

            with patch("paper_automation.batch_stages.sd_main", side_effect=fake_sd_main):
                results = run_sciencedirect_stage(input_path, root, BatchOptions())

        self.assertEqual(results[0].status, "downloaded")
        self.assertEqual(results[0].file, str(external.resolve()))

    def test_sciencedirect_relative_pdf_cannot_escape_pdf_directory(self) -> None:
        import csv

        from paper_automation.batch_stages import BatchOptions, run_sciencedirect_stage, write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            report_dir = root / "sciencedirect"
            pdf_dir = report_dir / "pdfs"
            pdf_dir.mkdir(parents=True)
            escaped_pdf = report_dir / "escaped.pdf"
            escaped_pdf.write_bytes(b"%PDF-1.7\nescaped fixture")
            input_path = write_stage_input(
                [{"task_id": "paper-0001", "doi": "10.1016/a", "title": "A"}],
                root / "input.csv",
            )

            def fake_sd_main(argv: list[str]) -> int:
                with (report_dir / "pdf_download_report.csv").open("w", newline="", encoding="utf-8-sig") as handle:
                    fields = ["doi", "pii", "title", "status", "file", "reason", "manual_pdf_url", "manual_status", "manual_reason"]
                    writer = csv.DictWriter(handle, fieldnames=fields)
                    writer.writeheader()
                    writer.writerow({"doi": "10.1016/a", "title": "A", "status": "success", "file": "../escaped.pdf"})
                return 0

            original_cwd = Path.cwd()
            os.chdir(pdf_dir)
            try:
                with patch("paper_automation.batch_stages.sd_main", side_effect=fake_sd_main):
                    results = run_sciencedirect_stage(input_path, root, BatchOptions())
            finally:
                os.chdir(original_cwd)

        self.assertEqual(results[0].status, "failed")
        self.assertEqual(results[0].reason, "invalid_pdf")

    def test_relative_pdf_symlink_is_rejected_before_resolution(self) -> None:
        from paper_automation.batch_stages import _resolve_report_pdf

        with tempfile.TemporaryDirectory() as tmp:
            pdf_dir = Path(tmp) / "pdfs"
            pdf_dir.mkdir()
            linked = pdf_dir / "linked.pdf"
            target = pdf_dir / "target.pdf"
            target.write_bytes(b"%PDF-1.7\ntarget fixture")
            original_resolve = Path.resolve

            def fake_resolve(path: Path, *args, **kwargs) -> Path:
                if path == linked:
                    return target
                return original_resolve(path, *args, **kwargs)

            def fake_is_symlink(path: Path) -> bool:
                return path == linked

            with patch.object(Path, "resolve", autospec=True, side_effect=fake_resolve), patch.object(
                Path,
                "is_symlink",
                autospec=True,
                side_effect=fake_is_symlink,
            ):
                resolved = _resolve_report_pdf("linked.pdf", pdf_dir)

        self.assertIsNone(resolved)

    def test_sciencedirect_absolute_pdf_path_is_canonicalized(self) -> None:
        import csv

        from paper_automation.batch_stages import BatchOptions, run_sciencedirect_stage, write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            valid_pdf = root / "actual" / "absolute.pdf"
            valid_pdf.parent.mkdir()
            valid_pdf.write_bytes(b"%PDF-1.7\nabsolute fixture")
            (root / "alias").mkdir()
            reported_pdf = root / "alias" / ".." / "actual" / "absolute.pdf"
            input_path = write_stage_input(
                [{"task_id": "paper-0001", "doi": "10.1016/a", "title": "A"}],
                root / "input.csv",
            )

            def fake_sd_main(argv: list[str]) -> int:
                report_dir = root / "sciencedirect"
                report_dir.mkdir()
                with (report_dir / "pdf_download_report.csv").open("w", newline="", encoding="utf-8-sig") as handle:
                    fields = ["doi", "pii", "title", "status", "file", "reason", "manual_pdf_url", "manual_status", "manual_reason"]
                    writer = csv.DictWriter(handle, fieldnames=fields)
                    writer.writeheader()
                    writer.writerow({"doi": "10.1016/a", "title": "A", "status": "success", "file": str(reported_pdf)})
                return 0

            with patch("paper_automation.batch_stages.sd_main", side_effect=fake_sd_main):
                results = run_sciencedirect_stage(input_path, root, BatchOptions())

        self.assertEqual(results[0].status, "downloaded")
        self.assertEqual(results[0].file, str(valid_pdf.resolve()))

    def test_absolute_pdf_resolve_returns_final_reparse_target_semantics(self) -> None:
        from paper_automation.batch_stages import _resolve_report_pdf

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            alias_file = root / "junction-alias" / "paper.pdf"
            final_file = root / "real-target" / "paper.pdf"
            final_file.parent.mkdir()
            final_file.write_bytes(b"%PDF-1.7\nreparse target fixture")
            original_resolve = Path.resolve

            def fake_resolve(path: Path, *args, **kwargs) -> Path:
                if path == alias_file:
                    return final_file
                return original_resolve(path, *args, **kwargs)

            with patch.object(Path, "resolve", autospec=True, side_effect=fake_resolve):
                resolved = _resolve_report_pdf(str(alias_file), root / "unused-base")

        self.assertEqual(resolved, final_file)

    @unittest.skipUnless(os.name == "nt", "Windows junction semantics")
    def test_sciencedirect_absolute_junction_path_resolves_to_final_pdf(self) -> None:
        import csv
        import subprocess

        from paper_automation.batch_stages import BatchOptions, run_sciencedirect_stage, write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target_dir = root / "junction-target"
            target_dir.mkdir()
            final_pdf = target_dir / "paper.pdf"
            final_pdf.write_bytes(b"%PDF-1.7\njunction target fixture")
            junction_dir = root / "junction-alias"
            created = subprocess.run(
                ["cmd.exe", "/d", "/c", "mklink", "/J", str(junction_dir), str(target_dir)],
                capture_output=True,
                text=True,
                check=False,
            )
            if created.returncode != 0:
                self.skipTest("current Windows environment cannot create a directory junction")

            try:
                input_path = write_stage_input(
                    [{"task_id": "paper-0001", "doi": "10.1016/a", "title": "A"}],
                    root / "input.csv",
                )

                def fake_sd_main(args: list[str]) -> int:
                    report = root / "sciencedirect" / "pdf_download_report.csv"
                    report.parent.mkdir(parents=True)
                    with report.open("w", newline="", encoding="utf-8-sig") as handle:
                        fields = ["doi", "pii", "title", "status", "file", "reason", "manual_pdf_url", "manual_status", "manual_reason"]
                        writer = csv.DictWriter(handle, fieldnames=fields)
                        writer.writeheader()
                        writer.writerow(
                            {
                                "doi": "10.1016/a",
                                "title": "A",
                                "status": "success",
                                "file": str(junction_dir / "paper.pdf"),
                            }
                        )
                    return 0

                with patch("paper_automation.batch_stages.sd_main", side_effect=fake_sd_main):
                    results = run_sciencedirect_stage(input_path, root, BatchOptions())
            finally:
                if junction_dir.exists():
                    junction_dir.rmdir()

        self.assertEqual(results[0].status, "downloaded")
        self.assertEqual(results[0].file, str(final_pdf.resolve()))

    def test_non_elsevier_adapter_keeps_missing_blank_rows_in_place(self) -> None:
        import csv
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions, run_non_elsevier_stage, write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_path = write_stage_input(
                [{"task_id": "paper-0001"}, {"task_id": "paper-0002"}],
                root / "input.csv",
            )

            def fake_workflow(*args, **kwargs):
                report = root / "non_elsevier_institutional" / "institutional_pdf_download_report.csv"
                report.parent.mkdir(parents=True, exist_ok=True)
                with report.open("w", newline="", encoding="utf-8-sig") as handle:
                    fields = ["row_number", "doi", "title", "journal", "year", "publisher", "adapter", "status", "file", "reason", "landing_url", "final_landing_url", "pdf_url", "metadata_source"]
                    writer = csv.DictWriter(handle, fieldnames=fields)
                    writer.writeheader()
                    writer.writerow({"row_number": "3", "status": "unsupported_publisher", "reason": "no_supported_adapter"})
                    writer.writerow({"row_number": "1", "status": "failed", "reason": "bogus_row_number"})
                return SimpleNamespace(report_path=str(report))

            with patch("paper_automation.batch_stages.run_institutional_workflow", side_effect=fake_workflow):
                results = run_non_elsevier_stage(input_path, root, BatchOptions())

        self.assertEqual([result.task_id for result in results], ["paper-0001", "paper-0002"])
        self.assertEqual(results[0].reason, "missing_stage_report_row")
        self.assertEqual(results[1].status, "unsupported_publisher")

    def test_non_elsevier_missing_report_path_values_are_safe_failures(self) -> None:
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions, run_non_elsevier_stage, write_stage_input

        variants = {
            "none": SimpleNamespace(report_path=None),
            "empty": SimpleNamespace(report_path=""),
            "missing_attribute": SimpleNamespace(),
        }
        for name, workflow_result in variants.items():
            with self.subTest(name=name), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                input_path = write_stage_input(
                    [{"task_id": "paper-0001", "doi": "10.1038/a", "title": "A"}],
                    root / "input.csv",
                )
                with patch(
                    "paper_automation.batch_stages.run_institutional_workflow",
                    return_value=workflow_result,
                ):
                    caught = None
                    try:
                        results = run_non_elsevier_stage(input_path, root, BatchOptions())
                    except Exception as exc:
                        caught = exc
                        results = []

                self.assertIsNone(caught, f"unexpected {type(caught).__name__}: {caught}")
                self.assertEqual(results[0].status, "failed")
                self.assertEqual(results[0].reason, "stage_report_missing")

    def test_oa_adapter_marks_normalized_duplicate_doi_without_second_download(self) -> None:
        import csv
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions, run_oa_stage

        with tempfile.TemporaryDirectory() as tmp:
            run_dir = Path(tmp) / "oa_run"
            valid_pdf = run_dir / "pdfs" / "owner.pdf"
            valid_pdf.parent.mkdir(parents=True)
            valid_pdf.write_bytes(b"%PDF-1.7\nowner fixture")
            manifest = run_dir / "metadata" / "final_manifest.csv"
            manifest.parent.mkdir(parents=True)
            received_lines: list[str] = []

            def fake_workflow(input_text: str, *args, **kwargs):
                received_lines.extend(input_text.splitlines())
                with manifest.open("w", newline="", encoding="utf-8-sig") as handle:
                    writer = csv.DictWriter(handle, fieldnames=["source_index", "doi", "title", "download_status", "file", "reason"])
                    writer.writeheader()
                    writer.writerow({"source_index": "1", "doi": "10.1000/duplicate", "title": "Owner", "download_status": "downloaded", "file": "owner.pdf"})
                return SimpleNamespace(manifest_csv=str(manifest), output_dir=str(run_dir))

            with patch("paper_automation.batch_stages.run_workflow", side_effect=fake_workflow):
                results = run_oa_stage(
                    [
                        {"task_id": "paper-0001", "doi": "https://doi.org/10.1000/duplicate", "title": "Owner"},
                        {"task_id": "paper-0002", "doi": "10.1000/DUPLICATE", "title": "Duplicate"},
                    ],
                    run_dir.parent,
                    BatchOptions(),
                )

        self.assertEqual(received_lines, ["https://doi.org/10.1000/duplicate"])
        self.assertEqual([result.task_id for result in results], ["paper-0001", "paper-0002"])
        self.assertEqual([result.status for result in results], ["downloaded", "duplicate"])
        self.assertEqual(results[1].reason, "duplicate_stage_input")
        self.assertEqual(results[1].source, "oa")
        self.assertEqual(results[1].file, "")

    def test_oa_adapter_marks_normalized_duplicate_title(self) -> None:
        import csv
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions, run_oa_stage

        with tempfile.TemporaryDirectory() as tmp:
            run_dir = Path(tmp) / "oa_run"
            manifest = run_dir / "metadata" / "final_manifest.csv"
            manifest.parent.mkdir(parents=True)
            received_lines: list[str] = []

            def fake_workflow(input_text: str, *args, **kwargs):
                received_lines.extend(input_text.splitlines())
                with manifest.open("w", newline="", encoding="utf-8-sig") as handle:
                    writer = csv.DictWriter(handle, fieldnames=["source_index", "doi", "title", "download_status", "file", "reason"])
                    writer.writeheader()
                    writer.writerow({"source_index": "1", "title": "Normalized Title", "download_status": "failed", "reason": "no_legal_open_pdf"})
                return SimpleNamespace(manifest_csv=str(manifest), output_dir=str(run_dir))

            with patch("paper_automation.batch_stages.run_workflow", side_effect=fake_workflow):
                results = run_oa_stage(
                    [
                        {"task_id": "paper-0001", "doi": "", "title": "Normalized   Title"},
                        {"task_id": "paper-0002", "doi": "", "title": " normalized title "},
                    ],
                    run_dir.parent,
                    BatchOptions(),
                )

        self.assertEqual(received_lines, ["Normalized   Title"])
        self.assertEqual([result.status for result in results], ["failed", "duplicate"])
        self.assertEqual(results[1].reason, "duplicate_stage_input")

    def test_sciencedirect_adapter_marks_normalized_duplicate_doi_without_second_download(self) -> None:
        import csv

        from paper_automation.batch_stages import BatchOptions, run_sciencedirect_stage, write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            valid_pdf = root / "sciencedirect" / "pdfs" / "owner.pdf"
            valid_pdf.parent.mkdir(parents=True)
            valid_pdf.write_bytes(b"%PDF-1.7\nowner fixture")
            input_path = write_stage_input(
                [
                    {"task_id": "paper-0001", "doi": "https://doi.org/10.1016/duplicate", "title": "Owner"},
                    {"task_id": "paper-0002", "doi": "10.1016/DUPLICATE", "title": "Duplicate"},
                ],
                root / "input.csv",
            )
            received_task_ids: list[str] = []

            def fake_sd_main(argv: list[str]) -> int:
                with Path(argv[argv.index("--input") + 1]).open("r", newline="", encoding="utf-8-sig") as handle:
                    received_task_ids.extend(row["task_id"] for row in csv.DictReader(handle))
                report = root / "sciencedirect" / "pdf_download_report.csv"
                with report.open("w", newline="", encoding="utf-8-sig") as handle:
                    fields = ["doi", "pii", "title", "status", "file", "reason", "manual_pdf_url", "manual_status", "manual_reason"]
                    writer = csv.DictWriter(handle, fieldnames=fields)
                    writer.writeheader()
                    writer.writerow({"doi": "10.1016/duplicate", "title": "Owner", "status": "success", "file": "owner.pdf"})
                return 0

            with patch("paper_automation.batch_stages.sd_main", side_effect=fake_sd_main):
                results = run_sciencedirect_stage(input_path, root, BatchOptions())

        self.assertEqual(received_task_ids, ["paper-0001"])
        self.assertEqual([result.task_id for result in results], ["paper-0001", "paper-0002"])
        self.assertEqual([result.status for result in results], ["downloaded", "duplicate"])
        self.assertEqual(results[1].reason, "duplicate_stage_input")
        self.assertEqual(results[1].source, "sciencedirect")
        self.assertEqual(results[1].file, "")

    def test_non_elsevier_adapter_marks_normalized_duplicate_doi_without_second_download(self) -> None:
        import csv
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions, run_non_elsevier_stage, write_stage_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            valid_pdf = root / "non_elsevier_institutional" / "pdfs" / "owner.pdf"
            valid_pdf.parent.mkdir(parents=True)
            valid_pdf.write_bytes(b"%PDF-1.7\nowner fixture")
            input_path = write_stage_input(
                [
                    {"task_id": "paper-0001", "doi": "https://doi.org/10.1038/duplicate", "title": "Owner"},
                    {"task_id": "paper-0002", "doi": "10.1038/DUPLICATE", "title": "Duplicate"},
                ],
                root / "input.csv",
            )
            received_task_ids: list[str] = []

            def fake_workflow(*args, **kwargs):
                with Path(kwargs["input_path"]).open("r", newline="", encoding="utf-8-sig") as handle:
                    received_task_ids.extend(row["task_id"] for row in csv.DictReader(handle))
                report = root / "non_elsevier_institutional" / "institutional_pdf_download_report.csv"
                with report.open("w", newline="", encoding="utf-8-sig") as handle:
                    fields = ["row_number", "doi", "title", "journal", "year", "publisher", "adapter", "status", "file", "reason", "landing_url", "final_landing_url", "pdf_url", "metadata_source"]
                    writer = csv.DictWriter(handle, fieldnames=fields)
                    writer.writeheader()
                    writer.writerow({"row_number": "2", "doi": "10.1038/duplicate", "title": "Owner", "status": "pdf_downloaded", "file": "owner.pdf"})
                return SimpleNamespace(report_path=str(report))

            with patch("paper_automation.batch_stages.run_institutional_workflow", side_effect=fake_workflow):
                results = run_non_elsevier_stage(input_path, root, BatchOptions())

        self.assertEqual(received_task_ids, ["paper-0001"])
        self.assertEqual([result.task_id for result in results], ["paper-0001", "paper-0002"])
        self.assertEqual([result.status for result in results], ["downloaded", "duplicate"])
        self.assertEqual(results[1].reason, "duplicate_stage_input")
        self.assertEqual(results[1].source, "non_elsevier")
        self.assertEqual(results[1].file, "")


class FakeBatchGateway:
    """Offline gateway that preserves the workflow contract for Task 4 tests."""

    def __init__(self, initial_updates=None, retry_updates=None) -> None:
        self.initial_updates = initial_updates
        self.retry_updates = retry_updates
        self.initial_rows: list[dict] = []
        self.retry_rows: list[dict] = []
        self.retry_options = []

    def run_initial(self, rows, paths, options):
        self.initial_rows = [dict(row) for row in rows]
        if self.initial_updates is not None:
            return self.initial_updates(rows)
        return [{**row, "status": "no_open_pdf", "source": "oa", "file": "", "reason": "no_open_pdf"} for row in rows]

    def run_retry(self, rows, paths, options):
        self.retry_rows.append([dict(row) for row in rows])
        self.retry_options.append(options)
        if self.retry_updates is not None:
            return self.retry_updates(rows)
        return [{**row, "status": "no_entitlement", "source": "institutional", "file": "", "reason": "retry_exhausted"} for row in rows]

    @property
    def retry_calls(self) -> int:
        return len(self.retry_rows)


class BatchRunTests(unittest.TestCase):
    def _normalized_rows(self, pdf: Path) -> list[dict]:
        return [
            {"task_id": "paper-0001", "source_index": "1", "input_doi": "10.1000/a", "doi": "10.1000/a", "title": "A", "fixture_pdf": str(pdf)},
            {"task_id": "paper-0002", "source_index": "2", "input_doi": "10.1000/b", "doi": "10.1000/b", "title": "B"},
            {"task_id": "paper-0003", "source_index": "3", "input_doi": "10.1000/c", "doi": "10.1000/c", "title": "C"},
        ]

    def test_start_writes_only_manual_rows_to_retry_and_other_failures_to_zotero(self) -> None:
        import csv

        from paper_automation.batch_workflow import start_batch

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            valid_pdf = root / "fixture.pdf"
            valid_pdf.write_bytes(b"%PDF-1.7\nfixture")

            def initial(rows):
                return [
                    {**rows[0], "status": "oa_downloaded", "source": "oa", "file": rows[0]["fixture_pdf"], "reason": ""},
                    {**rows[1], "status": "captcha_required", "source": "institutional", "file": "", "reason": "captcha_required"},
                    {**rows[2], "status": "no_open_pdf", "source": "oa", "file": "", "reason": "no_open_pdf"},
                ]

            result = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=root,
                gateway=FakeBatchGateway(initial_updates=initial),
                normalizer=lambda **_kwargs: self._normalized_rows(valid_pdf),
                now=datetime(2026, 7, 10, 17, 0, 0),
            )
            with result.paths.manual_retry.open("r", encoding="utf-8-sig") as handle:
                retry_rows = list(csv.DictReader(handle))
            with result.paths.zotero_fallback.open("r", encoding="utf-8-sig") as handle:
                fallback_rows = list(csv.DictReader(handle))

            self.assertEqual(len(list(result.paths.pdfs.glob("*.pdf"))), 1)
            self.assertEqual([row["task_id"] for row in retry_rows], ["paper-0002"])
            self.assertEqual([row["task_id"] for row in fallback_rows], ["paper-0003"])
            self.assertEqual(result.success_count, 1)

    def test_resume_runs_manual_retry_only_once_and_restores_saved_options(self) -> None:
        from paper_automation.batch_stages import BatchOptions
        from paper_automation.batch_workflow import load_batch_state, resume_batch, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            valid_pdf = root / "fixture.pdf"
            valid_pdf.write_bytes(b"%PDF-1.7\nfixture")
            options = BatchOptions(email="reader@example.edu", cookies=str(root / "cookies.json"), debug_port=9444)
            gateway = FakeBatchGateway(initial_updates=lambda rows: [
                {**rows[0], "status": "captcha_required", "source": "institutional", "file": "", "reason": "captcha_required"},
                {**rows[1], "status": "duplicate", "source": "oa", "file": "", "reason": "duplicate_input"},
                {**rows[2], "status": "duplicate", "source": "oa", "file": "", "reason": "duplicate_input"},
            ])
            started = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=root,
                options=options,
                gateway=gateway,
                normalizer=lambda **_kwargs: self._normalized_rows(valid_pdf),
            )
            resume_batch(started.paths.root, gateway=gateway)
            resume_batch(started.paths.root, gateway=gateway)

            state = load_batch_state(started.paths.root)
            self.assertEqual(len(gateway.retry_rows), 1)
            self.assertEqual(gateway.retry_rows[0][0]["task_id"], "paper-0001")
            self.assertEqual(gateway.retry_options[0], options)
            self.assertTrue(state["manual_retry_used"])
            self.assertEqual(state["options"]["cookies"], str(root / "cookies.json"))

    def test_second_resume_returns_without_revalidating_deleted_pending_file(self) -> None:
        from paper_automation.batch_workflow import resume_batch, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            gateway = FakeBatchGateway(initial_updates=lambda rows: [
                {**rows[0], "status": "captcha_required", "source": "institutional", "file": "", "reason": "captcha_required"},
            ])
            started = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=Path(tmp),
                gateway=gateway,
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                ],
            )
            first = resume_batch(started.paths.root, gateway=gateway)
            started.paths.manual_retry.unlink()
            second = resume_batch(started.paths.root, gateway=gateway)

        self.assertEqual(len(gateway.retry_rows), 1)
        self.assertEqual(second.total_count, first.total_count)

    def test_resume_preclaim_failures_do_not_consume_manual_retry(self) -> None:
        import csv

        from paper_automation.batch_workflow import load_batch_state, resume_batch, save_batch_state, start_batch

        def start_manual(root: Path):
            return start_batch(
                input_text="fixture",
                input_path=None,
                output_root=root,
                gateway=FakeBatchGateway(initial_updates=lambda rows: [
                    {**rows[0], "status": "captcha_required", "source": "institutional", "file": "", "reason": "captcha_required"},
                ]),
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                ],
            )

        for case in ("missing_file", "invalid_fields", "unknown_task", "stale_manual_status", "damaged_options"):
            with self.subTest(case=case), tempfile.TemporaryDirectory() as tmp:
                started = start_manual(Path(tmp))
                state = load_batch_state(started.paths.root)
                if case == "missing_file":
                    started.paths.manual_retry.unlink()
                elif case == "invalid_fields":
                    started.paths.manual_retry.write_text(
                        "task_id,status\npaper-0001,captcha_required\n",
                        encoding="utf-8-sig",
                    )
                elif case == "unknown_task":
                    with started.paths.manual_retry.open("r", newline="", encoding="utf-8-sig") as handle:
                        reader = csv.DictReader(handle)
                        rows = list(reader)
                        fieldnames = reader.fieldnames
                    rows[0]["task_id"] = "paper-9999"
                    with started.paths.manual_retry.open("w", newline="", encoding="utf-8-sig") as handle:
                        writer = csv.DictWriter(handle, fieldnames=fieldnames)
                        writer.writeheader()
                        writer.writerows(rows)
                elif case == "stale_manual_status":
                    state["rows"][0].update(status="no_open_pdf", reason="no_open_pdf")
                    save_batch_state(started.paths, state)
                else:
                    state["options"]["debug_port"] = 70000
                    save_batch_state(started.paths, state)

                gateway = FakeBatchGateway()
                with self.assertRaises((ValueError, FileNotFoundError)):
                    resume_batch(started.paths.root, gateway=gateway)

                self.assertFalse(load_batch_state(started.paths.root)["manual_retry_used"])
                self.assertEqual(gateway.retry_rows, [])

    def test_manual_retry_csv_must_exactly_match_all_current_manual_rows(self) -> None:
        import csv

        from paper_automation.batch_workflow import load_batch_state, resume_batch, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            started = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=Path(tmp),
                gateway=FakeBatchGateway(initial_updates=lambda rows: [
                    {**row, "status": "captcha_required", "source": "institutional", "file": "", "reason": "captcha_required"}
                    for row in rows
                ]),
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                    {"task_id": "paper-0002", "doi": "10.1000/b", "title": "B", "status": "pending"},
                ],
            )
            with started.paths.manual_retry.open("r", newline="", encoding="utf-8-sig") as handle:
                reader = csv.DictReader(handle)
                rows = list(reader)
                fieldnames = reader.fieldnames
            with started.paths.manual_retry.open("w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.DictWriter(handle, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerow(rows[0])

            gateway = FakeBatchGateway()
            with self.assertRaisesRegex(ValueError, "manual_retry_task_ids_mismatch"):
                resume_batch(started.paths.root, gateway=gateway)

            self.assertFalse(load_batch_state(started.paths.root)["manual_retry_used"])
            self.assertEqual(gateway.retry_rows, [])

    def test_empty_manual_retry_set_does_not_claim_rewrite_or_call_gateway(self) -> None:
        from paper_automation.batch_workflow import load_batch_state, resume_batch, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            gateway = FakeBatchGateway(initial_updates=lambda rows: [
                {**rows[0], "status": "no_open_pdf", "source": "oa", "file": "", "reason": "no_open_pdf"},
            ])
            started = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=Path(tmp),
                gateway=gateway,
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                ],
            )
            before = {
                path: path.read_bytes()
                for path in (started.paths.manual_retry, started.paths.zotero_fallback)
            }

            resumed = resume_batch(started.paths.root, gateway=gateway)

            self.assertFalse(load_batch_state(started.paths.root)["manual_retry_used"])
            self.assertEqual(gateway.retry_rows, [])
            self.assertEqual(
                {path: path.read_bytes() for path in before},
                before,
            )
            self.assertEqual(resumed.zotero_fallback_count, 1)

    def test_empty_manual_retry_csv_rejects_manual_state_without_claiming(self) -> None:
        from paper_automation.batch_workflow import NORMALIZED_FIELDS, load_batch_state, resume_batch, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            started = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=Path(tmp),
                gateway=FakeBatchGateway(initial_updates=lambda rows: [
                    {**rows[0], "status": "captcha_required", "source": "institutional", "file": "", "reason": "captcha_required"},
                ]),
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                ],
            )
            started.paths.manual_retry.write_text(
                ",".join(NORMALIZED_FIELDS) + "\n",
                encoding="utf-8-sig",
            )

            gateway = FakeBatchGateway()
            with self.assertRaisesRegex(ValueError, "manual_retry_task_ids_mismatch"):
                resume_batch(started.paths.root, gateway=gateway)

            self.assertFalse(load_batch_state(started.paths.root)["manual_retry_used"])
            self.assertEqual(gateway.retry_rows, [])

    def test_manual_retry_csv_status_must_match_the_current_state_status(self) -> None:
        import csv

        from paper_automation.batch_workflow import load_batch_state, resume_batch, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            started = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=Path(tmp),
                gateway=FakeBatchGateway(initial_updates=lambda rows: [
                    {**rows[0], "status": "captcha_required", "source": "institutional", "file": "", "reason": "captcha_required"},
                ]),
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                ],
            )
            with started.paths.manual_retry.open("r", newline="", encoding="utf-8-sig") as handle:
                reader = csv.DictReader(handle)
                rows = list(reader)
                fieldnames = reader.fieldnames
            rows[0]["status"] = "auth_required"
            rows[0]["reason"] = "auth_required"
            with started.paths.manual_retry.open("w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.DictWriter(handle, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerows(rows)

            with self.assertRaisesRegex(ValueError, "manual_retry_status_mismatch"):
                resume_batch(started.paths.root, gateway=FakeBatchGateway())

            self.assertFalse(load_batch_state(started.paths.root)["manual_retry_used"])

    def test_claim_callbacks_validate_state_before_used_and_can_cancel_atomically(self) -> None:
        from paper_automation.batch_workflow import claim_manual_retry, create_batch_paths, load_batch_state, save_batch_state

        with tempfile.TemporaryDirectory() as tmp:
            paths = create_batch_paths(Path(tmp))
            state = {
                "version": 1,
                "run_dir": str(paths.root),
                "manual_retry_used": False,
                "options": {
                    "email": "",
                    "cookies": "",
                    "browser_exe": "",
                    "login_wait_seconds": 120,
                    "debug_port": 9222,
                    "throttle_seconds": 0.0,
                },
                "rows": [],
            }
            save_batch_state(paths, state)
            calls = []

            claimed, _ = claim_manual_retry(
                paths.root,
                validate_state=lambda _state: calls.append("state"),
                validate_before_claim=lambda _state: calls.append("preclaim") or False,
            )

            self.assertFalse(claimed)
            self.assertEqual(calls, ["state", "preclaim"])
            self.assertFalse(load_batch_state(paths.root)["manual_retry_used"])

            state["manual_retry_used"] = True
            save_batch_state(paths, state)
            calls.clear()
            claimed, _ = claim_manual_retry(
                paths.root,
                validate_state=lambda _state: calls.append("state"),
                validate_before_claim=lambda _state: calls.append("preclaim"),
            )
            self.assertFalse(claimed)
            self.assertEqual(calls, ["state"])

    def test_resume_validates_corrupt_run_dir_even_after_retry_was_used(self) -> None:
        from paper_automation.batch_workflow import load_batch_state, resume_batch, save_batch_state, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            started = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=Path(tmp),
                gateway=FakeBatchGateway(initial_updates=lambda rows: [
                    {**rows[0], "status": "auth_required", "source": "institutional", "file": "", "reason": "auth_required"},
                ]),
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                ],
            )
            resume_batch(started.paths.root, gateway=FakeBatchGateway())
            state = load_batch_state(started.paths.root)
            state["run_dir"] = "..\\escaped-run"
            save_batch_state(started.paths, state)

            with self.assertRaisesRegex(ValueError, "invalid_batch_state_run_dir"):
                resume_batch(started.paths.root, gateway=FakeBatchGateway())

    def test_options_reject_cookie_secrets_and_invalid_numeric_values_before_state_write(self) -> None:
        from paper_automation.batch_stages import BatchOptions
        from paper_automation.batch_workflow import start_batch

        invalid_options = [
            BatchOptions(cookies="Cookie: session=topsecret"),
            BatchOptions(cookies="session=topsecret; token=abc"),
            BatchOptions(cookies='{"session":"topsecret"}'),
            BatchOptions(cookies="https://example.test/cookies.json"),
            BatchOptions(cookies="cookies.json\r\nCookie: secret=abc"),
            BatchOptions(login_wait_seconds=True),
            BatchOptions(login_wait_seconds=-1),
            BatchOptions(login_wait_seconds=1.5),
            BatchOptions(login_wait_seconds=float("nan")),
            BatchOptions(login_wait_seconds=float("inf")),
            BatchOptions(debug_port=False),
            BatchOptions(debug_port=-1),
            BatchOptions(debug_port=0),
            BatchOptions(debug_port=65536),
            BatchOptions(debug_port=float("nan")),
            BatchOptions(debug_port=float("inf")),
            BatchOptions(throttle_seconds=True),
            BatchOptions(throttle_seconds=-0.1),
            BatchOptions(throttle_seconds=float("nan")),
            BatchOptions(throttle_seconds=float("inf")),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for index, options in enumerate(invalid_options, start=1):
                with self.subTest(index=index, options=options):
                    output = root / str(index)
                    with self.assertRaises(ValueError):
                        start_batch(
                            input_text="fixture",
                            input_path=None,
                            output_root=output,
                            options=options,
                            gateway=FakeBatchGateway(),
                            normalizer=lambda **_kwargs: [
                                {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                            ],
                        )
                    state_files = list(output.rglob("batch_state.json")) if output.exists() else []
                    self.assertTrue(all(b"topsecret" not in path.read_bytes() for path in state_files))

    def test_options_allow_nonexistent_windows_and_relative_cookie_json_paths(self) -> None:
        from paper_automation.batch_stages import BatchOptions
        from paper_automation.batch_workflow import load_batch_state, start_batch

        for cookie_path in (r"C:\research\cookies.json", r"exports\cookies.json", "cookies.json", ""):
            with self.subTest(cookie_path=cookie_path), tempfile.TemporaryDirectory() as tmp:
                result = start_batch(
                    input_text="fixture",
                    input_path=None,
                    output_root=Path(tmp),
                    options=BatchOptions(
                        cookies=cookie_path,
                        login_wait_seconds=0,
                        debug_port=65535,
                        throttle_seconds=0.0,
                    ),
                    gateway=FakeBatchGateway(initial_updates=lambda rows: [
                        {**rows[0], "status": "duplicate", "source": "oa", "file": "", "reason": "duplicate_input"},
                    ]),
                    normalizer=lambda **_kwargs: [
                        {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                    ],
                )
                self.assertEqual(load_batch_state(result.paths.root)["options"]["cookies"], cookie_path)

    def test_state_writer_rejects_nan_without_replacing_previous_snapshot(self) -> None:
        from paper_automation.batch_workflow import create_batch_paths, load_batch_state, save_batch_state

        with tempfile.TemporaryDirectory() as tmp:
            paths = create_batch_paths(Path(tmp))
            save_batch_state(paths, {"version": 1, "value": "safe"})
            with self.assertRaises(ValueError):
                save_batch_state(paths, {"version": 1, "value": float("nan")})

            self.assertEqual(load_batch_state(paths.root), {"version": 1, "value": "safe"})

    def test_invalid_resume_option_override_does_not_consume_claim(self) -> None:
        from paper_automation.batch_stages import BatchOptions
        from paper_automation.batch_workflow import load_batch_state, resume_batch, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            started = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=Path(tmp),
                gateway=FakeBatchGateway(initial_updates=lambda rows: [
                    {**rows[0], "status": "auth_required", "source": "institutional", "file": "", "reason": "auth_required"},
                ]),
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                ],
            )
            with self.assertRaises(ValueError):
                resume_batch(
                    started.paths.root,
                    gateway=FakeBatchGateway(),
                    options=BatchOptions(cookies="Cookie: session=secret"),
                )

            self.assertFalse(load_batch_state(started.paths.root)["manual_retry_used"])

    def test_normalize_input_supports_markdown_csv_and_xlsx_without_changing_source(self) -> None:
        import csv

        from openpyxl import Workbook
        from paper_automation.batch_stages import BatchOptions
        from paper_automation.batch_workflow import create_batch_paths, normalize_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            markdown = root / "papers.md"
            markdown.write_text("- DOI: https://doi.org/10.1000/markdown\n", encoding="utf-8")
            csv_path = root / "papers.csv"
            with csv_path.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=["doi", "title"])
                writer.writeheader()
                writer.writerow({"doi": "10.1000/csv", "title": "CSV paper"})
            xlsx_path = root / "papers.xlsx"
            workbook = Workbook()
            workbook.active.append(["doi", "title"])
            workbook.active.append(["10.1000/xlsx", "XLSX paper"])
            workbook.save(xlsx_path)

            for index, input_path in enumerate((markdown, csv_path, xlsx_path), start=1):
                before = input_path.read_bytes()
                paths = create_batch_paths(root, run_name=f"run-{index}")
                rows = normalize_input(input_text=None, input_path=input_path, paths=paths, options=BatchOptions())
                self.assertEqual(len(rows), 1)
                self.assertTrue(rows[0]["doi"].startswith("10.1000/"))
                self.assertEqual(input_path.read_bytes(), before)
                self.assertTrue(paths.normalized_input.read_bytes().startswith(b"\xef\xbb\xbf"))

    def test_normalize_expands_multi_doi_before_gateway_and_keeps_deterministic_ids(self) -> None:
        from paper_automation.batch_workflow import start_batch

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            gateway = FakeBatchGateway()
            start_batch(
                input_text="10.1016/first and https://doi.org/10.1038/second",
                input_path=None,
                output_root=root,
                gateway=gateway,
            )

        self.assertEqual([row["task_id"] for row in gateway.initial_rows], ["paper-0001", "paper-0002"])
        self.assertEqual([row["doi"] for row in gateway.initial_rows], ["10.1016/first", "10.1038/second"])

    def test_normalize_expands_multi_doi_records_in_all_supported_file_types_without_mutation(self) -> None:
        import csv

        from openpyxl import Workbook

        from paper_automation.batch_stages import BatchOptions
        from paper_automation.batch_workflow import create_batch_paths, normalize_input

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            doi_cell = "10.1000/first; https://doi.org/10.1000/second"
            inputs = []

            csv_path = root / "papers.csv"
            with csv_path.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=["doi", "title"])
                writer.writeheader()
                writer.writerow({"doi": doi_cell, "title": "Two DOI table record"})
            inputs.append(csv_path)

            for suffix in (".xlsx", ".xlsm"):
                workbook = Workbook()
                workbook.active.append(["doi", "title"])
                workbook.active.append([doi_cell, "Two DOI workbook record"])
                workbook_path = root / f"papers{suffix}"
                workbook.save(workbook_path)
                inputs.append(workbook_path)

            for suffix in (".txt", ".md"):
                text_path = root / f"papers{suffix}"
                text_path.write_text(doi_cell + "\n", encoding="utf-8")
                inputs.append(text_path)

            for index, input_path in enumerate(inputs, start=1):
                with self.subTest(suffix=input_path.suffix):
                    original = input_path.read_bytes()
                    paths = create_batch_paths(root, run_name=f"multi-{index}")
                    rows = normalize_input(
                        input_text=None,
                        input_path=input_path,
                        paths=paths,
                        options=BatchOptions(),
                    )

                    self.assertEqual(
                        [(row["task_id"], row["doi"]) for row in rows],
                        [("paper-0001", "10.1000/first"), ("paper-0002", "10.1000/second")],
                    )
                    self.assertEqual(input_path.read_bytes(), original)

    def test_tabular_title_reference_doi_is_not_treated_as_an_explicit_target(self) -> None:
        import csv
        from types import SimpleNamespace

        from openpyxl import Workbook

        from paper_automation.batch_stages import BatchOptions
        from paper_automation.batch_workflow import create_batch_paths, normalize_input
        from sd_institutional_skill import IntakeRow

        title = "Target paper cites DOI 10.5555/reference"
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            inputs = []
            for suffix in (".csv", ".xlsx", ".xlsm"):
                for has_doi_column in (False, True):
                    path = root / f"reference-{'blank-doi' if has_doi_column else 'title-only'}{suffix}"
                    headers = ["doi", "title"] if has_doi_column else ["title", "authors"]
                    values = ["", title] if has_doi_column else [title, "Researcher"]
                    if suffix == ".csv":
                        with path.open("w", newline="", encoding="utf-8") as handle:
                            writer = csv.writer(handle)
                            writer.writerow(headers)
                            writer.writerow(values)
                    else:
                        workbook = Workbook()
                        workbook.active.append(headers)
                        workbook.active.append(values)
                        workbook.save(path)
                    inputs.append(path)

            intake_row = IntakeRow(
                source="fixture",
                row_number=2,
                input_doi="10.5555/reference",
                doi="10.5555/reference",
                input_title=title,
                title=title,
                raw_value=title,
                status="valid",
                reason="",
            )
            for index, input_path in enumerate(inputs, start=1):
                with self.subTest(name=input_path.name):
                    original = input_path.read_bytes()
                    paths = create_batch_paths(root, run_name=f"reference-{index}")
                    with patch(
                        "sd_institutional_skill.build_intake",
                        return_value=SimpleNamespace(all_rows=[intake_row]),
                    ):
                        rows = normalize_input(
                            input_text=None,
                            input_path=input_path,
                            paths=paths,
                            options=BatchOptions(),
                        )

                    self.assertEqual(len(rows), 1)
                    self.assertEqual(rows[0]["doi"], "")
                    self.assertEqual(rows[0]["status"], "metadata_uncertain")
                    self.assertEqual(rows[0]["reason"], "doi_not_from_doi_column")
                    self.assertEqual(input_path.read_bytes(), original)

    def test_tabular_title_only_metadata_resolution_can_supply_a_valid_doi(self) -> None:
        import csv
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions
        from paper_automation.batch_workflow import create_batch_paths, normalize_input
        from sd_institutional_skill import IntakeRow

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_path = root / "title-only.csv"
            with input_path.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                writer.writerow(["title", "authors"])
                writer.writerow(["Resolved target paper", "Researcher"])
            intake_row = IntakeRow(
                source=str(input_path),
                row_number=2,
                input_doi="",
                doi="10.5555/resolved-target",
                input_title="Resolved target paper",
                title="Resolved target paper",
                raw_value="Resolved target paper",
                status="valid",
                reason="",
            )
            paths = create_batch_paths(root, run_name="resolved-title")
            with patch(
                "sd_institutional_skill.build_intake",
                return_value=SimpleNamespace(all_rows=[intake_row]),
            ):
                rows = normalize_input(
                    input_text=None,
                    input_path=input_path,
                    paths=paths,
                    options=BatchOptions(),
                )

        self.assertEqual(rows[0]["doi"], "10.5555/resolved-target")
        self.assertEqual(rows[0]["status"], "pending")

    def test_normalize_marks_only_valid_doi_intake_rows_pending(self) -> None:
        from types import SimpleNamespace

        from paper_automation.batch_stages import BatchOptions
        from paper_automation.batch_workflow import create_batch_paths, normalize_input
        from sd_institutional_skill import IntakeRow

        intake_rows = [
            IntakeRow(
                source="fixture",
                row_number=1,
                input_doi="10.1000/low",
                doi="10.1000/low",
                title="Low confidence title",
                raw_value="10.1000/low",
                status="needs_review",
                reason="metadata_confidence_below_threshold:0.400",
            ),
            IntakeRow(
                source="fixture",
                row_number=2,
                input_doi="10.1000/invalid",
                doi="10.1000/invalid",
                title="Invalid intake row",
                raw_value="10.1000/invalid",
                status="invalid",
                reason="invalid_input",
            ),
            IntakeRow(
                source="fixture",
                row_number=3,
                input_doi="",
                doi="10.1000/resolved",
                input_title="Resolved title only paper",
                title="Resolved title only paper",
                raw_value="Resolved title only paper",
                status="valid",
                reason="",
            ),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            paths = create_batch_paths(Path(tmp))
            with patch(
                "sd_institutional_skill.build_intake",
                return_value=SimpleNamespace(all_rows=intake_rows),
            ):
                rows = normalize_input(
                    input_text="fixture",
                    input_path=None,
                    paths=paths,
                    options=BatchOptions(),
                )

        self.assertEqual([row["status"] for row in rows], ["metadata_uncertain", "metadata_uncertain", "pending"])
        self.assertEqual(rows[0]["reason"], "metadata_confidence_below_threshold:0.400")
        self.assertEqual(rows[1]["reason"], "invalid_input")
        self.assertEqual(rows[2]["doi"], "10.1000/resolved")

    def test_real_low_confidence_title_is_metadata_uncertain_and_never_sent_to_gateway_stages(self) -> None:
        from paper_automation.batch_stages import BatchOptions
        from paper_automation.batch_workflow import DefaultStageGateway, create_batch_paths, normalize_input

        with tempfile.TemporaryDirectory() as tmp:
            paths = create_batch_paths(Path(tmp))
            rows = normalize_input(
                input_text="Titanium",
                input_path=None,
                paths=paths,
                options=BatchOptions(),
            )
            with patch("paper_automation.batch_workflow.run_oa_stage") as oa_stage, patch(
                "paper_automation.batch_workflow.run_sciencedirect_stage"
            ) as sd_stage, patch("paper_automation.batch_workflow.run_non_elsevier_stage") as other_stage:
                updates = DefaultStageGateway().run_initial(rows, paths, BatchOptions())

        self.assertEqual(rows[0]["status"], "metadata_uncertain")
        self.assertEqual(updates, [])
        oa_stage.assert_not_called()
        sd_stage.assert_not_called()
        other_stage.assert_not_called()

    def test_duplicate_rows_are_terminal_and_excluded_from_pending_files(self) -> None:
        import csv

        from paper_automation.batch_workflow import start_batch

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            gateway = FakeBatchGateway(initial_updates=lambda rows: [
                {**rows[0], "status": "duplicate", "source": "oa", "file": "", "reason": "duplicate_input"},
            ])
            result = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=root,
                gateway=gateway,
                normalizer=lambda **_kwargs: [{"task_id": "paper-0001", "doi": "10.1000/a", "title": "A"}],
            )
            for path in (result.paths.manual_retry, result.paths.zotero_fallback):
                with path.open("r", encoding="utf-8-sig") as handle:
                    self.assertEqual(list(csv.DictReader(handle)), [])

    def test_default_gateway_routes_doi_url_to_sciencedirect_after_oa_failure(self) -> None:
        from paper_automation.batch_stages import BatchOptions, StageResult
        from paper_automation.batch_workflow import DefaultStageGateway, create_batch_paths

        with tempfile.TemporaryDirectory() as tmp:
            paths = create_batch_paths(Path(tmp))
            rows = [{"task_id": "paper-0001", "doi": "https://doi.org/10.1016/j.actamat.1", "title": "A", "status": "pending"}]
            oa_result = [StageResult("paper-0001", "10.1016/j.actamat.1", "A", "no_open_pdf", "", "no_open_pdf", "oa")]
            sd_result = [StageResult("paper-0001", "10.1016/j.actamat.1", "A", "failed", "", "no_entitlement", "sciencedirect")]
            with patch("paper_automation.batch_workflow.run_oa_stage", return_value=oa_result), patch(
                "paper_automation.batch_workflow.run_sciencedirect_stage", return_value=sd_result
            ) as sciencedirect, patch("paper_automation.batch_workflow.run_non_elsevier_stage") as non_elsevier:
                updates = DefaultStageGateway().run_initial(rows, paths, BatchOptions())

        self.assertEqual([update["task_id"] for update in updates], ["paper-0001"])
        sciencedirect.assert_called_once()
        non_elsevier.assert_not_called()

    def test_default_retry_gateway_processes_only_explicit_manual_retry_rows(self) -> None:
        from paper_automation.batch_stages import BatchOptions
        from paper_automation.batch_workflow import DefaultStageGateway, create_batch_paths

        with tempfile.TemporaryDirectory() as tmp:
            paths = create_batch_paths(Path(tmp))
            rows = [
                {"task_id": "paper-0001", "doi": "10.1016/a", "title": "A", "status": "auth_required", "reason": "auth_required"},
                {"task_id": "paper-0002", "doi": "10.1038/b", "title": "B", "status": "no_open_pdf", "reason": "no_open_pdf"},
            ]
            with patch("paper_automation.batch_workflow.run_sciencedirect_stage", return_value=[]) as sciencedirect, patch(
                "paper_automation.batch_workflow.run_non_elsevier_stage", return_value=[]
            ) as non_elsevier:
                DefaultStageGateway().run_retry(rows, paths, BatchOptions())

        sciencedirect.assert_called_once()
        non_elsevier.assert_not_called()

    def test_successful_absolute_adapter_pdf_is_copied_without_moving_source(self) -> None:
        from paper_automation.batch_workflow import start_batch

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "absolute.pdf"
            source.write_bytes(b"%PDF-1.7\nabsolute fixture")
            gateway = FakeBatchGateway(initial_updates=lambda rows: [
                {**rows[0], "status": "downloaded", "source": "oa", "file": str(source), "reason": ""},
            ])
            result = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=root,
                gateway=gateway,
                normalizer=lambda **_kwargs: [{"task_id": "paper-0001", "doi": "10.1000/a", "title": "A"}],
            )

            self.assertTrue(source.is_file())
            self.assertEqual(source.read_bytes(), b"%PDF-1.7\nabsolute fixture")
            self.assertEqual(Path(result.paths.root / "pdfs" / Path(result.paths.pdfs.glob("*.pdf").__next__().name)).read_bytes(), source.read_bytes())

    def test_successful_files_use_task5_compatible_canonical_statuses(self) -> None:
        from paper_automation.batch_workflow import load_batch_state, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            oa_pdf = root / "oa.pdf"
            institutional_pdf = root / "institutional.pdf"
            oa_pdf.write_bytes(b"%PDF-1.7\noa fixture")
            institutional_pdf.write_bytes(b"%PDF-1.7\ninstitutional fixture")
            gateway = FakeBatchGateway(initial_updates=lambda rows: [
                {**rows[0], "status": "downloaded", "source": "oa", "file": str(oa_pdf), "reason": ""},
                {**rows[1], "status": "downloaded", "source": "sciencedirect", "file": str(institutional_pdf), "reason": ""},
            ])
            result = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=root,
                gateway=gateway,
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                    {"task_id": "paper-0002", "doi": "10.1016/b", "title": "B", "status": "pending"},
                ],
            )
            state = load_batch_state(result.paths.root)

        self.assertEqual([row["status"] for row in state["rows"]], ["oa_downloaded", "institutional_downloaded"])
        self.assertNotIn("downloaded", {row["status"] for row in state["rows"]})
        self.assertEqual(result.success_count, 2)

    def test_unknown_download_source_and_invalid_pdf_are_isolated_as_row_failures(self) -> None:
        import csv

        from paper_automation.batch_workflow import load_batch_state, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            valid_pdf = root / "valid.pdf"
            html_pdf = root / "login.pdf"
            valid_pdf.write_bytes(b"%PDF-1.7\nvalid fixture")
            html_pdf.write_text("<html>login</html>", encoding="utf-8")
            gateway = FakeBatchGateway(initial_updates=lambda rows: [
                {**rows[0], "status": "downloaded", "source": "mystery", "file": str(valid_pdf), "reason": ""},
                {**rows[1], "status": "downloaded", "source": "oa", "file": str(html_pdf), "reason": ""},
                {**rows[2], "status": "downloaded", "source": "institutional", "file": str(root / "missing.pdf"), "reason": ""},
            ])
            result = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=root,
                gateway=gateway,
                normalizer=lambda **_kwargs: [
                    {"task_id": f"paper-{index:04d}", "doi": f"10.1000/{index}", "title": str(index), "status": "pending"}
                    for index in range(1, 4)
                ],
            )
            state = load_batch_state(result.paths.root)
            with result.paths.zotero_fallback.open("r", encoding="utf-8-sig") as handle:
                fallback = list(csv.DictReader(handle))

        self.assertEqual(
            [row["status"] for row in state["rows"]],
            ["invalid_download_source", "not_pdf_response", "not_pdf_response"],
        )
        self.assertTrue(all(row["file"] == "" for row in state["rows"]))
        self.assertEqual([row["task_id"] for row in fallback], ["paper-0001", "paper-0002", "paper-0003"])

    def test_success_file_symlink_is_a_row_failure_without_copying(self) -> None:
        from paper_automation.batch_workflow import load_batch_state, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "linked.pdf"
            source.write_bytes(b"%PDF-1.7\nsymlink fixture")
            gateway = FakeBatchGateway(initial_updates=lambda rows: [
                {**rows[0], "status": "downloaded", "source": "oa", "file": str(source), "reason": ""},
            ])
            original_is_symlink = Path.is_symlink

            def fake_is_symlink(path):
                return path == source or original_is_symlink(path)

            with patch.object(Path, "is_symlink", new=fake_is_symlink):
                result = start_batch(
                    input_text="fixture",
                    input_path=None,
                    output_root=root,
                    gateway=gateway,
                    normalizer=lambda **_kwargs: [
                        {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                    ],
                )
            state = load_batch_state(result.paths.root)

        self.assertEqual(state["rows"][0]["status"], "not_pdf_response")
        self.assertEqual(list(result.paths.pdfs.glob("*.pdf")), [])

    def test_empty_gateway_updates_mark_only_required_rows_missing(self) -> None:
        import csv

        from paper_automation.batch_workflow import load_batch_state, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            result = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=Path(tmp),
                gateway=FakeBatchGateway(initial_updates=lambda rows: []),
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                    {"task_id": "paper-0002", "doi": "", "title": "Low confidence", "status": "metadata_uncertain", "reason": "low_confidence"},
                    {"task_id": "paper-0003", "doi": "10.1000/c", "title": "C", "status": "duplicate", "reason": "duplicate_input"},
                ],
            )
            state = load_batch_state(result.paths.root)
            with result.paths.zotero_fallback.open("r", encoding="utf-8-sig") as handle:
                fallback = list(csv.DictReader(handle))

        self.assertEqual(
            [row["status"] for row in state["rows"]],
            ["missing_stage_update", "metadata_uncertain", "duplicate"],
        )
        self.assertEqual([row["task_id"] for row in fallback], ["paper-0001", "paper-0002"])
        self.assertEqual(result.zotero_fallback_count, 2)

    def test_partial_gateway_updates_mark_each_omitted_required_row_missing(self) -> None:
        from paper_automation.batch_workflow import load_batch_state, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            result = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=Path(tmp),
                gateway=FakeBatchGateway(initial_updates=lambda rows: [
                    {**rows[0], "status": "no_open_pdf", "source": "oa", "file": "", "reason": "no_open_pdf"},
                ]),
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                    {"task_id": "paper-0002", "doi": "10.1000/b", "title": "B", "status": "pending"},
                ],
            )
            state = load_batch_state(result.paths.root)

        self.assertEqual([row["status"] for row in state["rows"]], ["no_open_pdf", "missing_stage_update"])

    def test_start_gateway_exception_becomes_durable_row_failure_and_reports(self) -> None:
        from paper_automation.batch_workflow import load_batch_state, start_batch

        class RaisingGateway(FakeBatchGateway):
            def run_initial(self, rows, paths, options):
                raise RuntimeError("offline gateway failure")

        with tempfile.TemporaryDirectory() as tmp:
            result = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=Path(tmp),
                gateway=RaisingGateway(),
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                    {"task_id": "paper-0002", "doi": "", "title": "Review", "status": "metadata_uncertain", "reason": "review"},
                ],
            )
            state = load_batch_state(result.paths.root)

            self.assertTrue(result.paths.zotero_fallback.is_file())
            self.assertTrue((result.paths.reports / "batch_status.json").is_file())

        self.assertEqual(state["rows"][0]["status"], "gateway_exception_RuntimeError")
        self.assertEqual(state["rows"][1]["status"], "metadata_uncertain")

    def test_resume_gateway_exception_enters_fallback_after_one_claim(self) -> None:
        from paper_automation.batch_workflow import load_batch_state, resume_batch, start_batch

        class RaisingRetryGateway(FakeBatchGateway):
            def run_retry(self, rows, paths, options):
                raise ConnectionError("offline retry failure")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            started = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=root,
                gateway=FakeBatchGateway(initial_updates=lambda rows: [
                    {**rows[0], "status": "captcha_required", "source": "institutional", "file": "", "reason": "captcha_required"},
                ]),
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                ],
            )
            resumed = resume_batch(started.paths.root, gateway=RaisingRetryGateway())
            state = load_batch_state(started.paths.root)

        self.assertEqual(state["rows"][0]["status"], "gateway_exception_ConnectionError")
        self.assertTrue(state["manual_retry_used"])
        self.assertEqual(resumed.zotero_fallback_count, 1)
        self.assertEqual(resumed.manual_retry_count, 0)

    def test_oa_callback_state_survives_later_institutional_exception(self) -> None:
        from paper_automation.batch_stages import StageResult
        from paper_automation.batch_workflow import DefaultStageGateway, load_batch_state, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            oa_pdf = root / "oa.pdf"
            oa_pdf.write_bytes(b"%PDF-1.7\noa callback fixture")
            oa_results = [
                StageResult("paper-0001", "10.1000/a", "A", "downloaded", str(oa_pdf), "", "oa"),
                StageResult("paper-0002", "10.1038/b", "B", "no_open_pdf", "", "no_open_pdf", "oa"),
            ]
            with patch("paper_automation.batch_workflow.run_oa_stage", return_value=oa_results), patch(
                "paper_automation.batch_workflow.run_non_elsevier_stage",
                side_effect=RuntimeError("institutional stage stopped"),
            ):
                result = start_batch(
                    input_text="fixture",
                    input_path=None,
                    output_root=root,
                    gateway=DefaultStageGateway(),
                    normalizer=lambda **_kwargs: [
                        {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                        {"task_id": "paper-0002", "doi": "10.1038/b", "title": "B", "status": "pending"},
                    ],
                )
            state = load_batch_state(result.paths.root)

        self.assertEqual(state["rows"][0]["status"], "oa_downloaded")
        self.assertEqual(state["rows"][1]["status"], "gateway_exception_RuntimeError")

    def test_callback_success_is_not_rolled_back_by_stale_final_return(self) -> None:
        from paper_automation.batch_workflow import load_batch_state, start_batch

        class CallbackThenStaleGateway:
            def run_initial(self, rows, paths, options, *, on_updates):
                on_updates([
                    {**rows[0], "status": "oa_downloaded", "source": "oa", "file": rows[0]["fixture_pdf"], "reason": ""},
                ])
                return [
                    {**rows[0], "status": "no_open_pdf", "source": "oa", "file": "", "reason": "stale_final_return"},
                ]

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "oa.pdf"
            source.write_bytes(b"%PDF-1.7\nidempotent callback fixture")
            result = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=root,
                gateway=CallbackThenStaleGateway(),
                normalizer=lambda **_kwargs: [
                    {
                        "task_id": "paper-0001",
                        "doi": "10.1000/a",
                        "title": "A",
                        "status": "pending",
                        "fixture_pdf": str(source),
                    },
                ],
            )
            state = load_batch_state(result.paths.root)

        self.assertEqual(state["rows"][0]["status"], "oa_downloaded")
        self.assertEqual(state["rows"][0]["reason"], "")

    def test_callback_only_updates_are_not_replaced_by_missing_stage_update(self) -> None:
        from paper_automation.batch_workflow import load_batch_state, start_batch

        class CallbackOnlyGateway:
            def run_initial(self, rows, paths, options, *, on_updates):
                on_updates([
                    {**rows[0], "status": "captcha_required", "source": "institutional", "file": "", "reason": "captcha_required"},
                    {**rows[1], "status": "no_open_pdf", "source": "oa", "file": "", "reason": "no_open_pdf"},
                    {**rows[2], "status": "institutional_downloaded", "source": "institutional", "file": rows[2]["fixture_pdf"], "reason": ""},
                ])
                return []

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pdf = root / "institutional.pdf"
            pdf.write_bytes(b"%PDF-1.7\ncallback-only fixture")
            result = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=root,
                gateway=CallbackOnlyGateway(),
                normalizer=lambda **_kwargs: [
                    {"task_id": "paper-0001", "doi": "10.1000/a", "title": "A", "status": "pending"},
                    {"task_id": "paper-0002", "doi": "10.1000/b", "title": "B", "status": "pending"},
                    {"task_id": "paper-0003", "doi": "10.1000/c", "title": "C", "status": "pending", "fixture_pdf": str(pdf)},
                ],
            )
            state = load_batch_state(result.paths.root)

        self.assertEqual(
            [row["status"] for row in state["rows"]],
            ["captcha_required", "no_open_pdf", "institutional_downloaded"],
        )

    def test_unknown_or_duplicate_gateway_task_ids_are_rejected(self) -> None:
        from paper_automation.batch_workflow import start_batch

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for updates in (
                lambda rows: [{**rows[0], "task_id": "unknown", "status": "failed", "source": "oa", "file": "", "reason": "x"}],
                lambda rows: [{**rows[0], "status": "failed", "source": "oa", "file": "", "reason": "x"}] * 2,
            ):
                with self.subTest(updates=updates):
                    with self.assertRaisesRegex(ValueError, "gateway_task_id"):
                        start_batch(
                            input_text="fixture",
                            input_path=None,
                            output_root=root,
                            gateway=FakeBatchGateway(initial_updates=updates),
                        normalizer=lambda **_kwargs: [{"task_id": "paper-0001", "doi": "10.1000/a", "title": "A"}],
                    )

    def test_start_writes_fixed_normalized_schema_for_injected_normalizer(self) -> None:
        import csv

        from paper_automation.batch_workflow import NORMALIZED_FIELDS, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            result = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=Path(tmp),
                gateway=FakeBatchGateway(),
                normalizer=lambda **_kwargs: [{
                    "task_id": "paper-0001",
                    "source_index": "7",
                    "input_doi": "https://doi.org/10.1000/a",
                    "input_title": "Original A",
                    "doi": "10.1000/a",
                    "title": "A",
                }],
            )
            with result.paths.normalized_input.open("r", newline="", encoding="utf-8-sig") as handle:
                reader = csv.DictReader(handle)
                rows = list(reader)

        self.assertEqual(reader.fieldnames, NORMALIZED_FIELDS)
        self.assertEqual(rows[0]["input_doi"], "https://doi.org/10.1000/a")
        self.assertEqual(rows[0]["input_title"], "Original A")

    def test_concurrent_resume_claims_retry_only_once(self) -> None:
        import threading

        from paper_automation.batch_workflow import resume_batch, start_batch

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            entered = threading.Event()
            release = threading.Event()

            class BlockingGateway(FakeBatchGateway):
                def run_retry(self, rows, paths, options):
                    self.retry_rows.append([dict(row) for row in rows])
                    entered.set()
                    if not release.wait(5.0):
                        raise TimeoutError("test_retry_release_timeout")
                    return [{**row, "status": "no_entitlement", "source": "institutional", "file": "", "reason": "retry_exhausted"} for row in rows]

            gateway = BlockingGateway(initial_updates=lambda rows: [
                {**rows[0], "status": "auth_required", "source": "institutional", "file": "", "reason": "auth_required"},
            ])
            started = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=root,
                gateway=gateway,
                normalizer=lambda **_kwargs: [{"task_id": "paper-0001", "doi": "10.1000/a", "title": "A"}],
            )
            workers = [threading.Thread(target=resume_batch, args=(started.paths.root,), kwargs={"gateway": gateway}) for _ in range(2)]
            workers[0].start()
            self.assertTrue(entered.wait(5.0))
            workers[1].start()
            workers[1].join(5.0)
            release.set()
            workers[0].join(5.0)

        self.assertEqual(len(gateway.retry_rows), 1)
        self.assertEqual(len(gateway.retry_rows[0]), 1)


class BatchFinalizeTests(unittest.TestCase):
    def _paths_with_state(self, root: Path, rows: list[dict]):
        from paper_automation.batch_workflow import create_batch_paths, save_batch_state

        paths = create_batch_paths(root, now=datetime(2026, 7, 11, 9, 0, 0))
        save_batch_state(
            paths,
            {
                "version": 1,
                "run_dir": str(paths.root),
                "manual_retry_used": True,
                "options": {
                    "email": "",
                    "cookies": "",
                    "browser_exe": "",
                    "login_wait_seconds": 0,
                    "debug_port": 9222,
                    "throttle_seconds": 0.0,
                },
                "rows": rows,
            },
        )
        return paths

    @staticmethod
    def _row(task_id: str, status: str = "no_open_pdf", **extra) -> dict:
        row = {
            "task_id": task_id,
            "source_index": task_id[-1:],
            "input_doi": f"10.1000/{task_id}",
            "input_title": f"Input {task_id}",
            "doi": f"10.1000/{task_id}",
            "title": f"Title {task_id}",
            "authors": "Reader Example",
            "journal": "Journal",
            "year": "2026",
            "publisher": "Publisher",
            "status": status,
            "source": "oa",
            "file": "",
            "reason": "initial_failure",
        }
        row.update(extra)
        return row

    @staticmethod
    def _write_zotero_csv(path: Path, rows: list[dict], fieldnames=None) -> None:
        import csv

        columns = fieldnames or [
            "task_id",
            "zotero_item_id",
            "attachment_path",
            "status",
            "reason",
        ]
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns)
            writer.writeheader()
            writer.writerows(rows)

    def test_finalize_reconciles_safely_and_writes_consistent_reports(self) -> None:
        import csv

        from openpyxl import load_workbook

        from paper_automation.batch_workflow import FINAL_MANIFEST_FIELDS, finalize_batch, load_batch_state

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            attachment = root / "zotero.pdf"
            attachment.write_bytes(b"%PDF-1.7\nZotero fixture")
            attachment_hash = hashlib.sha256(attachment.read_bytes()).hexdigest()
            html = root / "login.html"
            html.write_text("<html>login</html>", encoding="utf-8")
            original = self._row("paper-0001", "oa_downloaded", source="oa", file="already.pdf", reason="")
            paths = self._paths_with_state(
                root,
                [original, self._row("paper-0002"), self._row("paper-0003")],
            )
            zotero_csv = root / "zotero_results.csv"
            self._write_zotero_csv(
                zotero_csv,
                [
                    {"task_id": "paper-0001", "zotero_item_id": "ignored", "attachment_path": str(attachment), "status": "downloaded", "reason": ""},
                    {"task_id": "paper-0002", "zotero_item_id": "42", "attachment_path": str(attachment), "status": "existing_pdf", "reason": ""},
                    {"task_id": "paper-0003", "zotero_item_id": "43", "attachment_path": str(html), "status": "downloaded", "reason": ""},
                ],
            )

            result = finalize_batch(paths.root, zotero_csv)
            state = load_batch_state(paths.root)
            rows_by_id = {row["task_id"]: row for row in state["rows"]}
            with (paths.reports / "final_manifest.csv").open("r", newline="", encoding="utf-8-sig") as handle:
                csv_reader = csv.DictReader(handle)
                csv_rows = list(csv_reader)
            workbook = load_workbook(paths.reports / "final_manifest.xlsx", read_only=True, data_only=False)
            worksheet = workbook.active
            xlsx_rows = list(worksheet.iter_rows(values_only=True))
            workbook.close()
            attachment_intact = (
                attachment.exists()
                and hashlib.sha256(attachment.read_bytes()).hexdigest() == attachment_hash
            )

        self.assertEqual(rows_by_id["paper-0001"], original)
        self.assertEqual(rows_by_id["paper-0002"]["status"], "zotero_existing_pdf")
        self.assertEqual(rows_by_id["paper-0002"]["zotero_item_id"], "42")
        self.assertEqual(rows_by_id["paper-0003"]["status"], "not_pdf_response")
        self.assertEqual(rows_by_id["paper-0003"]["source"], "zotero")
        self.assertTrue(attachment_intact)
        self.assertEqual(result.success_count, 2)
        self.assertEqual(result.failed_count, 1)
        self.assertEqual(csv_reader.fieldnames, FINAL_MANIFEST_FIELDS)
        self.assertEqual(xlsx_rows[0], tuple(FINAL_MANIFEST_FIELDS))
        self.assertEqual(
            [[str(value or "") for value in row] for row in xlsx_rows[1:]],
            [[row[field] for field in FINAL_MANIFEST_FIELDS] for row in csv_rows],
        )

    def test_finalize_rejects_invalid_zotero_csv_without_state_changes(self) -> None:
        from paper_automation.batch_workflow import finalize_batch

        cases = [
            ("missing_header", ["task_id", "status"], [{"task_id": "paper-0001", "status": "not_found"}], "zotero_results_fields_invalid"),
            ("unknown", None, [{"task_id": "paper-9999", "zotero_item_id": "", "attachment_path": "", "status": "not_found", "reason": "x"}], "zotero_result_task_id_unknown"),
            ("duplicate", None, [{"task_id": "paper-0001", "zotero_item_id": "", "attachment_path": "", "status": "not_found", "reason": "x"}] * 2, "zotero_result_task_id_duplicate"),
            ("empty", None, [{"task_id": "", "zotero_item_id": "", "attachment_path": "", "status": "not_found", "reason": "x"}], "zotero_result_task_id_missing"),
            ("duplicate_header", ["task_id", "zotero_item_id", "attachment_path", "status", "reason", "reason"], [{"task_id": "paper-0001", "zotero_item_id": "", "attachment_path": "", "status": "not_found", "reason": "x"}], "zotero_results_fields_invalid"),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for name, fieldnames, rows, error in cases:
                with self.subTest(name=name):
                    paths = self._paths_with_state(root / name, [self._row("paper-0001")])
                    zotero_csv = paths.working / "zotero_results.csv"
                    self._write_zotero_csv(zotero_csv, rows, fieldnames)
                    before = paths.state.read_bytes()
                    with self.assertRaisesRegex(ValueError, error):
                        finalize_batch(paths.root, zotero_csv)
                    self.assertEqual(paths.state.read_bytes(), before)

    def test_finalize_rejects_unsafe_attachments_and_deduplicates_repeated_pdf(self) -> None:
        from paper_automation.batch_workflow import finalize_batch, load_batch_state

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            valid_pdf = root / "attachment.pdf"
            valid_pdf.write_bytes(b"%PDF-1.7\nshared attachment")
            missing = root / "missing.pdf"
            directory = root / "attachment-directory"
            directory.mkdir()
            html = root / "attachment.html"
            html.write_text("<html>not a PDF</html>", encoding="utf-8")
            paths = self._paths_with_state(
                root,
                [self._row(f"paper-000{index}") for index in range(1, 7)],
            )
            zotero_csv = paths.working / "zotero_results.csv"
            self._write_zotero_csv(
                zotero_csv,
                [
                    {"task_id": "paper-0001", "zotero_item_id": "1", "attachment_path": str(valid_pdf), "status": "downloaded", "reason": ""},
                    {"task_id": "paper-0002", "zotero_item_id": "2", "attachment_path": str(valid_pdf), "status": "existing_pdf", "reason": ""},
                    {"task_id": "paper-0003", "zotero_item_id": "3", "attachment_path": "https://example.invalid/paper.pdf", "status": "downloaded", "reason": ""},
                    {"task_id": "paper-0004", "zotero_item_id": "4", "attachment_path": str(html), "status": "downloaded", "reason": ""},
                    {"task_id": "paper-0005", "zotero_item_id": "5", "attachment_path": str(missing), "status": "downloaded", "reason": ""},
                    {"task_id": "paper-0006", "zotero_item_id": "6", "attachment_path": str(directory), "status": "downloaded", "reason": ""},
                ],
            )
            finalize_batch(paths.root, zotero_csv)
            state = load_batch_state(paths.root)
            pdf_count = len(list(paths.pdfs.glob("*.pdf")))

        rows_by_id = {row["task_id"]: row for row in state["rows"]}
        self.assertEqual(rows_by_id["paper-0001"]["file"], rows_by_id["paper-0002"]["file"])
        self.assertEqual(pdf_count, 1)
        for task_id in ("paper-0003", "paper-0004", "paper-0005", "paper-0006"):
            self.assertEqual(rows_by_id[task_id]["status"], "not_pdf_response")
            self.assertEqual(rows_by_id[task_id]["file"], "")

    def test_finalize_preserves_metadata_uncertain_audit_and_non_success_results(self) -> None:
        from paper_automation.batch_workflow import finalize_batch, load_batch_state

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            attachment = root / "uncertain.pdf"
            attachment.write_bytes(b"%PDF-1.7\nuncertain")
            paths = self._paths_with_state(
                root,
                [
                    self._row("paper-0001", "metadata_uncertain", reason="title_low_confidence"),
                    self._row("paper-0002"),
                    self._row("paper-0003"),
                ],
            )
            zotero_csv = paths.working / "zotero_results.csv"
            self._write_zotero_csv(
                zotero_csv,
                [
                    {"task_id": "paper-0001", "zotero_item_id": "42", "attachment_path": str(attachment), "status": "existing_pdf", "reason": ""},
                    {"task_id": "paper-0002", "zotero_item_id": "", "attachment_path": "", "status": "not_found", "reason": "not in collection"},
                ],
            )
            result = finalize_batch(paths.root, zotero_csv)
            state = load_batch_state(paths.root)

        rows_by_id = {row["task_id"]: row for row in state["rows"]}
        self.assertEqual(rows_by_id["paper-0001"]["status"], "zotero_existing_pdf")
        self.assertIn("metadata_uncertain", rows_by_id["paper-0001"]["reason"])
        self.assertEqual(rows_by_id["paper-0002"]["status"], "not_found")
        self.assertEqual(rows_by_id["paper-0002"]["source"], "zotero")
        self.assertEqual(rows_by_id["paper-0002"]["file"], "")
        self.assertEqual(rows_by_id["paper-0003"]["status"], "no_open_pdf")
        self.assertEqual(result.total_count, 3)
        self.assertEqual(result.success_count, 1)
        self.assertEqual(result.failed_count, 2)

    def test_final_reports_escape_formula_text_and_summarize_failures(self) -> None:
        import csv

        from openpyxl import load_workbook

        from paper_automation.batch_workflow import write_final_reports

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            rows = [
                self._row("paper-0001", "zotero_downloaded", title="=DANGEROUS", reason=""),
                self._row("paper-0002", "no_open_pdf", reason="=FORMULA"),
                self._row("paper-0003", "duplicate", reason="duplicate_input"),
            ]
            paths = self._paths_with_state(root, rows)
            write_final_reports(paths, rows)
            with (paths.reports / "failed.csv").open("r", newline="", encoding="utf-8-sig") as handle:
                failed_rows = list(csv.DictReader(handle))
            workbook = load_workbook(paths.reports / "final_manifest.xlsx", read_only=True, data_only=False)
            values = list(workbook.active.iter_rows(values_only=False))
            formula_cell = values[1][5]
            workbook.close()
            summary = (paths.reports / "run_summary.txt").read_text(encoding="utf-8")

        self.assertEqual([row["task_id"] for row in failed_rows], ["paper-0002"])
        self.assertEqual(formula_cell.value, "=DANGEROUS")
        self.assertEqual(formula_cell.data_type, "s")
        self.assertIn("input_count: 3", summary)
        self.assertIn("success_count: 1", summary)
        self.assertIn("failure_count: 1", summary)
        self.assertIn("no_open_pdf: 1", summary)
        self.assertIn("paper-0002\tno_open_pdf\t=FORMULA", summary)
        self.assertIn("duplicate_terminal_rows_excluded: 1", summary)

    def test_finalize_rejects_symlink_attachment_and_empty_results_leave_state_pending(self) -> None:
        from paper_automation import batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            attachment = root / "attachment.pdf"
            attachment.write_bytes(b"%PDF-1.7\nfixture")
            paths = self._paths_with_state(root, [self._row("paper-0001"), self._row("paper-0002")])
            zotero_csv = paths.working / "zotero_results.csv"
            self._write_zotero_csv(
                zotero_csv,
                [{"task_id": "paper-0001", "zotero_item_id": "1", "attachment_path": str(attachment), "status": "downloaded", "reason": ""}],
            )
            original_is_symlink = workflow.Path.is_symlink

            def symlink_only_attachment(path):
                return path == attachment or original_is_symlink(path)

            with patch.object(workflow.Path, "is_symlink", new=symlink_only_attachment):
                workflow.finalize_batch(paths.root, zotero_csv)
            state_after_symlink = workflow.load_batch_state(paths.root)
            self._write_zotero_csv(zotero_csv, [])
            result = workflow.finalize_batch(paths.root, zotero_csv)
            state_after_empty = workflow.load_batch_state(paths.root)

        self.assertEqual(state_after_symlink["rows"][0]["status"], "not_pdf_response")
        self.assertEqual(state_after_symlink["rows"][0]["file"], "")
        self.assertEqual(state_after_empty["rows"][1]["status"], "no_open_pdf")
        self.assertEqual(result.success_count, 0)
        self.assertEqual(result.failed_count, 2)

    def test_final_report_write_error_is_diagnostic(self) -> None:
        from paper_automation import batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            paths = self._paths_with_state(Path(tmp), [self._row("paper-0001")])
            with patch.object(workflow, "_write_final_manifest_xlsx", side_effect=OSError("disk full")):
                with self.assertRaisesRegex(RuntimeError, "final_report_write_failed:OSError:disk full"):
                    workflow.write_final_reports(paths, [self._row("paper-0001")])

    def test_zotero_status_whitelist_rejects_internal_and_unknown_before_state_change(self) -> None:
        from paper_automation.batch_workflow import finalize_batch

        invalid_statuses = [
            "oa_downloaded",
            "institutional_downloaded",
            "zotero_existing_pdf",
            "zotero_downloaded",
            "duplicate",
            "success",
            "anything_else",
        ]
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for index, status in enumerate(invalid_statuses):
                with self.subTest(status=status):
                    paths = self._paths_with_state(root / str(index), [self._row("paper-0001")])
                    zotero_csv = paths.working / "zotero_results.csv"
                    self._write_zotero_csv(
                        zotero_csv,
                        [{"task_id": "paper-0001", "zotero_item_id": "", "attachment_path": "", "status": status, "reason": "invalid"}],
                    )
                    before = paths.state.read_bytes()
                    with self.assertRaisesRegex(ValueError, "^zotero_result_status_invalid$"):
                        finalize_batch(paths.root, zotero_csv)
                    self.assertEqual(paths.state.read_bytes(), before)

    def test_zotero_failure_status_whitelist_is_preserved(self) -> None:
        from paper_automation.batch_workflow import finalize_batch, load_batch_state

        statuses = [
            "no_pdf",
            "not_found",
            "metadata_uncertain",
            "zotero_unavailable",
            "no_attachment",
            "download_failed",
            "zotero_api_unavailable",
            "user_cancelled",
            "job_expired",
            "job_id_conflict",
            "plugin_error",
        ]
        with tempfile.TemporaryDirectory() as tmp:
            paths = self._paths_with_state(
                Path(tmp),
                [self._row(f"paper-{index:04d}") for index in range(1, len(statuses) + 1)],
            )
            zotero_csv = paths.working / "zotero_results.csv"
            self._write_zotero_csv(
                zotero_csv,
                [
                    {
                        "task_id": f"paper-{index:04d}",
                        "zotero_item_id": "",
                        "attachment_path": "",
                        "status": status,
                        "reason": f"reason-{status}",
                    }
                    for index, status in enumerate(statuses, start=1)
                ],
            )
            finalize_batch(paths.root, zotero_csv)
            state = load_batch_state(paths.root)

        self.assertEqual([row["status"] for row in state["rows"]], statuses)
        self.assertTrue(all(row["source"] == "zotero" for row in state["rows"]))

    def test_zotero_csv_requires_exact_order_and_rejects_blank_or_malformed_rows_atomically(self) -> None:
        from paper_automation.batch_workflow import ZOTERO_RESULT_FIELDS, finalize_batch

        exact = list(ZOTERO_RESULT_FIELDS)
        cases = [
            ("extra_header", [*exact, "extra"], [{"task_id": "paper-0001", "status": "not_found", "reason": "x", "extra": "x"}], None, "zotero_results_fields_invalid"),
            ("reordered_header", [exact[1], exact[0], *exact[2:]], [{"task_id": "paper-0001", "status": "not_found", "reason": "x"}], None, "zotero_results_fields_invalid"),
            ("blank_line", None, None, ",".join(exact) + "\n\n", "zotero_results_row_invalid"),
            ("whitespace_row", None, None, ",".join(exact) + "\n , , , , \n", "zotero_results_row_invalid"),
            ("too_few_cells", None, None, ",".join(exact) + "\npaper-0001,,,not_found\n", "zotero_results_row_invalid"),
            ("too_many_cells", None, None, ",".join(exact) + "\npaper-0001,,,not_found,x,extra\n", "zotero_results_row_invalid"),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for name, fieldnames, rows, raw_csv, error in cases:
                with self.subTest(name=name):
                    paths = self._paths_with_state(root / name, [self._row("paper-0001")])
                    zotero_csv = paths.working / "zotero_results.csv"
                    if raw_csv is None:
                        self._write_zotero_csv(zotero_csv, rows, fieldnames)
                    else:
                        zotero_csv.write_text(raw_csv, encoding="utf-8-sig")
                    before = paths.state.read_bytes()
                    with self.assertRaisesRegex(ValueError, f"^{error}$"):
                        finalize_batch(paths.root, zotero_csv)
                    self.assertEqual(paths.state.read_bytes(), before)

    def test_xlsx_closes_then_fsyncs_before_replace_and_remains_readable(self) -> None:
        import openpyxl

        from paper_automation import batch_workflow as workflow

        events = []
        real_workbook = openpyxl.Workbook
        real_fsync = workflow.os.fsync
        real_replace = workflow.os.replace

        class TrackingWorkbook:
            def __init__(self, *args, **kwargs):
                self.inner = real_workbook(*args, **kwargs)

            def create_sheet(self, *args, **kwargs):
                return self.inner.create_sheet(*args, **kwargs)

            def save(self, path):
                events.append("save")
                return self.inner.save(path)

            def close(self):
                events.append("close")
                return self.inner.close()

        def tracking_fsync(descriptor):
            events.append("fsync")
            return real_fsync(descriptor)

        def tracking_replace(source, destination):
            events.append("replace")
            return real_replace(source, destination)

        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "final_manifest.xlsx"
            with patch.object(openpyxl, "Workbook", TrackingWorkbook), patch.object(
                workflow.os, "fsync", side_effect=tracking_fsync
            ), patch.object(workflow.os, "replace", side_effect=tracking_replace):
                workflow._write_final_manifest_xlsx(target, [self._row("paper-0001")])
            workbook = openpyxl.load_workbook(target, read_only=True)
            header = tuple(cell.value for cell in next(workbook.active.iter_rows()))
            workbook.close()

        self.assertEqual(events, ["save", "close", "fsync", "replace"])
        self.assertEqual(header, tuple(workflow.FINAL_MANIFEST_FIELDS))

    def test_xlsx_closes_workbook_when_save_raises(self) -> None:
        import openpyxl

        from paper_automation import batch_workflow as workflow

        events = []
        real_workbook = openpyxl.Workbook

        class FailingWorkbook:
            def __init__(self, *args, **kwargs):
                self.inner = real_workbook(*args, **kwargs)

            def create_sheet(self, *args, **kwargs):
                return self.inner.create_sheet(*args, **kwargs)

            def save(self, path):
                events.append("save")
                self.inner.save(path)
                raise OSError("synthetic save failure")

            def close(self):
                events.append("close")
                return self.inner.close()

        with tempfile.TemporaryDirectory() as tmp, patch.object(openpyxl, "Workbook", FailingWorkbook):
            with self.assertRaisesRegex(OSError, "synthetic save failure"):
                workflow._write_final_manifest_xlsx(Path(tmp) / "final_manifest.xlsx", [])

        self.assertEqual(events, ["save", "close"])

    def test_zotero_attachment_requires_absolute_path_and_accepts_unicode_spaces(self) -> None:
        from paper_automation import batch_workflow as workflow

        with self.assertRaisesRegex(ValueError, "^zotero_attachment_not_absolute$"):
            workflow._local_zotero_attachment("relative folder/paper.pdf")

        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "中文 附件" / "论文 文件.pdf"
            source.parent.mkdir()
            source.write_bytes(b"%PDF-1.7\nunicode path fixture")
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            resolved = workflow._local_zotero_attachment(str(source))
            source_unchanged = hashlib.sha256(source.read_bytes()).hexdigest() == before

        self.assertTrue(resolved.is_absolute())
        self.assertEqual(resolved, source.resolve())
        self.assertTrue(source_unchanged)

    @unittest.skipUnless(os.name == "nt", "Windows junction semantics")
    def test_zotero_attachment_rejects_real_junction_ancestor(self) -> None:
        import subprocess

        from paper_automation import batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target_dir = root / "real target"
            target_dir.mkdir()
            source = target_dir / "paper.pdf"
            source.write_bytes(b"%PDF-1.7\njunction fixture")
            junction_dir = root / "junction alias"
            created = subprocess.run(
                ["cmd.exe", "/d", "/c", "mklink", "/J", str(junction_dir), str(target_dir)],
                capture_output=True,
                text=True,
                check=False,
            )
            if created.returncode != 0:
                self.skipTest("current Windows environment cannot create a directory junction")
            try:
                with self.assertRaisesRegex(ValueError, "^zotero_attachment_reparse_point$"):
                    workflow._local_zotero_attachment(str(junction_dir / "paper.pdf"))
            finally:
                if junction_dir.exists():
                    junction_dir.rmdir()

    def test_report_csv_escapes_formula_like_data_only(self) -> None:
        import csv

        from paper_automation import batch_workflow as workflow

        values = ["=SUM(1,2)", "+cmd", "-1+2", "@lookup", "   =leading", "ordinary", "  ordinary", ""]
        expected = ["'=SUM(1,2)", "'+cmd", "'-1+2", "'@lookup", "'   =leading", "ordinary", "  ordinary", ""]
        rows = [{"value": value} for value in values]
        original_rows = [dict(row) for row in rows]
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "safe.csv"
            workflow._write_report_csv(target, ["value"], rows)
            with target.open("r", newline="", encoding="utf-8-sig") as handle:
                reader = csv.reader(handle)
                records = list(reader)

        self.assertEqual(records[0], ["value"])
        self.assertEqual([record[0] for record in records[1:]], expected)
        self.assertEqual(rows, original_rows)

    def test_finalize_reloads_locked_state_and_preserves_concurrent_project_success(self) -> None:
        import threading

        from paper_automation import batch_workflow as workflow

        parsed_old_view = threading.Event()
        release_finalize = threading.Event()
        finalize_errors = []
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            project_pdf = root / "project.pdf"
            project_pdf.write_bytes(b"%PDF-1.7\nproject success fixture")
            zotero_pdf = root / "zotero.pdf"
            zotero_pdf.write_bytes(b"%PDF-1.7\nzotero stale fixture")
            paths = self._paths_with_state(root, [self._row("paper-0001")])
            zotero_csv = paths.working / "zotero_results.csv"
            self._write_zotero_csv(
                zotero_csv,
                [{"task_id": "paper-0001", "zotero_item_id": "42", "attachment_path": str(zotero_pdf), "status": "downloaded", "reason": ""}],
            )
            real_reader = workflow._read_zotero_results
            reader_calls = 0

            def controlled_reader(path, rows):
                nonlocal reader_calls
                result = real_reader(path, rows)
                reader_calls += 1
                if reader_calls == 1:
                    parsed_old_view.set()
                    if not release_finalize.wait(10.0):
                        raise TimeoutError("test_finalize_release_timeout")
                return result

            def run_finalize():
                try:
                    workflow.finalize_batch(paths.root, zotero_csv)
                except BaseException as exc:
                    finalize_errors.append(exc)

            caller_state = workflow.load_batch_state(paths.root)
            with patch.object(workflow, "_read_zotero_results", side_effect=controlled_reader):
                worker = threading.Thread(target=run_finalize)
                worker.start()
                self.assertTrue(parsed_old_view.wait(10.0))
                workflow._apply_stage_updates(
                    caller_state,
                    [{**caller_state["rows"][0], "status": "oa_downloaded", "source": "oa", "file": str(project_pdf), "reason": ""}],
                    paths,
                )
                release_finalize.set()
                worker.join(10.0)
            self.assertFalse(worker.is_alive())
            final_state = workflow.load_batch_state(paths.root)
            copied_hash = hashlib.sha256(Path(final_state["rows"][0]["file"]).read_bytes()).hexdigest()
            project_hash = hashlib.sha256(project_pdf.read_bytes()).hexdigest()

        self.assertEqual(finalize_errors, [])
        self.assertEqual(reader_calls, 2)
        self.assertEqual(final_state["rows"][0]["status"], "oa_downloaded")
        self.assertEqual(final_state["rows"][0]["source"], "oa")
        self.assertEqual(caller_state["rows"][0]["status"], "oa_downloaded")
        self.assertEqual(copied_hash, project_hash)

    def test_latest_locked_state_wins_when_finalize_report_overlaps_stage_update(self) -> None:
        import csv
        import threading
        from contextlib import contextmanager

        from paper_automation import batch_workflow as workflow

        finalize_in_report = threading.Event()
        release_finalize = threading.Event()
        stage_attempting_lock = threading.Event()
        stage_finished = threading.Event()
        state_mutex = threading.Lock()
        finalize_report_holds_state_lock = []
        errors = []
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            project_pdf = root / "project-latest.pdf"
            project_pdf.write_bytes(b"%PDF-1.7\nlatest project success")
            paths = self._paths_with_state(root / "run", [self._row("paper-0001")])
            zotero_csv = paths.working / "zotero_results.csv"
            self._write_zotero_csv(
                zotero_csv,
                [{"task_id": "paper-0001", "zotero_item_id": "", "attachment_path": "", "status": "not_found", "reason": "stale zotero view"}],
            )
            real_writer = workflow.write_final_reports

            def controlled_writer(report_paths, rows):
                if threading.current_thread().name == "finalize-F1":
                    finalize_report_holds_state_lock.append(state_mutex.locked())
                    finalize_in_report.set()
                    if not release_finalize.wait(10.0):
                        raise TimeoutError("release_finalize_timeout")
                return real_writer(report_paths, rows)

            @contextmanager
            def tracking_state_lock(*args, **kwargs):
                if threading.current_thread().name == "stage-F2":
                    stage_attempting_lock.set()
                with state_mutex:
                    yield

            def run_finalize():
                try:
                    workflow.finalize_batch(paths.root, zotero_csv)
                except BaseException as exc:
                    errors.append(exc)

            def run_stage_update():
                try:
                    caller_state = workflow.load_batch_state(paths.root)
                    workflow._apply_stage_updates(
                        caller_state,
                        [{**caller_state["rows"][0], "status": "oa_downloaded", "source": "oa", "file": str(project_pdf), "reason": ""}],
                        paths,
                    )
                    workflow._write_latest_state_outputs(
                        paths,
                        caller_state,
                        pending_manual_retry_used=None,
                    )
                except BaseException as exc:
                    errors.append(exc)
                finally:
                    stage_finished.set()

            with patch.object(workflow, "write_final_reports", side_effect=controlled_writer), patch.object(
                workflow, "batch_state_lock", side_effect=tracking_state_lock
            ):
                finalizer = threading.Thread(target=run_finalize, name="finalize-F1")
                finalizer.start()
                self.assertTrue(finalize_in_report.wait(10.0))
                stage = threading.Thread(target=run_stage_update, name="stage-F2")
                stage.start()
                self.assertTrue(stage_attempting_lock.wait(10.0))
                if finalize_report_holds_state_lock == [True]:
                    release_finalize.set()
                    self.assertTrue(stage_finished.wait(10.0))
                else:
                    self.assertTrue(stage_finished.wait(10.0))
                    release_finalize.set()
                finalizer.join(10.0)
                stage.join(10.0)

            self.assertFalse(finalizer.is_alive())
            self.assertFalse(stage.is_alive())
            disk_state = workflow.load_batch_state(paths.root)
            with (paths.reports / "final_manifest.csv").open(
                "r", newline="", encoding="utf-8-sig"
            ) as handle:
                manifest_rows = list(csv.DictReader(handle))
            with (paths.reports / "batch_status.csv").open(
                "r", newline="", encoding="utf-8-sig"
            ) as handle:
                status_rows = list(csv.DictReader(handle))

        self.assertEqual(errors, [])
        self.assertEqual(finalize_report_holds_state_lock, [True])
        self.assertEqual(disk_state["rows"][0]["status"], "oa_downloaded")
        expected_manifest = {
            field: str(disk_state["rows"][0].get(field, "") or "")
            for field in workflow.FINAL_MANIFEST_FIELDS
        }
        self.assertEqual(manifest_rows[0], expected_manifest)
        self.assertEqual(
            {field: manifest_rows[0][field] for field in workflow.NORMALIZED_FIELDS},
            status_rows[0],
        )

    def test_zotero_attachment_change_after_first_hash_is_rejected(self) -> None:
        from paper_automation import batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "attachment.pdf"
            source.write_bytes(b"%PDF-1.7\noriginal safe attachment")
            external_bytes = b"%PDF-1.7\nreplacement external attachment"
            paths = self._paths_with_state(root / "run", [self._row("paper-0001")])
            real_sha256 = workflow._sha256
            hash_calls = 0

            def replace_after_first_hash(path):
                nonlocal hash_calls
                digest = real_sha256(path)
                hash_calls += 1
                if hash_calls == 1:
                    source.write_bytes(external_bytes)
                return digest

            with patch.object(workflow, "_sha256", side_effect=replace_after_first_hash):
                with self.assertRaisesRegex(ValueError, "^zotero_attachment_changed$"):
                    workflow._copy_zotero_attachment(self._row("paper-0001"), str(source), paths)
            published = list(paths.pdfs.glob("*.pdf"))

        self.assertGreaterEqual(hash_calls, 2)
        self.assertEqual(published, [])

    def test_zotero_attachment_reparse_injected_after_first_hash_is_rejected(self) -> None:
        from paper_automation import batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "attachment.pdf"
            source.write_bytes(b"%PDF-1.7\noriginal safe attachment")
            paths = self._paths_with_state(root / "run", [self._row("paper-0001")])
            real_validator = workflow._local_zotero_attachment
            validation_calls = 0

            def reparse_on_second_validation(path):
                nonlocal validation_calls
                validation_calls += 1
                if validation_calls == 2:
                    raise ValueError("zotero_attachment_reparse_point")
                return real_validator(path)

            with patch.object(workflow, "_local_zotero_attachment", side_effect=reparse_on_second_validation):
                with self.assertRaisesRegex(ValueError, "^zotero_attachment_reparse_point$"):
                    workflow._copy_zotero_attachment(self._row("paper-0001"), str(source), paths)
            published = list(paths.pdfs.glob("*.pdf"))

        self.assertEqual(validation_calls, 2)
        self.assertEqual(published, [])

    def test_zotero_attachment_replaced_after_second_hash_is_not_published(self) -> None:
        from paper_automation import batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "attachment.pdf"
            original_bytes = b"%PDF-1.7\noriginal second-check attachment"
            replacement_bytes = b"%PDF-1.7\nreplacement after second check"
            source.write_bytes(original_bytes)
            paths = self._paths_with_state(root / "run", [self._row("paper-0001")])
            real_sha256 = workflow._sha256
            hash_calls = 0

            def replace_after_second_hash(path):
                nonlocal hash_calls
                digest = real_sha256(path)
                hash_calls += 1
                if hash_calls == 2:
                    source.write_bytes(replacement_bytes)
                return digest

            with patch.object(workflow, "_sha256", side_effect=replace_after_second_hash):
                with self.assertRaisesRegex(ValueError, "^zotero_attachment_changed$"):
                    workflow._copy_zotero_attachment(self._row("paper-0001"), str(source), paths)
            published_payloads = [candidate.read_bytes() for candidate in paths.pdfs.glob("*.pdf")]

        self.assertGreaterEqual(hash_calls, 2)
        self.assertNotIn(replacement_bytes, published_payloads)
        self.assertEqual(published_payloads, [])

    def test_zotero_attachment_changed_to_symlink_after_second_hash_is_rejected(self) -> None:
        from contextlib import nullcontext

        from paper_automation import batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "attachment.pdf"
            external = root / "external.pdf"
            source.write_bytes(b"%PDF-1.7\noriginal before symlink")
            external_bytes = b"%PDF-1.7\nexternal symlink target"
            external.write_bytes(external_bytes)
            probe = root / "symlink-probe.pdf"
            real_is_symlink = workflow.Path.is_symlink
            try:
                probe.symlink_to(external)
            except OSError:
                real_symlink_supported = False
            else:
                real_symlink_supported = True
                probe.unlink()
            paths = self._paths_with_state(root / "run", [self._row("paper-0001")])
            real_sha256 = workflow._sha256
            hash_calls = 0
            mocked_symlink_active = False

            def replace_with_symlink_after_second_hash(path):
                nonlocal hash_calls, mocked_symlink_active
                digest = real_sha256(path)
                hash_calls += 1
                if hash_calls == 2:
                    if real_symlink_supported:
                        source.unlink()
                        source.symlink_to(external)
                    else:
                        mocked_symlink_active = True
                return digest

            def injected_is_symlink(path):
                return (
                    mocked_symlink_active
                    and Path(path) == source
                ) or real_is_symlink(path)

            symlink_patch = (
                nullcontext()
                if real_symlink_supported
                else patch.object(workflow.Path, "is_symlink", new=injected_is_symlink)
            )
            with symlink_patch, patch.object(
                workflow, "_sha256", side_effect=replace_with_symlink_after_second_hash
            ):
                with self.assertRaisesRegex(
                    ValueError, "^(zotero_attachment_reparse_point|zotero_attachment_changed)$"
                ):
                    workflow._copy_zotero_attachment(self._row("paper-0001"), str(source), paths)
            published_payloads = [candidate.read_bytes() for candidate in paths.pdfs.glob("*.pdf")]

        self.assertGreaterEqual(hash_calls, 2)
        self.assertNotIn(external_bytes, published_payloads)
        self.assertEqual(published_payloads, [])

    def test_zotero_attachment_same_handle_snapshot_preserves_source(self) -> None:
        from paper_automation import batch_workflow as workflow

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "stable attachment.pdf"
            source_bytes = b"%PDF-1.7\nstable same-handle attachment"
            source.write_bytes(source_bytes)
            source_hash = hashlib.sha256(source_bytes).hexdigest()
            paths = self._paths_with_state(root / "run", [self._row("paper-0001")])

            published = workflow._copy_zotero_attachment(
                self._row("paper-0001"), str(source), paths
            )

            self.assertEqual(published.read_bytes(), source_bytes)
            self.assertTrue(source.exists())
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), source_hash)
            self.assertFalse(list(paths.pdfs.glob(".pdf_snapshot_*.tmp")))

    def test_report_generation_failure_leaves_existing_six_file_set_unchanged(self) -> None:
        from paper_automation import batch_workflow as workflow

        names = [
            "final_manifest.csv",
            "final_manifest.xlsx",
            "failed.csv",
            "run_summary.txt",
            "batch_status.csv",
            "batch_status.json",
        ]
        with tempfile.TemporaryDirectory() as tmp:
            paths = self._paths_with_state(Path(tmp), [self._row("paper-0001")])
            old_bytes = {name: f"old::{name}".encode("utf-8") for name in names}
            for name, content in old_bytes.items():
                (paths.reports / name).write_bytes(content)
            with patch.object(workflow, "_write_final_manifest_xlsx", side_effect=OSError("synthetic xlsx failure")):
                with self.assertRaisesRegex(RuntimeError, "final_report_write_failed:OSError:synthetic xlsx failure"):
                    workflow.write_final_reports(paths, [self._row("paper-0001")])
            after = {name: (paths.reports / name).read_bytes() for name in names}
            transient = list(paths.reports.glob(".final_reports_*"))

        self.assertEqual(after, old_bytes)
        self.assertEqual(transient, [])

    def test_report_publish_failure_rolls_back_entire_six_file_set(self) -> None:
        from paper_automation import batch_workflow as workflow

        names = [
            "final_manifest.csv",
            "final_manifest.xlsx",
            "failed.csv",
            "run_summary.txt",
            "batch_status.csv",
            "batch_status.json",
        ]
        with tempfile.TemporaryDirectory() as tmp:
            paths = self._paths_with_state(Path(tmp), [self._row("paper-0001")])
            old_names = {names[1], names[3], names[5]}
            old_bytes = {
                name: f"old::{name}".encode("utf-8")
                for name in old_names
            }
            for name, content in old_bytes.items():
                (paths.reports / name).write_bytes(content)
            real_replace = workflow.os.replace
            publish_count = 0
            failed_once = False

            def fail_third_publish(source, destination):
                nonlocal publish_count, failed_once
                source_path = Path(source)
                destination_path = Path(destination)
                is_publish = (
                    source_path.parent.name.startswith(".final_reports_staging_")
                    and destination_path.parent == paths.reports
                    and destination_path.name in names
                )
                if is_publish:
                    publish_count += 1
                    if publish_count == 3 and not failed_once:
                        failed_once = True
                        raise OSError("synthetic publish failure")
                return real_replace(source, destination)

            with patch.object(workflow.os, "replace", side_effect=fail_third_publish):
                with self.assertRaisesRegex(RuntimeError, "final_report_write_failed:OSError:synthetic publish failure"):
                    workflow.write_final_reports(paths, [self._row("paper-0001")])
            after = {name: (paths.reports / name).read_bytes() for name in old_names}
            absent_after = {
                name for name in names if name not in old_names and not (paths.reports / name).exists()
            }
            transient = list(paths.reports.glob(".final_reports_*"))

        self.assertEqual(publish_count, 3)
        self.assertEqual(after, old_bytes)
        self.assertEqual(absent_after, set(names) - old_names)
        self.assertEqual(transient, [])


class BatchEndToEndTests(unittest.TestCase):
    def test_start_resume_finalize_produces_one_manifest_and_valid_pdfs(self) -> None:
        import csv
        from io import BytesIO

        from openpyxl import load_workbook

        from paper_automation.batch_workflow import (
            finalize_batch,
            is_valid_pdf,
            resume_batch,
            start_batch,
        )

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            project_pdf = root / "project.pdf"
            project_payload = b"%PDF-1.7\nproject fixture"
            project_pdf.write_bytes(project_payload)
            zotero_pdf = root / "zotero.pdf"
            zotero_payload = b"%PDF-1.7\nzotero fixture"
            zotero_pdf.write_bytes(zotero_payload)
            source_hashes = {
                project_pdf: hashlib.sha256(project_payload).hexdigest(),
                zotero_pdf: hashlib.sha256(zotero_payload).hexdigest(),
            }
            normalized = [
                {
                    "task_id": "paper-0001",
                    "source_index": "1",
                    "input_doi": "10.1000/a",
                    "doi": "10.1000/a",
                    "title": "A",
                    "fixture_pdf": str(project_pdf),
                },
                {
                    "task_id": "paper-0002",
                    "source_index": "2",
                    "input_doi": "10.1000/b",
                    "doi": "10.1000/b",
                    "title": "B",
                },
                {
                    "task_id": "paper-0003",
                    "source_index": "3",
                    "input_doi": "10.1000/c",
                    "doi": "10.1000/c",
                    "title": "C",
                },
            ]

            def initial_updates(rows):
                return [
                    {
                        **rows[0],
                        "status": "oa_downloaded",
                        "source": "oa",
                        "file": rows[0]["fixture_pdf"],
                        "reason": "",
                    },
                    *[
                        {
                            **row,
                            "status": "captcha_required",
                            "source": "institutional",
                            "file": "",
                            "reason": "captcha_required",
                        }
                        for row in rows[1:]
                    ],
                ]

            gateway = FakeBatchGateway(
                initial_updates=initial_updates,
                retry_updates=lambda rows: [
                    {
                        **row,
                        "status": "no_entitlement",
                        "source": "institutional",
                        "file": "",
                        "reason": "retry_exhausted",
                    }
                    for row in rows
                ],
            )
            started = start_batch(
                input_text="fixture",
                input_path=None,
                output_root=root,
                gateway=gateway,
                normalizer=lambda **_kwargs: normalized,
                now=datetime(2026, 7, 11, 6, 0, 0),
            )
            resume_batch(started.paths.root, gateway=gateway)

            with started.paths.zotero_results.open("w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=[
                        "task_id",
                        "zotero_item_id",
                        "attachment_path",
                        "status",
                        "reason",
                    ],
                )
                writer.writeheader()
                writer.writerow(
                    {
                        "task_id": "paper-0002",
                        "zotero_item_id": "42",
                        "attachment_path": str(zotero_pdf),
                        "status": "downloaded",
                        "reason": "",
                    }
                )
                writer.writerow(
                    {
                        "task_id": "paper-0003",
                        "zotero_item_id": "43",
                        "attachment_path": "",
                        "status": "no_pdf",
                        "reason": "no_available_pdf",
                    }
                )

            finalized = finalize_batch(started.paths.root, started.paths.zotero_results)
            report_names = (
                "final_manifest.csv",
                "final_manifest.xlsx",
                "failed.csv",
                "run_summary.txt",
                "batch_status.csv",
                "batch_status.json",
            )
            reports_before_rerun = {
                name: (started.paths.reports / name).read_bytes()
                for name in report_names
            }
            state_before_rerun = started.paths.state.read_bytes()
            delivered_before_rerun = {
                path.name: path.read_bytes() for path in started.paths.pdfs.glob("*.pdf")
            }
            rerun = finalize_batch(started.paths.root, started.paths.zotero_results)
            reports_after_rerun = {
                name: (started.paths.reports / name).read_bytes()
                for name in report_names
            }
            state_after_rerun = started.paths.state.read_bytes()
            delivered_after_rerun = {
                path.name: path.read_bytes() for path in started.paths.pdfs.glob("*.pdf")
            }
            with (started.paths.reports / "final_manifest.csv").open(
                "r",
                newline="",
                encoding="utf-8-sig",
            ) as handle:
                final_rows = list(csv.DictReader(handle))

            self.assertEqual(gateway.retry_calls, 1)
            self.assertEqual(finalized.success_count, 2)
            self.assertEqual(finalized.failed_count, 1)
            self.assertEqual(rerun.success_count, 2)
            self.assertEqual(rerun.failed_count, 1)
            delivered = list(started.paths.pdfs.glob("*.pdf"))
            self.assertEqual(len(delivered), 2)
            self.assertTrue(all(is_valid_pdf(path) for path in delivered))
            self.assertEqual(
                {
                    hashlib.sha256(payload).hexdigest()
                    for payload in delivered_after_rerun.values()
                },
                {
                    hashlib.sha256(project_payload).hexdigest(),
                    hashlib.sha256(zotero_payload).hexdigest(),
                },
            )
            self.assertTrue((started.paths.reports / "final_manifest.xlsx").exists())
            self.assertIn(
                "no_available_pdf",
                (started.paths.reports / "run_summary.txt").read_text(encoding="utf-8"),
            )
            expected_task_ids = ["paper-0001", "paper-0002", "paper-0003"]
            self.assertEqual(len(final_rows), 3)
            self.assertCountEqual(
                [row["task_id"] for row in final_rows],
                expected_task_ids,
            )
            byte_stable_reports = set(report_names) - {"final_manifest.xlsx"}
            for name in byte_stable_reports:
                with self.subTest(idempotent_report=name):
                    self.assertEqual(
                        reports_before_rerun[name],
                        reports_after_rerun[name],
                    )

            def xlsx_rows(payload: bytes) -> list[tuple]:
                workbook = load_workbook(
                    BytesIO(payload),
                    read_only=True,
                    data_only=False,
                )
                try:
                    return list(workbook.active.iter_rows(values_only=True))
                finally:
                    workbook.close()

            self.assertEqual(
                xlsx_rows(reports_before_rerun["final_manifest.xlsx"]),
                xlsx_rows(reports_after_rerun["final_manifest.xlsx"]),
            )
            self.assertEqual(state_before_rerun, state_after_rerun)
            self.assertEqual(delivered_before_rerun, delivered_after_rerun)
            self.assertEqual(
                hashlib.sha256(project_pdf.read_bytes()).hexdigest(),
                source_hashes[project_pdf],
            )
            self.assertEqual(
                hashlib.sha256(zotero_pdf.read_bytes()).hexdigest(),
                source_hashes[zotero_pdf],
            )


class BatchCliTests(unittest.TestCase):
    @staticmethod
    def _ps_quote(value: object) -> str:
        return "'" + str(value).replace("'", "''") + "'"

    def _result(
        self,
        root: Path,
        *,
        success: int = 2,
        failed: int = 3,
        manual: int = 0,
        fallback: int = 0,
    ):
        from paper_automation.batch_workflow import BatchPaths, BatchRunResult

        paths = BatchPaths(
            root=root,
            pdfs=root / "pdfs",
            reports=root / "reports",
            working=root / "working",
            state=root / "working" / "batch_state.json",
            normalized_input=root / "working" / "normalized_input.csv",
            manual_retry=root / "working" / "manual_retry.csv",
            zotero_fallback=root / "working" / "zotero_fallback.csv",
            zotero_results=root / "working" / "zotero_results.csv",
        )
        return BatchRunResult(
            paths=paths,
            total_count=5,
            success_count=success,
            failed_count=failed,
            manual_retry_count=manual,
            zotero_fallback_count=fallback,
        )

    def test_start_calls_workflow_with_all_safe_options_and_prints_resume_command(self) -> None:
        from contextlib import redirect_stderr, redirect_stdout
        from io import StringIO

        from paper_batch import main

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "含 空格 O'Brien 的批次"
            secret = "cookie-secret-must-not-appear"
            stdout, stderr = StringIO(), StringIO()
            result = self._result(root, manual=1, fallback=2)
            with patch("paper_batch.start_batch", return_value=result) as start:
                with redirect_stdout(stdout), redirect_stderr(stderr):
                    exit_code = main([
                        "start", "--text", "10.1000/example", "--out", str(root),
                        "--run-name", "实验批次", "--email", "student@example.edu",
                        "--cookies", secret, "--browser-exe", "browser.exe",
                        "--login-wait-seconds", "12", "--debug-port", "9444",
                        "--throttle-seconds", "0.5",
                    ])

        self.assertEqual(exit_code, 0)
        self.assertEqual(start.call_args.kwargs["input_text"], "10.1000/example")
        self.assertIsNone(start.call_args.kwargs["input_path"])
        self.assertEqual(start.call_args.kwargs["output_root"], str(root))
        self.assertEqual(start.call_args.kwargs["run_name"], "实验批次")
        options = start.call_args.kwargs["options"]
        self.assertEqual(options.email, "student@example.edu")
        self.assertEqual(options.cookies, secret)
        self.assertEqual(options.browser_exe, "browser.exe")
        self.assertEqual(options.login_wait_seconds, 12)
        self.assertEqual(options.debug_port, 9444)
        self.assertEqual(options.throttle_seconds, 0.5)
        output = stdout.getvalue()
        for label in (
            "运行目录：", "最终 PDF 目录：", "总计：5", "成功：2", "失败：3",
            "待人工重试：1", "待 Zotero 回退：2", "人工重试清单：",
            "Zotero 回退清单：", "报告目录：",
        ):
            self.assertIn(label, output)
        script = (PROJECT_ROOT / "paper_batch.py").resolve()
        expected = (
            f"& {self._ps_quote(Path(sys.executable).resolve())} {self._ps_quote(script)} "
            f"'resume' '--run-dir' {self._ps_quote(root)}"
        )
        self.assertIn(expected, output)
        self.assertTrue(str(script).startswith(str(PROJECT_ROOT.resolve())))
        self.assertIn("O''Brien", output)
        self.assertNotIn(secret, output)
        self.assertNotIn(secret, stderr.getvalue())

    def test_resume_prints_zotero_finalize_command_when_no_manual_rows_remain(self) -> None:
        from contextlib import redirect_stdout
        from io import StringIO

        from paper_batch import main

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "resume O'Brien run"
            result = self._result(root, fallback=2)
            stdout = StringIO()
            with patch("paper_batch.resume_batch", return_value=result) as resume:
                with redirect_stdout(stdout):
                    exit_code = main(["resume", "--run-dir", str(root)])

        self.assertEqual(exit_code, 0)
        self.assertEqual(resume.call_args.args, (str(root),))
        output = stdout.getvalue()
        self.assertIn("请在 Zotero 中处理回退条目", output)
        script = (PROJECT_ROOT / "paper_batch.py").resolve()
        self.assertIn(
            f"& {self._ps_quote(Path(sys.executable).resolve())} {self._ps_quote(script)} "
            f"'finalize' '--run-dir' {self._ps_quote(root)} "
            f"'--zotero-results' {self._ps_quote(root / 'working' / 'zotero_results.csv')}",
            output,
        )

    def test_start_without_pending_rows_creates_header_only_zotero_results(self) -> None:
        from contextlib import redirect_stdout
        from io import StringIO

        from paper_batch import main

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "completed run"
            result = self._result(root)
            stdout = StringIO()
            with patch("paper_batch.start_batch", return_value=result):
                with redirect_stdout(stdout):
                    exit_code = main(["start", "--text", "10.1000/example", "--out", str(root)])

            results_path = root / "working" / "zotero_results.csv"
            raw = results_path.read_bytes()
            text = results_path.read_text(encoding="utf-8-sig")

        self.assertEqual(exit_code, 0)
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(
            text.splitlines(),
            ["task_id,zotero_item_id,attachment_path,status,reason"],
        )
        self.assertIn(str(results_path.resolve()), stdout.getvalue())
        self.assertIn("task_id,zotero_item_id,attachment_path,status,reason", stdout.getvalue())
        self.assertIn("'finalize' '--run-dir'", stdout.getvalue())

    def test_start_without_pending_rows_does_not_overwrite_existing_zotero_results(self) -> None:
        from contextlib import redirect_stdout
        from io import StringIO

        from paper_batch import main

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "existing results"
            results_path = root / "working" / "zotero_results.csv"
            results_path.parent.mkdir(parents=True)
            original = b"existing-content-must-stay"
            results_path.write_bytes(original)
            stdout = StringIO()
            with patch("paper_batch.start_batch", return_value=self._result(root)):
                with redirect_stdout(stdout):
                    exit_code = main(["start", "--text", "10.1000/example"])
            after = results_path.read_bytes()
            retry_paths = list(results_path.parent.glob("zotero_results_retry_*.csv"))
            retry_raw = retry_paths[0].read_bytes() if len(retry_paths) == 1 else b""
            retry_text = (
                retry_paths[0].read_text(encoding="utf-8-sig")
                if len(retry_paths) == 1
                else ""
            )
            output = stdout.getvalue()

        self.assertEqual(exit_code, 0)
        self.assertEqual(after, original)
        self.assertEqual(len(retry_paths), 1)
        retry_path = retry_paths[0]
        self.assertTrue(retry_raw.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(
            retry_text.splitlines(),
            ["task_id,zotero_item_id,attachment_path,status,reason"],
        )
        self.assertIn(str(retry_path.resolve()), output)
        self.assertNotIn(
            f"'--zotero-results' {self._ps_quote(results_path.resolve())}",
            output,
        )

    def test_start_without_pending_rows_reuses_valid_header_only_canonical(self) -> None:
        import csv
        from contextlib import redirect_stdout
        from io import StringIO

        from paper_automation.batch_workflow import ZOTERO_RESULT_FIELDS
        from paper_batch import main

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "valid canonical"
            results_path = root / "working" / "zotero_results.csv"
            results_path.parent.mkdir(parents=True)
            with results_path.open("x", newline="", encoding="utf-8-sig") as handle:
                csv.writer(handle).writerow(ZOTERO_RESULT_FIELDS)
            original = results_path.read_bytes()
            stdout = StringIO()
            with patch("paper_batch.start_batch", return_value=self._result(root)):
                with redirect_stdout(stdout):
                    exit_code = main(["start", "--text", "10.1000/example"])

            retry_paths = list(results_path.parent.glob("zotero_results_retry_*.csv"))
            after = results_path.read_bytes()
            output = stdout.getvalue()

        self.assertEqual(exit_code, 0)
        self.assertEqual(after, original)
        self.assertEqual(retry_paths, [])
        self.assertIn(
            f"'--zotero-results' {self._ps_quote(results_path.resolve())}",
            output,
        )

    def test_start_without_pending_rows_rejects_canonical_with_blank_record(self) -> None:
        from contextlib import redirect_stdout
        from io import StringIO

        from paper_batch import main

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "blank canonical record"
            results_path = root / "working" / "zotero_results.csv"
            results_path.parent.mkdir(parents=True)
            original = (
                b"\xef\xbb\xbf"
                b"task_id,zotero_item_id,attachment_path,status,reason\r\n\r\n"
            )
            results_path.write_bytes(original)
            stdout = StringIO()
            with patch("paper_batch.start_batch", return_value=self._result(root)):
                with redirect_stdout(stdout):
                    exit_code = main(["start", "--text", "10.1000/example"])

            retry_paths = list(results_path.parent.glob("zotero_results_retry_*.csv"))
            after = results_path.read_bytes()
            output = stdout.getvalue()

        self.assertEqual(exit_code, 0)
        self.assertEqual(after, original)
        self.assertEqual(len(retry_paths), 1)
        self.assertIn(str(retry_paths[0].resolve()), output)

    def test_start_without_pending_rows_avoids_same_second_retry_collision(self) -> None:
        from contextlib import redirect_stdout
        from io import StringIO

        from paper_batch import main

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "same second collision"
            results_path = root / "working" / "zotero_results.csv"
            results_path.parent.mkdir(parents=True)
            original = b"stale-canonical"
            results_path.write_bytes(original)
            with patch("paper_batch.start_batch", return_value=self._result(root)), patch(
                "paper_batch._retry_timestamp",
                return_value="20260711_050000",
                create=True,
            ):
                with redirect_stdout(StringIO()):
                    first_exit = main(["start", "--text", "10.1000/example"])
                with redirect_stdout(StringIO()):
                    second_exit = main(["start", "--text", "10.1000/example"])

            retry_names = sorted(path.name for path in results_path.parent.glob(
                "zotero_results_retry_*.csv"
            ))
            after = results_path.read_bytes()

        self.assertEqual((first_exit, second_exit), (0, 0))
        self.assertEqual(after, original)
        self.assertEqual(
            retry_names,
            [
                "zotero_results_retry_20260711_050000.csv",
                "zotero_results_retry_20260711_050000_2.csv",
            ],
        )

    def test_fallback_rows_do_not_create_zotero_results(self) -> None:
        from contextlib import redirect_stdout
        from io import StringIO

        from paper_batch import main

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "fallback run"
            stdout = StringIO()
            with patch("paper_batch.start_batch", return_value=self._result(root, fallback=1)):
                with redirect_stdout(stdout):
                    exit_code = main(["start", "--text", "10.1000/example"])
            exists = (root / "working" / "zotero_results.csv").exists()

        self.assertEqual(exit_code, 0)
        self.assertFalse(exists)

    def test_header_file_write_error_returns_two_without_secret_or_traceback(self) -> None:
        from contextlib import redirect_stderr, redirect_stdout
        from io import StringIO

        from paper_batch import main

        secret = "cookie-secret-in-path-must-not-leak"
        stdout, stderr = StringIO(), StringIO()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "write failure"
            with patch("paper_batch.start_batch", return_value=self._result(root)), patch(
                "paper_batch._ensure_header_only_zotero_results",
                side_effect=OSError(secret),
            ):
                with redirect_stdout(stdout), redirect_stderr(stderr):
                    exit_code = main(["start", "--text", "10.1000/example"])

        self.assertEqual(exit_code, 2)
        self.assertNotIn("Traceback", stderr.getvalue())
        self.assertNotIn(secret, stdout.getvalue())
        self.assertNotIn(secret, stderr.getvalue())

    def test_finalize_with_unresolved_rows_prints_recoverable_not_complete(self) -> None:
        from contextlib import redirect_stdout
        from io import StringIO

        from paper_batch import main

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "final run"
            results_csv = root / "working" / "zotero_results.csv"
            result = self._result(root, fallback=2)
            stdout = StringIO()
            with patch("paper_batch.finalize_batch", return_value=result) as finalize:
                with redirect_stdout(stdout):
                    exit_code = main([
                        "finalize", "--run-dir", str(root), "--zotero-results", str(results_csv),
                    ])

        self.assertEqual(exit_code, 0)
        self.assertEqual(finalize.call_args.args, (str(root), str(results_csv)))
        output = stdout.getvalue()
        self.assertNotIn("批次已完成", output)
        self.assertIn("报告已更新", output)
        self.assertIn("批次未完成且可恢复", output)
        self.assertIn("未解决数量：3", output)
        self.assertIn(f"最终 PDF 目录：{root / 'pdfs'}", output)
        self.assertNotIn("'resume' '--run-dir'", output)

    def test_finalize_without_unresolved_rows_prints_completion(self) -> None:
        from contextlib import redirect_stdout
        from io import StringIO

        from paper_batch import main

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "complete final run"
            results_csv = root / "working" / "zotero_results.csv"
            result = self._result(root, success=5, failed=0, fallback=0)
            stdout = StringIO()
            with patch("paper_batch.finalize_batch", return_value=result):
                with redirect_stdout(stdout):
                    exit_code = main([
                        "finalize", "--run-dir", str(root), "--zotero-results", str(results_csv),
                    ])

        self.assertEqual(exit_code, 0)
        output = stdout.getvalue()
        self.assertIn("批次已完成", output)
        self.assertNotIn("批次未完成且可恢复", output)

    def test_start_rejects_both_input_sources(self) -> None:
        from contextlib import redirect_stderr
        from io import StringIO

        from paper_batch import main

        with redirect_stderr(StringIO()):
            with self.assertRaises(SystemExit) as raised:
                main(["start", "--text", "10.1000/example", "--input", "papers.csv"])
        self.assertEqual(raised.exception.code, 2)

    def test_workflow_errors_return_two_without_traceback_or_cookie_secret(self) -> None:
        from contextlib import redirect_stderr, redirect_stdout
        from io import StringIO

        from paper_batch import main

        secret = "cookie-secret-must-not-appear"
        stdout, stderr = StringIO(), StringIO()
        with patch("paper_batch.start_batch", side_effect=ValueError(secret)):
            with redirect_stdout(stdout), redirect_stderr(stderr):
                exit_code = main([
                    "start", "--text", "10.1000/example", "--cookies", secret,
                ])

        self.assertEqual(exit_code, 2)
        self.assertIn("错误码 2", stderr.getvalue())
        self.assertNotIn("Traceback", stderr.getvalue())
        self.assertNotIn(secret, stdout.getvalue())
        self.assertNotIn(secret, stderr.getvalue())

    def test_known_workflow_error_codes_get_fixed_hints_without_secret_leak(self) -> None:
        from contextlib import redirect_stderr, redirect_stdout
        from io import StringIO

        from paper_batch import main

        cases = {
            "zotero_results_fields_invalid": "Zotero 结果 CSV 表头",
            "zotero_results_file_missing": "Zotero 结果文件",
            "zotero_result_status_invalid": "status",
            "invalid_batch_state": "批次状态",
            "cookies_must_be_path": "Cookie JSON 文件路径",
            "empty_input": "至少一条文献",
            "manual_retry_status_mismatch": "人工重试清单",
        }
        secret = "cookie=value;session=must-not-leak"
        for code, hint in cases.items():
            with self.subTest(code=code):
                stdout, stderr = StringIO(), StringIO()
                with patch("paper_batch.start_batch", side_effect=ValueError(f"{code}:{secret}")):
                    with redirect_stdout(stdout), redirect_stderr(stderr):
                        exit_code = main(["start", "--text", "10.1000/example"])
                self.assertEqual(exit_code, 2)
                self.assertIn(code, stderr.getvalue())
                self.assertIn(hint, stderr.getvalue())
                self.assertNotIn(secret, stdout.getvalue())
                self.assertNotIn(secret, stderr.getvalue())

    def test_generated_powershell_help_command_runs_from_another_directory(self) -> None:
        import subprocess

        from paper_batch import _powershell_command

        command = _powershell_command("--help")
        with tempfile.TemporaryDirectory() as tmp:
            completed = subprocess.run(
                ["powershell.exe", "-NoProfile", "-Command", command],
                cwd=tmp,
                capture_output=True,
                timeout=20,
                check=False,
            )

        self.assertEqual(
            completed.returncode,
            0,
            completed.stderr.decode(errors="replace"),
        )
        self.assertTrue(command.startswith("& '"))
        self.assertIn(self._ps_quote((PROJECT_ROOT / "paper_batch.py").resolve()), command)

    def test_help_is_chinese_and_describes_supported_safe_workflow(self) -> None:
        from contextlib import redirect_stdout
        from io import StringIO

        from paper_batch import main

        stdout = StringIO()
        with redirect_stdout(stdout):
            with self.assertRaises(SystemExit) as raised:
                main(["--help"])
        self.assertEqual(raised.exception.code, 0)
        output = stdout.getvalue()
        for text in ("TXT/MD/CSV/XLSX/XLSM", "只重试一次", "合法 OA/授权访问", "非破坏复制", "start", "resume", "finalize", "zotero"):
            self.assertIn(text, output)

    def test_zotero_waiting_returns_three_and_prints_one_action(self) -> None:
        from contextlib import redirect_stdout
        from io import StringIO

        from paper_automation.zotero_bridge import BridgeBatch, BridgeJob, BridgeRunResult
        from paper_batch import main

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "zotero waiting"
            job = BridgeJob(
                job_id="11111111-1111-4111-8111-111111111111",
                payload_sha256="a" * 64,
                request_path=root / "bridge" / "inbox" / "11111111-1111-4111-8111-111111111111.json",
                result_path=root / "bridge" / "outbox" / "11111111-1111-4111-8111-111111111111.result.json",
                run_dir=root,
                chunk_index=1,
                chunk_count=1,
            )
            waiting = BridgeRunResult(
                status="awaiting_confirmation",
                bridge=BridgeBatch(
                    run_id=root.name,
                    manifest_path=root / "working" / "zotero_bridge_jobs.json",
                    jobs=(job,),
                ),
                zotero_results=None,
                batch_result=None,
            )
            stdout = StringIO()
            with patch("paper_batch.run_zotero_bridge", return_value=waiting, create=True) as run:
                with redirect_stdout(stdout):
                    exit_code = main([
                        "zotero", "--run-dir", str(root), "--wait-seconds", "0",
                    ])

        self.assertEqual(exit_code, 3)
        self.assertEqual(run.call_args.args, (str(root),))
        self.assertEqual(run.call_args.kwargs, {"library_id": 1, "wait_seconds": 0})
        self.assertIn("请在 Zotero 中确认一次", stdout.getvalue())
        self.assertNotIn("'resume'", stdout.getvalue())

    def test_zotero_wait_seconds_outside_range_returns_two(self) -> None:
        from contextlib import redirect_stderr
        from io import StringIO

        from paper_batch import main

        with tempfile.TemporaryDirectory() as tmp:
            run_dir = Path(tmp) / "missing"
            for seconds in ("-1", "86401"):
                with self.subTest(seconds=seconds):
                    stderr = StringIO()
                    with redirect_stderr(stderr):
                        exit_code = main([
                            "zotero", "--run-dir", str(run_dir), "--wait-seconds", seconds,
                        ])
                    self.assertEqual(exit_code, 2)
                    self.assertIn("bridge_wait_seconds_invalid", stderr.getvalue())
                    self.assertNotIn("Traceback", stderr.getvalue())

    def test_zotero_invalid_result_returns_two_without_payload_leak(self) -> None:
        from contextlib import redirect_stderr, redirect_stdout
        from io import StringIO

        from paper_batch import main

        secret = '{"plugin_result":"must-not-leak"}'
        stdout, stderr = StringIO(), StringIO()
        with patch(
            "paper_batch.run_zotero_bridge",
            side_effect=ValueError(f"bridge_result_fields_invalid:{secret}"),
            create=True,
        ):
            with redirect_stdout(stdout), redirect_stderr(stderr):
                exit_code = main(["zotero", "--run-dir", "existing run"])

        self.assertEqual(exit_code, 2)
        self.assertIn("bridge_result_fields_invalid", stderr.getvalue())
        self.assertNotIn("Traceback", stderr.getvalue())
        self.assertNotIn(secret, stdout.getvalue())
        self.assertNotIn(secret, stderr.getvalue())
